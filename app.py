import heapq
import os
import re
import time
import json
import hashlib
import uuid
import io
import html
import math
from datetime import datetime
from collections import Counter
import matplotlib.pyplot as plt
import numpy as np
import pdfplumber
import docx
import plotly.graph_objects as go
import plotly.express as px
import streamlit as st
from fpdf import FPDF
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
from textblob import TextBlob
import nltk
from wordcloud import WordCloud
import pandas as pd

# Import authentication, shared HTML/escaping utilities, PDF export, and modules
import auth
import summarizer
import workspaces
import admin_panel
import pdf_export
from utils import (
    clean_plain_text,
    normalize_html,
    render_html,
    safe_text,
    truncate_title_word_boundary,
    unescape_text,
)

# ----------------------------------------------------
# NLTK Resource Initialization (Streamlit Cloud Safe)
# ----------------------------------------------------
for _resource in ("punkt_tab", "punkt", "averaged_perceptron_tagger_eng", "averaged_perceptron_tagger"):
    try:
        nltk.download(_resource, quiet=True)
    except Exception:
        pass

# ----------------------------------------------------
# Page Configuration
# ----------------------------------------------------
st.set_page_config(
    page_title="PaperIQ – AI Research Intelligence Platform",
    page_icon="📄",
    layout="wide",
    initial_sidebar_state="expanded"
)

# ----------------------------------------------------
# ----------------------------------------------------
# ----------------------------------------------------
# Centralized Score Calculation & Verification
# ----------------------------------------------------
def calculate_overall_score(language, coherence, argumentation, academic_style, readability):
    """
    Authoritative single calculation for PaperIQ overall composite score.
    Weights:
      Language Quality: 25%
      Structural Coherence: 25%
      Argumentation: 20%
      Academic Style: 15%
      Readability: 15%
    """
    def _validate_score(val):
        if val is None:
            return 0.0
        try:
            return max(0.0, min(100.0, float(val)))
        except (TypeError, ValueError):
            return 0.0

    lang_val = _validate_score(language)
    coh_val = _validate_score(coherence)
    arg_val = _validate_score(argumentation)
    sty_val = _validate_score(academic_style)
    rd_val = _validate_score(readability)

    comp = (
        lang_val * 0.25 +
        coh_val * 0.25 +
        arg_val * 0.20 +
        sty_val * 0.15 +
        rd_val * 0.15
    )
    return round(max(0.0, min(100.0, comp)), 1)

def get_overall_score(res_or_scores):
    """
    Authoritative single getter for overall score across all pages, reports, and DB records.
    Guarantees no conflicting score values anywhere in the application.
    """
    if res_or_scores is None:
        return 0.0
    if isinstance(res_or_scores, (int, float)):
        try:
            return round(max(0.0, min(100.0, float(res_or_scores))), 1)
        except (TypeError, ValueError):
            return 0.0
    if isinstance(res_or_scores, dict):
        if "overall_score" in res_or_scores and res_or_scores["overall_score"] is not None:
            try:
                return round(float(res_or_scores["overall_score"]), 1)
            except (TypeError, ValueError):
                pass
        if "score" in res_or_scores and res_or_scores["score"] is not None:
            try:
                return round(float(res_or_scores["score"]), 1)
            except (TypeError, ValueError):
                pass
        scores = res_or_scores.get("scores", res_or_scores)
        if isinstance(scores, dict):
            if "overall_score" in scores and scores["overall_score"] is not None:
                try:
                    return round(float(scores["overall_score"]), 1)
                except (TypeError, ValueError):
                    pass
            if "Composite" in scores and scores["Composite"] is not None:
                try:
                    return round(float(scores["Composite"]), 1)
                except (TypeError, ValueError):
                    pass
            # If component scores are present, calculate authoritative weighted score
            if any(k in scores for k in ("Language Quality", "Language", "Structural Coherence", "Coherence", "language_quality", "structural_coherence")):
                lang = scores.get("Language Quality", scores.get("Language", scores.get("language_quality", 0.0)))
                coh = scores.get("Structural Coherence", scores.get("Coherence", scores.get("structural_coherence", 0.0)))
                arg = scores.get("Argumentation", scores.get("Reasoning", scores.get("argumentation", 0.0)))
                sty = scores.get("Academic Style", scores.get("Sophistication", scores.get("academic_style", 0.0)))
                rd = scores.get("Readability", scores.get("readability", 0.0))
                return calculate_overall_score(lang, coh, arg, sty, rd)
    return 0.0

# ----------------------------------------------------
# Centralized Filename Normalizer & Title Resolver
# ----------------------------------------------------
def get_clean_base_filename(filename):
    """
    Sanitizes raw filenames, strips path traversal, duplicate extensions (.pdf_report.pdf),
    embedded report headers ('PaperIQ Analysis Report: ...'), score/grade fragments, and copy numbers.
    Returns the core manuscript identifier.
    """
    if not filename:
        return "Research_Paper"
    raw = str(filename).strip()
    
    # If the string contains an embedded file name like Gangadharan...pdf, extract it first!
    m_file = re.search(r"([A-Za-z0-9_\-\.]{4,}\.(?:pdf|docx|txt))", raw, re.I)
    if m_file:
        raw = m_file.group(1)
        
    # Strip any report/score clutter
    raw = re.sub(r"^(paperiq\s+analysis\s+report\s*[:\-–—]?\s*)+", "", raw, flags=re.I)
    raw = re.sub(r"^(paperiq[_\-\s]+report[_\-\s]*)+", "", raw, flags=re.I)
    raw = re.sub(r"^(report[_\-\s]*)+", "", raw, flags=re.I)
    raw = re.sub(r"(final\s+)?composite\s+score\s*[:\-–—]?\s*[\d.]+(\s*[\/\\]\s*\d+)?", "", raw, flags=re.I)
    raw = re.sub(r"score\s*[:\-–—]?\s*[\d.]+(\s*[\/\\]\s*\d+)?", "", raw, flags=re.I)
    raw = re.sub(r"grade\s*[:\-–—]?\s*[A-F][+\-]?", "", raw, flags=re.I)
    
    # Iteratively strip duplicate extensions and report suffixes
    base = raw
    changed = True
    while changed:
        changed = False
        lower = base.lower().strip()
        for ext in (".pdf", ".docx", ".txt", ".doc"):
            if lower.endswith(ext):
                base = base.strip()[:-len(ext)]
                lower = base.lower().strip()
                changed = True
        m = re.search(r"([_\-.]?(paperiq[_\-]?report|report))([_\-]?\d+)?$", base, re.I)
        if m:
            base = base[:m.start()]
            changed = True
            
    clean = re.sub(r"[^\w\-.]", "_", base).strip(" ._-")
    clean = re.sub(r"_+", "_", clean)
    if not clean or clean.lower() in ("report", "paperiq", "analysis", "paperiq_report", "research_paper"):
        clean = "Research_Paper"
    return clean

def format_clean_display_title(base):
    """Formats raw filename base into a professional human-readable title without generated suffixes."""
    base = unescape_text(base) if base else ""
    if not base or base == "Research_Paper":
        return "Research Paper"
    # Match author-year code patterns like Gangadharan45122023JEAI111549 -> Gangadharan et al. – Research Paper
    m = re.match(r"^([A-Z][a-z]{3,})\d+[A-Za-z0-9]*$", base)
    if m:
        return f"{m.group(1)} et al. – Research Paper"
    # Replace underscores and hyphens with spaces
    spaced = base.replace("_", " ").replace("-", " ")
    # Clean up double spaces
    spaced = re.sub(r"\s+", " ", spaced).strip()
    if spaced.isupper() or spaced.islower():
        spaced = spaced.title()
    return spaced

def get_clean_report_filename(filename):
    """
    Guarantees strictly formatted: PaperIQ_Report_<clean_slug>.pdf with only ONE prefix/suffix.
    Never produces repeated suffixes (e.g. .pdf_report.pdf_Report.pdf).
    """
    clean_base = get_clean_base_filename(filename)
    m = re.match(r"^([A-Z][a-z]{3,})\d+[A-Za-z0-9]*$", clean_base)
    if m:
        slug = f"{m.group(1)}_et_al"
    else:
        slug = re.sub(r"[^\w\-]", "_", clean_base).strip("_")
        slug = re.sub(r"_+", "_", slug)
    if not slug or slug.lower() in ("report", "paperiq", "analysis", "paperiq_report", "research_paper"):
        slug = "Research_Paper"
    return f"PaperIQ_Report_{slug}.pdf"

# ----------------------------------------------------
# Multi-Format Text Extraction (PDF • DOCX • TXT)
# ----------------------------------------------------
_EXTRACTED_PDF_TITLES = {}

def compute_text_counts(text: str):
    """
    Phase 1.2: Single source of truth for word count and sentence count on cleaned text.
    """
    return summarizer.compute_text_counts(text)

def clean_extracted_text(text):
    """
    Phase 1.1: Clean extracted text pipeline run BEFORE sectioning, scoring, and summarizing.
    """
    return summarizer.clean_extracted_text(text)

def extract_page_words_column_aware(page):
    """
    Extracts text from a single pdfplumber page while respecting two-column layouts.
    Prevents left and right columns from merging horizontally across the page gutter.
    """
    w = float(page.width or 600)
    h = float(page.height or 800)
    mid = w / 2.0

    raw_words = page.extract_words(x_tolerance=1.5, y_tolerance=3)
    if not raw_words:
        return page.extract_text(x_tolerance=2, y_tolerance=3) or ""

    sorted_by_top = sorted(raw_words, key=lambda x: (x["top"], x["x0"]))
    lines = []
    for wd in sorted_by_top:
        if not lines or abs(wd["top"] - lines[-1]["top_avg"]) > 3.8:
            lines.append({"top_avg": wd["top"], "words": [wd]})
        else:
            grp = lines[-1]["words"]
            grp.append(wd)
            lines[-1]["top_avg"] = sum(x["top"] for x in grp) / len(grp)

    classified_lines = []
    for ln in lines:
        lw = sorted(ln["words"], key=lambda x: x["x0"])
        y = ln["top_avg"]
        line_str = " ".join(x["text"] for x in lw).strip()

        if y < 65 or y > h - 55:
            classified_lines.append(("full", y, lw))
            continue

        if re.match(r"^[_=\-—]{10,}$", line_str) or re.match(r"^(Table|Fig\.|Figure|Source:)\s*", line_str, re.I):
            classified_lines.append(("full", y, lw))
            continue

        crosses = any(wd["x0"] < (mid - 5) and wd["x1"] > (mid + 5) for wd in lw)
        if not crosses:
            for idx in range(len(lw) - 1):
                if lw[idx]["x1"] <= (mid + 5) and lw[idx + 1]["x0"] >= (mid - 5):
                    gap = lw[idx + 1]["x0"] - lw[idx]["x1"]
                    if gap < 11.0:
                        crosses = True
                        break

        classified_lines.append(("full" if crosses else "col", y, lw))

    left_col_lines = sum(1 for t, _, lw in classified_lines if t == "col" and any(x["x1"] <= mid + 6 for x in lw))
    right_col_lines = sum(1 for t, _, lw in classified_lines if t == "col" and any(x["x0"] >= mid - 6 for x in lw))

    if left_col_lines < 4 or right_col_lines < 4:
        return "\n".join(" ".join(x["text"] for x in lw) for _, _, lw in classified_lines)

    blocks = []
    for t, y, lw in classified_lines:
        if not blocks or blocks[-1]["type"] != t:
            blocks.append({"type": t, "lines": [(y, lw)]})
        else:
            blocks[-1]["lines"].append((y, lw))

    output_lines = []
    for blk in blocks:
        if blk["type"] == "full":
            for _, lw in blk["lines"]:
                output_lines.append(" ".join(x["text"] for x in lw))
        else:
            left_lines = []
            right_lines = []
            for _, lw in blk["lines"]:
                l_words = [x for x in lw if (x["x0"] + x["x1"]) / 2.0 < mid]
                r_words = [x for x in lw if (x["x0"] + x["x1"]) / 2.0 >= mid]
                if l_words:
                    left_lines.append(" ".join(x["text"] for x in l_words))
                if r_words:
                    right_lines.append(" ".join(x["text"] for x in r_words))
            output_lines.extend(left_lines)
            output_lines.extend(right_lines)

    return "\n".join(output_lines)

def extract_title_from_pdf_object(pdf) -> str:
    """
    Phase 1.3: Detect paper title from:
    (1) PDF metadata title if sane, else
    (2) largest-font / top-of-page-1 title lines (excluding journal name, volume/issue,
        DOI, 'Original Research Article', running headers).
    """
    meta = getattr(pdf, "metadata", None) or {}
    meta_title = str(meta.get("Title") or meta.get("title") or "").strip()
    if (
        meta_title
        and len(meta_title) >= 12
        and len(meta_title.split()) >= 3
        and not meta_title.endswith(",")
        and not summarizer.is_citation_or_header_title_line(meta_title)
        and not summarizer._is_author_byline_line(meta_title)
    ):
        mt_low = meta_title.lower()
        bad_meta = [
            "untitled", "microsoft word", ".doc", ".pdf", "journal of", "volume ",
            "issue ", "article no", "original research", "paperiq", "analysis report"
        ]
        if not any(b in mt_low for b in bad_meta):
            return re.sub(r"\s+", " ", meta_title).strip(" ,;:-.")

    if not getattr(pdf, "pages", None):
        if hasattr(pdf, "read") or isinstance(pdf, (str, bytes, os.PathLike)):
            try:
                if hasattr(pdf, "seek"):
                    pdf.seek(0)
                with pdfplumber.open(pdf) as opened_pdf:
                    return extract_title_from_pdf_object(opened_pdf)
            except Exception:
                return ""
        return ""
    page1 = pdf.pages[0]
    p1_h = float(page1.height or 800)
    words = page1.extract_words(x_tolerance=1.5, y_tolerance=3, extra_attrs=["size"])
    if not words:
        return ""

    sorted_by_top = sorted(words, key=lambda x: (x["top"], x["x0"]))
    lines = []
    for wd in sorted_by_top:
        if not lines or abs(wd["top"] - lines[-1]["top_avg"]) > 4.0:
            lines.append({"top_avg": wd["top"], "words": [wd]})
        else:
            grp = lines[-1]["words"]
            grp.append(wd)
            lines[-1]["top_avg"] = sum(x["top"] for x in grp) / len(grp)

    header_noise_re = re.compile(
        r"(journal\s+of|international\s+journal|transactions\s+on|proceedings\s+of|"
        r"\bvolume\b|\bvol\.|\bissue\b|article\s+no|issn|isbn|\bdoi\b|http|www\.|"
        r"original\s+research\s+article|review\s+article|short\s+communication|"
        r"research\s+article|received\s*:|accepted\s*:|published\s*:|copyright|"
        r"all\s+rights\s+reserved|department\s+of|university|institute\s+of|college\s+of|"
        r"school\s+of|academy\s+of|past\s+name:)",
        re.IGNORECASE
    )

    candidate_lines = []
    for ln in lines:
        y = ln["top_avg"]
        if y > p1_h * 0.55:
            break
        lw = sorted(ln["words"], key=lambda x: x["x0"])
        text_line = " ".join(x["text"] for x in lw).strip()
        if not text_line:
            continue
        low = text_line.lower()
        if low.startswith("abstract") or low.startswith("keywords") or low.startswith("key words") or low.startswith("1. introduction"):
            break
        if (
            header_noise_re.search(text_line)
            or "@" in text_line
            or summarizer.is_citation_or_header_title_line(text_line)
            or summarizer._is_author_byline_line(text_line)
        ):
            continue
        if len(re.findall(r"[A-Za-z]{2,}", text_line)) == 0:
            continue
        max_sz = max(float(x.get("size", 10.0)) for x in lw)
        candidate_lines.append({"y": y, "text": text_line, "size": max_sz})

    if not candidate_lines:
        return ""

    best_size = max(c["size"] for c in candidate_lines)
    title_parts = []
    for c in candidate_lines:
        if abs(c["size"] - best_size) <= 0.8:
            title_parts.append(c["text"])
        elif title_parts:
            break

    full_title = re.sub(r"\s+", " ", " ".join(title_parts)).strip().strip(" ,;:-.")
    full_title, _ = summarizer.split_title_and_author_byline(full_title)
    full_title = full_title.strip(" ,;:-.")
    if (
        len(full_title) >= 10
        and len(full_title.split()) >= 2
        and not full_title.endswith(",")
        and not summarizer.is_citation_or_header_title_line(full_title)
        and not summarizer._is_author_byline_line(full_title)
    ):
        return full_title
    return ""

def extract_text_from_file(file):
    """
    Robust multi-format extraction for PDF, DOCX, and TXT documents.
    Detects scanned image-only PDFs, password encryption, and file limits.
    Performs column-aware PDF extraction, title detection, and running header/footer filtering.
    Enforces platform settings (max_upload_size_mb, allowed_formats) and logs processing_events.
    """
    t0 = time.perf_counter()
    raw_name = getattr(file, "name", "document.pdf") or "document.pdf"
    filename = raw_name.lower()
    text = ""
    is_scanned = False
    pdf_title = ""

    cur_u = None
    cur_role = "Student"
    cur_pg = "Analyze Paper"
    try:
        cur_u = st.session_state.get("username") or (st.session_state.get("user_data") or {}).get("username")
        cur_role = (st.session_state.get("user_data") or {}).get("role") or "Student"
        cur_pg = st.session_state.get("current_page") or "Analyze Paper"
    except Exception:
        pass

    file_ext = raw_name.rsplit(".", 1)[-1].upper() if "." in raw_name else "UNKNOWN"
    file_size_bytes = 0
    if hasattr(file, "size") and file.size is not None:
        file_size_bytes = int(file.size)
    elif hasattr(file, "getvalue"):
        try:
            raw_b = file.getvalue()
            if isinstance(raw_b, (bytes, bytearray)):
                file_size_bytes = len(raw_b)
        except Exception:
            file_size_bytes = 0
    file_size_kb = round(file_size_bytes / 1024.0, 2)

    def _log_proc(status_val: str, reason_val: str = ""):
        dur_ms = round((time.perf_counter() - t0) * 1000.0, 1)
        try:
            auth.log_processing_event(
                username=cur_u,
                file_type=file_ext,
                file_size_kb=file_size_kb,
                stage="extraction",
                status=status_val,
                failure_reason=reason_val,
                duration_ms=dur_ms,
            )
        except Exception:
            pass

    try:
        max_mb = int(auth.get_platform_setting("max_upload_size_mb", 30) or 30)
    except Exception:
        max_mb = 30
    max_bytes = max_mb * 1024 * 1024

    try:
        allowed_str = str(auth.get_platform_setting("allowed_formats", "PDF,DOCX,TXT") or "PDF,DOCX,TXT")
        allowed_exts = tuple("." + f.strip().lower().lstrip(".") for f in allowed_str.split(",") if f.strip())
        if not allowed_exts:
            allowed_exts = (".pdf", ".docx", ".txt")
    except Exception:
        allowed_exts = (".pdf", ".docx", ".txt")

    if not filename.endswith(allowed_exts):
        _log_proc("failed", "unsupported type")
        return "__ERROR_UNSUPPORTED_FORMAT__"

    # Check file size limit and empty file (0 bytes)
    if hasattr(file, "size") and file.size is not None:
        if file.size > max_bytes:
            _log_proc("failed", "too large")
            return "__ERROR_FILE_TOO_LARGE__"
        if file.size == 0:
            _log_proc("failed", "empty document")
            return "__ERROR_EMPTY_DOCUMENT__"
    elif hasattr(file, "getvalue"):
        raw_bytes = file.getvalue()
        if isinstance(raw_bytes, (bytes, bytearray)):
            if len(raw_bytes) > max_bytes:
                _log_proc("failed", "too large")
                return "__ERROR_FILE_TOO_LARGE__"
            if len(raw_bytes) == 0:
                _log_proc("failed", "empty document")
                return "__ERROR_EMPTY_DOCUMENT__"

    try:
        if hasattr(file, "seek"):
            try:
                file.seek(0)
            except Exception:
                pass

        if filename.endswith(".pdf"):
            try:
                with pdfplumber.open(file) as pdf:
                    try:
                        pdf_title = extract_title_from_pdf_object(pdf)
                        if pdf_title:
                            _EXTRACTED_PDF_TITLES[raw_name] = pdf_title
                            _EXTRACTED_PDF_TITLES[get_clean_base_filename(raw_name)] = pdf_title
                    except Exception:
                        pdf_title = ""

                    total_images = 0
                    page_texts = []
                    for page in pdf.pages:
                        extracted = extract_page_words_column_aware(page)
                        if not extracted:
                            extracted = page.extract_text(x_tolerance=2, y_tolerance=3, layout=False) or ""
                        page_texts.append(extracted or "")
                        if getattr(page, "images", None):
                            total_images += len(page.images)

                    if len(page_texts) >= 3:
                        candidate_headers = Counter()
                        candidate_footers = Counter()

                        for p_txt in page_texts:
                            p_lines = [l.strip() for l in p_txt.splitlines() if l.strip()]
                            if len(p_lines) >= 3:
                                for hl in p_lines[:2]:
                                    if len(hl) < 120:
                                        candidate_headers[hl.lower()] += 1
                                for fl in p_lines[-2:]:
                                    if len(fl) < 120:
                                        candidate_footers[fl.lower()] += 1

                        threshold = max(2, int(len(page_texts) * 0.35))
                        repeated_headers = {h for h, c in candidate_headers.items() if c >= threshold}
                        repeated_footers = {f for f, c in candidate_footers.items() if c >= threshold}
                        bad_lines = repeated_headers | repeated_footers

                        cleaned_pages = []
                        for p_txt in page_texts:
                            p_lines = p_txt.splitlines()
                            filtered = [l for l in p_lines if l.strip().lower() not in bad_lines]
                            cleaned_pages.append("\n".join(filtered))
                        text = "\n\n".join(cleaned_pages)
                    else:
                        text = "\n\n".join(page_texts)

                    if not text.strip():
                        is_scanned = True
            except Exception as pdf_err:
                err_str = str(pdf_err).lower()
                if any(w in err_str for w in ["password", "encrypted", "decrypt"]):
                    _log_proc("failed", "password-protected")
                    return "__ERROR_PASSWORD_PROTECTED__"
                _log_proc("failed", "corrupted file")
                return "__ERROR_CORRUPTED_PDF__"

            if is_scanned and not text.strip():
                _log_proc("failed", "scanned PDF, no text layer")
                return "__ERROR_SCANNED_PDF__"

        elif filename.endswith(".docx"):
            try:
                doc = docx.Document(file)
                paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]
                table_text = []
                for table in doc.tables:
                    for row in table.rows:
                        row_vals = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_vals:
                            table_text.append(" | ".join(row_vals))
                combined_docx = paragraphs + table_text
                text = "\n\n".join(combined_docx)
            except Exception:
                _log_proc("failed", "corrupted file")
                return "__ERROR_CORRUPTED_DOCX__"

        elif filename.endswith(".txt"):
            content = file.getvalue() if hasattr(file, "getvalue") else (file.read() if hasattr(file, "read") else file)
            if isinstance(content, bytes):
                for enc in ("utf-8", "latin-1", "cp1252"):
                    try:
                        text = content.decode(enc)
                        break
                    except Exception:
                        pass
                if not text:
                    text = content.decode("utf-8", errors="ignore")
            else:
                text = str(content)

    except Exception as ex:
        try:
            auth.log_app_error(ex, page=cur_pg, username=cur_u, role=cur_role)
        except Exception:
            pass
        _log_proc("failed", "corrupted file")
        return ""
    finally:
        if hasattr(file, "seek"):
            try:
                file.seek(0)
            except Exception:
                pass

    cleaned = clean_extracted_text(text)
    if not cleaned:
        _log_proc("failed", "empty document")
        return "__ERROR_EMPTY_DOCUMENT__"
    if pdf_title and not cleaned.startswith(pdf_title):
        cleaned = f"{pdf_title}\n\n{cleaned}"
    _log_proc("success", "")
    return cleaned

def extract_text_from_pdf(file):
    """Backwards compatibility alias for PDF files."""
    return extract_text_from_file(file)

def clean_text(text):
    if not text:
        return ""
    return clean_extracted_text(text)

# ----------------------------------------------------
# Section Detection & Heading Cleaning Pipeline
# ----------------------------------------------------
def is_report_noise_line(line):
    if not line:
        return True
    s = str(line).strip().lower()
    report_markers = [
        "paperiq", "analysis report", "composite score", "writing & structure score",
        "document metadata", "document name:", "evaluation scores", "grade:",
        "academic writing tone", "reading difficulty", "potential research areas",
        "corpus & lexical", "responsible ai", "analytical scoring framework",
        "final composite score", "score:"
    ]
    return any(marker in s for marker in report_markers)

def is_journal_or_metadata_noise(line):
    s = line.strip().lower()
    if len(s) < 2:
        return True
    if re.match(r"^[\d\s.,\-]+$", s):
        return True
    noise_tokens = [
        "volume", "issue", "issn", "isbn", "doi:", "http:", "https:",
        "all rights reserved", "authorized licensed", "ieee", "acm", "springer",
        "elsevier", "conference on", "proceedings of", "transactions on",
        "copyright", "©", "published by", "pp.", "pages ",
        "journal of", "international journal", "original research article",
        "review article", "short communication", "article no", "past name:",
        "authors' contributions", "article information", "open peer review"
    ]
    return any(token in s for token in noise_tokens)

def sanitize_detected_title(raw_title: str) -> str:
    """Strips accidental journal header prefixes or author suffixes from a title string,
    and unescapes any HTML entities (e.g., &#x005F;, &amp;) once."""
    if not raw_title:
        return ""
    t = re.sub(r"\s+", " ", unescape_text(raw_title)).strip()
    if summarizer._is_author_byline_line(t) or t.endswith(","):
        return ""
    t = re.sub(r"^\(Past\s+name:[^)]+\)\s*", "", t, flags=re.I)
    t = re.sub(r"^Journal\s+of\s+[A-Za-z\s]+International\s+", "", t, flags=re.I)
    t = re.sub(r"^Original\s+Research\s+Article\s+", "", t, flags=re.I)
    t, _ = summarizer.split_title_and_author_byline(t)
    # Strip trailing author list if merged (e.g. "Growth Performance of Nutri-Cereals in India Gangadharan S. a*...")
    t = re.sub(r"\s+[A-Z][a-z]+\s+[A-Z]\.\s*[a-z]?\*?\s*,.*$", "", t)
    t = t.strip(" ,;:-.")
    if summarizer._is_author_byline_line(t):
        return ""
    return t

def extract_paper_title_from_text(text):
    """
    Phase 1.3: Extracts the actual scholarly paper title from the document text.
    Inspects initial lines before Abstract, skipping journal metadata, author blocks,
    and generated PaperIQ report headers.
    """
    if not text:
        return None

    sample = text[:1800]
    m_docname = re.search(r"Document Name:\s*([^\n\r]+)", sample, re.I)
    if m_docname:
        clean_base = get_clean_base_filename(m_docname.group(1))
        return format_clean_display_title(clean_base)

    # First check summarizer's citation-aware title extractor
    sum_title = summarizer.extract_paper_title_from_text(text)
    if (
        sum_title
        and not is_report_noise_line(sum_title)
        and not summarizer.is_citation_or_header_title_line(sum_title)
        and not summarizer._is_author_byline_line(sum_title)
    ):
        clean_st = sanitize_detected_title(sum_title)
        if clean_st:
            return clean_st

    lines = [l.strip() for l in text.split("\n")[:30] if l.strip()]
    candidate_lines = []
    for line in lines:
        l_lower = line.lower()
        if any(l_lower.startswith(prefix) for prefix in [
            "abstract", "keywords", "key words", "1. introduction", "1 introduction", "i. introduction", "contents"
        ]):
            break
        if summarizer.is_citation_or_header_title_line(line) or summarizer._is_author_byline_line(line):
            continue
        line_clean = sanitize_detected_title(line)
        if not line_clean:
            continue
        lc_lower = line_clean.lower()
        if (
            is_journal_or_metadata_noise(line_clean)
            or is_report_noise_line(line_clean)
            or summarizer.is_citation_or_header_title_line(line_clean)
            or summarizer._is_author_byline_line(line_clean)
        ):
            continue
        if any(x in lc_lower for x in ["@", "http", "doi:", "issn", "isbn", "volume", "issue", "pages", "copyright", "all rights reserved"]):
            continue
        if any(aff in lc_lower for aff in ["department of", "university", "institute of", "college", "faculty of", "school of", "center for", "laboratory", "academy of"]):
            continue
        # Skip author list lines (contain multiple commas or initials like "S. a*")
        if re.search(r"\b[A-Z]\.\s*[a-z]?\*?\s*,", line_clean) or line_clean.count(",") >= 2:
            if candidate_lines:
                break
            continue

        words = line_clean.split()
        if candidate_lines and 1 <= len(words) <= 5 and len(line_clean) <= 45 and not line_clean.endswith("."):
            candidate_lines.append(line_clean)
            break
        if len(words) < 3 or len(line_clean) < 12:
            continue
        if 12 <= len(line_clean) <= 160 and len(words) <= 22:
            candidate_lines.append(line_clean)
            if len(candidate_lines) >= 2:
                break
    if candidate_lines:
        full_title = sanitize_detected_title(" ".join(candidate_lines))
        if 12 <= len(full_title) <= 180 and len(full_title.split()) >= 3:
            if (
                not is_report_noise_line(full_title)
                and not summarizer.is_citation_or_header_title_line(full_title)
                and not summarizer._is_author_byline_line(full_title)
            ):
                return full_title
    return None

def get_clean_display_title(filename, text=None, stored_title=None, file=None):
    """
    Phase 1.3: Central display title resolver following priority:
    1. PDF metadata / largest-font title from Page 1 (if available).
    2. Pre-stored or extracted actual paper title from manuscript text.
    3. Fallback to cleaned filename ONLY when nothing is found.
    Always returns clean plain text with HTML entities unescaped once.
    """
    filename = unescape_text(filename) if filename else filename
    base = get_clean_base_filename(filename)

    # 1. Check PDF title cache or live PDF object
    if filename and filename in _EXTRACTED_PDF_TITLES:
        cached_t = sanitize_detected_title(_EXTRACTED_PDF_TITLES[filename])
        if cached_t and not summarizer.is_citation_or_header_title_line(cached_t) and not summarizer._is_author_byline_line(cached_t):
            return cached_t
    if base and base in _EXTRACTED_PDF_TITLES:
        cached_b = sanitize_detected_title(_EXTRACTED_PDF_TITLES[base])
        if cached_b and not summarizer.is_citation_or_header_title_line(cached_b) and not summarizer._is_author_byline_line(cached_b):
            return cached_b
    if file is not None and getattr(file, "name", "").lower().endswith(".pdf"):
        try:
            if hasattr(file, "seek"):
                file.seek(0)
            with pdfplumber.open(file) as pdf:
                t_pdf = sanitize_detected_title(extract_title_from_pdf_object(pdf))
                if t_pdf and not summarizer.is_citation_or_header_title_line(t_pdf) and not summarizer._is_author_byline_line(t_pdf):
                    _EXTRACTED_PDF_TITLES[getattr(file, "name", filename)] = t_pdf
                    _EXTRACTED_PDF_TITLES[base] = t_pdf
                    return t_pdf
        except Exception:
            pass
        finally:
            if hasattr(file, "seek"):
                try:
                    file.seek(0)
                except Exception:
                    pass

    # 2. Check text-extracted title (preferred over legacy stored_title if stored_title had journal header noise or byline)
    if text:
        extracted = extract_paper_title_from_text(text)
        if extracted and not is_report_noise_line(extracted) and not summarizer.is_citation_or_header_title_line(extracted) and not summarizer._is_author_byline_line(extracted):
            return unescape_text(extracted)

    # 3. Check stored_title
    if stored_title:
        st_clean = sanitize_detected_title(unescape_text(stored_title).strip())
        fallback_disp = format_clean_display_title(base)
        if (
            len(st_clean) > 5
            and st_clean != fallback_disp
            and not is_report_noise_line(st_clean)
            and not summarizer.is_citation_or_header_title_line(st_clean)
            and not summarizer._is_author_byline_line(st_clean)
            and not re.search(r"score\s*:\s*[\d.]+", st_clean, re.I)
        ):
            return st_clean

    # 4. Fallback to cleaned filename ONLY when nothing is found
    return unescape_text(format_clean_display_title(base))

def load_analysis_into_session(item, source="history"):
    """
    Authoritatively loads an analysis history record into active session state.
    Ensures complete synchronization of results, filenames, titles, authoritative score,
    and records the active analysis source ('uploaded', 'history', or 'repository').
    """
    if not item:
        return False
    cur_u = st.session_state.get("username") or st.session_state.get("user_data", {}).get("username")
    item_u = item.get("username")
    rec_id_check = item.get("id") or item.get("analysis_id")
    if source == "history" and cur_u:
        if item_u and str(item_u).strip() != str(cur_u).strip():
            auth.record_audit_event(
                username=cur_u,
                action_type="UNAUTHORIZED_ACCESS",
                target=str(rec_id_check),
                details=f"Cross-user load blocked: user '{cur_u}' tried to view analysis of '{item_u}'",
                status="Denied"
            )
            st.error("Access Denied: You do not own this analysis record.")
            return False
        if not item_u and rec_id_check and not auth.verify_analysis_ownership(rec_id_check, cur_u):
            auth.record_audit_event(
                username=cur_u,
                action_type="UNAUTHORIZED_ACCESS",
                target=str(rec_id_check),
                details=f"Cross-user load blocked via DB ownership check for user '{cur_u}'",
                status="Denied"
            )
            st.error("Access Denied: You do not own this analysis record.")
            return False
    res = {}
    if item.get("results_json"):
        try:
            res = json.loads(item["results_json"])
        except Exception:
            res = {}

    authoritative_score = get_overall_score(item.get("score") if item.get("score") is not None else res)
    if authoritative_score == 0.0 and res:
        authoritative_score = get_overall_score(res)

    if not res:
        res = {
            "scores": {
                "Language Quality": item.get("language_score") or authoritative_score,
                "Structural Coherence": item.get("coherence_score") or authoritative_score,
                "Argumentation": item.get("argumentation_score") or authoritative_score,
                "Academic Style": item.get("academic_style_score") or authoritative_score,
                "Readability": item.get("readability_score") or authoritative_score,
                "Composite": authoritative_score,
                "overall_score": authoritative_score
            },
            "domain": item.get("domain", "General Academic"),
            "full_text": ""
        }

    if "scores" not in res:
        res["scores"] = {}
    res["scores"]["Composite"] = authoritative_score
    res["scores"]["overall_score"] = authoritative_score
    res["overall_score"] = authoritative_score

    authoritative_filename = item.get("original_filename") or item.get("filename") or "Paper.pdf"
    clean_base = get_clean_base_filename(authoritative_filename)
    display_title = get_clean_display_title(
        clean_base,
        text=res.get("full_text"),
        stored_title=item.get("display_title") or res.get("title")
    )
    res["title"] = display_title
    res["display_title"] = display_title

    if "stats" not in res or not isinstance(res.get("stats"), dict):
        counts = compute_text_counts(res.get("full_text", ""))
        res["stats"] = {
            "reading_difficulty": "Standard Academic",
            "reading_difficulty_desc": "Appropriate scholarly syntactic complexity",
            "word_count": counts["word_count"],
            "sentence_count": counts["sentence_count"]
        }

    clean_report_fn = item.get("report_filename") or get_clean_report_filename(authoritative_filename)

    # Fully update session state
    st.session_state["results"] = res
    st.session_state["active_analysis"] = res
    st.session_state["analysis_result"] = res
    st.session_state["analysis_results"] = res
    st.session_state["analysis_complete"] = True
    st.session_state["filename"] = authoritative_filename
    st.session_state["analyzed_filename"] = authoritative_filename
    st.session_state["paper_title"] = display_title
    st.session_state["overall_score"] = authoritative_score
    st.session_state["current_analysis_id"] = item.get("analysis_id") or str(item.get("id"))
    st.session_state["report_filename"] = clean_report_fn
    st.session_state["analysis_just_completed"] = authoritative_filename
    st.session_state["active_analysis_source"] = source
    st.session_state["chat_history"] = []
    
    pdf_bytes = item.get("pdf") or item.get("pdf_report_bytes")
    if pdf_bytes:
        st.session_state["pdf_bytes"] = bytes(pdf_bytes)

def clear_active_analysis_session():
    """
    Completely resets all active analysis state, uploaded document metadata,
    executive summaries, scores, and cached page data in session state.
    Does NOT affect saved history records in the database.
    """
    keys_to_clear = [
        "results",
        "active_analysis",
        "analysis_result",
        "analysis_results",
        "analysis_complete",
        "filename",
        "analyzed_filename",
        "paper_title",
        "overall_score",
        "current_analysis_id",
        "report_filename",
        "analysis_just_completed",
        "pdf_bytes",
        "chat_history",
        "summary_result",
        "summary_doc_text",
        "summary_doc_name",
        "summary_length",
        "active_analysis_source"
    ]
    for k in keys_to_clear:
        if k in st.session_state:
            st.session_state[k] = None
    st.session_state["analysis_complete"] = False
    st.session_state["chat_history"] = []
    st.session_state["summary_result"] = None
    st.session_state["summary_doc_text"] = ""
    st.session_state["summary_doc_name"] = ""
    st.session_state["summary_length"] = "Medium"
    st.session_state["active_analysis_source"] = None
    # Reset file uploader widget state by incrementing key
    st.session_state["uploader_key"] = st.session_state.get("uploader_key", 0) + 1

def clear_active_analysis():
    """Canonical alias for clearing active analysis session state."""
    clear_active_analysis_session()

def clear_chat_context():
    """Resets the Ask PaperIQ chat history and tracking state."""
    st.session_state["chat_history"] = []
    st.session_state["last_chat_doc_id"] = None

def clear_comparison_state():
    """Resets the Compare Papers selection and cached comparative state."""
    for k in ["compare_p1", "compare_p2", "compare_p3", "selected_compare_papers", "comparison_result", "comparison_data"]:
        if k in st.session_state:
            st.session_state[k] = None

def is_valid_analysis(res) -> bool:
    """Checks whether an object represents a complete, valid PaperIQ analysis result."""
    return isinstance(res, dict) and bool(res) and "scores" in res and "overall_score" in res

def has_any_valid_analysis() -> bool:
    """
    Checks whether active session state contains a valid PaperIQ analysis result.
    Does NOT automatically restore or hydrate old historical or deleted records
    from the database without explicit user action.
    """
    res = st.session_state.get("results") or st.session_state.get("active_analysis")
    return is_valid_analysis(res)

def has_matching_analysis_for_file(file_name: str) -> bool:
    """Checks whether the active session state or user repository contains a valid analysis matching the specified filename."""
    if not file_name:
        return False
        
    # Check if active session state matches
    if has_any_valid_analysis():
        saved_fn = st.session_state.get("filename")
        if saved_fn and (
            saved_fn == file_name 
            or get_clean_base_filename(saved_fn) == get_clean_base_filename(file_name)
        ):
            return True

    # Check user history for a matching record
    user_data = st.session_state.get("user_data") or {}
    username = user_data.get("username") or st.session_state.get("username") or st.session_state.get("current_user")
    if username:
        user_hist = st.session_state.get("analysis_history", {}).get(username)
        if not user_hist:
            try:
                user_hist = auth.get_user_history(username)
                if user_hist:
                    st.session_state.setdefault("analysis_history", {})[username] = user_hist
            except Exception:
                user_hist = []
        for rec in (user_hist or []):
            rec_fn = rec.get("filename") or rec.get("original_filename")
            if rec_fn and (
                rec_fn == file_name 
                or get_clean_base_filename(rec_fn) == get_clean_base_filename(file_name)
            ):
                load_analysis_into_session(rec, source="history")
                return True

    return False

def is_generated_report_text(sentence):
    """Detects whether a sentence originates from PaperIQ generated report metadata rather than original manuscript."""
    s_lower = sentence.lower()
    report_markers = [
        "academic writing tone", "evaluation - formality", "objectivity:", "scholarly hedging",
        "promotional stance", "author framing", "paperiq", "analytical scoring framework",
        "responsible ai notice", "language quality", "structural coherence",
        "composite writing score", "reading difficulty", "executive summary",
        "citations detected", "lexical composition", "sentences exceed"
    ]
    return any(marker in s_lower for marker in report_markers)

def clean_heading_spacing(title):
    t = str(title).strip()
    # Separate jammed all-caps words only when title is a single concatenated string without spaces
    if ' ' not in t and len(t) > 18 and t.isupper():
        tokens = [
            'EMBEDDING', 'GENERATION', 'DATASET', 'METHODOLOGY', 'PROPOSED',
            'FRAMEWORK', 'EXPERIMENTAL', 'EXPERIMENT', 'EVALUATION', 'RESULTS',
            'DISCUSSION', 'CONCLUSION', 'RELATED', 'WORK', 'SYSTEM', 'MODEL',
            'ARCHITECTURE', 'ALGORITHM', 'ANALYSIS', 'AND'
        ]
        pattern = '|'.join(tokens)
        t = re.sub(pattern, lambda m: m.group(0) + ' ', t).strip()
    t = re.sub(r'([a-z])([A-Z])', r'\1 \2', t)
    t = re.sub(r'\s+', ' ', t).strip()
    return t.title() if t.isupper() else t

def extract_sections(text):
    """
    Phase 1.4: Robust academic section detector supporting numbered (1., 1.1, I., II.),
    all-caps, title-case, and combined headings (e.g. '3. Results and Discussion').
    Preserves individual reference lines inside References and keeps numbered subsections
    inside their parent section unless they represent canonical standalone sections.
    """
    if not text:
        return {}
    lines = text.split("\n")
    sections = {}
    current_header = "Preamble"
    current_content = []
    in_references = False

    def _save_section(hdr, content_list):
        if not content_list and hdr == "Preamble":
            return
        is_ref_sec = any(k in hdr.lower() for k in ["reference", "bibliography", "works cited"])
        sep = "\n" if is_ref_sec else " "
        body = sep.join(c for c in content_list if c).strip()
        if body or hdr != "Preamble":
            if hdr in sections and sections[hdr]:
                sections[hdr] = (sections[hdr] + sep + body).strip()
            else:
                sections[hdr] = body

    for line in lines:
        line_s = line.strip()
        if not line_s:
            continue

        if not in_references and is_journal_or_metadata_noise(line_s):
            current_content.append(line_s)
            continue

        is_header = False
        candidate = line_s
        inline_body = ""

        colon_match = re.match(r"^([A-Z][A-Za-z\s]{2,28}):\s+(.*)$", line_s)
        if not in_references and colon_match:
            prefix_up = colon_match.group(1).strip().upper()
            if prefix_up in {
                "ABSTRACT", "INTRODUCTION", "BACKGROUND", "LITERATURE REVIEW", "RELATED WORK",
                "METHODOLOGY", "METHODS", "MATERIALS AND METHODS", "RESULTS", "DISCUSSION",
                "RESULTS AND DISCUSSION", "LIMITATIONS", "CONCLUSION", "CONCLUSIONS", "FUTURE WORK"
            }:
                is_header = True
                candidate = colon_match.group(1).strip()
                inline_body = colon_match.group(2).strip()

        if not is_header and summarizer.is_top_level_section_heading(line_s, in_references):
            is_header = True
            candidate = line_s.rstrip(".:- ").strip()
            if re.search(r"\b(?:REFERENCES|BIBLIOGRAPHY|WORKS\s+CITED)\b", candidate, re.I):
                in_references = True

        if is_header:
            cleaned_title = clean_heading_spacing(candidate)
            _save_section(current_header, current_content)
            current_header = cleaned_title
            current_content = [inline_body] if inline_body else []
        else:
            current_content.append(line_s)

    _save_section(current_header, current_content)
    return sections

def extract_reference_entries(ref_text: str):
    """
    Phase 1.4: Parses the References section into clean individual reference items.
    Never summarizes references into prose.
    """
    if not ref_text or not ref_text.strip():
        return []
    text = ref_text.strip()
    # Split by newline first if multi-line
    raw_lines = [l.strip() for l in text.split("\n") if l.strip()]
    if len(raw_lines) <= 1:
        # Split inline numbered references like "1. Author... 2. Author..." or "[1] Author..."
        parts = re.split(r"(?:^|\s+)(?=(?:\[\d+\]|\d+\.)\s+[A-Z])", text)
        raw_lines = [p.strip() for p in parts if p.strip()]

    entries = []
    for item in raw_lines:
        if re.match(r"^(?:REFERENCES|BIBLIOGRAPHY|WORKS\s+CITED)$", item, re.I):
            continue
        if len(item) < 6:
            continue
        entries.append(re.sub(r"\s+", " ", item).strip())
    return entries

def summarize_text(text, num_sentences=3):
    if not text:
        return "No content to summarize."
    sents = summarizer.split_into_sentences(text)
    if not sents:
        blob = TextBlob(text)
        sents = [str(s).strip() for s in blob.sentences if str(s).strip()]
    if len(sents) <= num_sentences:
        return " ".join(sents) if sents else text
    word_frequencies = {}
    stop_words = set([
        "the", "is", "in", "and", "to", "of", "a", "for", "on", "with",
        "as", "by", "at", "this", "that", "it", "from", "an", "be", "are", "was"
    ])
    for word in re.findall(r"\b[A-Za-z]{3,}\b", text.lower()):
        if word not in stop_words:
            word_frequencies[word] = word_frequencies.get(word, 0) + 1
    if not word_frequencies:
        return " ".join(sents[:num_sentences])
    max_frequency = max(word_frequencies.values())
    for word in word_frequencies:
        word_frequencies[word] = word_frequencies[word] / max_frequency
    sentence_scores = {}
    for idx, sent in enumerate(sents):
        for word in re.findall(r"\b[A-Za-z]{3,}\b", sent.lower()):
            if word in word_frequencies:
                sentence_scores[idx] = sentence_scores.get(idx, 0.0) + word_frequencies[word]
    top_indices = sorted(heapq.nlargest(num_sentences, sentence_scores, key=sentence_scores.get))
    return " ".join([sents[i] for i in top_indices])

def get_important_sentences(text, num_sentences=3):
    sents = summarizer.split_into_sentences(text)
    if not sents:
        sents = [s.strip() for s in re.split(r"(?<=[.!?]) +", text) if s.strip()]
    if len(sents) <= num_sentences:
        return sents
    word_freq = {}
    words = re.findall(r"\w+", text.lower())
    for word in words:
        word_freq[word] = word_freq.get(word, 0) + 1
    sentence_scores = {}
    for sentence in sents:
        for word in re.findall(r"\w+", sentence.lower()):
            if word in word_freq:
                sentence_scores[sentence] = sentence_scores.get(sentence, 0) + word_freq[word]
    important = heapq.nlargest(num_sentences, sentence_scores, key=sentence_scores.get)
    return important

# ----------------------------------------------------
# Document Scoring & Readability Engine
# ----------------------------------------------------
def count_syllables(word):
    word = word.lower()
    vowels = "aeiou"
    count = 0
    if not word:
        return 0
    if word[0] in vowels:
        count += 1
    for index in range(1, len(word)):
        if word[index] in vowels and word[index - 1] not in vowels:
            count += 1
    if word.endswith("e"):
        count -= 1
    if count <= 0:
        count = 1
    return count

def calculate_academic_readability(text):
    """
    Computes academic readability normalized for scholarly literature using the
    single source of truth for word and sentence counts.
    """
    counts = compute_text_counts(text)
    words = counts["words"]
    total_words = counts["word_count"]
    total_sentences = counts["sentence_count"]
    if total_sentences == 0 or total_words == 0:
        return 70.0, "Moderate Academic", "Sufficient prose for evaluation."

    syllables = sum(count_syllables(w) for w in words)
    asl = total_words / max(1, total_sentences)
    asw = syllables / max(1, total_words)
    raw_flesch = 206.835 - (1.015 * asl) - (84.6 * asw)

    if raw_flesch >= 55:
        norm_score = min(100.0, 75.0 + (raw_flesch - 55) * 0.5)
        difficulty = "Accessible Academic"
        desc = "Clear thematic accessibility with accessible sentence syntax."
    elif raw_flesch >= 35:
        norm_score = 80.0 + ((raw_flesch - 35) / 20.0) * 15.0
        difficulty = "Standard Academic"
        desc = "Well-balanced scholarly prose with conventional sentence density."
    elif raw_flesch >= 20:
        norm_score = 65.0 + ((raw_flesch - 20) / 15.0) * 15.0
        difficulty = "Dense Academic"
        desc = "Dense syntactic construction typical of peer-reviewed journal papers."
    elif raw_flesch >= 0:
        norm_score = 50.0 + (raw_flesch / 20.0) * 15.0
        difficulty = "Highly Complex"
        desc = "Elaborate multi-clause propositions requiring concentrated reviewer attention."
    else:
        norm_score = max(35.0, 50.0 + (raw_flesch * 0.25))
        difficulty = "Very Dense / Convoluted"
        desc = "Extremely long compound clauses; recommend segmenting into concise statements."

    return round(float(norm_score), 2), difficulty, desc

def calculate_readability(text):
    """Backwards compatibility alias returning normalized readability score."""
    score, _, _ = calculate_academic_readability(text)
    return score

def get_grade(score):
    """Returns PaperIQ Evaluation Grade (A, B, C, D, E)."""
    if score >= 85:
        return "A"
    elif score >= 70:
        return "B"
    elif score >= 55:
        return "C"
    elif score >= 40:
        return "D"
    else:
        return "E"

def interpret_score(score, metric_name):
    """Provides validated academic rating and contextual interpretation based on real score."""
    if score >= 85:
        rating = "Excellent"
        color = "#10B981"  # Emerald
    elif score >= 75:
        rating = "Very Good"
        color = "#3B82F6"  # Blue
    elif score >= 65:
        rating = "Good"
        color = "#8B5CF6"  # Purple
    elif score >= 50:
        rating = "Fair"
        color = "#F59E0B"  # Amber
    else:
        rating = "Needs Improvement"
        color = "#EF4444"  # Red

    descriptions = {
        "Language Quality": {
            "Excellent": "Strong academic vocabulary, clear sentence structure, and confident tone.",
            "Very Good": "Fluent academic writing with sound grammar and balanced syntax.",
            "Good": "Generally clear language with minor phrasing inconsistencies.",
            "Fair": "Passable language quality; some informal phrases or syntax irregularities.",
            "Needs Improvement": "Syntax and vocabulary require revision for scholarly rigor."
        },
        "Structural Coherence": {
            "Excellent": "Seamless logical transitions between paragraphs and concepts.",
            "Very Good": "Consistent flow of ideas with effective use of discourse markers.",
            "Good": "Understandable narrative; a few abrupt section shifts.",
            "Fair": "Paragraph linkages are somewhat disjointed; needs better transitional flow.",
            "Needs Improvement": "Fragmented arguments; lacks overarching structural narrative."
        },
        "Argumentation": {
            "Excellent": "Rigorously backed propositions with robust evidence and causal explanations.",
            "Very Good": "Solid analytical deductions with clear supporting data.",
            "Good": "Valid arguments present, though some claims lack direct evidentiary support.",
            "Fair": "Causal links are loosely asserted; requires stronger empirical citations.",
            "Needs Improvement": "Unsupported claims; argumentative chain needs fundamental grounding."
        },
        "Academic Style": {
            "Excellent": "Rich lexical diversity with specialized scholarly terminology.",
            "Very Good": "Substantial technical depth and appropriate domain-specific nomenclature.",
            "Good": "Adequate vocabulary with moderate variation of conceptual terms.",
            "Fair": "Repetitive vocabulary; relies on generic rather than precise academic verbs.",
            "Needs Improvement": "Limited lexical variety; significant repetition of basic phrasing."
        },
        "Readability": {
            "Excellent": "Optimal readability for scholarly evaluation without undue density.",
            "Very Good": "Well-balanced sentence rhythm and clear thematic accessibility.",
            "Good": "Readable academic prose; occasional dense clauses.",
            "Fair": "Dense syntactic structures typical of specialized research.",
            "Needs Improvement": "Extremely convoluted phrasing; sentences exceed typical readability thresholds."
        }
    }
    alias_map = {
        "Language": "Language Quality",
        "Coherence": "Structural Coherence",
        "Reasoning": "Argumentation",
        "Sophistication": "Academic Style"
    }
    canonical = alias_map.get(metric_name, metric_name)
    interp = descriptions.get(canonical, {}).get(rating, "Analytical metric calculated from document characteristics.")
    return rating, color, interp

def get_dimension_metric_explanation(dim_name: str, res: dict) -> str:
    """
    Phase 2: Returns a brief, data-grounded explanation showing which measured metric(s)
    drove each dimension score.
    """
    if not isinstance(res, dict):
        return ""
    stats = res.get("stats") or {}
    explanations = stats.get("dimension_explanations") or {}
    if dim_name in explanations:
        return explanations[dim_name]

    tone = res.get("tone") or {}
    cits = (res.get("citation_analysis") or {}).get("total_citations", stats.get("citation_count", 0))
    asl = stats.get("avg_sentence_len", 20.0)
    awl = stats.get("avg_word_len", 5.0)
    ttr = stats.get("vocab_diversity", 0.4)
    cwr = stats.get("complex_word_ratio", 0.25)
    diff = stats.get("reading_difficulty", "Standard Academic")
    secs = res.get("sections") or {}
    core_sec_cnt = sum(1 for cs in ["abstract", "introduction", "method", "result", "conclusion"] if any(cs in s.lower() for s in secs.keys()))
    trans_cnt = stats.get("transition_count", 0)
    arg_cnt = stats.get("argument_markers", 0)

    if dim_name == "Language Quality":
        return f"Driven by {asl:.1f}-word avg sentence length, {awl:.1f}-char avg word length, and {ttr:.2f} lexical diversity."
    elif dim_name == "Structural Coherence":
        return f"Driven by {core_sec_cnt}/5 core academic sections detected and {trans_cnt} discourse transitions."
    elif dim_name == "Argumentation":
        return f"Driven by {cits} in-text citations and {arg_cnt} causal/evidentiary reasoning markers."
    elif dim_name == "Academic Style":
        return f"Driven by {cwr * 100:.0f}% complex vocabulary ratio and {tone.get('objectivity_score', 80):.0f}/100 objectivity."
    elif dim_name == "Readability":
        return f"Driven by {diff} syntax profile ({asl:.1f} words/sentence)."
    return ""

def analyze_academic_writing_tone(text, sections):
    """
    Reworked Academic Writing Tone Analysis replacing generic positive/negative sentiment.
    Evaluates formality, objectivity, hedging, claim strength, promotional tone, and author framing.
    """
    text_lower = text.lower()
    words = re.findall(r"\b[a-zA-Z]+\b", text_lower)
    word_count = max(1, len(words))
    
    # 1. Formality (Contraction avoidance & formal diction)
    contractions = ["don't", "can't", "won't", "didn't", "it's", "they're", "we've", "isn't", "aren't"]
    contraction_count = sum(len(re.findall(r"\b" + re.escape(c) + r"\b", text_lower)) for c in contractions)
    formality_score = min(100.0, max(0.0, 100.0 - (contraction_count * 5)))
    formality_status = "Formal Academic" if contraction_count == 0 else ("Generally Formal" if contraction_count <= 2 else "Contains Colloquial Contractions")
    
    # 2. Objectivity (Absence of emotive / subjective adjectives)
    blob = TextBlob(text)
    subjectivity = round(float(blob.sentiment.subjectivity), 2)
    objectivity_score = round(max(0.0, min(100.0, (1.0 - subjectivity) * 100)), 1)
    if objectivity_score >= 70:
        objectivity_status = "High Objectivity (Empirically Centered)"
    elif objectivity_score >= 50:
        objectivity_status = "Balanced (Moderate Subjective Commentary)"
    else:
        objectivity_status = "Subjective / Interpretive Framing"
        
    # 3. Hedging (Cautious scholarly assertions / epistemic modality)
    hedging_terms = [
        "suggests", "indicates", "appears to", "may indicate", "potential",
        "plausible", "tentatively", "presumably", "partially", "likely",
        "might", "could", "to our knowledge"
    ]
    hedge_count = sum(len(re.findall(r"\b" + re.escape(h) + r"\b", text_lower)) for h in hedging_terms)
    hedge_density = round((hedge_count / word_count) * 1000, 2)
    if hedge_density >= 4.0:
        hedging_status = "Robust Scholarly Hedging (Cautious Claims)"
    elif hedge_density >= 1.5:
        hedging_status = "Balanced Hedging"
    else:
        hedging_status = "Sparse Hedging (Claims Stated Directly)"
        
    # 4. Claim Strength (Definitive vs balanced claims)
    strong_claim_terms = [
        "proves", "proven", "undeniable", "incontestable", "irrefutable",
        "absolute truth", "completely solves", "flawless", "perfect solution"
    ]
    strong_claim_count = sum(len(re.findall(r"\b" + re.escape(sc) + r"\b", text_lower)) for sc in strong_claim_terms)
    claim_strength_status = "Balanced Academic Assertions" if strong_claim_count == 0 else f"{strong_claim_count} Absolute Assertions Detected"
    
    # 5. Promotional / Hype Language
    promo_terms = [
        "revolutionary", "groundbreaking", "game-changing", "miraculous",
        "unprecedented leap", "unbelievable", "world-changing", "disruptive breakthrough"
    ]
    promo_count = sum(len(re.findall(r"\b" + re.escape(p) + r"\b", text_lower)) for p in promo_terms)
    promo_status = "Neutral Academic Stance (No Marketing Hype)" if promo_count == 0 else f"Promotional Tone Detected ({promo_count} hype phrases)"
    
    # 6. First-Person Framing
    first_person_singular = len(re.findall(r"\b(i|my|me|mine)\b", text_lower))
    first_person_plural = len(re.findall(r"\b(we|our|us|ours)\b", text_lower))
    if first_person_singular > 0:
        fp_status = f"Personal Voice Detected ({first_person_singular} single, {first_person_plural} collective)"
    elif first_person_plural > 0:
        fp_status = f"Author-Collective Framing ('we'/'our' used {first_person_plural} times)"
    else:
        fp_status = "Third-Person Passive Voice (Impersonal Academic)"
        
    return {
        "formality_score": formality_score,
        "formality_status": formality_status,
        "contraction_count": contraction_count,
        "objectivity_score": objectivity_score,
        "objectivity_status": objectivity_status,
        "subjectivity_raw": subjectivity,
        "hedging_count": hedge_count,
        "hedging_density": hedge_density,
        "hedging_status": hedging_status,
        "claim_strength_status": claim_strength_status,
        "strong_claim_count": strong_claim_count,
        "promotional_count": promo_count,
        "promotional_status": promo_status,
        "first_person_singular": first_person_singular,
        "first_person_plural": first_person_plural,
        "first_person_status": fp_status
    }

def analyze_full_document(text):
    t0 = time.perf_counter()
    cur_u = None
    cur_role = "Student"
    cur_pg = "Analyze Paper"
    try:
        cur_u = st.session_state.get("username") or (st.session_state.get("user_data") or {}).get("username")
        cur_role = (st.session_state.get("user_data") or {}).get("role") or "Student"
        cur_pg = st.session_state.get("current_page") or "Analyze Paper"
    except Exception:
        pass
    size_kb = round(len((text or "").encode("utf-8", errors="ignore")) / 1024.0, 2)

    if not text or len(text.strip()) < 50:
        try:
            auth.log_processing_event(
                username=cur_u, file_type="TEXT", file_size_kb=size_kb,
                stage="analysis", status="failed", failure_reason="insufficient text",
                duration_ms=round((time.perf_counter() - t0) * 1000.0, 1)
            )
        except Exception:
            pass
        return None

    try:
        cleaned = clean_extracted_text(text)
        if not cleaned or len(cleaned.strip()) < 50:
            try:
                auth.log_processing_event(
                    username=cur_u, file_type="TEXT", file_size_kb=size_kb,
                    stage="analysis", status="failed", failure_reason="insufficient text",
                    duration_ms=round((time.perf_counter() - t0) * 1000.0, 1)
                )
            except Exception:
                pass
            return None

        counts = compute_text_counts(cleaned)
        words = counts["words"]
        sentences = counts["sentences"]
        word_count = counts["word_count"]
        sentence_count = counts["sentence_count"]
        if sentence_count == 0 or word_count < 15:
            try:
                auth.log_processing_event(
                    username=cur_u, file_type="TEXT", file_size_kb=size_kb,
                    stage="analysis", status="failed", failure_reason="insufficient text",
                    duration_ms=round((time.perf_counter() - t0) * 1000.0, 1)
                )
            except Exception:
                pass
            return None

        blob = TextBlob(cleaned)
        avg_sentence_len = word_count / max(1, sentence_count)
        avg_word_len = float(np.mean([len(w) for w in words])) if words else 5.0

        sections_data = extract_sections(cleaned)
        tone_data = analyze_academic_writing_tone(cleaned, sections_data)

        citation_patterns = [r"\(\d{4}\)", r"\[\d+\]", r"\(\w+ et al\., \d{4}\)"]
        citation_count = sum(len(re.findall(p, cleaned)) for p in citation_patterns)

        # 1. Language Quality (0-100)
        len_penalty = min(20.0, abs(float(avg_sentence_len) - 22.0) * 0.9)
        word_adj = min(15.0, max(-10.0, (float(avg_word_len) - 5.0) * 7.5))
        contractions_penalty = min(20.0, float(tone_data.get("contraction_count", 0)) * 3.5)
        ttr = len(set([w.lower() for w in words])) / max(1, word_count)
        ttr_adj = min(10.0, max(-5.0, (ttr - 0.35) * 25.0))
        language_score = min(98.0, max(30.0, 76.0 - len_penalty + word_adj - contractions_penalty + ttr_adj))

        # 2. Structural Coherence (0-100)
        transitions = [
            "however", "therefore", "thus", "consequently", "furthermore",
            "meanwhile", "moreover", "in addition", "in contrast", "subsequently",
            "accordingly", "specifically", "conversely", "notably"
        ]
        transition_count = sum(cleaned.lower().count(t) for t in transitions)
        trans_density = transition_count / max(1, sentence_count)
        density_score = min(40.0, trans_density * 160.0)
        canonical_sections = ["abstract", "introduction", "method", "result", "conclusion"]
        sec_matches = sum(1 for cs in canonical_sections if any(cs in s.lower() for s in sections_data.keys()))
        sec_bonus = min(24.0, sec_matches * 5.0)
        coherence_score = min(98.0, max(30.0, 36.0 + density_score + sec_bonus))

        # 3. Argumentation (0-100)
        argumentation_keywords = [
            "because", "since", "implies", "due to", "as a result", "evidence",
            "demonstrates", "proves", "corroborates", "substantiates", "suggests",
            "indicates", "validates", "confirms", "illustrates"
        ]
        arg_count = sum(cleaned.lower().count(k) for k in argumentation_keywords)
        arg_density = arg_count / max(1, sentence_count)
        arg_density_score = min(42.0, arg_density * 170.0)
        cit_bonus = min(25.0, float(citation_count) * 2.5)
        argumentation_score = min(98.0, max(30.0, 32.0 + arg_density_score + cit_bonus))

        # 4. Academic Style (0-100)
        complex_words = [w for w in words if len(w) > 6]
        comp_ratio = len(complex_words) / max(1, word_count)
        style_base = 45.0 + min(36.0, max(0.0, (comp_ratio - 0.18) * 150.0))
        obj_score = float(tone_data.get("objectivity_score", 70.0))
        obj_adj = min(12.0, max(-10.0, (obj_score - 60.0) * 0.3))
        promo_penalty = min(20.0, float(tone_data.get("promotional_count", 0)) * 4.0)
        style_score = min(98.0, max(30.0, style_base + obj_adj - promo_penalty))

        # 5. Readability (0-100 normalized)
        readability_score, reading_difficulty, reading_desc = calculate_academic_readability(cleaned)

        final_score = calculate_overall_score(
            language_score, coherence_score, argumentation_score, style_score, readability_score
        )

        vocab_div = round(ttr, 2) if word_count else 0
        comp_word_ratio = round(comp_ratio, 2) if word_count else 0

        dimension_explanations = {
            "Language Quality": f"Driven by {avg_sentence_len:.1f}-word avg sentence length, {avg_word_len:.1f}-char avg word length, and {vocab_div:.2f} lexical diversity.",
            "Structural Coherence": f"Driven by {sec_matches}/5 core academic sections detected and {transition_count} discourse transitions.",
            "Argumentation": f"Driven by {citation_count} in-text citations and {arg_count} causal/evidentiary reasoning markers.",
            "Academic Style": f"Driven by {comp_word_ratio * 100:.0f}% complex vocabulary ratio and {obj_score:.0f}/100 objectivity.",
            "Readability": f"Driven by {reading_difficulty} syntax profile ({avg_sentence_len:.1f} words/sentence)."
        }

        stats = {
            "word_count": word_count,
            "sentence_count": sentence_count,
            "avg_sentence_len": round(float(avg_sentence_len), 2),
            "avg_word_len": round(float(avg_word_len), 2),
            "vocab_diversity": vocab_div,
            "complex_word_ratio": comp_word_ratio,
            "reading_difficulty": reading_difficulty,
            "reading_difficulty_desc": reading_desc,
            "transition_count": transition_count,
            "detected_core_sections": sec_matches,
            "argument_markers": arg_count,
            "citation_count": citation_count,
            "dimension_explanations": dimension_explanations,
        }

        try:
            auth.log_processing_event(
                username=cur_u, file_type="TEXT", file_size_kb=size_kb,
                stage="analysis", status="success", failure_reason="",
                duration_ms=round((time.perf_counter() - t0) * 1000.0, 1)
            )
        except Exception:
            pass

        return {
            "scores": {
                "Language Quality": round(float(language_score), 2),
                "Structural Coherence": round(float(coherence_score), 2),
                "Argumentation": round(float(argumentation_score), 2),
                "Academic Style": round(float(style_score), 2),
                "Readability": round(float(readability_score), 2),
                "Composite": round(float(final_score), 1),
                "overall_score": round(float(final_score), 1),
                "Language": round(float(language_score), 2),
                "Coherence": round(float(coherence_score), 2),
                "Reasoning": round(float(argumentation_score), 2),
                "Sophistication": round(float(style_score), 2),
            },
            "overall_score": round(float(final_score), 1),
            "stats": stats,
            "tone": tone_data,
            "sentiment": round(float(blob.sentiment.polarity), 2),
            "subjectivity": round(float(blob.sentiment.subjectivity), 2),
            "sections": sections_data,
        }
    except Exception as ex:
        try:
            auth.log_app_error(ex, page=cur_pg, username=cur_u, role=cur_role)
            auth.log_processing_event(
                username=cur_u, file_type="TEXT", file_size_kb=size_kb,
                stage="analysis", status="failed", failure_reason="analysis error",
                duration_ms=round((time.perf_counter() - t0) * 1000.0, 1)
            )
        except Exception:
            pass
        return None

# ----------------------------------------------------
# Advanced NLP Insights & Suggestion Engines
# ----------------------------------------------------
ACADEMIC_VOCAB_SUGGESTIONS = [
    ("show", "demonstrate / exhibit", "Conveys rigorous empirical evidence rather than casual observation."),
    ("use", "utilize / employ", "More precise terminology in methodology and experimental design."),
    ("get", "obtain / derive", "Appropriate scholarly phrasing for results, data, and outcomes."),
    ("big", "substantial / significant", "Quantifies magnitude with academic specificity."),
    ("make", "formulate / construct", "Clearer action verb for models, theories, and protocols."),
    ("find", "identify / determine", "Reflects systematic investigation rather than incidental discovery."),
    ("good", "favorable / effective", "Specifies evaluation criteria rather than generic subjective value."),
    ("bad", "suboptimal / adverse", "Maintains objective, neutral academic tone in critique."),
    ("look at", "examine / investigate", "Standard terminology for literature reviews and methodology."),
    ("give", "provide / furnish", "Formal phrasing for presenting evidence or resources."),
    ("a lot of", "numerous / considerable", "Replaces informal quantifiers with precise scholarly vocabulary."),
    ("start", "initiate / commence", "Professional phrasing for workflows and experimental phases."),
    ("end", "conclude / terminate", "Conventional academic phrasing for processes and experiments."),
    ("tell", "indicate / signify", "More accurate description of data implications."),
]

def extract_vocabulary_improvements(text):
    text_lower = text.lower()
    results = []
    for original, suggestion, reason in ACADEMIC_VOCAB_SUGGESTIONS:
        matches = list(re.finditer(r"\b" + re.escape(original) + r"\b", text_lower))
        if matches:
            results.append({
                "original": original,
                "suggestion": suggestion,
                "reason": reason,
                "occurrences": len(matches)
            })
    return results

def detect_paper_issues(text, sections, scores, stats, citations):
    issues = []
    if sections and len(sections) >= 3:
        has_methodology = any("method" in s.lower() for s in sections.keys())
        if not has_methodology:
            issues.append({
                "severity": "🔴 High",
                "category": "Structure",
                "problem": "Missing or unlabelled Methodology section",
                "issue": "Missing or unlabelled Methodology section",
                "why_it_matters": "Peer reviewers and readers need to understand the empirical procedure, dataset, or theoretical framework used in the research.",
                "recommendation": "Add a dedicated Methodology section describing the research procedure, dataset, tools, and evaluation process."
            })

    total_cits = citations.get("total_citations", 0)
    has_ref_keyword = any(k in text.lower() for k in ["references", "bibliography", "works cited"])
    if total_cits == 0:
        if not has_ref_keyword:
            issues.append({
                "severity": "🟠 High",
                "category": "Citations",
                "problem": "No in-text citations or references section detected",
                "issue": "No in-text citations or references section detected",
                "why_it_matters": "Scholarly research requires rigorous attribution to situate claims within the broader academic corpus.",
                "recommendation": "Incorporate standard academic citations (e.g., [1] or Author, Year) and append a complete References section."
            })
        else:
            issues.append({
                "severity": "🟡 Moderate",
                "category": "Citations",
                "problem": "In-text citation markers could not be reliably extracted",
                "issue": "In-text citation markers could not be reliably extracted",
                "why_it_matters": "References heading was detected, but standard bracket or parenthetical markers were not located in the body prose.",
                "recommendation": "Ensure citation callouts adhere to standard conventions such as IEEE [1] or APA (Smith et al., 2023)."
            })
    elif total_cits < 5:
        issues.append({
            "severity": "🟡 Moderate",
            "category": "Citations",
            "problem": f"Sparse reference citations ({total_cits} citations detected)",
            "issue": f"Sparse reference citations ({total_cits} citations detected)",
            "why_it_matters": "A low citation count suggests related work and comparative baselines may be insufficiently contextualized.",
            "recommendation": "Incorporate authoritative peer-reviewed references to contextualize your findings within existing literature."
        })

    sentences = summarizer.split_into_sentences(text)
    if not sentences:
        sentences = [s.strip() for s in re.split(r"[.!?]+", text) if s.strip()]
    long_sentences = [s for s in sentences if len(s.split()) > 35 and not is_generated_report_text(s)]
    if len(long_sentences) >= 3:
        issues.append({
            "severity": "🟡 Medium",
            "category": "Readability",
            "problem": f"{len(long_sentences)} sentences exceed 35 words",
            "issue": f"{len(long_sentences)} sentences exceed 35 words",
            "why_it_matters": "Convoluted sentences increase cognitive load and obscure logical relationships between assertions.",
            "recommendation": "Segment convoluted sentences into concise, logical propositions to reduce cognitive load on reviewers."
        })

    transitions = ["however", "therefore", "thus", "consequently", "furthermore", "meanwhile", "moreover", "in addition"]
    trans_count = sum(text.lower().count(t) for t in transitions)
    sentence_count = stats.get("sentence_count", 1)
    if sentence_count > 10 and (trans_count / max(1, sentence_count)) < 0.08:
        issues.append({
            "severity": "🟡 Medium",
            "category": "Structural Coherence",
            "problem": "Low density of transitional discourse markers",
            "issue": "Low density of transitional discourse markers",
            "why_it_matters": "Discourse markers help readers navigate logical relationships and transitions between consecutive paragraphs.",
            "recommendation": "Use transition words (e.g., 'Consequently', 'Furthermore', 'In contrast') to guide logical transitions between paragraphs."
        })

    vocab_div = stats.get("vocab_diversity", 0.5)
    word_count = stats.get("word_count", 0)
    if vocab_div < 0.22 and word_count > 400:
        issues.append({
            "severity": "🔵 Low",
            "category": "Academic Style",
            "problem": "Low vocabulary diversity ratio",
            "issue": "Low vocabulary diversity ratio",
            "why_it_matters": "Repetitive non-technical phrasing diminishes scholarly rigor and precision.",
            "recommendation": "Expand scholarly terminology and avoid repeating the same non-technical words across consecutive sections."
        })
    return issues

def compute_section_sentiments(sections):
    results = []
    for title, content in sections.items():
        if title == "Preamble":
            continue
        if content and len(content.strip()) > 30:
            blob = TextBlob(content)
            polarity = round(blob.sentiment.polarity, 3)
            subjectivity = round(blob.sentiment.subjectivity, 3)
            short_title = title if len(title) <= 22 else title[:19] + "..."
            if polarity > 0.05:
                tag = "Assertive 🟢"
                interp = "Optimistic / Assertive academic framing"
            elif polarity < -0.05:
                tag = "Critical 🔴"
                interp = "Critical / Problem-focused framing"
            else:
                tag = "Objective ⚪"
                interp = "Objective empirical reporting"
            subj_desc = "Objective" if subjectivity < 0.35 else ("Balanced" if subjectivity <= 0.6 else "Subjective")
            results.append({
                "full_title": title,
                "short_title": short_title,
                "polarity": polarity,
                "subjectivity": subjectivity,
                "sentiment_tag": tag,
                "subjectivity_desc": subj_desc,
                "interpretation": interp
            })
    return results

def get_reliable_academic_sections(sections):
    """
    Phase 1.4: Maps detected sections to recognized scholarly headings:
    Abstract, Introduction, Literature Review / Background, Methodology,
    Results and Discussion (combined), Results, Discussion, Limitations,
    Conclusion, Future Work, Acknowledgements, References.
    Shows only sections that actually exist in the paper.
    """
    if not sections or not isinstance(sections, dict):
        return {}

    matched = {}
    used_titles = set()

    # Check if there is a combined Results and Discussion section first
    for title, content in sections.items():
        t_low = title.lower()
        if "result" in t_low and "discussion" in t_low and len(content.strip().split()) >= 10:
            matched["Results and Discussion"] = content.strip()
            used_titles.add(title)
            break

    canonical_specs = [
        ("Abstract", ["abstract"]),
        ("Introduction", ["introduction"]),
        ("Literature Review / Background", ["literature review", "related work", "prior work", "background"]),
        ("Methodology", ["methodology", "methods", "materials and methods", "materials & methods", "system model", "experimental design", "experimental setup", "proposed method"]),
        ("Results", ["results", "findings", "experimental results", "evaluation"]),
        ("Discussion", ["discussion"]),
        ("Limitations", ["limitations", "threats to validity", "scope and limitations"]),
        ("Conclusion", ["conclusion", "conclusions", "concluding remarks", "summary of work"]),
        ("Future Work", ["future work", "future directions"]),
        ("Acknowledgements", ["acknowledgements", "acknowledgments", "competing interests", "conflict of interest"]),
        ("References", ["references", "bibliography", "works cited"]),
    ]

    ordered = {}
    for canon_name, keywords in canonical_specs:
        if canon_name in ["Results", "Discussion"] and "Results and Discussion" in matched:
            if canon_name == "Results" and "Results and Discussion" not in ordered:
                ordered["Results and Discussion"] = matched["Results and Discussion"]
            continue
        for title, content in sections.items():
            if title in used_titles or title == "Preamble":
                continue
            t_low = title.lower()
            if any(k in t_low for k in keywords):
                min_words = 5 if canon_name in ["References", "Acknowledgements"] else 10
                if len(content.strip().split()) >= min_words:
                    ordered[canon_name] = content.strip()
                    used_titles.add(title)
                    break

    if len(ordered) < 2:
        return {}
    return ordered

def get_missing_standard_sections(reliable_secs: dict):
    """
    Phase 1.4: Returns standard academic sections that were not detected in the manuscript,
    so the UI can display a neutral 'Not detected in this manuscript' note.
    """
    if not reliable_secs:
        return []
    missing = []
    standard_checks = [
        ("Abstract", ["Abstract"]),
        ("Introduction", ["Introduction"]),
        ("Literature Review / Background", ["Literature Review / Background", "Related Work"]),
        ("Methodology", ["Methodology"]),
        ("Results", ["Results", "Results and Discussion"]),
        ("Discussion", ["Discussion", "Results and Discussion"]),
        ("Limitations", ["Limitations"]),
        ("Conclusion", ["Conclusion"]),
        ("References", ["References"]),
    ]
    for label, satisfying_keys in standard_checks:
        if not any(k in reliable_secs for k in satisfying_keys):
            missing.append(label)
    return missing

def _phrase_specific_gap_opportunity(sentence_text: str, marker: str) -> str:
    """Phrase a specific research opportunity grounded in the terms of the quoted sentence, never a generic template."""
    stop = {
        "the", "and", "for", "with", "that", "this", "from", "were", "was", "are", "have", "has", "had",
        "not", "but", "however", "despite", "lack", "limitation", "limitations", "limited", "study",
        "paper", "work", "research", "further", "future", "into", "over", "under", "between", "across",
        "their", "these", "those", "which", "while", "where", "when", "such", "only", "also", "more",
    }
    tokens = [
        w for w in re.findall(r"\b[A-Za-z][A-Za-z0-9\-]{3,}\b", sentence_text)
        if w.lower() not in stop
    ]
    seen_t = []
    for t in tokens:
        if t.lower() not in [x.lower() for x in seen_t]:
            seen_t.append(t)
        if len(seen_t) >= 4:
            break
    focus_phrase = ", ".join(seen_t[:3]) if seen_t else "the reported scope"
    if "lack of" in marker or "scarce" in marker or "limited data" in marker:
        return f"Develop expanded datasets or empirical benchmarks specifically addressing {focus_phrase}."
    if "future" in marker or "further research" in marker or "remains" in marker:
        return f"Investigate {focus_phrase} as prioritized by the authors' stated open inquiry."
    return f"Extend experimental evaluation and methodological coverage across {focus_phrase}."


def extract_structured_research_gaps(text, sections=None):
    """
    Extracts document-grounded research gaps, open problems, or author-stated limitations
    strictly from explicit gap/limitation/'lack of'/'further research' statements.
    """
    sentences = summarizer.split_into_sentences(text)
    gap_markers = [
        ("limitation", "Author-stated study limitation", "Methodological or sample boundary noted by authors", "High"),
        ("limited to", "Scope restriction stated by authors", "Evaluation is bounded to specific conditions or cohorts", "High"),
        ("limited by", "Constraint noted in manuscript", "Empirical or resource constraint identified in study", "High"),
        ("lack of", "Explicitly stated resource or data gap", "Absence of data, coverage, or baseline noted in text", "High"),
        ("further research", "Author-identified need for further research", "Authors explicitly call for additional investigation", "High"),
        ("future work", "Author-identified open research inquiry", "Authors explicitly recommend future extension", "High"),
        ("remains unclear", "Unresolved scientific question", "Mechanism or relationship not yet established", "High"),
        ("remains unexplored", "Unexplored research area", "Topic or condition not yet empirically evaluated", "High"),
        ("not addressed", "Unaddressed aspect in current scope", "Aspect explicitly left outside current scope", "Moderate"),
        ("scarce", "Scarcity of prior evidence or data", "Limited existing literature or data on this topic", "Moderate"),
        ("drawback", "Identified methodological drawback", "Known trade-off or constraint in the approach", "Moderate"),
    ]

    # Map sentence to section if sections dict is provided
    sec_lookup = []
    if isinstance(sections, dict):
        for s_name, s_body in sections.items():
            if s_name.lower() in ("references", "bibliography", "acknowledgements"):
                continue
            sec_Lookup_txt = summarizer.unescape_text(s_body) if hasattr(summarizer, "unescape_text") else str(s_body)
            sec_lookup.append((s_name, sec_Lookup_txt.lower()))

    found_gaps = []
    seen_excerpts = set()

    for s in sentences:
        if is_generated_report_text(s):
            continue
        s_clean = summarizer.strip_leading_heading_fragment(s)
        s_lower = s_clean.lower()
        if 35 <= len(s_clean) <= 320:
            for marker, lim_desc, gap_desc, conf in gap_markers:
                if re.search(r"\b" + re.escape(marker) + r"\b", s_lower):
                    if s_clean not in seen_excerpts:
                        seen_excerpts.add(s_clean)
                        source_sec = "Manuscript Text"
                        for sec_name, sec_low in sec_lookup:
                            if s_lower[:45] in sec_low:
                                source_sec = sec_name
                                break
                        opp = _phrase_specific_gap_opportunity(s_clean, marker)
                        found_gaps.append({
                            "evidence": s_clean,
                            "quote": s_clean,
                            "section": source_sec,
                            "limitation": lim_desc,
                            "why_gap": gap_desc,
                            "opportunity": opp,
                            "confidence": conf,
                        })
                        break
        if len(found_gaps) >= 5:
            break

    return found_gaps


def extract_research_gaps(text):
    """Backwards compatibility alias returning string list."""
    structured = extract_structured_research_gaps(text)
    if not structured:
        return ["Not found in this manuscript"]
    return [g["evidence"] for g in structured]


def extract_key_contributions_structured(text):
    """
    Extracts author-stated contributions or empirical findings using robust sentence splitting.
    """
    sentences = summarizer.split_into_sentences(text)
    explicit_markers = [
        "we propose", "in this paper", "this work presents", "our contribution",
        "we contribute", "main contribution", "key contribution", "primary contribution",
        "we demonstrate", "we develop", "we introduce", "first time", "our framework",
        "we design", "we present", "in this study", "this study evaluates", "this study investigates",
    ]
    inferred_markers = [
        "results demonstrate", "achieves", "achieved", "outperforms", "outperformed",
        "findings indicate", "findings show", "results show", "results indicate",
        "significant improvement", "experiment shows", "experiments show", "observed that",
        "recorded", "increased by", "growth rate of",
    ]

    contributions = []
    seen = set()

    for s in sentences:
        if is_generated_report_text(s):
            continue
        s_clean = summarizer.strip_leading_heading_fragment(s)
        s_lower = s_clean.lower()
        if any(k in s_lower for k in explicit_markers) and 30 <= len(s_clean) <= 300:
            if s_clean not in seen:
                seen.add(s_clean)
                contributions.append({
                    "type": "Author-Stated Contribution",
                    "text": s_clean,
                    "confidence": "High",
                })
        if len(contributions) >= 3:
            break

    if len(contributions) < 2:
        for s in sentences:
            if is_generated_report_text(s):
                continue
            s_clean = summarizer.strip_leading_heading_fragment(s)
            s_lower = s_clean.lower()
            if any(k in s_lower for k in inferred_markers) and 30 <= len(s_clean) <= 300:
                if s_clean not in seen:
                    seen.add(s_clean)
                    contributions.append({
                        "type": "Inferred from Empirical Findings",
                        "text": s_clean,
                        "confidence": "Moderate",
                    })
            if len(contributions) >= 3:
                break

    return contributions


def extract_key_contributions(text):
    """Backwards compatibility alias returning string list."""
    structured = extract_key_contributions_structured(text)
    if not structured:
        return ["Not found in this manuscript"]
    return [c["text"] for c in structured]


def extract_overly_long_sentences(text, max_words=35, limit=2):
    """Identifies concrete overly long sentences using the unified decimal-safe sentence splitter."""
    sentences = summarizer.split_into_sentences(text)
    long_ones = []
    for s in sentences:
        if is_generated_report_text(s):
            continue
        s_clean = summarizer.strip_leading_heading_fragment(s)
        words = s_clean.split()
        if len(words) > max_words and not any(ch in s_clean for ch in ["http", "www", "doi"]):
            long_ones.append((s_clean, len(words)))
            if len(long_ones) >= limit:
                break
    return long_ones


def extract_future_directions(text):
    """Extracts ONLY explicit author-stated future research directions from the manuscript."""
    sentences = summarizer.split_into_sentences(text)
    future_keywords = [
        "future work", "future research", "future investigations", "in the future",
        "next step", "promising avenue", "remains for future", "plan to explore",
        "further study", "extensions of this work", "will investigate", "future studies",
        "open direction", "further research", "should be explored", "can be extended",
    ]
    found = []
    for s in sentences:
        if is_generated_report_text(s):
            continue
        s_clean = summarizer.strip_leading_heading_fragment(s)
        s_lower = s_clean.lower()
        if any(k in s_lower for k in future_keywords) and 25 <= len(s_clean) <= 300:
            if s_clean not in found:
                found.append(s_clean)
            if len(found) >= 4:
                break
    if not found:
        return ["Not stated by the authors"]
    return found


def extract_methodology_and_findings(text, sections=None):
    """Extracts clean, complete sentences for methodology, empirical findings, and data/materials."""
    meth_summary = ""
    findings_summary = ""
    data_summary = ""

    def _pick_clean_sentences(block_text: str, count: int = 2) -> str:
        sents = [
            summarizer.strip_leading_heading_fragment(s)
            for s in summarizer.split_into_sentences(block_text)
        ]
        valid = [
            s for s in sents
            if len(s) >= 28 and len(s.split()) >= 6 and not is_generated_report_text(s)
        ]
        if not valid:
            return ""
        joined = " ".join(valid[:count]).strip()
        if joined and joined[-1] not in ".!?":
            joined += "."
        return joined

    if sections and isinstance(sections, dict):
        for sec_name, sec_text in sections.items():
            sec_lower = sec_name.lower()
            if any(k in sec_lower for k in ["method", "approach", "framework", "architecture", "experimental setup"]):
                if not meth_summary and len(sec_text.strip()) > 35:
                    meth_summary = _pick_clean_sentences(sec_text, 2)
            if any(k in sec_lower for k in ["result", "finding", "experiment", "evaluation", "discussion"]):
                if not findings_summary and len(sec_text.strip()) > 35:
                    findings_summary = _pick_clean_sentences(sec_text, 2)
            if any(k in sec_lower for k in ["dataset", "material", "cohort", "data collection", "data and"]):
                if not data_summary and len(sec_text.strip()) > 30:
                    data_summary = _pick_clean_sentences(sec_text, 2)

    sentences = [summarizer.strip_leading_heading_fragment(s) for s in summarizer.split_into_sentences(text)]
    if not meth_summary:
        m_keys = [
            "we implement", "we evaluate", "experimental setup", "we trained",
            "proposed architecture", "methodology", "baseline", "randomized",
            "regression", "compound annual growth", "analytical framework", "formulate",
        ]
        for s in sentences:
            if is_generated_report_text(s):
                continue
            if any(k in s.lower() for k in m_keys) and 30 <= len(s) <= 280:
                meth_summary = s if s[-1] in ".!?" else s + "."
                break
    if not findings_summary:
        f_keys = [
            "results indicate", "results show", "outperforms", "achieves", "achieved",
            "findings demonstrate", "empirically", "accuracy", "growth rate of",
            "increased by", "statistically significant", "recorded",
        ]
        for s in sentences:
            if is_generated_report_text(s):
                continue
            if any(k in s.lower() for k in f_keys) and 30 <= len(s) <= 280:
                findings_summary = s if s[-1] in ".!?" else s + "."
                break
    if not data_summary:
        d_keys = [
            "dataset", "cohort", "sample of", "materials", "telemetry", "field trial",
            "cultivars", "benchmark", "participants", "secondary data", "time series data",
            "collected from", "corpus",
        ]
        for s in sentences:
            if is_generated_report_text(s):
                continue
            if any(k in s.lower() for k in d_keys) and 28 <= len(s) <= 280:
                data_summary = s if s[-1] in ".!?" else s + "."
                break

    if not meth_summary:
        meth_summary = "Not found in this manuscript"
    if not findings_summary:
        findings_summary = "Not found in this manuscript"
    if not data_summary:
        data_summary = "Not found in this manuscript"

        
    return {"methodology": meth_summary, "findings": findings_summary, "data_and_materials": data_summary}

ACADEMIC_DOMAIN_TAXONOMY = {
    "Agricultural & Food Sciences": [
        "agriculture", "agricultural", "crop", "crops", "farming", "farm", "farms", "soil", "yield",
        "cultivar", "cultivars", "cultivation", "harvest", "livestock", "agronomy", "irrigation",
        "fertilizer", "fertilizers", "pest", "pesticide", "grain", "wheat", "rice", "maize",
        "horticulture", "forestry", "agroforestry", "plant pathology", "food production", "seed",
        "germination", "agrochemical", "fodder", "silage", "pasture", "canopy", "tillage",
        "biochar", "crop yield", "plant growth", "soil health", "food security"
    ],
    "Biological & Life Sciences": [
        "biology", "biological", "genetics", "genome", "genomic", "cellular", "organism", "organisms",
        "ecology", "biodiversity", "species", "molecular biology", "physiology", "ecosystem", "dna",
        "rna", "protein", "proteins", "evolution", "botany", "zoology", "microbiology",
        "photosynthesis", "enzyme", "enzymes", "bacterial", "membrane", "metabolic", "pathogen"
    ],
    "Medicine & Health Sciences": [
        "clinical", "patient", "patients", "hospital", "therapy", "therapeutic", "pathology", "medical",
        "diagnosis", "diagnostic", "treatment", "treatments", "surgery", "surgical", "oncology",
        "pharmacological", "pharmacology", "epidemiology", "pharmacy", "vaccine", "cardiovascular",
        "physician", "morbidity", "mortality rate", "biomedical", "disease prevention", "healthcare"
    ],
    "Computer Science & Artificial Intelligence": [
        "algorithm", "algorithms", "neural network", "machine learning", "deep learning",
        "artificial intelligence", "nlp", "computer vision", "transformer", "transformers",
        "embeddings", "classification model", "reinforcement learning", "supervised learning",
        "unsupervised learning", "convolutional", "backpropagation", "benchmark dataset",
        "large language model", "latent space", "semantic segmentation", "hyperparameter"
    ],
    "Cybersecurity & Cryptography": [
        "cybersecurity", "cryptography", "cryptographic", "cyber attack", "vulnerability",
        "vulnerabilities", "malware", "ransomware", "intrusion detection", "encryption",
        "decryption", "adversarial attack", "phishing", "firewall", "zero-day",
        "authentication protocol", "exploit", "threat actor", "security penetration"
    ],
    "Networks & Distributed Systems": [
        "distributed systems", "cloud computing", "peer-to-peer", "consensus protocol",
        "byzantine", "latency", "throughput", "routing protocol", "kubernetes",
        "fault tolerance", "cluster computing", "blockchain", "load balancing", "serverless",
        "microservices", "network topology", "distributed consensus"
    ],
    "Physical Sciences & Mathematics": [
        "physics", "quantum", "thermodynamics", "astrophysics", "electromagnetism", "optics",
        "relativity", "particle physics", "mathematics", "mathematical theorem", "topology",
        "calculus", "differential equation", "algebraic", "numerical analysis",
        "probability distribution", "gravitational", "condensed matter", "hamiltonian"
    ],
    "Chemical & Materials Sciences": [
        "chemistry", "chemical", "chemical synthesis", "polymer", "polymers", "catalyst",
        "catalysts", "reagent", "molecular structure", "spectroscopy", "nanoparticles",
        "nanomaterials", "electrochemistry", "crystallography", "metallurgy", "inorganic chemistry",
        "organic synthesis", "crystallographic", "chemical bond", "thermogravimetric"
    ],
    "Earth & Environmental Sciences": [
        "environmental science", "climate change", "hydrology", "geology", "geological",
        "meteorology", "atmospheric science", "carbon emissions", "oceanography",
        "environmental pollution", "ecosystem preservation", "geospatial", "remote sensing",
        "precipitation", "temperature anomalies", "greenhouse gas", "groundwater"
    ],
    "Engineering & Robotics": [
        "mechanical engineering", "robotics", "actuator", "actuators", "kinematics",
        "control systems", "robot", "electrical grid", "structural engineering",
        "signal processing", "telecommunications", "embedded systems", "sensors",
        "microcontroller", "finite element", "aerospace", "mechatronics"
    ],
    "Economics & Business": [
        "economics", "econometric", "macroeconomic", "microeconomic", "financial markets",
        "inflation", "gdp", "market volatility", "portfolio optimization", "asset pricing",
        "monetary policy", "fiscal policy", "trade volume", "consumer behavior", "interest rate"
    ],
    "Social Sciences & Humanities": [
        "sociology", "psychology", "pedagogy", "educational curriculum", "linguistics",
        "anthropology", "political science", "philosophy", "ethics", "qualitative research",
        "ethnography", "cognitive development", "historical analysis", "governance policy"
    ]
}

def classify_research_domain(text):
    """
    Rigorously classifies document research domain using frequency-weighted token modeling.
    Distinguishes primary domain, secondary candidates, and confidence levels.
    Prevents cross-domain pollution (e.g. agricultural papers mistakenly assigned cloud or cybersecurity).
    """
    if not text:
        return {
            "primary": "Domain Uncertain",
            "secondary": [],
            "all_candidates": ["Domain Uncertain"],
            "confidence": "Uncertain",
            "scores": {}
        }
        
    text_lower = text.lower()
    domain_scores = {}
    
    for domain, terms in ACADEMIC_DOMAIN_TAXONOMY.items():
        score = 0
        for term in terms:
            pattern = r"\b" + re.escape(term) + r"\b"
            matches = len(re.findall(pattern, text_lower))
            if matches > 0:
                weight = 3 if " " in term else 1
                score += matches * weight
        if score > 0:
            domain_scores[domain] = score
            
    if not domain_scores:
        return {
            "primary": "General Academic",
            "secondary": [],
            "all_candidates": ["General Academic"],
            "confidence": "Uncertain",
            "scores": {}
        }
        
    sorted_domains = sorted(domain_scores.items(), key=lambda x: x[1], reverse=True)
    top_domain, top_score = sorted_domains[0]
    second_domain, second_score = sorted_domains[1] if len(sorted_domains) > 1 else (None, 0)
    
    if top_score < 3:
        return {
            "primary": "Domain Uncertain",
            "secondary": [],
            "all_candidates": ["Domain Uncertain"],
            "confidence": "Uncertain",
            "scores": domain_scores
        }
        
    if top_score >= 8 and (second_score == 0 or (top_score / max(1, second_score)) >= 1.7):
        confidence = "High"
    elif top_score >= 4:
        confidence = "Moderate"
    else:
        confidence = "Uncertain"
        
    secondary = []
    for d, s in sorted_domains[1:]:
        if s >= 4 and s >= 0.40 * top_score and len(secondary) < 2:
            secondary.append(d)
            
    all_candidates = [top_domain] + secondary
    
    return {
        "primary": top_domain,
        "secondary": secondary,
        "all_candidates": all_candidates,
        "confidence": confidence,
        "scores": domain_scores
    }

def extract_potential_research_areas(text):
    """Backwards compatibility wrapper returning candidate research domains."""
    clf = classify_research_domain(text)
    return clf["all_candidates"]

def extract_keywords_and_domain(text, top_n=12):
    blob = TextBlob(text)
    stop_words = set([
        "the", "is", "in", "and", "to", "of", "a", "for", "on", "with", "as",
        "by", "at", "this", "that", "it", "from", "an", "be", "are", "was",
        "which", "were", "been", "have", "has", "using", "used", "paper"
    ])
    word_freq = {}
    for word in blob.words:
        word_l = word.lower()
        if word_l.isalpha() and word_l not in stop_words and len(word_l) > 3:
            word_freq[word_l] = word_freq.get(word_l, 0) + 1
    sorted_words = sorted(word_freq.items(), key=lambda x: x[1], reverse=True)
    keywords = [w[0] for w in sorted_words[:top_n]]
    
    clf = classify_research_domain(text)
    return keywords, clf["all_candidates"]

def detect_publisher(text):
    publishers = {
        "IEEE": ["ieee", "ieeexplore"],
        "Springer Nature": ["springer", "springerlink", "nature"],
        "Elsevier / ScienceDirect": ["elsevier", "sciencedirect"],
        "Wiley Online Library": ["wiley"],
        "Taylor & Francis": ["taylor", "francis"],
        "ACM": ["acm", "association for computing machinery"],
        "MDPI": ["mdpi"],
        "Oxford University Press": ["oxford university press"],
    }
    text_lower = text.lower()
    for publisher, keywords in publishers.items():
        for word in keywords:
            if word in text_lower:
                return publisher
    return "Independent Academic Manuscript / Pre-print"

def analyze_citations(text):
    citation_patterns = [r"\(\d{4}\)", r"\[\d+\]", r"\(\w+ et al\., \d{4}\)"]
    citation_count = sum(len(re.findall(p, text)) for p in citation_patterns)
    total_words = len(text.split())
    citation_density = round(citation_count / total_words, 4) if total_words > 0 else 0.0
    impact_score = min(100.0, citation_count * 2.5)
    years = [int(y) for y in re.findall(r"(19\d{2}|20\d{2})", text) if 1970 <= int(y) <= 2026]
    avg_year = round(sum(years) / len(years), 1) if years else 0.0
    return {
        "total_citations": citation_count,
        "citation_density": citation_density,
        "impact_score": round(impact_score, 1),
        "average_year": avg_year,
    }

def calculate_semantic_strength(text):
    try:
        blob = TextBlob(text)
        nouns = [word.lower() for word, tag in blob.tags if tag.startswith("NN")]
        unique_nouns = len(set(nouns))
        total_nouns = len(nouns)
        if total_nouns == 0:
            return 0.0
        return round(min(100.0, (unique_nouns / total_nouns) * 150), 2)
    except Exception:
        words = text.split()
        return round(min(100.0, (len(set(words)) / max(1, len(words))) * 100), 2)

def novelty_score(text):
    words = re.findall(r"\b\w+\b", text.lower())
    if not words:
        return 0.0
    return round((len(set(words)) / len(words)) * 100, 2)

def generate_research_feedback(results):
    feedback = []
    scores = results.get("scores", {})
    comp = scores.get("Composite", 0)
    citations = results.get("citation_analysis", {}).get("total_citations", 0)
    sections = results.get("sections", {})
    if comp < 70:
        feedback.append("Overall academic writing and structure require targeted revision.")
    if citations < 5:
        feedback.append("Incorporate more scholarly citations to anchor empirical findings.")
    has_methodology = any("method" in k.lower() for k in sections.keys())
    if not has_methodology:
        feedback.append("Methodology section is absent or insufficiently labeled.")
    if not feedback:
        feedback.append("Paper demonstrates balanced academic structure and sound prose quality.")
    return feedback

def recommend_journal(domain):
    """Backwards-compatibility alias returning inferred research areas."""
    if isinstance(domain, list) and domain:
        return ", ".join(domain[:2])
    return str(domain)

SCHOLARLY_STOPWORDS = {
    'work', 'paper', 'study', 'research', 'author', 'authors', 'article',
    'approach', 'present', 'presents', 'presented', 'describe', 'describes',
    'described', 'discuss', 'discusses', 'discussed', 'use', 'uses', 'used',
    'using', 'does', 'did', 'done', 'what', 'how', 'why', 'when', 'which',
    'where', 'who', 'whom', 'tell', 'explain', 'give', 'show', 'shows',
    'in', 'on', 'at', 'by', 'for', 'with', 'about', 'against', 'between',
    'into', 'through', 'during', 'before', 'after', 'above', 'below', 'to',
    'from', 'up', 'down', 'the', 'a', 'an', 'is', 'are', 'was', 'were',
    'be', 'been', 'being', 'have', 'has', 'had', 'having', 'do', 'does'
}

INTENT_MAP = {
    'evaluation': {
        'stems': ['evaluat', 'metric', 'benchmark', 'validat', 'assess', 'measur', 'experiment', 'test', 'accura', 'f1', 'baseline', 'compar', 'cross-validation', 'signific', 't-test'],
        'sections': ['evaluation', 'experiment', 'experiments', 'assessment', 'results and discussion', 'results', 'methodology', 'materials and methods', 'validation', 'empirical', 'setup', 'ablation']
    },
    'methodology': {
        'stems': ['method', 'algorithm', 'architect', 'procedur', 'design', 'pipelin', 'framework', 'formul', 'techniqu', 'system model', 'implement', 'material', 'statist', 'analys', 'model', 'regression', 'cagr'],
        'sections': ['methodology', 'materials and methods', 'methods', 'proposed', 'architecture', 'system model', 'design', 'algorithm']
    },
    'findings': {
        'stems': ['find', 'result', 'outcom', 'observ', 'achiev', 'discover', 'demonstrat', 'reveal', 'perform', 'gain', 'growth', 'trend', 'increas', 'decreas', 'signific', 'conclud'],
        'sections': ['results and discussion', 'results', 'findings', 'discoveries', 'discussion', 'conclusion', 'conclusions', 'abstract', 'experiments']
    },
    'dataset': {
        'stems': ['dataset', 'data', 'corpus', 'sample', 'materi', 'subject', 'particip', 'source', 'benchmark data', 'benchmark', 'secondary data', 'period', 'collect'],
        'sections': ['methodology', 'materials and methods', 'methods', 'dataset', 'datasets', 'data collection', 'materials', 'experimental setup', 'experiment', 'evaluation']
    },
    'limitations': {
        'stems': ['limit', 'threat', 'weakness', 'shortcom', 'constraint', 'drawback', 'challeng', 'caveat', 'lack', 'bottleneck', 'barrier'],
        'sections': ['limitations', 'threats to validity', 'constraints', 'discussion', 'results and discussion', 'conclusion', 'conclusions']
    },
    'future_work': {
        'stems': ['future', 'next step', 'promis', 'exten', 'direction', 'prospect', 'avenue', 'recommend', 'further', 'policy', 'implicat'],
        'sections': ['future work', 'conclusion', 'conclusions', 'discussion', 'results and discussion']
    },
    'objective': {
        'stems': ['object', 'aim', 'goal', 'purpos', 'motivat', 'propos', 'focus', 'investigat', 'tackl', 'studi', 'determine', 'conduct', 'examine', 'analys', 'analyz'],
        'sections': ['abstract', 'introduction']
    },
    'research_gap': {
        'stems': ['gap', 'unresolved', 'open question', 'addressed', 'prior work', 'lack', 'need for', 'deficienc', 'declin', 'problem'],
        'sections': ['introduction', 'background', 'literature review', 'related work', 'abstract']
    },
    'contributions': {
        'stems': ['contribut', 'novel', 'propose', 'key insight', 'advance', 'first time', 'present', 'core', 'studi', 'analyz'],
        'sections': ['abstract', 'introduction', 'contributions', 'conclusion', 'conclusions']
    },
    'theoretical_framework': {
        'stems': ['theor', 'framework', 'foundat', 'principl', 'concept', 'formulat', 'formalism', 'model', 'function', 'growth rate'],
        'sections': ['theory', 'background', 'methodology', 'materials and methods', 'theoretical framework', 'system model', 'introduction']
    },
    'conclusion': {
        'stems': ['conclus', 'summar', 'conclud', 'takeaway', 'in summary', 'concluding remarks'],
        'sections': ['conclusion', 'conclusions', 'concluding remarks', 'summary', 'discussion', 'abstract', 'results and discussion']
    }
}

def _extract_clean_qa_sentences(text):
    """
    Phase 1.5 helper: Splits section/document text into complete, well-formed sentences
    using abbreviation-safe splitting, removing broken mid-sentence fragments, inline
    Keywords tails, inline numbered subsection headers, and author/grant boilerplate.
    """
    if not text:
        return []
    cleaned_block = re.split(r"\bKeywords\s*:", str(text), maxsplit=1, flags=re.IGNORECASE)[0]
    cleaned_block = re.sub(
        r"(?:^|\s)\d+(?:\.\d+)+\s+[A-Z][A-Za-z\s,\-&]{3,65}?(?=\s+(?:The|A|An|In|During|Compound|Overall|Although|Even|However|Further|From|Table|Fig)\b)",
        " ",
        cleaned_block,
    )
    raw_sents = summarizer.split_into_sentences(cleaned_block)
    valid_sents = []
    for s in raw_sents:
        s = re.sub(r"\s+", " ", s).strip()
        s = re.sub(r"^\d+(?:\.\d+)+\s+[A-Z][A-Za-z\s,\-&]{3,55}?(?=\s+[A-Z][a-z])", "", s).strip()
        # Strip accidental broken leading fragments like "e., " or "i.e., " or leading punctuation
        s = re.sub(r"^(?:i\.?e\.?|e\.?g\.?|viz\.?)\s*,\s*", "", s, flags=re.IGNORECASE).strip()
        s = re.sub(r"^[,;:\-\–\—\)\]]+\s*", "", s).strip()
        if len(s) < 20 or len(s.split()) < 4:
            continue
        s_low = s.lower()
        if any(noise in s_low for noise in [
            "all authors contributed equally",
            "supported by grant",
            "authors have declared that no competing interests",
            "open peer review history",
            "corresponding author:"
        ]):
            continue
        # Ensure starts at a clean capitalized boundary
        if s and s[0].islower():
            s = s[0].upper() + s[1:]
        if s and s[-1] not in ".!?":
            s = s.rstrip(",;:") + "."
        valid_sents.append(s)
    return valid_sents

def grounded_paper_qa(question, full_text, sections):
    """
    Phase 1.5: Answers questions strictly grounded in the analyzed document.
    - Chunks CLEANED text by section with overlap, keeping section metadata.
    - Routes suggested and free-text questions to relevant sections first.
    - Ranks candidate chunks by TF-IDF + BM25/keyword relevance.
    - Returns 2-4 complete, well-formed sentences plus a clearly labeled
      'Supporting evidence (from <Section>)' snippet.
    - If relevance score is below threshold, responds:
      'I could not find this information in the paper.'
    """
    if not question or not full_text:
        return "Please ask a question regarding the uploaded document.", []

    q_clean = question.strip()
    q_lower = q_clean.lower()

    # Extract substantive tokens from question (removing scholarly stopwords)
    q_words = re.findall(r"\b[a-z0-9]{2,}\b", q_lower)
    substantive_tokens = [w for w in q_words if len(w) >= 3 and w not in SCHOLARLY_STOPWORDS]

    # Detect Question Intent(s)
    active_intents = []
    for intent_name, config in INTENT_MAP.items():
        for stem in config['stems']:
            if (" " in stem and stem in q_lower) or (" " not in stem and any(w.startswith(stem) for w in q_words)):
                active_intents.append(intent_name)
                break

    # Build overlapping section chunks from cleaned sentences
    passages = []
    skip_sec_keywords = ['preamble', 'reference', 'bibliograph', 'competing interest', 'conflict of interest', 'acknowledg', 'author contribution']
    asking_about_refs_or_meta = any(k in q_lower for k in ['reference', 'citation', 'cite', 'acknowledg', 'funding', 'grant', 'competing', 'conflict'])

    if sections:
        for sec_name, sec_text in sections.items():
            if not sec_text or len(str(sec_text).strip()) < 20:
                continue
            sec_low = sec_name.lower()
            if not asking_about_refs_or_meta and any(sk in sec_low for sk in skip_sec_keywords):
                continue
            raw_sents = _extract_clean_qa_sentences(sec_text)
            if not raw_sents:
                continue
            if len(raw_sents) <= 3:
                passages.append({
                    "section": sec_name,
                    "text": " ".join(raw_sents),
                    "sentences": raw_sents
                })
            else:
                # Overlapping window of 3 sentences with step=1 for high recall at sentence boundaries
                for i in range(0, len(raw_sents) - 1, 1):
                    chunk_sents = raw_sents[i:i+3]
                    if len(chunk_sents) < 2 and i > 0:
                        continue
                    passages.append({
                        "section": sec_name,
                        "text": " ".join(chunk_sents),
                        "sentences": chunk_sents
                    })

    if not passages:
        raw_sents = _extract_clean_qa_sentences(full_text)
        for i in range(0, max(1, len(raw_sents)), 2):
            chunk_sents = raw_sents[i:i+3]
            if chunk_sents:
                passages.append({
                    "section": "Document Body",
                    "text": " ".join(chunk_sents),
                    "sentences": chunk_sents
                })

    if not passages:
        return (
            "I could not find this information in the paper. The uploaded paper does not provide enough information to answer this question confidently.",
            []
        )

    passage_texts = [f"{p['section']}. {p['text']}" for p in passages]

    # 1. TF-IDF Cosine Similarity across passages
    try:
        vectorizer = TfidfVectorizer(stop_words="english", ngram_range=(1, 2))
        tfidf = vectorizer.fit_transform(passage_texts + [q_clean])
        vectors = tfidf.toarray()
        cosine_sims = cosine_similarity([vectors[-1]], vectors[:-1])[0]
    except Exception:
        cosine_sims = np.zeros(len(passages))

    # 2. BM25 lexical scoring across passages
    doc_tokens_list = [re.findall(r"\b[a-z0-9]{3,}\b", pt.lower()) for pt in passage_texts]
    avgdl = max(1.0, sum(len(dt) for dt in doc_tokens_list) / max(1, len(doc_tokens_list)))
    N_docs = len(doc_tokens_list)
    k1, b_param = 1.5, 0.75
    bm25_scores = np.zeros(N_docs)
    if substantive_tokens:
        for idx_d, d_toks in enumerate(doc_tokens_list):
            dl = len(d_toks)
            score_b = 0.0
            for q_tok in substantive_tokens:
                q_stem = q_tok[:5] if len(q_tok) >= 5 else q_tok
                tf = sum(1 for t in d_toks if t == q_tok or t.startswith(q_stem))
                if tf > 0:
                    df = sum(1 for dt in doc_tokens_list if any(t == q_tok or t.startswith(q_stem) for t in dt))
                    idf = math.log(1.0 + (N_docs - df + 0.5) / (df + 0.5))
                    score_b += idf * ((tf * (k1 + 1.0)) / (tf + k1 * (1.0 - b_param + b_param * (dl / avgdl))))
            bm25_scores[idx_d] = score_b
        max_bm25 = float(np.max(bm25_scores)) if len(bm25_scores) > 0 else 0.0
        if max_bm25 > 0:
            bm25_scores = bm25_scores / max_bm25

    scored_passages = []
    for idx, p in enumerate(passages):
        p_body_low = p["text"].lower()
        full_p_text = f"{p['section']}. {p['text']}".lower()
        sec_lower = p["section"].lower()

        base_score = float(cosine_sims[idx]) if idx < len(cosine_sims) else 0.0
        bm25_val = float(bm25_scores[idx]) if idx < len(bm25_scores) else 0.0

        # Substantive token match (with stem flexibility)
        token_matches = 0
        body_token_matches = 0
        for tok in substantive_tokens:
            tok_stem = tok[:5] if len(tok) >= 5 else tok
            if tok in full_p_text or tok_stem in full_p_text:
                token_matches += 1
            if tok in p_body_low or tok_stem in p_body_low:
                body_token_matches += 1

        token_ratio = token_matches / max(1, len(substantive_tokens))

        # Section affinity and intent routing boost
        sec_boost = 0.0
        intent_match = False
        for intent in active_intents:
            target_secs = INTENT_MAP[intent]['sections']
            target_stems = INTENT_MAP[intent]['stems']
            has_stem_in_text = any(ts in p_body_low for ts in target_stems)

            # Priority section routing
            matched_sec_rank = None
            for rank_i, ts in enumerate(target_secs):
                if ts in sec_lower:
                    matched_sec_rank = rank_i
                    break

            if matched_sec_rank is not None:
                rank_bonus = max(0.25, 0.65 - (matched_sec_rank * 0.06))
                if has_stem_in_text or body_token_matches > 0:
                    sec_boost += rank_bonus
                else:
                    sec_boost += 0.25
                intent_match = True

            if any(ts in sec_lower for ts in target_stems):
                sec_boost += 0.45
                intent_match = True

            if has_stem_in_text:
                sec_boost += 0.35
                intent_match = True

            # Special domain refinement for 'findings': prefer this study's empirical outcomes over prior-work literature citations
            if intent == 'findings':
                prior_lit_citations = len(re.findall(r"(?:\[\d+\]\s*(?:reported|studied|observed|found)|et\s+al\.\s*\[\d+\])", p_body_low))
                if prior_lit_citations > 0 and not any(k in p_body_low for k in ["this study", "present study", "concluded", "overall"]):
                    sec_boost -= 0.25 * prior_lit_citations
                if any(k in p_body_low for k in ["results revealed", "study revealed", "achieved", "outperforming", "key finding", "concluded", "registered", "significant", "growth rate"]):
                    sec_boost += 0.30

            # Special refinement for 'limitations' when no dedicated Limitations section exists:
            # boost passages discussing constraints, lack of varieties/policies, or challenges
            if intent == 'limitations':
                if any(k in p_body_low for k in ["limitation", "constraint", "lack of", "shortcoming", "bottleneck", "threat", "quadratic", "unmodeled", "challenge", "decline", "drastic"]):
                    sec_boost += 0.45
                    intent_match = True

        # Penalize noise sections
        if any(noise in sec_lower for noise in ['author', 'acknowledg', 'funding', 'reference', 'grant', 'conflict', 'competing', 'biograph', 'about the author', 'preamble']):
            sec_boost -= 5.0

        final_score = (0.35 * base_score) + (0.25 * bm25_val) + (0.25 * token_ratio) + sec_boost

        if token_matches == 0 and not intent_match:
            final_score = 0.0

        scored_passages.append((final_score, p, base_score, token_matches, intent_match))

    scored_passages.sort(key=lambda x: x[0], reverse=True)
    best_score, best_p, best_base, best_tokens, best_intent = scored_passages[0]

    # Strict relevance threshold: must have genuine evidence match
    if best_score < 0.25 or (best_tokens == 0 and not best_intent) or (not best_intent and best_base < 0.12 and best_tokens < max(1, len(substantive_tokens) // 2)):
        return (
            "I could not find this information in the paper. The uploaded paper does not provide enough information to answer this question confidently.",
            []
        )

    best_sec = best_p["section"]
    ans_sents = list(best_p["sentences"])

    # Ensure 2-4 complete, well-formed sentences when the manuscript has additional relevant sentences
    if len(ans_sents) < 2:
        for cand_score, cand_p, _, _, _ in scored_passages[1:]:
            if cand_score < 0.20:
                break
            for s in cand_p["sentences"]:
                if s not in ans_sents:
                    ans_sents.append(s)
                    if len(ans_sents) >= 2:
                        break
            if len(ans_sents) >= 2:
                break

    ans_sents = ans_sents[:4]
    ans_text = " ".join(ans_sents)

    # Clean excerpt: pick the single sentence in best_p that best matches the query intent/tokens
    best_sent = best_p["sentences"][0] if best_p["sentences"] else ans_sents[0]
    best_sent_score = -1
    for sent in best_p["sentences"]:
        sent_l = sent.lower()
        score_s = sum(2 for tok in substantive_tokens if tok in sent_l or tok[:5] in sent_l)
        for intent in active_intents:
            score_s += sum(3 for ts in INTENT_MAP[intent]['stems'] if ts in sent_l)
        if score_s > best_sent_score:
            best_sent_score = score_s
            best_sent = sent

    excerpt = best_sent.strip()

    sources = [{
        "section": best_sec,
        "excerpt": excerpt,
        "label": f"Supporting evidence (from {best_sec})"
    }]

    formatted_answer = f"According to {best_sec}:\n\n{ans_text}\n\nSupporting evidence (from {best_sec}): \"{excerpt}\""
    return formatted_answer, sources

def semantic_answer(question, full_text, sections):
    """Backwards compatibility alias returning answer text only."""
    answer, _ = grounded_paper_qa(question, full_text, sections)
    return answer

# ----------------------------------------------------
# PDF Report Generator (Strict Pure Bytes Return)
# ----------------------------------------------------
def clean_for_pdf(text):
    if not text:
        return ""
    return clean_plain_text(text)


def _prepare_analysis_pdf_record(res, filename="Paper.pdf"):
    """Assemble a complete dictionary for pdf_export.build_analysis_pdf from an app.py analysis result."""
    res = res or {}
    clean_base = get_clean_base_filename(filename)
    clean_disp = get_clean_display_title(clean_base, text=res.get("full_text"), stored_title=res.get("title") or res.get("display_title"))
    scores = res.get("scores", {}) if isinstance(res.get("scores"), dict) else {}
    comp = get_overall_score(res)
    grade = get_grade(comp)

    stats = res.get("stats", {}) if isinstance(res.get("stats"), dict) else {}
    pdf_counts = compute_text_counts(res.get("full_text", "")) if res.get("full_text") else None
    rep_words = pdf_counts["word_count"] if pdf_counts and pdf_counts["word_count"] > 0 else int(stats.get("word_count", 0) or res.get("word_count", 0) or 0)
    rep_sents = pdf_counts["sentence_count"] if pdf_counts and pdf_counts["sentence_count"] > 0 else int(stats.get("sentence_count", 0) or res.get("sentence_count", 0) or 0)

    # Dimension explanations
    dim_explanations = dict(res.get("dimension_explanations") or {})
    for m_key in ("Language Quality", "Structural Coherence", "Argumentation", "Academic Style", "Readability"):
        if m_key not in dim_explanations:
            try:
                dim_explanations[m_key] = get_dimension_metric_explanation(m_key, res)
            except Exception:
                pass

    # Section summaries
    sec_summaries = {}
    rel_pdf_secs = get_reliable_academic_sections(res.get("sections", {})) if isinstance(res.get("sections"), dict) else {}
    for title, content in rel_pdf_secs.items():
        if not content or len(str(content).strip()) < 25:
            continue
        if title == "References":
            ref_list = extract_reference_entries(content)
            sec_summaries[title] = f"Extracted {len(ref_list)} reference entries from manuscript bibliography."
        else:
            sec_summaries[title] = summarize_text(content, 2)

    cur_sum_text = res.get("summary") or ""
    if not cur_sum_text and "summary_result" in st.session_state:
        s_obj = st.session_state.get("summary_result")
        if isinstance(s_obj, dict) and s_obj.get("success"):
            cur_sum_text = s_obj.get("summary", "")

    # Key findings
    l_val = float(scores.get("Language Quality", scores.get("Language", comp)) or comp)
    c_val = float(scores.get("Structural Coherence", scores.get("Coherence", comp)) or comp)
    a_val = float(scores.get("Argumentation", scores.get("Reasoning", comp)) or comp)
    s_val = float(scores.get("Academic Style", scores.get("Sophistication", comp)) or comp)
    r_val = float(scores.get("Readability", comp) or comp)
    findings_list = []
    if l_val >= 70:
        findings_list.append(f"Strong academic language quality ({l_val:.1f}/100): Clear sentence structure with formal scholarly cadence.")
    else:
        findings_list.append(f"Consistent prose rhythm ({l_val:.1f}/100): Sentences maintain an understandable scholarly cadence.")
    if c_val >= 65:
        findings_list.append(f"Consistent structural organization ({c_val:.1f}/100): Logical transitions between sections and topic sentences.")
    if a_val >= 65:
        findings_list.append(f"Evidence-supported arguments ({a_val:.1f}/100): Causal propositions are linked to empirical observations.")
    if s_val >= 65:
        findings_list.append(f"Lexical sophistication ({s_val:.1f}/100): Effective use of domain-specific scholarly nomenclature.")
    if r_val >= 60:
        findings_list.append(f"Accessible academic readability ({r_val:.1f}/100): Phrasing is well-calibrated for scholarly peer review.")

    return {
        "paper_title": clean_disp,
        "title": clean_disp,
        "filename": filename,
        "domain": res.get("domain") or "General Academic",
        "overall_score": comp,
        "grade": grade,
        "scores": scores,
        "dimension_explanations": dim_explanations,
        "word_count": rep_words,
        "sentence_count": rep_sents,
        "avg_sentence_len": float(stats.get("avg_sentence_len", 0.0) or 0.0),
        "key_findings": findings_list,
        "issues": res.get("issues") or [],
        "executive_summary": cur_sum_text,
        "section_summaries": sec_summaries,
    }


def generate_pdf_report(res, filename):
    """Generate a complete Unicode PDF report (with embedded charts & Page X of Y) in memory."""
    try:
        record = _prepare_analysis_pdf_record(res, filename)
        return pdf_export.build_analysis_pdf(record)
    except Exception as ex:
        try:
            cur_u = st.session_state.get("username") or (st.session_state.get("user_data") or {}).get("username")
            cur_role = (st.session_state.get("user_data") or {}).get("role") or "Student"
            cur_pg = st.session_state.get("current_page") or "PDF Report"
            auth.log_app_error(ex, page=cur_pg, username=cur_u, role=cur_role)
        except Exception:
            pass
        raise


def get_cached_summary_pdf(arg1, mode=None, summary_data=None, metadata=None, paper_title=None):
    """Return cached summary PDF bytes from st.session_state so reruns do not regenerate."""
    cache = st.session_state.setdefault("_pdf_cache_summary", {})
    if isinstance(arg1, dict) and summary_data is None:
        summary_data = arg1
        resolved_title = paper_title or arg1.get("document_info", {}).get("title") or arg1.get("title") or "Research Paper"
        resolved_mode = mode or arg1.get("summary_type") or "Medium"
    else:
        resolved_title = paper_title or str(arg1 or "Research Paper")
        resolved_mode = mode or "Medium"
    key = f"{clean_plain_text(resolved_title)}|{resolved_mode}|{len(str(summary_data))}"
    if key not in cache:
        cache[key] = pdf_export.build_summary_pdf(resolved_title, resolved_mode, summary_data, metadata)
    return cache[key]


def get_cached_analysis_pdf(res, filename="Paper.pdf", cache_key=None):
    """Return cached analysis/feedback PDF bytes from st.session_state so reruns do not regenerate."""
    cache = st.session_state.setdefault("_pdf_cache_analysis", {})
    comp = get_overall_score(res) if isinstance(res, dict) else 0.0
    key = str(cache_key or f"{filename}|{comp:.1f}")
    if key not in cache:
        cache[key] = generate_pdf_report(res, filename)
    return cache[key]


def get_cached_comparison_pdf(comp_data, comparison_meta=None):
    """Return cached multi-paper comparison PDF bytes from st.session_state."""
    cache = st.session_state.setdefault("_pdf_cache_comparison", {})
    sig_parts = []
    papers_payload = []
    for idx, p in enumerate(comp_data or []):
        p_title = p.get("clean_title") or p.get("title") or p.get("filename") or f"Paper {idx + 1}"
        ovr = get_overall_score(p)
        sig_parts.append(f"{clean_plain_text(p_title)}:{ovr:.1f}")
        p_copy = dict(p)
        p_copy["title"] = clean_plain_text(p_title)
        p_copy["overall_score"] = ovr
        p_copy["composite"] = ovr
        papers_payload.append(p_copy)
    key = "|".join(sig_parts)
    if key not in cache:
        cache[key] = pdf_export.build_comparison_pdf(papers_payload, comparison_meta or {})
    return cache[key]


def get_cached_chat_pdf(paper_title, chat_history):
    """Return cached Ask PaperIQ chat transcript PDF bytes from st.session_state."""
    cache = st.session_state.setdefault("_pdf_cache_chat", {})
    key = f"{clean_plain_text(paper_title)}|{len(chat_history or [])}"
    if key not in cache:
        cache[key] = pdf_export.build_chat_pdf(paper_title, chat_history or [])
    return cache[key]

# ----------------------------------------------------
# UI Theme Architecture: Light, Dark & System Modes
# ----------------------------------------------------
def get_system_theme():
    """Detects system appearance preference (Windows registry, environment, or default Dark)."""
    try:
        import winreg
        key = winreg.OpenKey(winreg.HKEY_CURRENT_USER, r"Software\Microsoft\Windows\CurrentVersion\Themes\Personalize")
        val, _ = winreg.QueryValueEx(key, "AppsUseLightTheme")
        return "Light" if val == 1 else "Dark"
    except Exception:
        return "Dark"

def get_effective_theme(theme_pref=None):
    """Resolves 'System' to current system mode, or returns explicit 'Light'/'Dark'."""
    if theme_pref is None:
        theme_pref = st.session_state.get("theme_preference", "System")
    if theme_pref == "System":
        return get_system_theme()
    return theme_pref

THEME_TOKENS = {
    "Light": {
        # Backgrounds
        "bg_page": "#F8FAFC",
        "bg_content": "#F8FAFC",
        "bg_sidebar": "#FFFFFF",
        "bg_card": "#FFFFFF",
        "bg_card_elevated": "#F1F5F9",
        "bg_card_inner": "#F8FAFC",
        "bg_user_card": "#F1F5F9",
        
        # Typography (Strict WCAG AA >= 4.5:1 contrast)
        "text_primary": "#0F172A",
        "text_secondary": "#334155",
        "text_muted": "#64748B",
        "text_placeholder": "#94A3B8",
        
        # Borders & Dividers
        "border_color": "#E2E8F0",
        "border_subtle": "#CBD5E1",
        "border_divider": "#E2E8F0",
        
        # Inputs & Forms
        "input_bg": "#FFFFFF",
        "input_border": "#CBD5E1",
        "input_text": "#0F172A",
        
        # Buttons & Controls
        "btn_bg": "#FFFFFF",
        "btn_text": "#1E293B",
        "btn_border": "#CBD5E1",
        "btn_hover_bg": "#F1F5F9",
        
        # Primary Brand Accent (PaperIQ Purple)
        "accent_primary": "#7C3AED",
        "accent_hover": "#6D28D9",
        "accent_primary_subtle": "rgba(124, 58, 237, 0.08)",
        "accent_primary_border": "rgba(124, 58, 237, 0.25)",
        "accent_purple": "#7C3AED",
        "accent_purple_text": "#6D28D9",
        "accent_purple_subtle": "rgba(124, 58, 237, 0.10)",
        "focus_outline": "#7C3AED",
        
        # Status Semantic Colors (Tested WCAG AA >= 4.5:1)
        "success_bg": "rgba(16, 185, 129, 0.12)",
        "success_border": "rgba(16, 185, 129, 0.35)",
        "success_text": "#047857",
        
        "warning_bg": "rgba(245, 158, 11, 0.12)",
        "warning_border": "rgba(245, 158, 11, 0.35)",
        "warning_text": "#B45309",
        
        "error_bg": "rgba(239, 68, 68, 0.12)",
        "error_border": "rgba(239, 68, 68, 0.35)",
        "error_text": "#B91C1C",
        
        "info_bg": "rgba(14, 165, 233, 0.10)",
        "info_border": "rgba(14, 165, 233, 0.30)",
        "info_text": "#0369A1",
        
        # Badges & Pills
        "badge_domain_bg": "rgba(124, 58, 237, 0.08)",
        "badge_domain_border": "rgba(124, 58, 237, 0.22)",
        "badge_domain_text": "#6D28D9",
        "badge_role_bg": "rgba(124, 58, 237, 0.12)",
        "badge_role_border": "rgba(124, 58, 237, 0.25)",
        "badge_role_text": "#6D28D9",
        
        # Chat
        "chat_user_bg": "#F1F5F9",
        "chat_user_border": "#E2E8F0",
        "chat_user_text": "#0F172A",
        "chat_bot_bg": "#FFFFFF",
        "chat_bot_border": "#E2E8F0",
        "chat_bot_text": "#1E293B",
        
        # Tables & DataFrames
        "table_header_bg": "#F8FAFC",
        "table_header_text": "#0F172A",
        "table_body_bg": "#FFFFFF",
        "table_row_alt": "#F8FAFC",
        "table_border": "#E2E8F0",
        "table_text": "#1E293B",
        
        # Charts
        "chart_bg": "#FFFFFF",
        "chart_plot_bg": "#F8FAFC",
        "chart_font": "#0F172A",
        "chart_grid": "#E2E8F0",
        "chart_line": "#CBD5E1",
        "chart_tick": "#475569",
        "chart_radar_line": "#7C3AED",
        "chart_radar_fill": "rgba(124, 58, 237, 0.22)",
        
        # Disabled Controls
        "disabled_bg": "#F1F5F9",
        "disabled_text": "#94A3B8",
        "disabled_border": "#E2E8F0",
    },
    "Dark": {
        # Backgrounds
        "bg_page": "#0B0F19",
        "bg_content": "#0B0F19",
        "bg_sidebar": "#0D1117",
        "bg_card": "#111827",
        "bg_card_elevated": "#1E293B",
        "bg_card_inner": "#1E293B",
        "bg_user_card": "#111827",
        
        # Typography
        "text_primary": "#F8FAFC",
        "text_secondary": "#E2E8F0",
        "text_muted": "#94A3B8",
        "text_placeholder": "#64748B",
        
        # Borders & Dividers
        "border_color": "#1F2937",
        "border_subtle": "#374151",
        "border_divider": "#1F2937",
        
        # Inputs & Forms
        "input_bg": "#111827",
        "input_border": "#374151",
        "input_text": "#F8FAFC",
        
        # Buttons & Controls
        "btn_bg": "#1E293B",
        "btn_text": "#F1F5F9",
        "btn_border": "#374151",
        "btn_hover_bg": "#334155",
        
        # Primary Brand Accent (PaperIQ Purple)
        "accent_primary": "#7C3AED",
        "accent_hover": "#6D28D9",
        "accent_primary_subtle": "rgba(124, 58, 237, 0.20)",
        "accent_primary_border": "rgba(124, 58, 237, 0.40)",
        "accent_purple": "#A78BFA",
        "accent_purple_text": "#C4B5FD",
        "accent_purple_subtle": "rgba(139, 92, 246, 0.20)",
        "focus_outline": "#8B5CF6",
        
        # Status Semantic Colors
        "success_bg": "rgba(16, 185, 129, 0.15)",
        "success_border": "rgba(16, 185, 129, 0.35)",
        "success_text": "#34D399",
        
        "warning_bg": "rgba(245, 158, 11, 0.15)",
        "warning_border": "rgba(245, 158, 11, 0.35)",
        "warning_text": "#FBBF24",
        
        "error_bg": "rgba(239, 68, 68, 0.15)",
        "error_border": "rgba(239, 68, 68, 0.35)",
        "error_text": "#F87171",
        
        "info_bg": "rgba(14, 165, 233, 0.15)",
        "info_border": "rgba(14, 165, 233, 0.30)",
        "info_text": "#38BDF8",
        
        # Badges & Pills
        "badge_domain_bg": "rgba(139, 92, 246, 0.20)",
        "badge_domain_border": "rgba(139, 92, 246, 0.35)",
        "badge_domain_text": "#C4B5FD",
        "badge_role_bg": "rgba(139, 92, 246, 0.20)",
        "badge_role_border": "rgba(139, 92, 246, 0.40)",
        "badge_role_text": "#C4B5FD",
        
        # Chat
        "chat_user_bg": "#1E293B",
        "chat_user_border": "#374151",
        "chat_user_text": "#F8FAFC",
        "chat_bot_bg": "#111827",
        "chat_bot_border": "#1F2937",
        "chat_bot_text": "#E2E8F0",
        
        # Tables & DataFrames
        "table_header_bg": "#1E293B",
        "table_header_text": "#F8FAFC",
        "table_body_bg": "#111827",
        "table_row_alt": "#162032",
        "table_border": "#1F2937",
        "table_text": "#E2E8F0",
        
        # Charts
        "chart_bg": "#111827",
        "chart_plot_bg": "#111827",
        "chart_font": "#E2E8F0",
        "chart_grid": "#1F2937",
        "chart_line": "#374151",
        "chart_tick": "#94A3B8",
        "chart_radar_line": "#A78BFA",
        "chart_radar_fill": "rgba(139, 92, 246, 0.28)",
        
        # Disabled Controls
        "disabled_bg": "#1F2937",
        "disabled_text": "#64748B",
        "disabled_border": "#374151",
    }
}

def get_theme_tokens(theme_pref=None):
    """Returns semantic theme token dictionary for the active theme."""
    eff = get_effective_theme(theme_pref)
    return THEME_TOKENS.get(eff, THEME_TOKENS["Dark"])

def get_chart_theme(theme_pref=None):
    """Returns Plotly layout color palette customized for the active theme."""
    eff = get_effective_theme(theme_pref)
    T = THEME_TOKENS.get(eff, THEME_TOKENS["Dark"])
    return {
        "paper_bgcolor": T["chart_bg"],
        "plot_bgcolor": T["chart_plot_bg"],
        "font_color": T["chart_font"],
        "grid_color": T["chart_grid"],
        "line_color": T["chart_line"],
        "tick_color": T["chart_tick"],
        "polar_bg": T["chart_plot_bg"],
        "legend_bg": T["chart_bg"],
        "legend_border": T["chart_line"],
        "radar_line": T["chart_radar_line"],
        "radar_fill": T["chart_radar_fill"]
    }

def inject_theme_css(theme_pref="System"):
    """
    Injects comprehensive CSS rules tailored for Light, Dark, or System mode.
    Maintains clean layout, high contrast, readable cards, and professional purple branding.
    """
    eff = get_effective_theme(theme_pref)
    T = THEME_TOKENS.get(eff, THEME_TOKENS["Dark"])
    is_light = (eff == "Light")
    
    # CSS custom properties
    css_vars = f"""
    :root, .stApp {{
        --bg-page: {T['bg_page']};
        --bg-content: {T['bg_content']};
        --bg-sidebar: {T['bg_sidebar']};
        --bg-card: {T['bg_card']};
        --bg-card-elevated: {T['bg_card_elevated']};
        --bg-card-inner: {T['bg_card_inner']};
        --bg-user-card: {T['bg_user_card']};
        --text-primary: {T['text_primary']};
        --text-secondary: {T['text_secondary']};
        --text-muted: {T['text_muted']};
        --text-placeholder: {T['text_placeholder']};
        --border-color: {T['border_color']};
        --border-subtle: {T['border_subtle']};
        --border-divider: {T['border_divider']};
        --input-bg: {T['input_bg']};
        --input-border: {T['input_border']};
        --input-text: {T['input_text']};
        --btn-bg: {T['btn_bg']};
        --btn-text: {T['btn_text']};
        --btn-border: {T['btn_border']};
        --btn-hover-bg: {T['btn_hover_bg']};
        --accent-primary: {T['accent_primary']};
        --accent-hover: {T['accent_hover']};
        --accent-purple: {T['accent_purple']};
        --accent-purple-text: {T['accent_purple_text']};
        --accent-purple-subtle: {T['accent_purple_subtle']};
        --success-bg: {T['success_bg']};
        --success-border: {T['success_border']};
        --success-text: {T['success_text']};
        --warning-bg: {T['warning_bg']};
        --warning-border: {T['warning_border']};
        --warning-text: {T['warning_text']};
        --error-bg: {T['error_bg']};
        --error-border: {T['error_border']};
        --error-text: {T['error_text']};
        --info-bg: {T['info_bg']};
        --info-border: {T['info_border']};
        --info-text: {T['info_text']};
    }}
    """
    
    # Common layout styles
    base_css = """
    header[data-testid="stHeader"] {
        background: transparent !important;
    }
    .block-container {
        padding-top: 1.8rem !important;
        padding-bottom: 3rem !important;
        max-width: 1250px !important;
    }
    section[data-testid="stSidebar"] {
        overflow-x: hidden !important;
        overflow-y: auto !important;
    }
    section[data-testid="stSidebar"] > div:first-child,
    div[data-testid="stSidebarContent"] {
        padding-top: 0.75rem !important;
        overflow: visible !important;
        height: auto !important;
    }
    div[data-testid="stSidebarUserContent"] {
        padding-top: 0.75rem !important;
        padding-bottom: 1.5rem !important;
        overflow: visible !important;
    }
    div[data-testid="stSidebarHeader"] {
        height: auto !important;
        min-height: 2rem !important;
        padding-bottom: 0 !important;
    }
    .sidebar-logo-card {
        margin-top: 4px !important;
    }
    """
    
    if is_light:
        theme_css = f"""
        /* Main Light Background & Typography */
        .stApp {{
            background-color: {T['bg_page']} !important;
            font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, "Helvetica Neue", Arial, sans-serif !important;
            color: {T['text_primary']} !important;
        }}
        h1, h2, h3, h4, h5, h6 {{
            color: {T['text_primary']} !important;
            font-weight: 700 !important;
            letter-spacing: -0.02em !important;
        }}
        .stApp p, .stApp span, .stApp label, .stApp li {{
            color: {T['text_secondary']} !important;
        }}
        .stApp b, .stApp strong {{
            color: {T['text_primary']} !important;
            font-weight: 700 !important;
        }}
        .stApp a {{
            color: {T['accent_purple']} !important;
        }}
        .stApp hr {{
            border-color: {T['border_divider']} !important;
        }}

        /* Sidebar: Clean White with Crisp Hierarchy */
        section[data-testid="stSidebar"] {{
            background-color: {T['bg_sidebar']} !important;
            border-right: 1px solid {T['border_color']} !important;
            padding-top: 1.2rem !important;
        }}
        section[data-testid="stSidebar"] h1,
        section[data-testid="stSidebar"] h2,
        section[data-testid="stSidebar"] h3 {{
            color: {T['text_primary']} !important;
        }}
        section[data-testid="stSidebar"] p,
        section[data-testid="stSidebar"] span,
        section[data-testid="stSidebar"] label {{
            color: {T['text_secondary']} !important;
        }}
        .sidebar-logo-card {{
            padding: 10px 4px 16px 4px;
            border-bottom: 1px solid {T['border_divider']};
            margin-bottom: 16px;
        }}
        .sidebar-logo-title {{
            font-size: 20px;
            font-weight: 800;
            background: linear-gradient(135deg, #7C3AED, #4F46E5);
            -webkit-background-clip: text;
            -webkit-text-fill-color: transparent;
            letter-spacing: 0.05em;
            margin: 0;
        }}
        .sidebar-logo-sub {{
            font-size: 11px;
            color: {T['text_muted']} !important;
            text-transform: uppercase;
            letter-spacing: 0.08em;
            margin-top: 2px;
        }}
        .sidebar-user-card {{
            background: {T['bg_user_card']};
            border: 1px solid {T['border_color']};
            border-radius: 10px;
            padding: 12px;
            margin-bottom: 16px;
        }}
        .sidebar-user-role {{
            display: inline-block;
            background: {T['badge_role_bg']};
            color: {T['badge_role_text']} !important;
            border: 1px solid {T['badge_role_border']};
            font-size: 11px;
            font-weight: 600;
            padding: 2px 8px;
            border-radius: 12px;
            margin-top: 4px;
        }}

        /* Sidebar Navigation Buttons */
        section[data-testid="stSidebar"] div.stButton > button {{
            width: 100% !important;
            text-align: left !important;
            justify-content: flex-start !important;
            background-color: transparent !important;
            border: 1px solid transparent !important;
            color: {T['text_secondary']} !important;
            border-radius: 8px !important;
            padding: 0.55rem 0.85rem !important;
            margin-bottom: 3px !important;
            font-size: 14px !important;
            font-weight: 500 !important;
            box-shadow: none !important;
            display: flex !important;
            align-items: center !important;
            transition: all 0.15s ease-in-out !important;
        }}
        section[data-testid="stSidebar"] div.stButton > button p,
        section[data-testid="stSidebar"] div.stButton > button span,
        section[data-testid="stSidebar"] div.stButton > button div {{
            color: inherit !important;
            font-weight: inherit !important;
            text-align: left !important;
            margin: 0 !important;
        }}
        section[data-testid="stSidebar"] div.stButton > button:hover {{
            background-color: {T['btn_hover_bg']} !important;
            color: {T['text_primary']} !important;
            border-color: {T['border_color']} !important;
        }}
        section[data-testid="stSidebar"] div.stButton > button:hover p,
        section[data-testid="stSidebar"] div.stButton > button:hover span {{
            color: {T['text_primary']} !important;
        }}
        section[data-testid="stSidebar"] div.stButton > button[kind="primary"] {{
            background: linear-gradient(90deg, rgba(124, 58, 237, 0.14) 0%, rgba(99, 102, 241, 0.08) 100%) !important;
            color: {T['accent_purple_text']} !important;
            border: 1px solid #C4B5FD !important;
            border-left: 4px solid {T['accent_primary']} !important;
            font-weight: 700 !important;
            box-shadow: 0 2px 6px rgba(124, 58, 237, 0.12) !important;
        }}
        section[data-testid="stSidebar"] div.stButton > button[kind="primary"] p,
        section[data-testid="stSidebar"] div.stButton > button[kind="primary"] span {{
            color: {T['accent_purple_text']} !important;
            font-weight: 700 !important;
        }}

        /* Cards & Containers */
        .iq-card {{
            background-color: {T['bg_card']} !important;
            border: 1px solid {T['border_color']} !important;
            border-radius: 12px !important;
            padding: 20px !important;
            margin-bottom: 16px !important;
            box-shadow: 0 2px 10px rgba(0, 0, 0, 0.05) !important;
        }}
        .iq-card h3, .iq-card h4 {{
            margin-top: 0 !important;
            color: {T['text_primary']} !important;
        }}
        div[data-testid="stVerticalBlockBorderWrapper"] {{
            background-color: {T['bg_card']} !important;
            border: 1px solid {T['border_color']} !important;
            border-radius: 12px !important;
            padding: 14px 18px !important;
            margin-bottom: 12px !important;
            box-shadow: 0 2px 8px rgba(0, 0, 0, 0.04) !important;
        }}
        div[data-testid="stVerticalBlockBorderWrapper"] > div {{
            background-color: transparent !important;
        }}
        .journal-hero-card {{
            background: {T['bg_card']} !important;
            border: 1px solid {T['border_color']} !important;
            border-radius: 12px !important;
            padding: 18px 22px !important;
            margin-bottom: 18px !important;
            box-shadow: 0 2px 8px rgba(0, 0, 0, 0.04) !important;
        }}
        .journal-hero-card h4 {{
            color: {T['text_primary']} !important;
            margin: 0 !important;
        }}

        /* Metric Displays */
        [data-testid="stMetric"] {{
            background-color: {T['bg_card']} !important;
            border: 1px solid {T['border_color']} !important;
            border-radius: 12px !important;
            padding: 16px 18px !important;
            box-shadow: 0 2px 6px rgba(0, 0, 0, 0.04) !important;
        }}
        [data-testid="stMetricLabel"] p {{
            color: {T['text_muted']} !important;
            font-size: 12px !important;
            font-weight: 600 !important;
            text-transform: uppercase !important;
            letter-spacing: 0.06em !important;
        }}
        [data-testid="stMetricValue"] div {{
            color: {T['text_primary']} !important;
            font-size: 28px !important;
            font-weight: 800 !important;
        }}

        /* Score Cards */
        .score-card {{
            background-color: {T['bg_card']};
            border: 1px solid {T['border_color']};
            border-radius: 12px;
            padding: 18px;
            height: 100%;
            box-shadow: 0 2px 6px rgba(0, 0, 0, 0.04);
        }}
        .score-card-title {{
            font-size: 12px;
            font-weight: 700;
            color: {T['text_muted']} !important;
            text-transform: uppercase;
            letter-spacing: 0.06em;
        }}
        .score-card-val {{
            font-size: 28px;
            font-weight: 800;
            color: {T['text_primary']} !important;
            margin: 4px 0;
        }}
        .score-card-interp {{
            font-size: 12px;
            color: {T['text_secondary']} !important;
            line-height: 1.4;
            margin-top: 6px;
        }}

        /* Buttons */
        .stButton > button,
        .stButton > button:not([kind="primary"]),
        [data-testid="stBaseButton-secondary"] {{
            background-color: {T['btn_bg']} !important;
            color: {T['btn_text']} !important;
            border: 1px solid {T['btn_border']} !important;
            border-radius: 8px !important;
            font-weight: 600 !important;
            padding: 0.5rem 1rem !important;
            box-shadow: 0 1px 3px rgba(0, 0, 0, 0.06) !important;
            transition: all 0.2s ease !important;
        }}
        .stButton > button *,
        .stButton > button:not([kind="primary"]) *,
        [data-testid="stBaseButton-secondary"] * {{
            color: {T['btn_text']} !important;
        }}
        .stButton > button:hover,
        .stButton > button:not([kind="primary"]):hover,
        [data-testid="stBaseButton-secondary"]:hover {{
            background-color: {T['btn_hover_bg']} !important;
            border-color: {T['accent_primary']} !important;
            color: {T['accent_primary']} !important;
        }}
        .stButton > button:hover *,
        .stButton > button:not([kind="primary"]):hover *,
        [data-testid="stBaseButton-secondary"]:hover * {{
            color: {T['accent_primary']} !important;
        }}
        .stButton > button[kind="primary"],
        [data-testid="stBaseButton-primary"] {{
            background: linear-gradient(135deg, #7C3AED, #4F46E5) !important;
            color: #FFFFFF !important;
            border: none !important;
            box-shadow: 0 4px 14px rgba(124, 58, 237, 0.3) !important;
        }}
        .stButton > button[kind="primary"] *,
        [data-testid="stBaseButton-primary"] * {{
            color: #FFFFFF !important;
        }}
        .stButton > button[kind="primary"]:hover,
        [data-testid="stBaseButton-primary"]:hover {{
            background: linear-gradient(135deg, #6D28D9, #4338CA) !important;
            box-shadow: 0 6px 18px rgba(124, 58, 237, 0.4) !important;
        }}
        .stDownloadButton > button {{
            background: linear-gradient(135deg, #4F46E5, #2563EB) !important;
            color: #FFFFFF !important;
            border: none !important;
            border-radius: 8px !important;
            font-weight: 600 !important;
            padding: 0.5rem 1rem !important;
            box-shadow: 0 4px 12px rgba(79, 70, 229, 0.25) !important;
        }}
        .stDownloadButton > button * {{
            color: #FFFFFF !important;
        }}

        /* Popover Trigger Buttons (Explicit High-Specificity Contrast Repair) */
        .stApp [data-testid="stPopoverButton"],
        .stApp [data-testid="stPopover"] button,
        .stApp .stPopover button,
        .stApp div[data-testid="stPopover"] button,
        .stApp [data-testid="stPopover"] [data-testid="stBaseButton-secondary"],
        .stApp .stPopover [data-testid="stBaseButton-secondary"],
        [data-testid="stPopoverButton"],
        [data-testid="stPopover"] button,
        .stPopover button {{
            background-color: {T['btn_bg']} !important;
            color: {T['btn_text']} !important;
            border: 1.5px solid {T['btn_border']} !important;
            border-radius: 8px !important;
            font-weight: 600 !important;
            font-size: 14px !important;
            min-height: 42px !important;
            padding: 0.5rem 1rem !important;
            box-shadow: 0 1px 3px rgba(0, 0, 0, 0.08) !important;
            transition: all 0.2s ease !important;
            display: inline-flex !important;
            align-items: center !important;
            justify-content: center !important;
            cursor: pointer !important;
            width: 100% !important;
        }}
        .stApp [data-testid="stPopoverButton"] *,
        .stApp [data-testid="stPopover"] button *,
        .stApp .stPopover button *,
        .stApp div[data-testid="stPopover"] button *,
        [data-testid="stPopoverButton"] *,
        [data-testid="stPopover"] button *,
        .stPopover button * {{
            color: {T['btn_text']} !important;
            font-weight: 600 !important;
            fill: currentColor !important;
            stroke: currentColor !important;
        }}
        .stApp [data-testid="stPopoverButton"] svg,
        .stApp [data-testid="stPopover"] button svg,
        .stApp .stPopover button svg,
        .stApp [data-testid="stPopoverButton"] [data-testid="stIconMaterial"],
        .stApp [data-testid="stPopoverButton"] span[class*="DynamicIcon"],
        .stApp [data-testid="stPopoverButton"] div[class*="StyledPopoverExpansionIcon"] {{
            color: {T['btn_text']} !important;
            fill: currentColor !important;
            stroke: currentColor !important;
            display: inline-flex !important;
            align-items: center !important;
        }}
        .stApp [data-testid="stPopoverButton"]:hover,
        .stApp [data-testid="stPopover"] button:hover,
        .stApp .stPopover button:hover,
        .stApp div[data-testid="stPopover"] button:hover,
        [data-testid="stPopoverButton"]:hover,
        [data-testid="stPopover"] button:hover,
        .stPopover button:hover {{
            background-color: {T['btn_hover_bg']} !important;
            border-color: {T['accent_primary']} !important;
            color: {T['accent_primary']} !important;
            box-shadow: 0 2px 6px rgba(124, 58, 237, 0.15) !important;
        }}
        .stApp [data-testid="stPopoverButton"]:hover *,
        .stApp [data-testid="stPopover"] button:hover *,
        .stApp .stPopover button:hover *,
        .stApp div[data-testid="stPopover"] button:hover *,
        [data-testid="stPopoverButton"]:hover *,
        [data-testid="stPopover"] button:hover *,
        .stPopover button:hover * {{
            color: {T['accent_primary']} !important;
            fill: currentColor !important;
            stroke: currentColor !important;
        }}
        .stApp [data-testid="stPopoverButton"]:hover svg,
        .stApp [data-testid="stPopover"] button:hover svg,
        .stApp .stPopover button:hover svg,
        .stApp [data-testid="stPopoverButton"]:hover [data-testid="stIconMaterial"],
        .stApp [data-testid="stPopoverButton"]:hover span[class*="DynamicIcon"] {{
            color: {T['accent_primary']} !important;
            fill: currentColor !important;
            stroke: currentColor !important;
        }}

        /* Focus rings and disabled buttons */
        .stButton > button:focus-visible,
        .stDownloadButton > button:focus-visible,
        [data-testid="stFileUploaderDropzone"] button:focus-visible,
        .stApp [data-testid="stPopoverButton"]:focus-visible,
        .stApp [data-testid="stPopover"] button:focus-visible,
        [data-testid="stPopoverButton"]:focus-visible,
        [data-testid="stPopover"] button:focus-visible {{
            outline: 2px solid {T['accent_primary']} !important;
            outline-offset: 2px !important;
            box-shadow: 0 0 0 3px rgba(124, 58, 237, 0.25) !important;
        }}
        .stButton > button:disabled,
        .stButton > button[disabled],
        .stDownloadButton > button:disabled,
        .stDownloadButton > button[disabled],
        [data-testid="stFileUploaderDropzone"] button:disabled,
        [data-testid="stBaseButton-secondary"]:disabled,
        [data-testid="stBaseButton-secondary"][disabled],
        .stApp [data-testid="stPopoverButton"]:disabled,
        .stApp [data-testid="stPopoverButton"][disabled],
        .stApp [data-testid="stPopover"] button:disabled,
        .stApp [data-testid="stPopover"] button[disabled],
        [data-testid="stPopoverButton"]:disabled,
        [data-testid="stPopover"] button:disabled {{
            background-color: {T['disabled_bg']} !important;
            color: {T['disabled_text']} !important;
            border: 1px solid {T['disabled_border']} !important;
            cursor: not-allowed !important;
            opacity: 0.65 !important;
            box-shadow: none !important;
        }}
        .stButton > button:disabled *,
        .stButton > button[disabled] *,
        .stDownloadButton > button:disabled *,
        [data-testid="stFileUploaderDropzone"] button:disabled *,
        [data-testid="stBaseButton-secondary"]:disabled *,
        [data-testid="stBaseButton-secondary"][disabled] *,
        .stApp [data-testid="stPopoverButton"]:disabled *,
        .stApp [data-testid="stPopoverButton"][disabled] *,
        .stApp [data-testid="stPopover"] button:disabled *,
        [data-testid="stPopoverButton"]:disabled *,
        [data-testid="stPopover"] button:disabled * {{
            color: {T['disabled_text']} !important;
            fill: currentColor !important;
        }}

        /* File Uploader Container & Dropzone */
        [data-testid="stFileUploader"] {{
            background-color: {T['bg_card']} !important;
            border: 1px solid {T['border_color']} !important;
            border-radius: 12px !important;
            padding: 18px !important;
        }}
        [data-testid="stFileUploaderDropzone"] {{
            border: 2px dashed {T['border_subtle']} !important;
            border-radius: 10px !important;
            background-color: {T['bg_card_inner']} !important;
        }}
        [data-testid="stFileUploaderDropzoneInstructions"],
        [data-testid="stFileUploaderDropzoneInstructions"] span,
        [data-testid="stFileUploaderDropzoneInstructions"] small,
        [data-testid="stFileUploaderDropzone"] > div > span {{
            color: {T['text_secondary']} !important;
        }}

        /* File Uploader "Browse files" Button - Explicit high-specificity styles */
        [data-testid="stFileUploaderDropzone"] button,
        [data-testid="stFileUploaderDropzone"] [data-testid="stBaseButton-secondary"],
        section[data-testid="stFileUploadDropzone"] button,
        section[data-testid="stFileUploadDropzone"] [data-testid="stBaseButton-secondary"],
        [data-testid="stFileUploader"] section button {{
            background-color: {T['btn_bg']} !important;
            color: {T['btn_text']} !important;
            border: 1.5px solid {T['btn_border']} !important;
            border-radius: 8px !important;
            font-weight: 600 !important;
            padding: 0.45rem 1.1rem !important;
            box-shadow: 0 1px 3px rgba(0, 0, 0, 0.08) !important;
            transition: all 0.2s ease !important;
            cursor: pointer !important;
        }}
        [data-testid="stFileUploaderDropzone"] button *,
        [data-testid="stFileUploaderDropzone"] [data-testid="stBaseButton-secondary"] *,
        section[data-testid="stFileUploadDropzone"] button *,
        section[data-testid="stFileUploadDropzone"] [data-testid="stBaseButton-secondary"] * {{
            color: {T['btn_text']} !important;
            font-weight: 600 !important;
        }}
        [data-testid="stFileUploaderDropzone"] button:hover,
        [data-testid="stFileUploaderDropzone"] [data-testid="stBaseButton-secondary"]:hover,
        section[data-testid="stFileUploadDropzone"] button:hover,
        section[data-testid="stFileUploadDropzone"] [data-testid="stBaseButton-secondary"]:hover {{
            background-color: {T['btn_hover_bg']} !important;
            border-color: {T['accent_primary']} !important;
            color: {T['accent_primary']} !important;
            box-shadow: 0 2px 6px rgba(124, 58, 237, 0.15) !important;
        }}
        [data-testid="stFileUploaderDropzone"] button:hover *,
        [data-testid="stFileUploaderDropzone"] [data-testid="stBaseButton-secondary"]:hover * {{
            color: {T['accent_primary']} !important;
        }}

        /* Uploaded File Chip */
        [data-testid="stFileUploaderFile"] {{
            background-color: {T['bg_card_elevated']} !important;
            border: 1px solid {T['border_subtle']} !important;
            border-radius: 8px !important;
            color: {T['text_primary']} !important;
        }}
        [data-testid="stFileUploaderFile"] * {{
            color: {T['text_primary']} !important;
        }}
        [data-testid="stFileUploaderDeleteBtn"] button {{
            color: {T['text_muted']} !important;
            border: none !important;
            background: transparent !important;
        }}
        [data-testid="stFileUploaderDeleteBtn"] button:hover {{
            color: #EF4444 !important;
        }}

        /* Form Inputs */
        div[data-baseweb="select"] > div {{
            background-color: {T['input_bg']} !important;
            border-color: {T['input_border']} !important;
            color: {T['text_primary']} !important;
        }}
        div[data-baseweb="select"] * {{
            color: {T['text_primary']} !important;
        }}
        div[data-baseweb="popover"],
        div[data-baseweb="menu"],
        ul[role="listbox"] {{
            background-color: {T['bg_card']} !important;
            border: 1px solid {T['border_subtle']} !important;
            box-shadow: 0 4px 16px rgba(0, 0, 0, 0.1) !important;
        }}
        li[role="option"] {{
            background-color: {T['bg_card']} !important;
            color: {T['text_primary']} !important;
        }}
        li[role="option"]:hover,
        li[role="option"][aria-selected="true"] {{
            background-color: {T['bg_card_elevated']} !important;
            color: {T['accent_purple_text']} !important;
        }}
        div[data-testid="stTextInput"] input {{
            background-color: {T['input_bg']} !important;
            border: 1px solid {T['input_border']} !important;
            border-radius: 8px !important;
            color: {T['text_primary']} !important;
            padding: 8px 12px !important;
        }}
        div[data-testid="stTextInput"] input::placeholder {{
            color: {T['text_placeholder']} !important;
        }}
        div[data-testid="stTextInput"] input:focus {{
            border-color: {T['accent_primary']} !important;
            box-shadow: 0 0 0 2px rgba(124, 58, 237, 0.2) !important;
        }}
        div[data-testid="stRadio"] label, div[data-testid="stRadio"] p, div[data-testid="stRadio"] span {{
            color: {T['text_primary']} !important;
        }}
        div[data-testid="stCheckbox"] label[data-baseweb="checkbox"] {{
            display: inline-flex !important;
            align-items: center !important;
            gap: 8px !important;
            margin: 0 !important;
            cursor: pointer !important;
        }}
        div[data-testid="stCheckbox"] label[data-baseweb="checkbox"] > div,
        div[data-testid="stCheckbox"] [data-testid="stWidgetLabel"] {{
            margin: 0 !important;
            padding: 0 !important;
            min-height: 0 !important;
            display: inline-flex !important;
            align-items: center !important;
        }}
        div[data-testid="stCheckbox"] label, div[data-testid="stCheckbox"] p, div[data-testid="stCheckbox"] span {{
            color: {T['text_primary']} !important;
            margin: 0 !important;
            line-height: 1.35 !important;
        }}
        div[data-testid="stCheckbox"] label[data-baseweb="checkbox"] > span:first-child {{
            width: 18px !important;
            height: 18px !important;
            min-width: 18px !important;
            min-height: 18px !important;
            background-color: #FFFFFF !important;
            background-image: none !important;
            border: 1.5px solid #64748B !important;
            border-radius: 5px !important;
            margin: 0 !important;
            box-sizing: border-box !important;
            transition: all 0.18s ease !important;
        }}
        div[data-testid="stCheckbox"] label[data-baseweb="checkbox"]:hover > span:first-child {{
            border-color: #7C3AED !important;
            box-shadow: 0 0 0 2px rgba(124, 58, 237, 0.15) !important;
        }}
        div[data-testid="stCheckbox"] label[data-baseweb="checkbox"]:has(input:checked) > span:first-child,
        div[data-testid="stCheckbox"] label[data-baseweb="checkbox"]:has(input[aria-checked="true"]) > span:first-child {{
            background-color: #7C3AED !important;
            border: 1.5px solid #6D28D9 !important;
            background-image: url("data:image/svg+xml,%3Csvg xmlns='http://www.w3.org/2000/svg' viewBox='0 0 16 16' fill='none'%3E%3Cpath d='M3.5 8.2L6.6 11.3L12.8 4.8' stroke='%23FFFFFF' stroke-width='2.3' stroke-linecap='round' stroke-linejoin='round'/%3E%3C/svg%3E") !important;
            background-repeat: no-repeat !important;
            background-position: center !important;
            background-size: 13px 13px !important;
            box-shadow: 0 2px 6px rgba(124, 58, 237, 0.28) !important;
        }}

        /* Expanders */
        [data-testid="stExpander"] {{
            background-color: {T['bg_card']} !important;
            border: 1px solid {T['border_subtle']} !important;
            border-radius: 10px !important;
            box-shadow: 0 1px 3px rgba(0, 0, 0, 0.04) !important;
            margin-bottom: 12px !important;
        }}
        [data-testid="stExpander"] summary {{
            background-color: {T['bg_card_inner']} !important;
            color: {T['text_primary']} !important;
            font-weight: 600 !important;
            border-radius: 10px 10px 0 0 !important;
            padding: 10px 14px !important;
        }}
        [data-testid="stExpander"] summary * {{
            color: {T['text_primary']} !important;
        }}
        [data-testid="stExpander"] summary:hover {{
            background-color: {T['bg_card_elevated']} !important;
            color: {T['accent_purple']} !important;
        }}
        [data-testid="stExpander"] summary:hover * {{
            color: {T['accent_purple']} !important;
        }}
        [data-testid="stExpander"] [data-testid="stExpanderDetails"] {{
            background-color: {T['bg_card']} !important;
            color: {T['text_secondary']} !important;
            padding: 14px 18px !important;
        }}

        /* Tabs */
        .stTabs [data-baseweb="tab-list"] {{
            background-color: transparent !important;
            border-bottom: 2px solid {T['border_divider']} !important;
            gap: 8px !important;
        }}
        .stTabs [data-baseweb="tab"] {{
            background-color: transparent !important;
            color: {T['text_muted']} !important;
            font-weight: 600 !important;
            font-size: 14px !important;
            border: none !important;
            padding: 8px 16px !important;
        }}
        .stTabs [data-baseweb="tab"]:hover {{
            color: {T['text_primary']} !important;
        }}
        .stTabs [data-baseweb="tab"][aria-selected="true"] {{
            color: {T['accent_purple_text']} !important;
            border-bottom: 3px solid {T['accent_primary']} !important;
            font-weight: 700 !important;
        }}
        .stTabs [data-baseweb="tab-highlight"] {{
            background-color: {T['accent_primary']} !important;
        }}

        /* Chat Styling */
        .chat-msg-user {{
            background-color: {T['chat_user_bg']} !important;
            border: 1px solid {T['chat_user_border']} !important;
            border-radius: 10px !important;
            padding: 12px 16px !important;
            margin-bottom: 10px !important;
            color: {T['chat_user_text']} !important;
        }}
        .chat-msg-user * {{
            color: {T['chat_user_text']} !important;
        }}
        .chat-msg-bot {{
            background-color: {T['chat_bot_bg']} !important;
            border-left: 4px solid {T['accent_primary']} !important;
            border-top: 1px solid {T['chat_bot_border']} !important;
            border-right: 1px solid {T['chat_bot_border']} !important;
            border-bottom: 1px solid {T['chat_bot_border']} !important;
            border-radius: 10px !important;
            padding: 14px 18px !important;
            margin-bottom: 14px !important;
            color: {T['chat_bot_text']} !important;
            box-shadow: 0 2px 6px rgba(0, 0, 0, 0.04) !important;
        }}
        .chat-msg-bot * {{
            color: {T['chat_bot_text']} !important;
        }}

        /* Tables & DataFrames */
        [data-testid="stDataFrame"] {{
            border: 1px solid {T['border_color']} !important;
            border-radius: 10px !important;
            background-color: {T['bg_card']} !important;
        }}
        [data-testid="stTable"] {{
            background-color: {T['bg_card']} !important;
            color: {T['text_primary']} !important;
        }}
        [data-testid="stTable"] th {{
            background-color: {T['table_header_bg']} !important;
            color: {T['table_header_text']} !important;
            border-bottom: 2px solid {T['border_subtle']} !important;
            font-weight: 700 !important;
        }}
        [data-testid="stTable"] td {{
            background-color: {T['bg_card']} !important;
            color: {T['text_secondary']} !important;
            border-bottom: 1px solid {T['border_color']} !important;
        }}
        [data-testid="stTable"] tr:nth-child(even) td {{
            background-color: {T['table_row_alt']} !important;
        }}

        /* Badges */
        .domain-badge {{
            display: inline-block !important;
            background: {T['badge_domain_bg']} !important;
            color: {T['badge_domain_text']} !important;
            border: 1px solid {T['badge_domain_border']} !important;
            font-size: 11.5px !important;
            font-weight: 600 !important;
            padding: 2px 9px !important;
            border-radius: 6px !important;
            margin-right: 6px !important;
            margin-bottom: 4px !important;
        }}

        /* Popovers Body */
        [data-testid="stPopoverBody"],
        div[data-baseweb="popover"] {{
            background-color: {T['bg_card']} !important;
            border: 1px solid {T['border_color']} !important;
            border-radius: 12px !important;
            box-shadow: 0 8px 24px rgba(0, 0, 0, 0.16) !important;
            color: {T['text_primary']} !important;
            padding: 16px !important;
            max-width: 95vw !important;
        }}
        [data-testid="stPopoverBody"] *,
        div[data-baseweb="popover"] * {{
            color: {T['text_primary']} !important;
        }}
        [data-testid="stPopoverBody"] p,
        [data-testid="stPopoverBody"] span,
        [data-testid="stPopoverBody"] [data-testid="stCaptionContainer"] p {{
            color: {T['text_secondary']} !important;
        }}
        [data-testid="stPopoverBody"] pre,
        [data-testid="stPopoverBody"] code {{
            background-color: {T['bg_card_elevated']} !important;
            color: {T['text_primary']} !important;
            border: 1px solid {T['border_subtle']} !important;
        }}
        [data-testid="stPopoverBody"] button,
        [data-testid="stPopoverBody"] [data-testid="stCodeCopyButton"] {{
            background-color: {T['bg_card_elevated']} !important;
            color: {T['text_primary']} !important;
            border: 1px solid {T['border_subtle']} !important;
            border-radius: 6px !important;
        }}
        [data-testid="stPopoverBody"] button:hover,
        [data-testid="stPopoverBody"] [data-testid="stCodeCopyButton"]:hover {{
            background-color: {T['btn_hover_bg']} !important;
            color: {T['accent_primary']} !important;
            border-color: {T['accent_primary']} !important;
        }}

        /* Custom Responsive Comparison Table */
        .iq-comparison-table-wrapper {{
            overflow-x: auto;
            border: 1px solid {T['border_color']};
            border-radius: 12px;
            background-color: {T['bg_card']};
            margin: 16px 0 24px 0;
            box-shadow: 0 2px 8px rgba(0, 0, 0, 0.05);
        }}
        .iq-comp-table {{
            width: 100%;
            border-collapse: collapse;
            font-size: 13.5px;
            color: {T['text_primary']};
        }}
        .iq-comp-table th {{
            background-color: {T['table_header_bg']};
            color: {T['table_header_text']};
            padding: 14px 18px;
            font-weight: 700;
            border-bottom: 2px solid {T['border_subtle']};
            font-size: 12px;
            letter-spacing: 0.04em;
        }}
        .iq-comp-table td {{
            padding: 12px 18px;
            border-bottom: 1px solid {T['border_color']};
            color: {T['text_secondary']};
        }}
        .iq-comp-table tr:last-child td {{
            border-bottom: none;
            font-weight: 700;
            background-color: {T['bg_card_inner']};
        }}
        .iq-comp-table tr:nth-child(even) {{
            background-color: {T['table_row_alt']};
        }}
        .iq-comp-table tr:hover {{
            background-color: {T['bg_card_elevated']};
        }}
        .iq-comp-dim-name {{
            color: {T['text_primary']} !important;
            font-weight: 600;
        }}
        .iq-comp-val-cell {{
            text-align: center;
            font-weight: 700;
            font-size: 14px;
        }}
        .iq-comp-val-best {{
            color: {T['success_text']} !important;
            font-weight: 800;
        }}
        .iq-comp-val-pill {{
            display: inline-block;
            padding: 3px 10px;
            border-radius: 6px;
            font-weight: 700;
        }}

        /* Status Widgets */
        [data-testid="stStatusWidget"] {{
            background-color: {T['bg_card']} !important;
            border: 1px solid {T['border_subtle']} !important;
        }}
        [data-testid="stStatusWidget"] * {{
            color: {T['text_primary']} !important;
        }}

        /* Code Blocks */
        pre, code, [data-testid="stCode"] {{
            background-color: {T['bg_card_elevated']} !important;
            color: {T['text_primary']} !important;
            border-radius: 8px !important;
        }}

        /* Global Light Mode Overrides for Residual Hardcoded Dark Styles */
        .stApp [style*="color: #FFFFFF"], .stApp [style*="color:#FFFFFF"] {{ color: {T['text_primary']} !important; }}
        .stApp [style*="color: #F8FAFC"], .stApp [style*="color:#F8FAFC"] {{ color: {T['text_primary']} !important; }}
        .stApp [style*="color: #CBD5E1"], .stApp [style*="color:#CBD5E1"] {{ color: {T['text_secondary']} !important; }}
        .stApp [style*="color: #94A3B8"], .stApp [style*="color:#94A3B8"] {{ color: {T['text_muted']} !important; }}
        .stApp [style*="background: rgba(30, 41, 59"], .stApp [style*="background:rgba(30, 41, 59"],
        .stApp [style*="background: rgba(15, 23, 42"], .stApp [style*="background:rgba(15, 23, 42"] {{
            background: {T['bg_card_inner']} !important;
            border-color: {T['border_subtle']} !important;
            color: {T['text_primary']} !important;
        }}
        .stApp [style*="background: #111827"], .stApp [style*="background-color: #111827"] {{
            background: {T['bg_card']} !important;
            background-color: {T['bg_card']} !important;
            border-color: {T['border_color']} !important;
        }}
        .stApp [style*="border-color: #1F2937"], .stApp [style*="border: 1px solid #1F2937"],
        .stApp [style*="border: 1px solid #374151"] {{
            border-color: {T['border_color']} !important;
        }}
        .stApp [style*="border-bottom: 1px solid #1F2937"] {{
            border-bottom: 1px solid {T['border_divider']} !important;
        }}
        .stApp [style*="border-top: 1px solid #1F2937"] {{
            border-top: 1px solid {T['border_divider']} !important;
        }}
        .stApp [style*="background: rgba(139, 92, 246, 0.12)"],
        .stApp [style*="background: rgba(139, 92, 246, 0.2)"] {{
            background: {T['accent_purple_subtle']} !important;
            border-color: rgba(124, 58, 237, 0.25) !important;
            color: {T['accent_purple_text']} !important;
        }}
        .stApp [style*="color: #A78BFA"] {{ color: {T['accent_purple_text']} !important; }}
        .stApp [style*="color: #C4B5FD"] {{ color: {T['accent_purple_text']} !important; }}
        """
    else:
        # Dark Theme (Signature Deep Navy Analytical Workspace)
        theme_css = f"""
        /* Main Dark Navy Background & Typography */
        .stApp {{
            background-color: {T['bg_page']} !important;
            font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, "Helvetica Neue", Arial, sans-serif !important;
            color: {T['text_secondary']} !important;
        }}
        h1, h2, h3, h4, h5, h6 {{
            color: {T['text_primary']} !important;
            font-weight: 700 !important;
            letter-spacing: -0.02em !important;
        }}
        .stApp p, .stApp span, .stApp label, .stApp li {{
            color: {T['text_secondary']} !important;
        }}
        .stApp b, .stApp strong {{
            color: {T['text_primary']} !important;
            font-weight: 700 !important;
        }}

        /* Sidebar: Dark Navy */
        section[data-testid="stSidebar"] {{
            background-color: {T['bg_sidebar']} !important;
            border-right: 1px solid {T['border_color']} !important;
            padding-top: 1.2rem !important;
        }}
        section[data-testid="stSidebar"] h1,
        section[data-testid="stSidebar"] h2,
        section[data-testid="stSidebar"] h3 {{
            color: {T['text_primary']} !important;
        }}
        section[data-testid="stSidebar"] p,
        section[data-testid="stSidebar"] span,
        section[data-testid="stSidebar"] label {{
            color: {T['text_secondary']} !important;
        }}
        .sidebar-logo-card {{
            padding: 10px 4px 16px 4px;
            border-bottom: 1px solid {T['border_divider']};
            margin-bottom: 16px;
        }}
        .sidebar-logo-title {{
            font-size: 20px;
            font-weight: 800;
            background: linear-gradient(135deg, #A78BFA, #60A5FA);
            -webkit-background-clip: text;
            -webkit-text-fill-color: transparent;
            letter-spacing: 0.05em;
            margin: 0;
        }}
        .sidebar-logo-sub {{
            font-size: 11px;
            color: {T['text_muted']} !important;
            text-transform: uppercase;
            letter-spacing: 0.08em;
            margin-top: 2px;
        }}
        .sidebar-user-card {{
            background: {T['bg_user_card']};
            border: 1px solid {T['border_color']};
            border-radius: 10px;
            padding: 12px;
            margin-bottom: 16px;
        }}
        .sidebar-user-role {{
            display: inline-block;
            background: {T['badge_role_bg']};
            color: {T['badge_role_text']} !important;
            border: 1px solid {T['badge_role_border']};
            font-size: 11px;
            font-weight: 600;
            padding: 2px 8px;
            border-radius: 12px;
            margin-top: 4px;
        }}

        /* Sidebar Navigation Buttons */
        section[data-testid="stSidebar"] div.stButton > button {{
            width: 100% !important;
            text-align: left !important;
            justify-content: flex-start !important;
            background-color: transparent !important;
            border: 1px solid transparent !important;
            color: {T['text_muted']} !important;
            border-radius: 8px !important;
            padding: 0.55rem 0.85rem !important;
            margin-bottom: 3px !important;
            font-size: 14px !important;
            font-weight: 500 !important;
            box-shadow: none !important;
            display: flex !important;
            align-items: center !important;
            transition: all 0.15s ease-in-out !important;
        }}
        section[data-testid="stSidebar"] div.stButton > button p,
        section[data-testid="stSidebar"] div.stButton > button span,
        section[data-testid="stSidebar"] div.stButton > button div {{
            color: inherit !important;
            font-weight: inherit !important;
            text-align: left !important;
            margin: 0 !important;
        }}
        section[data-testid="stSidebar"] div.stButton > button:hover {{
            background-color: {T['border_color']} !important;
            color: {T['text_primary']} !important;
            border-color: {T['border_subtle']} !important;
        }}
        section[data-testid="stSidebar"] div.stButton > button:hover p,
        section[data-testid="stSidebar"] div.stButton > button:hover span {{
            color: {T['text_primary']} !important;
        }}
        section[data-testid="stSidebar"] div.stButton > button[kind="primary"] {{
            background: linear-gradient(90deg, rgba(124, 58, 237, 0.25) 0%, rgba(79, 70, 229, 0.15) 100%) !important;
            color: #FFFFFF !important;
            border: 1px solid #7C3AED !important;
            border-left: 4px solid #8B5CF6 !important;
            font-weight: 600 !important;
            box-shadow: 0 2px 8px rgba(124, 58, 237, 0.2) !important;
        }}
        section[data-testid="stSidebar"] div.stButton > button[kind="primary"] p,
        section[data-testid="stSidebar"] div.stButton > button[kind="primary"] span {{
            color: #FFFFFF !important;
            font-weight: 600 !important;
        }}

        /* Cards & Containers */
        .iq-card {{
            background-color: {T['bg_card']} !important;
            border: 1px solid {T['border_color']} !important;
            border-radius: 12px !important;
            padding: 20px !important;
            margin-bottom: 16px !important;
            box-shadow: 0 4px 12px rgba(0, 0, 0, 0.3) !important;
        }}
        .iq-card h3, .iq-card h4 {{
            margin-top: 0 !important;
            color: {T['text_primary']} !important;
        }}
        div[data-testid="stVerticalBlockBorderWrapper"] {{
            background-color: {T['bg_card']} !important;
            border: 1px solid {T['border_color']} !important;
            border-radius: 12px !important;
            padding: 14px 18px !important;
            margin-bottom: 12px !important;
        }}
        div[data-testid="stVerticalBlockBorderWrapper"] > div {{
            background-color: transparent !important;
        }}
        .journal-hero-card {{
            background: {T['bg_card']} !important;
            border: 1px solid {T['border_color']} !important;
            border-radius: 12px !important;
            padding: 18px 22px !important;
            margin-bottom: 18px !important;
        }}
        .journal-hero-card h4 {{
            color: {T['text_primary']} !important;
            margin: 0 !important;
        }}

        /* Metrics */
        [data-testid="stMetric"] {{
            background-color: {T['bg_card']} !important;
            border: 1px solid {T['border_color']} !important;
            border-radius: 12px !important;
            padding: 16px 18px !important;
            box-shadow: 0 2px 8px rgba(0, 0, 0, 0.25) !important;
        }}
        [data-testid="stMetricLabel"] p {{
            color: {T['text_muted']} !important;
            font-size: 12px !important;
            font-weight: 600 !important;
            text-transform: uppercase !important;
            letter-spacing: 0.06em !important;
        }}
        [data-testid="stMetricValue"] div {{
            color: {T['text_primary']} !important;
            font-size: 28px !important;
            font-weight: 800 !important;
        }}

        /* Score Cards */
        .score-card {{
            background-color: {T['bg_card']};
            border: 1px solid {T['border_color']};
            border-radius: 12px;
            padding: 18px;
            height: 100%;
        }}
        .score-card-title {{
            font-size: 12px;
            font-weight: 700;
            color: {T['text_muted']} !important;
            text-transform: uppercase;
            letter-spacing: 0.06em;
        }}
        .score-card-val {{
            font-size: 28px;
            font-weight: 800;
            color: {T['text_primary']} !important;
            margin: 4px 0;
        }}
        .score-card-interp {{
            font-size: 12px;
            color: {T['text_muted']} !important;
            line-height: 1.4;
            margin-top: 6px;
        }}

        /* Buttons */
        .stButton > button,
        .stButton > button:not([kind="primary"]),
        [data-testid="stBaseButton-secondary"] {{
            background-color: {T['btn_bg']} !important;
            color: {T['btn_text']} !important;
            border: 1px solid {T['btn_border']} !important;
            border-radius: 8px !important;
            font-weight: 600 !important;
            padding: 0.5rem 1rem !important;
            box-shadow: 0 1px 3px rgba(0, 0, 0, 0.2) !important;
            transition: all 0.2s ease !important;
        }}
        .stButton > button *,
        .stButton > button:not([kind="primary"]) *,
        [data-testid="stBaseButton-secondary"] * {{
            color: {T['btn_text']} !important;
        }}
        .stButton > button:hover,
        .stButton > button:not([kind="primary"]):hover,
        [data-testid="stBaseButton-secondary"]:hover {{
            background-color: {T['btn_hover_bg']} !important;
            border-color: {T['accent_purple']} !important;
            color: #FFFFFF !important;
        }}
        .stButton > button:hover *,
        .stButton > button:not([kind="primary"]):hover *,
        [data-testid="stBaseButton-secondary"]:hover * {{
            color: #FFFFFF !important;
        }}
        .stButton > button[kind="primary"],
        [data-testid="stBaseButton-primary"] {{
            background: linear-gradient(135deg, #7C3AED, #4F46E5) !important;
            color: #FFFFFF !important;
            border: none !important;
            box-shadow: 0 4px 14px rgba(124, 58, 237, 0.35) !important;
        }}
        .stButton > button[kind="primary"] *,
        [data-testid="stBaseButton-primary"] * {{
            color: #FFFFFF !important;
        }}
        .stButton > button[kind="primary"]:hover,
        [data-testid="stBaseButton-primary"]:hover {{
            background: linear-gradient(135deg, #6D28D9, #4338CA) !important;
            box-shadow: 0 6px 18px rgba(124, 58, 237, 0.45) !important;
        }}
        .stDownloadButton > button {{
            background: linear-gradient(135deg, #4F46E5, #2563EB) !important;
            color: #FFFFFF !important;
            border: none !important;
            border-radius: 8px !important;
            font-weight: 600 !important;
            padding: 0.5rem 1rem !important;
            box-shadow: 0 4px 14px rgba(79, 70, 229, 0.35) !important;
        }}
        .stDownloadButton > button * {{
            color: #FFFFFF !important;
        }}

        /* Popover Trigger Buttons (Explicit High-Specificity Contrast Repair) */
        .stApp [data-testid="stPopoverButton"],
        .stApp [data-testid="stPopover"] button,
        .stApp .stPopover button,
        .stApp div[data-testid="stPopover"] button,
        .stApp [data-testid="stPopover"] [data-testid="stBaseButton-secondary"],
        .stApp .stPopover [data-testid="stBaseButton-secondary"],
        [data-testid="stPopoverButton"],
        [data-testid="stPopover"] button,
        .stPopover button {{
            background-color: {T['btn_bg']} !important;
            color: {T['btn_text']} !important;
            border: 1.5px solid {T['btn_border']} !important;
            border-radius: 8px !important;
            font-weight: 600 !important;
            font-size: 14px !important;
            min-height: 42px !important;
            padding: 0.5rem 1rem !important;
            box-shadow: 0 1px 3px rgba(0, 0, 0, 0.25) !important;
            transition: all 0.2s ease !important;
            display: inline-flex !important;
            align-items: center !important;
            justify-content: center !important;
            cursor: pointer !important;
            width: 100% !important;
        }}
        .stApp [data-testid="stPopoverButton"] *,
        .stApp [data-testid="stPopover"] button *,
        .stApp .stPopover button *,
        .stApp div[data-testid="stPopover"] button *,
        [data-testid="stPopoverButton"] *,
        [data-testid="stPopover"] button *,
        .stPopover button * {{
            color: {T['btn_text']} !important;
            font-weight: 600 !important;
            fill: currentColor !important;
            stroke: currentColor !important;
        }}
        .stApp [data-testid="stPopoverButton"] svg,
        .stApp [data-testid="stPopover"] button svg,
        .stApp .stPopover button svg,
        .stApp [data-testid="stPopoverButton"] [data-testid="stIconMaterial"],
        .stApp [data-testid="stPopoverButton"] span[class*="DynamicIcon"],
        .stApp [data-testid="stPopoverButton"] div[class*="StyledPopoverExpansionIcon"] {{
            color: {T['btn_text']} !important;
            fill: currentColor !important;
            stroke: currentColor !important;
            display: inline-flex !important;
            align-items: center !important;
        }}
        .stApp [data-testid="stPopoverButton"]:hover,
        .stApp [data-testid="stPopover"] button:hover,
        .stApp .stPopover button:hover,
        .stApp div[data-testid="stPopover"] button:hover,
        [data-testid="stPopoverButton"]:hover,
        [data-testid="stPopover"] button:hover,
        .stPopover button:hover {{
            background-color: {T['btn_hover_bg']} !important;
            border-color: {T['accent_purple']} !important;
            color: {T['accent_purple']} !important;
            box-shadow: 0 2px 8px rgba(167, 139, 250, 0.2) !important;
        }}
        .stApp [data-testid="stPopoverButton"]:hover *,
        .stApp [data-testid="stPopover"] button:hover *,
        .stApp .stPopover button:hover *,
        .stApp div[data-testid="stPopover"] button:hover *,
        [data-testid="stPopoverButton"]:hover *,
        [data-testid="stPopover"] button:hover *,
        .stPopover button:hover * {{
            color: {T['accent_purple']} !important;
            fill: currentColor !important;
            stroke: currentColor !important;
        }}
        .stApp [data-testid="stPopoverButton"]:hover svg,
        .stApp [data-testid="stPopover"] button:hover svg,
        .stApp .stPopover button:hover svg,
        .stApp [data-testid="stPopoverButton"]:hover [data-testid="stIconMaterial"],
        .stApp [data-testid="stPopoverButton"]:hover span[class*="DynamicIcon"] {{
            color: {T['accent_purple']} !important;
            fill: currentColor !important;
            stroke: currentColor !important;
        }}

        /* Focus rings and disabled buttons */
        .stButton > button:focus-visible,
        .stDownloadButton > button:focus-visible,
        [data-testid="stFileUploaderDropzone"] button:focus-visible,
        .stApp [data-testid="stPopoverButton"]:focus-visible,
        .stApp [data-testid="stPopover"] button:focus-visible,
        [data-testid="stPopoverButton"]:focus-visible,
        [data-testid="stPopover"] button:focus-visible {{
            outline: 2px solid {T['focus_outline']} !important;
            outline-offset: 2px !important;
            box-shadow: 0 0 0 3px rgba(139, 92, 246, 0.25) !important;
        }}
        .stButton > button:disabled,
        .stButton > button[disabled],
        .stDownloadButton > button:disabled,
        .stDownloadButton > button[disabled],
        [data-testid="stFileUploaderDropzone"] button:disabled,
        [data-testid="stBaseButton-secondary"]:disabled,
        [data-testid="stBaseButton-secondary"][disabled],
        .stApp [data-testid="stPopoverButton"]:disabled,
        .stApp [data-testid="stPopoverButton"][disabled],
        .stApp [data-testid="stPopover"] button:disabled,
        .stApp [data-testid="stPopover"] button[disabled],
        [data-testid="stPopoverButton"]:disabled,
        [data-testid="stPopover"] button:disabled {{
            background-color: {T['disabled_bg']} !important;
            color: {T['disabled_text']} !important;
            border: 1px solid {T['disabled_border']} !important;
            cursor: not-allowed !important;
            opacity: 0.65 !important;
            box-shadow: none !important;
        }}
        .stButton > button:disabled *,
        .stButton > button[disabled] *,
        .stDownloadButton > button:disabled *,
        [data-testid="stFileUploaderDropzone"] button:disabled *,
        [data-testid="stBaseButton-secondary"]:disabled *,
        [data-testid="stBaseButton-secondary"][disabled] *,
        .stApp [data-testid="stPopoverButton"]:disabled *,
        .stApp [data-testid="stPopoverButton"][disabled] *,
        .stApp [data-testid="stPopover"] button:disabled *,
        [data-testid="stPopoverButton"]:disabled *,
        [data-testid="stPopover"] button:disabled * {{
            color: {T['disabled_text']} !important;
            fill: currentColor !important;
        }}

        /* File Uploader Container & Dropzone */
        [data-testid="stFileUploader"] {{
            background-color: {T['bg_card']} !important;
            border: 1px solid {T['border_color']} !important;
            border-radius: 12px !important;
            padding: 18px !important;
        }}
        [data-testid="stFileUploaderDropzone"] {{
            border: 2px dashed {T['border_subtle']} !important;
            border-radius: 10px !important;
            background-color: {T['bg_page']} !important;
        }}
        [data-testid="stFileUploaderDropzoneInstructions"],
        [data-testid="stFileUploaderDropzoneInstructions"] span,
        [data-testid="stFileUploaderDropzoneInstructions"] small,
        [data-testid="stFileUploaderDropzone"] > div > span {{
            color: {T['text_secondary']} !important;
        }}

        /* File Uploader "Browse files" Button - Explicit dark styles */
        [data-testid="stFileUploaderDropzone"] button,
        [data-testid="stFileUploaderDropzone"] [data-testid="stBaseButton-secondary"],
        section[data-testid="stFileUploadDropzone"] button,
        section[data-testid="stFileUploadDropzone"] [data-testid="stBaseButton-secondary"],
        [data-testid="stFileUploader"] section button {{
            background-color: {T['btn_bg']} !important;
            color: {T['btn_text']} !important;
            border: 1.5px solid {T['btn_border']} !important;
            border-radius: 8px !important;
            font-weight: 600 !important;
            padding: 0.45rem 1.1rem !important;
            box-shadow: 0 1px 3px rgba(0, 0, 0, 0.25) !important;
            transition: all 0.2s ease !important;
            cursor: pointer !important;
        }}
        [data-testid="stFileUploaderDropzone"] button *,
        [data-testid="stFileUploaderDropzone"] [data-testid="stBaseButton-secondary"] *,
        section[data-testid="stFileUploadDropzone"] button *,
        section[data-testid="stFileUploadDropzone"] [data-testid="stBaseButton-secondary"] * {{
            color: {T['btn_text']} !important;
            font-weight: 600 !important;
        }}
        [data-testid="stFileUploaderDropzone"] button:hover,
        [data-testid="stFileUploaderDropzone"] [data-testid="stBaseButton-secondary"]:hover,
        section[data-testid="stFileUploadDropzone"] button:hover,
        section[data-testid="stFileUploadDropzone"] [data-testid="stBaseButton-secondary"]:hover {{
            background-color: {T['btn_hover_bg']} !important;
            border-color: {T['accent_purple']} !important;
            color: #FFFFFF !important;
            box-shadow: 0 2px 8px rgba(124, 58, 237, 0.25) !important;
        }}
        [data-testid="stFileUploaderDropzone"] button:hover *,
        [data-testid="stFileUploaderDropzone"] [data-testid="stBaseButton-secondary"]:hover * {{
            color: #FFFFFF !important;
        }}

        /* Uploaded File Chip */
        [data-testid="stFileUploaderFile"] {{
            background-color: {T['bg_card_elevated']} !important;
            border: 1px solid {T['border_subtle']} !important;
            border-radius: 8px !important;
            color: {T['text_primary']} !important;
        }}
        [data-testid="stFileUploaderFile"] * {{
            color: {T['text_primary']} !important;
        }}
        [data-testid="stFileUploaderDeleteBtn"] button {{
            color: {T['text_muted']} !important;
            border: none !important;
            background: transparent !important;
        }}
        [data-testid="stFileUploaderDeleteBtn"] button:hover {{
            color: #EF4444 !important;
        }}

        /* Form Inputs */
        div[data-baseweb="select"] > div {{
            background-color: {T['input_bg']} !important;
            border-color: {T['input_border']} !important;
            color: {T['input_text']} !important;
        }}
        div[data-baseweb="select"] * {{
            color: {T['input_text']} !important;
        }}
        div[data-baseweb="popover"],
        div[data-baseweb="menu"],
        ul[role="listbox"] {{
            background-color: {T['bg_card']} !important;
            border: 1px solid {T['border_subtle']} !important;
            box-shadow: 0 4px 16px rgba(0, 0, 0, 0.4) !important;
        }}
        li[role="option"] {{
            background-color: {T['bg_card']} !important;
            color: {T['text_primary']} !important;
        }}
        li[role="option"]:hover,
        li[role="option"][aria-selected="true"] {{
            background-color: {T['bg_card_elevated']} !important;
            color: {T['accent_purple']} !important;
        }}
        div[data-testid="stTextInput"] input {{
            background-color: {T['input_bg']} !important;
            border: 1px solid {T['input_border']} !important;
            border-radius: 8px !important;
            color: {T['input_text']} !important;
            padding: 8px 12px !important;
        }}
        div[data-testid="stTextInput"] input:focus {{
            border-color: {T['focus_outline']} !important;
            box-shadow: 0 0 0 2px rgba(139, 92, 246, 0.3) !important;
        }}
        div[data-testid="stCheckbox"] label[data-baseweb="checkbox"] {{
            display: inline-flex !important;
            align-items: center !important;
            gap: 8px !important;
            margin: 0 !important;
            cursor: pointer !important;
        }}
        div[data-testid="stCheckbox"] label[data-baseweb="checkbox"] > div,
        div[data-testid="stCheckbox"] [data-testid="stWidgetLabel"] {{
            margin: 0 !important;
            padding: 0 !important;
            min-height: 0 !important;
            display: inline-flex !important;
            align-items: center !important;
        }}
        div[data-testid="stCheckbox"] label, div[data-testid="stCheckbox"] p, div[data-testid="stCheckbox"] span {{
            color: {T['text_primary']} !important;
            margin: 0 !important;
            line-height: 1.35 !important;
        }}
        div[data-testid="stCheckbox"] label[data-baseweb="checkbox"] > span:first-child {{
            width: 18px !important;
            height: 18px !important;
            min-width: 18px !important;
            min-height: 18px !important;
            background-color: #0F172A !important;
            background-image: none !important;
            border: 1.5px solid #94A3B8 !important;
            border-radius: 5px !important;
            margin: 0 !important;
            box-sizing: border-box !important;
            transition: all 0.18s ease !important;
            box-shadow: inset 0 1px 2px rgba(0, 0, 0, 0.35) !important;
        }}
        div[data-testid="stCheckbox"] label[data-baseweb="checkbox"]:hover > span:first-child {{
            border-color: #A78BFA !important;
            background-color: #1E293B !important;
            box-shadow: 0 0 0 2px rgba(139, 92, 246, 0.2) !important;
        }}
        div[data-testid="stCheckbox"] label[data-baseweb="checkbox"]:has(input:checked) > span:first-child,
        div[data-testid="stCheckbox"] label[data-baseweb="checkbox"]:has(input[aria-checked="true"]) > span:first-child {{
            background-color: #7C3AED !important;
            border: 1.5px solid #A78BFA !important;
            background-image: url("data:image/svg+xml,%3Csvg xmlns='http://www.w3.org/2000/svg' viewBox='0 0 16 16' fill='none'%3E%3Cpath d='M3.5 8.2L6.6 11.3L12.8 4.8' stroke='%23FFFFFF' stroke-width='2.3' stroke-linecap='round' stroke-linejoin='round'/%3E%3C/svg%3E") !important;
            background-repeat: no-repeat !important;
            background-position: center !important;
            background-size: 13px 13px !important;
            box-shadow: 0 2px 6px rgba(124, 58, 237, 0.35) !important;
        }}
        div[data-testid="stCheckbox"] label[data-baseweb="checkbox"]:has(input:checked):hover > span:first-child,
        div[data-testid="stCheckbox"] label[data-baseweb="checkbox"]:has(input[aria-checked="true"]):hover > span:first-child {{
            background-color: #6D28D9 !important;
            border-color: #C4B5FD !important;
            box-shadow: 0 0 0 3px rgba(139, 92, 246, 0.28) !important;
        }}

        /* Expanders */
        [data-testid="stExpander"] {{
            background-color: {T['bg_card']} !important;
            border: 1px solid {T['border_color']} !important;
            border-radius: 10px !important;
            margin-bottom: 12px !important;
        }}
        [data-testid="stExpander"] summary {{
            background-color: {T['bg_card']} !important;
            color: {T['text_primary']} !important;
            font-weight: 600 !important;
            border-radius: 10px 10px 0 0 !important;
            padding: 10px 14px !important;
        }}
        [data-testid="stExpander"] summary * {{
            color: {T['text_primary']} !important;
        }}
        [data-testid="stExpander"] summary:hover {{
            background-color: {T['bg_card_elevated']} !important;
            color: {T['accent_purple']} !important;
        }}
        [data-testid="stExpander"] summary:hover * {{
            color: {T['accent_purple']} !important;
        }}
        [data-testid="stExpander"] [data-testid="stExpanderDetails"] {{
            background-color: {T['bg_card']} !important;
            color: {T['text_secondary']} !important;
            padding: 14px 18px !important;
        }}

        /* Tabs */
        .stTabs [data-baseweb="tab-list"] {{
            background-color: transparent !important;
            border-bottom: 2px solid {T['border_divider']} !important;
            gap: 8px !important;
        }}
        .stTabs [data-baseweb="tab"] {{
            background-color: transparent !important;
            color: {T['text_muted']} !important;
            font-weight: 600 !important;
            font-size: 14px !important;
            border: none !important;
            padding: 8px 16px !important;
        }}
        .stTabs [data-baseweb="tab"]:hover {{
            color: {T['text_primary']} !important;
        }}
        .stTabs [data-baseweb="tab"][aria-selected="true"] {{
            color: {T['accent_purple']} !important;
            border-bottom: 3px solid {T['accent_purple']} !important;
            font-weight: 700 !important;
        }}
        .stTabs [data-baseweb="tab-highlight"] {{
            background-color: {T['accent_purple']} !important;
        }}

        /* Chat */
        .chat-msg-user {{
            background-color: {T['chat_user_bg']} !important;
            border: 1px solid {T['chat_user_border']} !important;
            border-radius: 10px !important;
            padding: 12px 16px !important;
            margin-bottom: 10px !important;
            color: {T['chat_user_text']} !important;
        }}
        .chat-msg-bot {{
            background-color: {T['chat_bot_bg']} !important;
            border-left: 4px solid {T['accent_purple']} !important;
            border-top: 1px solid {T['chat_bot_border']} !important;
            border-right: 1px solid {T['chat_bot_border']} !important;
            border-bottom: 1px solid {T['chat_bot_border']} !important;
            border-radius: 10px !important;
            padding: 14px 18px !important;
            margin-bottom: 14px !important;
            color: {T['chat_bot_text']} !important;
        }}

        /* Tables & DataFrames */
        [data-testid="stDataFrame"] {{
            border: 1px solid {T['border_color']} !important;
            border-radius: 10px !important;
            background-color: {T['bg_card']} !important;
        }}
        [data-testid="stTable"] {{
            background-color: {T['bg_card']} !important;
            color: {T['text_primary']} !important;
        }}
        [data-testid="stTable"] th {{
            background-color: {T['table_header_bg']} !important;
            color: {T['table_header_text']} !important;
            border-bottom: 2px solid {T['border_subtle']} !important;
            font-weight: 700 !important;
        }}
        [data-testid="stTable"] td {{
            background-color: {T['bg_card']} !important;
            color: {T['text_secondary']} !important;
            border-bottom: 1px solid {T['border_color']} !important;
        }}
        [data-testid="stTable"] tr:nth-child(even) td {{
            background-color: {T['table_row_alt']} !important;
        }}

        /* Badges */
        .domain-badge {{
            display: inline-block !important;
            background: {T['badge_domain_bg']} !important;
            color: {T['badge_domain_text']} !important;
            border: 1px solid {T['badge_domain_border']} !important;
            font-size: 11.5px !important;
            font-weight: 600 !important;
            padding: 2px 9px !important;
            border-radius: 6px !important;
            margin-right: 6px !important;
            margin-bottom: 4px !important;
        }}

        /* Popovers Body */
        [data-testid="stPopoverBody"],
        div[data-baseweb="popover"] {{
            background-color: {T['bg_card']} !important;
            border: 1px solid {T['border_color']} !important;
            border-radius: 12px !important;
            box-shadow: 0 8px 24px rgba(0, 0, 0, 0.4) !important;
            color: {T['text_primary']} !important;
            padding: 16px !important;
            max-width: 95vw !important;
        }}
        [data-testid="stPopoverBody"] *,
        div[data-baseweb="popover"] * {{
            color: {T['text_primary']} !important;
        }}
        [data-testid="stPopoverBody"] p,
        [data-testid="stPopoverBody"] span,
        [data-testid="stPopoverBody"] [data-testid="stCaptionContainer"] p {{
            color: {T['text_secondary']} !important;
        }}
        [data-testid="stPopoverBody"] pre,
        [data-testid="stPopoverBody"] code {{
            background-color: {T['bg_card_elevated']} !important;
            color: {T['text_primary']} !important;
            border: 1px solid {T['border_subtle']} !important;
        }}
        [data-testid="stPopoverBody"] button,
        [data-testid="stPopoverBody"] [data-testid="stCodeCopyButton"] {{
            background-color: {T['bg_card_elevated']} !important;
            color: {T['text_primary']} !important;
            border: 1px solid {T['border_subtle']} !important;
            border-radius: 6px !important;
        }}
        [data-testid="stPopoverBody"] button:hover,
        [data-testid="stPopoverBody"] [data-testid="stCodeCopyButton"]:hover {{
            background-color: {T['btn_hover_bg']} !important;
            color: {T['accent_purple']} !important;
            border-color: {T['accent_purple']} !important;
        }}

        /* Custom Responsive Comparison Table */
        .iq-comparison-table-wrapper {{
            overflow-x: auto;
            border: 1px solid {T['border_color']};
            border-radius: 12px;
            background-color: {T['bg_card']};
            margin: 16px 0 24px 0;
            box-shadow: 0 4px 12px rgba(0, 0, 0, 0.25);
        }}
        .iq-comp-table {{
            width: 100%;
            border-collapse: collapse;
            font-size: 13.5px;
            color: {T['text_primary']};
        }}
        .iq-comp-table th {{
            background-color: {T['table_header_bg']};
            color: {T['table_header_text']};
            padding: 14px 18px;
            font-weight: 700;
            border-bottom: 2px solid {T['border_subtle']};
            font-size: 12px;
            letter-spacing: 0.04em;
        }}
        .iq-comp-table td {{
            padding: 12px 18px;
            border-bottom: 1px solid {T['border_color']};
            color: {T['text_secondary']};
        }}
        .iq-comp-table tr:last-child td {{
            border-bottom: none;
            font-weight: 700;
            background-color: {T['bg_card_inner']};
        }}
        .iq-comp-table tr:nth-child(even) {{
            background-color: {T['table_row_alt']};
        }}
        .iq-comp-table tr:hover {{
            background-color: {T['bg_card_elevated']};
        }}
        .iq-comp-dim-name {{
            color: {T['text_primary']} !important;
            font-weight: 600;
        }}
        .iq-comp-val-cell {{
            text-align: center;
            font-weight: 700;
            font-size: 14px;
        }}
        .iq-comp-val-best {{
            color: {T['success_text']} !important;
            font-weight: 800;
        }}
        .iq-comp-val-pill {{
            display: inline-block;
            padding: 3px 10px;
            border-radius: 6px;
            font-weight: 700;
        }}
        """

    full_css = f"<style>{css_vars}\n{base_css}\n{theme_css}</style>"
    render_html(full_css)


# ----------------------------------------------------
# Session State Initialization
# ----------------------------------------------------
auth.create_users_table()

if "authenticated" not in st.session_state:
    st.session_state.authenticated = False
if "user_data" not in st.session_state:
    st.session_state.user_data = {}
if "analysis_history" not in st.session_state:
    st.session_state.analysis_history = {}
if "chat_history" not in st.session_state:
    st.session_state.chat_history = []
if "current_page" not in st.session_state:
    st.session_state.current_page = "🏠 Dashboard"
if "theme_preference" not in st.session_state:
    qp = st.query_params.get("theme")
    if qp and qp.lower() in ["light", "dark", "system"]:
        st.session_state.theme_preference = qp.capitalize()
    else:
        st.session_state.theme_preference = "System"
else:
    qp = st.query_params.get("theme")
    if qp and qp.lower() in ["light", "dark", "system"] and st.session_state.theme_preference != qp.capitalize():
        st.session_state.theme_preference = qp.capitalize()
if "summary_result" not in st.session_state:
    st.session_state.summary_result = None
if "summary_length" not in st.session_state:
    st.session_state.summary_length = "Medium"
if "summary_doc_text" not in st.session_state:
    st.session_state.summary_doc_text = ""
if "summary_doc_name" not in st.session_state:
    st.session_state.summary_doc_name = ""

# Inject active theme stylesheet
inject_theme_css(st.session_state.theme_preference)

def navigate_to(page_name, *extra_args):
    """Sets current page in session state with defensive argument handling."""
    if isinstance(page_name, (list, tuple)) and page_name:
        page_name = page_name[0]
    st.session_state["current_page"] = str(page_name)

# Reusable Plotly Layout Preset (Theme-Aware)
DARK_LAYOUT = get_chart_theme(st.session_state.theme_preference)

# ----------------------------------------------------
# Authentication View (Sign In & Create Account Redesign)
# ----------------------------------------------------
if not st.session_state.authenticated:
    eff_auth_theme = "Dark"
    if "auth_mode" not in st.session_state:
        st.session_state.auth_mode = "login"
    if "auth_role_choice" not in st.session_state:
        st.session_state.auth_role_choice = "🎓 STUDENT LOGIN"
        
    # Read vector logo SVG and encode as Base64 to prevent markdown indentation code leaks
    logo_img_html = ""
    try:
        import base64
        logo_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets", "paperiq_logo.svg")
        if os.path.exists(logo_path):
            with open(logo_path, "rb") as f:
                b64_logo = base64.b64encode(f.read()).decode("utf-8")
            logo_img_html = f'<img src="data:image/svg+xml;base64,{b64_logo}" width="44" height="36" alt="PaperIQ Logo" style="display: block; margin: 0 auto; filter: drop-shadow(0 4px 10px rgba(2, 132, 199, 0.22));" />'
    except Exception:
        logo_img_html = ""
    if not logo_img_html:
        logo_img_html = '<div style="font-size: 32px; text-align: center; margin-bottom: 2px;">📄</div>'

    # Background SVG (Base64 data URL)
    bg_svg_css = ""
    try:
        bg_fname = "login_bg.svg" if eff_auth_theme == "Light" else "login_bg_dark.svg"
        bg_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets", bg_fname)
        if os.path.exists(bg_path):
            with open(bg_path, "rb") as f:
                b64_bg = base64.b64encode(f.read()).decode("utf-8")
            bg_svg_css = f"url('data:image/svg+xml;base64,{b64_bg}')"
    except Exception:
        bg_svg_css = ""

    # Scoped Auth CSS with strict overflow control, centered layout, and pixel alignment
    if eff_auth_theme == "Light":
        auth_theme_css = f"""
        <style>
        /* Viewport styling for login view */
        html, body {{
            overflow-x: hidden !important;
            max-width: 100vw !important;
            margin: 0 !important;
            padding: 0 !important;
        }}
        [data-testid="stAppViewContainer"] {{
            height: 100vh !important;
            height: 100dvh !important;
            overflow-y: auto !important;
            overflow-x: hidden !important;
        }}
        [data-testid="stAppViewContainer"] > .main {{
            min-height: 100vh !important;
            min-height: 100dvh !important;
            display: flex !important;
            flex-direction: column !important;
            justify-content: center !important;
            align-items: center !important;
            padding: 0 !important;
            overflow-x: hidden !important;
        }}
        section[data-testid="stSidebar"] {{
            display: none !important;
        }}
        div[data-testid="collapsedControl"] {{
            display: none !important;
        }}
        header[data-testid="stHeader"] {{
            display: none !important;
            height: 0 !important;
            min-height: 0 !important;
            padding: 0 !important;
        }}
        /* Soft Aqua & Research Illustration Background */
        .stApp {{
            background: radial-gradient(circle at 50% 28%, #FFFFFF 0%, #EBF6F8 42%, #DBF1F5 75%, #CEEBF2 100%) {f", {bg_svg_css} center center / cover no-repeat fixed" if bg_svg_css else ""} !important;
            font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, "Helvetica Neue", Arial, sans-serif !important;
        }}
        /* Centered Responsive Container */
        .block-container {{
            max-width: 420px !important;
            width: 100% !important;
            padding: 0.5rem 1rem !important;
            margin: auto !important;
            box-sizing: border-box !important;
            overflow-x: hidden !important;
        }}
        /* White Authentication Card */
        div[data-testid="stVerticalBlockBorderWrapper"] {{
            background-color: #FFFFFF !important;
            border: 1px solid #D8E7EE !important;
            border-radius: 14px !important;
            padding: 16px 20px !important;
            box-shadow: 0 16px 40px rgba(14, 42, 71, 0.08), 0 2px 8px rgba(0, 0, 0, 0.03) !important;
            box-sizing: border-box !important;
            width: 100% !important;
        }}
        div[data-testid="stVerticalBlockBorderWrapper"] > div[data-testid="stVerticalBlock"] {{
            gap: 0.45rem !important;
        }}
        /* Segmented Student / Researcher Selector Track */
        .st-key-auth_role_radio,
        div[data-testid="stElementContainer"]:has(div[data-testid="stRadio"]) {{
            width: 100% !important;
            max-width: 100% !important;
            display: flex !important;
            justify-content: center !important;
            align-items: center !important;
            margin: 0 auto 12px auto !important;
            box-sizing: border-box !important;
        }}
        div[data-testid="stRadio"] {{
            width: 100% !important;
            max-width: 100% !important;
            margin: 0 !important;
            display: flex !important;
            justify-content: center !important;
            align-items: center !important;
            box-sizing: border-box !important;
        }}
        div[data-testid="stRadio"] div[role="radiogroup"] {{
            display: flex !important;
            flex-direction: row !important;
            flex-wrap: nowrap !important;
            width: 100% !important;
            max-width: 100% !important;
            background: #EEF5F8 !important;
            border: 1px solid #DCE7EE !important;
            border-radius: 28px !important;
            padding: 3px !important;
            gap: 4px !important;
            box-sizing: border-box !important;
            align-items: stretch !important;
            justify-content: center !important;
        }}
        div[data-testid="stRadio"] label[data-baseweb="radio"] {{
            flex: 1 1 50% !important;
            width: 50% !important;
            max-width: 50% !important;
            min-width: 0 !important;
            min-height: 38px !important;
            height: 38px !important;
            display: flex !important;
            align-items: center !important;
            justify-content: center !important;
            text-align: center !important;
            border-radius: 24px !important;
            padding: 0 8px !important;
            margin: 0 !important;
            cursor: pointer !important;
            border: 1.5px solid transparent !important;
            box-sizing: border-box !important;
            transition: all 0.2s ease !important;
        }}
        div[data-testid="stRadio"] label[data-baseweb="radio"] > div:first-child {{
            display: none !important;
        }}
        div[data-testid="stRadio"] label[data-baseweb="radio"] > div:last-child {{
            display: flex !important;
            align-items: center !important;
            justify-content: center !important;
            text-align: center !important;
            width: 100% !important;
            height: 100% !important;
            margin: 0 !important;
            padding: 0 !important;
        }}
        div[data-testid="stRadio"] label[data-baseweb="radio"] p,
        div[data-testid="stRadio"] label[data-baseweb="radio"] span {{
            display: flex !important;
            align-items: center !important;
            justify-content: center !important;
            text-align: center !important;
            font-size: 11.5px !important;
            font-weight: 700 !important;
            color: #62758A !important;
            letter-spacing: 0.02em !important;
            margin: 0 !important;
            line-height: 1 !important;
            white-space: nowrap !important;
            overflow: hidden !important;
            text-overflow: ellipsis !important;
        }}
        div[data-testid="stRadio"] label[data-baseweb="radio"]:has(input:checked) {{
            background: #FFFFFF !important;
            border: 1.5px solid #1A6ED8 !important;
            box-shadow: 0 2px 8px rgba(26, 110, 216, 0.2) !important;
        }}
        div[data-testid="stRadio"] label[data-baseweb="radio"]:has(input:checked) p,
        div[data-testid="stRadio"] label[data-baseweb="radio"]:has(input:checked) span {{
            color: #0E2A47 !important;
        }}
        /* Form Inputs (Full Container & Inner Elements) */
        div[data-testid="stTextInput"] {{
            margin-bottom: 6px !important;
            width: 100% !important;
        }}
        div[data-testid="stTextInput"] label p {{
            color: #172B45 !important;
            font-weight: 600 !important;
            font-size: 12.5px !important;
            margin-bottom: 2px !important;
        }}
        div[data-testid="stTextInput"] div[data-baseweb="input"],
        div[data-testid="stTextInput"] div[data-baseweb="base-input"] {{
            background-color: #FFFFFF !important;
            border: 1.5px solid #D6E1EA !important;
            border-radius: 9px !important;
            height: 42px !important;
            box-sizing: border-box !important;
            transition: all 0.2s ease !important;
        }}
        div[data-testid="stTextInput"] div[data-baseweb="input"]:focus-within {{
            border-color: #1A6ED8 !important;
            box-shadow: 0 0 0 3px rgba(26, 110, 216, 0.15) !important;
        }}
        div[data-testid="stTextInput"] input {{
            background-color: transparent !important;
            border: none !important;
            height: 40px !important;
            padding: 8px 12px !important;
            color: #172B45 !important;
            font-size: 13.5px !important;
            width: 100% !important;
            box-sizing: border-box !important;
        }}
        div[data-testid="stTextInput"] button[aria-label*="password"],
        div[data-testid="stTextInput"] div[data-baseweb="input"] button {{
            background-color: transparent !important;
            border: none !important;
            box-shadow: none !important;
            color: #62758A !important;
        }}
        div[data-testid="stTextInput"] div[data-baseweb="input"] button svg {{
            fill: #62758A !important;
        }}
        /* Suppress browser native password reveal duplicate icons */
        input[type="password"]::-ms-reveal,
        input[type="password"]::-ms-clear {{
            display: none !important;
            width: 0 !important;
            height: 0 !important;
        }}
        /* Selectbox for Academic Role */
        div[data-testid="stSelectbox"] {{
            margin-bottom: 6px !important;
            width: 100% !important;
        }}
        div[data-testid="stSelectbox"] label p {{
            color: #172B45 !important;
            font-weight: 600 !important;
            font-size: 12.5px !important;
            margin-bottom: 2px !important;
        }}
        div[data-baseweb="select"] > div {{
            background-color: #FFFFFF !important;
            border: 1.5px solid #D6E1EA !important;
            border-radius: 9px !important;
            min-height: 42px !important;
            height: 42px !important;
            color: #172B45 !important;
        }}
        div[data-baseweb="select"] span {{
            color: #172B45 !important;
            font-size: 13.5px !important;
        }}
        div[data-baseweb="select"] svg {{
            fill: #62758A !important;
        }}
        /* Dropdown Popover Listbox */
        div[data-baseweb="popover"] ul,
        div[data-baseweb="menu"] {{
            background-color: #FFFFFF !important;
            border: 1px solid #D6E1EA !important;
            border-radius: 9px !important;
            box-shadow: 0 10px 25px rgba(0, 0, 0, 0.1) !important;
        }}
        li[role="option"] {{
            background-color: #FFFFFF !important;
            color: #172B45 !important;
            font-size: 13.5px !important;
        }}
        li[role="option"]:hover,
        li[role="option"][aria-selected="true"] {{
            background-color: #F1F5F9 !important;
            color: #1A6ED8 !important;
        }}
        /* Remember Me & Checkbox */
        div[data-testid="stCheckbox"] {{
            margin: 0 !important;
            padding-top: 2px !important;
        }}
        div[data-testid="stCheckbox"] label,
        div[data-testid="stCheckbox"] label[data-baseweb="checkbox"] {{
            display: inline-flex !important;
            align-items: center !important;
            gap: 8px !important;
            margin: 0 !important;
            cursor: pointer !important;
        }}
        div[data-testid="stCheckbox"] label[data-baseweb="checkbox"] > div,
        div[data-testid="stCheckbox"] [data-testid="stWidgetLabel"] {{
            margin: 0 !important;
            padding: 0 !important;
            min-height: 0 !important;
            display: inline-flex !important;
            align-items: center !important;
        }}
        div[data-testid="stCheckbox"] label[data-baseweb="checkbox"] > span:first-child {{
            width: 18px !important;
            height: 18px !important;
            min-width: 18px !important;
            min-height: 18px !important;
            background-color: #FFFFFF !important;
            background-image: none !important;
            border: 1.5px solid #64748B !important;
            border-radius: 5px !important;
            margin: 0 !important;
            box-sizing: border-box !important;
            transition: all 0.18s ease !important;
        }}
        div[data-testid="stCheckbox"] label[data-baseweb="checkbox"]:hover > span:first-child {{
            border-color: #7C3AED !important;
            box-shadow: 0 0 0 2px rgba(124, 58, 237, 0.15) !important;
        }}
        div[data-testid="stCheckbox"] label[data-baseweb="checkbox"]:has(input:checked) > span:first-child,
        div[data-testid="stCheckbox"] label[data-baseweb="checkbox"]:has(input[aria-checked="true"]) > span:first-child {{
            background-color: #7C3AED !important;
            border: 1.5px solid #6D28D9 !important;
            background-image: url("data:image/svg+xml,%3Csvg xmlns='http://www.w3.org/2000/svg' viewBox='0 0 16 16' fill='none'%3E%3Cpath d='M3.5 8.2L6.6 11.3L12.8 4.8' stroke='%23FFFFFF' stroke-width='2.3' stroke-linecap='round' stroke-linejoin='round'/%3E%3C/svg%3E") !important;
            background-repeat: no-repeat !important;
            background-position: center !important;
            background-size: 13px 13px !important;
            box-shadow: 0 2px 6px rgba(124, 58, 237, 0.28) !important;
        }}
        div[data-testid="stCheckbox"] label p {{
            font-size: 12.5px !important;
            font-weight: 500 !important;
            color: #475569 !important;
            margin: 0 !important;
            line-height: 1.3 !important;
            white-space: nowrap !important;
        }}
        /* High specificity overrides for Auth Action & Switch Buttons */
        div.stElementContainer.st-key-btn_signin div.stButton > button,
        div.stElementContainer.st-key-btn_register div.stButton > button {{
            background: linear-gradient(135deg, #6D28D9 0%, #2563EB 100%) !important;
            background-color: #6D28D9 !important;
            color: #FFFFFF !important;
            border: none !important;
            border-radius: 9px !important;
            padding: 10px 16px !important;
            height: 42px !important;
            font-weight: 700 !important;
            font-size: 13.5px !important;
            letter-spacing: 0.04em !important;
            box-shadow: 0 4px 14px rgba(109, 40, 217, 0.35) !important;
            width: 100% !important;
            margin-top: 6px !important;
            margin-bottom: 2px !important;
            transition: all 0.2s ease !important;
        }}
        div.stElementContainer.st-key-btn_signin div.stButton > button:hover,
        div.stElementContainer.st-key-btn_register div.stButton > button:hover {{
            background: linear-gradient(135deg, #5B21B6 0%, #1D4ED8 100%) !important;
            box-shadow: 0 6px 18px rgba(109, 40, 217, 0.45) !important;
            transform: translateY(-1px) !important;
        }}
        div.stElementContainer.st-key-btn_signin div.stButton > button *,
        div.stElementContainer.st-key-btn_register div.stButton > button * {{
            color: #FFFFFF !important;
            font-weight: 700 !important;
        }}
        /* Secondary Demo Button */
        div.stElementContainer.st-key-btn_demo div.stButton > button {{
            background: #FFFFFF !important;
            background-color: #FFFFFF !important;
            color: #172B45 !important;
            border: 1.5px solid #D6E1EA !important;
            border-radius: 9px !important;
            padding: 9px 16px !important;
            height: 42px !important;
            font-weight: 600 !important;
            font-size: 13px !important;
            box-shadow: 0 2px 5px rgba(0, 0, 0, 0.03) !important;
            width: 100% !important;
            margin-top: 4px !important;
            transition: all 0.2s ease !important;
        }}
        div.stElementContainer.st-key-btn_demo div.stButton > button:hover {{
            background: #F4F8FA !important;
            background-color: #F4F8FA !important;
            border-color: #B8CEDC !important;
            color: #0E2A47 !important;
        }}
        div.stElementContainer.st-key-btn_demo div.stButton > button * {{
            color: #172B45 !important;
            font-weight: 600 !important;
        }}
        /* Create Account / Back to Sign In Switch Link */
        div.stElementContainer.st-key-btn_switch_signup div.stButton > button,
        div.stElementContainer.st-key-btn_switch_login div.stButton > button,
        div.stElementContainer.st-key-btn_back_from_admin div.stButton > button {{
            background: transparent !important;
            background-color: transparent !important;
            border: none !important;
            color: #1A6ED8 !important;
            font-size: 13px !important;
            font-weight: 700 !important;
            white-space: nowrap !important;
            box-shadow: none !important;
            padding: 0 !important;
            height: 32px !important;
            min-height: 32px !important;
            line-height: 32px !important;
            text-align: left !important;
            justify-content: flex-start !important;
            margin: 0 !important;
        }}
        div.stElementContainer.st-key-btn_switch_signup div.stButton > button:hover,
        div.stElementContainer.st-key-btn_switch_login div.stButton > button:hover,
        div.stElementContainer.st-key-btn_back_from_admin div.stButton > button:hover {{
            color: #0E4B9F !important;
            text-decoration: underline !important;
            background: transparent !important;
            background-color: transparent !important;
            box-shadow: none !important;
        }}
        div.stElementContainer.st-key-btn_switch_signup div.stButton > button *,
        div.stElementContainer.st-key-btn_switch_login div.stButton > button *,
        div.stElementContainer.st-key-btn_back_from_admin div.stButton > button * {{
            white-space: nowrap !important;
            font-size: 13px !important;
            font-weight: 700 !important;
            color: #1A6ED8 !important;
        }}
        div.stElementContainer.st-key-btn_switch_admin_portal div.stButton > button {{
            background: transparent !important;
            background-color: transparent !important;
            border: none !important;
            color: #62758A !important;
            font-size: 12.5px !important;
            font-weight: 600 !important;
            white-space: nowrap !important;
            box-shadow: none !important;
            padding: 0 !important;
            height: 32px !important;
            min-height: 32px !important;
            line-height: 32px !important;
            text-align: right !important;
            justify-content: flex-end !important;
            margin: 0 !important;
            transition: color 0.15s ease !important;
        }}
        div.stElementContainer.st-key-btn_switch_admin_portal div.stButton > button:hover {{
            color: #172B45 !important;
            text-decoration: underline !important;
            background: transparent !important;
            background-color: transparent !important;
            box-shadow: none !important;
        }}
        div.stElementContainer.st-key-btn_switch_admin_portal div.stButton > button * {{
            white-space: nowrap !important;
            font-size: 12.5px !important;
            font-weight: 600 !important;
            color: #62758A !important;
        }}
        @media (max-width: 480px) {{
            .block-container {{
                padding-left: 0.75rem !important;
                padding-right: 0.75rem !important;
                padding-top: 0.5rem !important;
            }}
            div[data-testid="stVerticalBlockBorderWrapper"] {{
                padding: 16px 14px !important;
                border-radius: 14px !important;
            }}
        }}
        </style>
        """
    else:
        auth_theme_css = f"""
        <style>
        /* Viewport styling for login view */
        html, body {{
            overflow-x: hidden !important;
            max-width: 100vw !important;
            margin: 0 !important;
            padding: 0 !important;
        }}
        [data-testid="stAppViewContainer"] {{
            height: 100vh !important;
            height: 100dvh !important;
            overflow-y: auto !important;
            overflow-x: hidden !important;
        }}
        [data-testid="stAppViewContainer"] > .main {{
            min-height: 100vh !important;
            min-height: 100dvh !important;
            display: flex !important;
            flex-direction: column !important;
            justify-content: center !important;
            align-items: center !important;
            padding: 0 !important;
            overflow-x: hidden !important;
        }}
        section[data-testid="stSidebar"] {{
            display: none !important;
        }}
        div[data-testid="collapsedControl"] {{
            display: none !important;
        }}
        header[data-testid="stHeader"] {{
            display: none !important;
            height: 0 !important;
            min-height: 0 !important;
            padding: 0 !important;
        }}
        .stApp {{
            background: radial-gradient(circle at 50% 28%, #111D30 0%, #0B1424 50%, #060B14 100%) {f", {bg_svg_css} center center / cover no-repeat fixed" if bg_svg_css else ""} !important;
            font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, "Helvetica Neue", Arial, sans-serif !important;
        }}
        .block-container {{
            max-width: 420px !important;
            width: 100% !important;
            padding: 0.5rem 1rem !important;
            margin: auto !important;
            box-sizing: border-box !important;
            overflow-x: hidden !important;
        }}
        div[data-testid="stVerticalBlockBorderWrapper"] {{
            background-color: #111E2E !important;
            border: 1px solid rgba(56, 189, 248, 0.25) !important;
            border-radius: 14px !important;
            padding: 16px 20px !important;
            box-shadow: 0 16px 40px rgba(0, 0, 0, 0.45) !important;
            box-sizing: border-box !important;
            width: 100% !important;
        }}
        div[data-testid="stVerticalBlockBorderWrapper"] > div[data-testid="stVerticalBlock"] {{
            gap: 0.45rem !important;
        }}
        /* Segmented Student / Researcher Selector Track */
        .st-key-auth_role_radio,
        div[data-testid="stElementContainer"]:has(div[data-testid="stRadio"]) {{
            width: 100% !important;
            max-width: 100% !important;
            display: flex !important;
            justify-content: center !important;
            align-items: center !important;
            margin: 0 auto 12px auto !important;
            box-sizing: border-box !important;
        }}
        div[data-testid="stRadio"] {{
            width: 100% !important;
            max-width: 100% !important;
            margin: 0 !important;
            display: flex !important;
            justify-content: center !important;
            align-items: center !important;
            box-sizing: border-box !important;
        }}
        div[data-testid="stRadio"] div[role="radiogroup"] {{
            display: flex !important;
            flex-direction: row !important;
            flex-wrap: nowrap !important;
            width: 100% !important;
            max-width: 100% !important;
            background: #0B1424 !important;
            border: 1px solid rgba(56, 189, 248, 0.2) !important;
            border-radius: 28px !important;
            padding: 3px !important;
            gap: 4px !important;
            box-sizing: border-box !important;
            align-items: stretch !important;
            justify-content: center !important;
        }}
        div[data-testid="stRadio"] label[data-baseweb="radio"] {{
            flex: 1 1 50% !important;
            width: 50% !important;
            max-width: 50% !important;
            min-width: 0 !important;
            min-height: 36px !important;
            height: 36px !important;
            display: flex !important;
            align-items: center !important;
            justify-content: center !important;
            text-align: center !important;
            border-radius: 24px !important;
            padding: 0 8px !important;
            margin: 0 !important;
            cursor: pointer !important;
            border: 1.5px solid transparent !important;
            box-sizing: border-box !important;
            transition: all 0.2s ease !important;
        }}
        div[data-testid="stRadio"] label[data-baseweb="radio"] > div:first-child {{
            display: none !important;
        }}
        div[data-testid="stRadio"] label[data-baseweb="radio"] > div:last-child {{
            display: flex !important;
            align-items: center !important;
            justify-content: center !important;
            text-align: center !important;
            width: 100% !important;
            height: 100% !important;
            margin: 0 !important;
            padding: 0 !important;
        }}
        div[data-testid="stRadio"] label[data-baseweb="radio"] p,
        div[data-testid="stRadio"] label[data-baseweb="radio"] span {{
            display: flex !important;
            align-items: center !important;
            justify-content: center !important;
            text-align: center !important;
            font-size: 11.5px !important;
            font-weight: 700 !important;
            color: #94A3B8 !important;
            letter-spacing: 0.02em !important;
            margin: 0 !important;
            line-height: 1 !important;
            white-space: nowrap !important;
            overflow: hidden !important;
            text-overflow: ellipsis !important;
        }}
        div[data-testid="stRadio"] label[data-baseweb="radio"]:has(input:checked) {{
            background: #111E2E !important;
            border: 1.5px solid #38BDF8 !important;
            box-shadow: 0 2px 8px rgba(56, 189, 248, 0.25) !important;
        }}
        div[data-testid="stRadio"] label[data-baseweb="radio"]:has(input:checked) p,
        div[data-testid="stRadio"] label[data-baseweb="radio"]:has(input:checked) span {{
            color: #FFFFFF !important;
        }}
        /* Form Inputs */
        div[data-testid="stTextInput"] {{
            margin-bottom: 4px !important;
            width: 100% !important;
        }}
        div[data-testid="stTextInput"] label p {{
            color: #E2E8F0 !important;
            font-weight: 600 !important;
            font-size: 12.5px !important;
            margin-bottom: 2px !important;
        }}
        div[data-testid="stTextInput"] div[data-baseweb="input"],
        div[data-testid="stTextInput"] div[data-baseweb="base-input"] {{
            background-color: #0B1424 !important;
            border: 1.5px solid #1E2E42 !important;
            border-radius: 9px !important;
            height: 38px !important;
            box-sizing: border-box !important;
            transition: all 0.2s ease !important;
        }}
        div[data-testid="stTextInput"] div[data-baseweb="input"]:focus-within {{
            border-color: #38BDF8 !important;
            box-shadow: 0 0 0 3px rgba(56, 189, 248, 0.2) !important;
        }}
        div[data-testid="stTextInput"] input {{
            background-color: transparent !important;
            border: none !important;
            height: 36px !important;
            padding: 6px 12px !important;
            color: #F8FAFC !important;
            font-size: 13.5px !important;
            width: 100% !important;
            box-sizing: border-box !important;
        }}
        div[data-testid="stTextInput"] button[aria-label*="password"],
        div[data-testid="stTextInput"] div[data-baseweb="input"] button {{
            background-color: transparent !important;
            border: none !important;
            box-shadow: none !important;
            color: #94A3B8 !important;
        }}
        div[data-testid="stTextInput"] div[data-baseweb="input"] button svg {{
            fill: #94A3B8 !important;
        }}
        input[type="password"]::-ms-reveal,
        input[type="password"]::-ms-clear {{
            display: none !important;
            width: 0 !important;
            height: 0 !important;
        }}
        /* Selectbox */
        div[data-testid="stSelectbox"] {{
            margin-bottom: 4px !important;
            width: 100% !important;
        }}
        div[data-testid="stSelectbox"] label p {{
            color: #E2E8F0 !important;
            font-weight: 600 !important;
            font-size: 12.5px !important;
            margin-bottom: 2px !important;
        }}
        div[data-baseweb="select"] > div {{
            background-color: #0B1424 !important;
            border: 1.5px solid #1E2E42 !important;
            border-radius: 9px !important;
            min-height: 38px !important;
            height: 38px !important;
            color: #F8FAFC !important;
        }}
        div[data-baseweb="select"] span {{
            color: #F8FAFC !important;
            font-size: 13.5px !important;
        }}
        div[data-baseweb="select"] svg {{
            fill: #94A3B8 !important;
        }}
        /* Dropdown Popover Listbox */
        div[data-baseweb="popover"] ul,
        div[data-baseweb="menu"] {{
            background-color: #0B1424 !important;
            border: 1px solid #1E2E42 !important;
            border-radius: 9px !important;
            box-shadow: 0 10px 25px rgba(0, 0, 0, 0.4) !important;
        }}
        li[role="option"] {{
            background-color: #0B1424 !important;
            color: #F8FAFC !important;
            font-size: 13.5px !important;
        }}
        li[role="option"]:hover,
        li[role="option"][aria-selected="true"] {{
            background-color: #1E2E42 !important;
            color: #38BDF8 !important;
        }}
        /* Remember Me & Checkbox */
        div[data-testid="stCheckbox"] {{
            margin: 0 !important;
            padding-top: 2px !important;
        }}
        div[data-testid="stCheckbox"] label,
        div[data-testid="stCheckbox"] label[data-baseweb="checkbox"] {{
            display: inline-flex !important;
            align-items: center !important;
            gap: 8px !important;
            margin: 0 !important;
            cursor: pointer !important;
        }}
        div[data-testid="stCheckbox"] label[data-baseweb="checkbox"] > div,
        div[data-testid="stCheckbox"] [data-testid="stWidgetLabel"] {{
            margin: 0 !important;
            padding: 0 !important;
            min-height: 0 !important;
            display: inline-flex !important;
            align-items: center !important;
        }}
        div[data-testid="stCheckbox"] label[data-baseweb="checkbox"] > span:first-child {{
            width: 18px !important;
            height: 18px !important;
            min-width: 18px !important;
            min-height: 18px !important;
            background-color: #0B1424 !important;
            background-image: none !important;
            border: 1.5px solid #94A3B8 !important;
            border-radius: 5px !important;
            margin: 0 !important;
            box-sizing: border-box !important;
            transition: all 0.18s ease !important;
            box-shadow: inset 0 1px 2px rgba(0, 0, 0, 0.35) !important;
        }}
        div[data-testid="stCheckbox"] label[data-baseweb="checkbox"]:hover > span:first-child {{
            border-color: #A78BFA !important;
            background-color: #162338 !important;
            box-shadow: 0 0 0 2px rgba(139, 92, 246, 0.22) !important;
        }}
        div[data-testid="stCheckbox"] label[data-baseweb="checkbox"]:has(input:checked) > span:first-child,
        div[data-testid="stCheckbox"] label[data-baseweb="checkbox"]:has(input[aria-checked="true"]) > span:first-child {{
            background-color: #7C3AED !important;
            border: 1.5px solid #A78BFA !important;
            background-image: url("data:image/svg+xml,%3Csvg xmlns='http://www.w3.org/2000/svg' viewBox='0 0 16 16' fill='none'%3E%3Cpath d='M3.5 8.2L6.6 11.3L12.8 4.8' stroke='%23FFFFFF' stroke-width='2.3' stroke-linecap='round' stroke-linejoin='round'/%3E%3C/svg%3E") !important;
            background-repeat: no-repeat !important;
            background-position: center !important;
            background-size: 13px 13px !important;
            box-shadow: 0 2px 6px rgba(124, 58, 237, 0.35) !important;
        }}
        div[data-testid="stCheckbox"] label[data-baseweb="checkbox"]:has(input:checked):hover > span:first-child,
        div[data-testid="stCheckbox"] label[data-baseweb="checkbox"]:has(input[aria-checked="true"]):hover > span:first-child {{
            background-color: #6D28D9 !important;
            border-color: #C4B5FD !important;
            box-shadow: 0 0 0 3px rgba(139, 92, 246, 0.28) !important;
        }}
        div[data-testid="stCheckbox"] label p {{
            font-size: 12.5px !important;
            font-weight: 500 !important;
            color: #CBD5E1 !important;
            margin: 0 !important;
            line-height: 1.3 !important;
            white-space: nowrap !important;
        }}
        /* Primary Action Buttons */
        div.stElementContainer.st-key-btn_signin div.stButton > button,
        div.stElementContainer.st-key-btn_register div.stButton > button,
        div.stElementContainer.st-key-btn_admin_signin div.stButton > button {{
            background: linear-gradient(135deg, #6D28D9 0%, #2563EB 100%) !important;
            background-color: #6D28D9 !important;
            color: #FFFFFF !important;
            border: none !important;
            border-radius: 9px !important;
            padding: 8px 16px !important;
            height: 40px !important;
            font-weight: 700 !important;
            font-size: 13.5px !important;
            letter-spacing: 0.04em !important;
            box-shadow: 0 4px 14px rgba(109, 40, 217, 0.35) !important;
            width: 100% !important;
            margin-top: 4px !important;
            margin-bottom: 2px !important;
            transition: all 0.2s ease !important;
        }}
        div.stElementContainer.st-key-btn_signin div.stButton > button:hover,
        div.stElementContainer.st-key-btn_register div.stButton > button:hover,
        div.stElementContainer.st-key-btn_admin_signin div.stButton > button:hover {{
            background: linear-gradient(135deg, #5B21B6 0%, #1D4ED8 100%) !important;
            box-shadow: 0 6px 18px rgba(109, 40, 217, 0.45) !important;
            transform: translateY(-1px) !important;
        }}
        div.stElementContainer.st-key-btn_signin div.stButton > button *,
        div.stElementContainer.st-key-btn_register div.stButton > button *,
        div.stElementContainer.st-key-btn_admin_signin div.stButton > button * {{
            color: #FFFFFF !important;
            font-weight: 700 !important;
        }}
        /* Secondary Demo Button */
        div.stElementContainer.st-key-btn_demo div.stButton > button {{
            background: #0B1424 !important;
            background-color: #0B1424 !important;
            color: #E2E8F0 !important;
            border: 1.5px solid #1E2E42 !important;
            border-radius: 9px !important;
            padding: 7px 16px !important;
            height: 38px !important;
            font-weight: 600 !important;
            font-size: 13px !important;
            width: 100% !important;
            margin-top: 2px !important;
            transition: all 0.2s ease !important;
        }}
        div.stElementContainer.st-key-btn_demo div.stButton > button:hover {{
            background: #111E2E !important;
            background-color: #111E2E !important;
            border-color: #38BDF8 !important;
            color: #FFFFFF !important;
        }}
        div.stElementContainer.st-key-btn_demo div.stButton > button * {{
            color: #E2E8F0 !important;
            font-weight: 600 !important;
        }}
        /* Switch Links */
        div.stElementContainer.st-key-btn_switch_signup div.stButton > button,
        div.stElementContainer.st-key-btn_switch_login div.stButton > button,
        div.stElementContainer.st-key-btn_back_from_admin div.stButton > button {{
            background: transparent !important;
            background-color: transparent !important;
            border: none !important;
            color: #38BDF8 !important;
            font-size: 13px !important;
            font-weight: 700 !important;
            white-space: nowrap !important;
            box-shadow: none !important;
            padding: 0 !important;
            height: 32px !important;
            min-height: 32px !important;
            line-height: 32px !important;
            text-align: left !important;
            justify-content: flex-start !important;
            margin: 0 !important;
        }}
        div.stElementContainer.st-key-btn_switch_signup div.stButton > button:hover,
        div.stElementContainer.st-key-btn_switch_login div.stButton > button:hover,
        div.stElementContainer.st-key-btn_back_from_admin div.stButton > button:hover {{
            color: #7DD3FC !important;
            text-decoration: underline !important;
            background: transparent !important;
            background-color: transparent !important;
            box-shadow: none !important;
        }}
        div.stElementContainer.st-key-btn_switch_signup div.stButton > button *,
        div.stElementContainer.st-key-btn_switch_login div.stButton > button *,
        div.stElementContainer.st-key-btn_back_from_admin div.stButton > button * {{
            white-space: nowrap !important;
            font-size: 13px !important;
            font-weight: 700 !important;
            color: #38BDF8 !important;
        }}
        div.stElementContainer.st-key-btn_switch_admin_portal div.stButton > button {{
            background: transparent !important;
            background-color: transparent !important;
            border: none !important;
            color: #94A3B8 !important;
            font-size: 12.5px !important;
            font-weight: 600 !important;
            white-space: nowrap !important;
            box-shadow: none !important;
            padding: 0 !important;
            height: 32px !important;
            min-height: 32px !important;
            line-height: 32px !important;
            text-align: right !important;
            justify-content: flex-end !important;
            margin: 0 !important;
            transition: color 0.15s ease !important;
        }}
        div.stElementContainer.st-key-btn_switch_admin_portal div.stButton > button:hover {{
            color: #E2E8F0 !important;
            text-decoration: underline !important;
            background: transparent !important;
            background-color: transparent !important;
            box-shadow: none !important;
        }}
        div.stElementContainer.st-key-btn_switch_admin_portal div.stButton > button * {{
            white-space: nowrap !important;
            font-size: 12.5px !important;
            font-weight: 600 !important;
            color: #94A3B8 !important;
        }}
        @media (max-width: 480px) {{
            .block-container {{
                padding-left: 0.75rem !important;
                padding-right: 0.75rem !important;
                padding-top: 0.5rem !important;
            }}
            div[data-testid="stVerticalBlockBorderWrapper"] {{
                padding: 16px 14px !important;
                border-radius: 14px !important;
            }}
        }}
        </style>
        """
    render_html(auth_theme_css)

    # Centered PaperIQ Branding Header
    render_html([
        '<div style="text-align: center; margin-bottom: 10px;">',
        f'<div style="display: flex; justify-content: center; align-items: center; margin-bottom: 4px;">{logo_img_html}</div>',
        f'<div style="font-size: 22px; font-weight: 800; letter-spacing: 0.16em; color: {"#0E2A47" if eff_auth_theme == "Light" else "#FFFFFF"}; margin: 0; line-height: 1.1;">PAPERIQ</div>',
        f'<div style="font-size: 12px; font-weight: 600; color: {"#2C4E72" if eff_auth_theme == "Light" else "#94A3B8"}; letter-spacing: 0.03em; margin-top: 2px;">AI-Powered Research Insight Analyzer</div>',
        f'<div style="font-size: 11px; font-style: italic; color: {"#5A7B9D" if eff_auth_theme == "Light" else "#64748B"}; margin-top: 2px;">Analyze. Understand. Improve.</div>',
        '</div>',
    ])

    if st.session_state.get("session_expired_notice", False):
        st.info("🕒 Your session expired due to inactivity. Please sign in again.")
        st.session_state["session_expired_notice"] = False

    try:
        login_maint_banner = str(auth.get_platform_setting("maintenance_banner", "") or "").strip()
        if login_maint_banner:
            st.info(f"📢 **Platform Notice:** {login_maint_banner}")
    except Exception:
        pass

    # Centered Authentication Card
    with st.container(border=True):
        if st.session_state.auth_mode == "login":
            # 1. Login Role Selector (Side-by-side, 50/50 width, no wrapping)
            role_options = ["🎓 STUDENT LOGIN", "🔬 RESEARCHER LOGIN"]
            sel_role = st.radio(
                "Login Role Selection",
                options=role_options,
                index=0 if "STUDENT" in st.session_state.auth_role_choice else 1,
                horizontal=True,
                label_visibility="collapsed",
                key="auth_role_radio"
            )
            st.session_state.auth_role_choice = sel_role
            is_student = "STUDENT" in sel_role

            # 2. Username Input (Full Width)
            login_user_input = st.text_input(
                "Username" if is_student else "Researcher Username",
                placeholder="Enter your username or Student ID" if is_student else "Enter your researcher username",
                key="auth_login_user"
            )

            # 3. Password Input (Full Width with Native Built-In Visibility Eye Toggle)
            login_pass_input = st.text_input(
                "Password",
                type="password",
                placeholder="Enter your password",
                key="auth_login_pass"
            )

            # 4. Remember Me & Forgot Password (Aligned on One Row)
            col_rem, col_forgot = st.columns([1.1, 1.1])
            with col_rem:
                st.checkbox("Remember Me", value=True, key="auth_remember_me")
            with col_forgot:
                render_html([
                    '<div style="text-align: right; padding-top: 3px; white-space: nowrap;">',
                    f'<a href="mailto:support@paperiq.edu?subject=PaperIQ%20Account%20Assistance" style="color: {"#1A6ED8" if eff_auth_theme == "Light" else "#38BDF8"}; font-size: 12.5px; font-weight: 600; text-decoration: none;">Forgot Password?</a>',
                    '</div>',
                ])

            # 5. Primary Sign In Action
            if st.button("LOG IN", type="primary", use_container_width=True, key="btn_signin"):
                if login_user_input and login_pass_input:
                    user = auth.login_user(login_user_input, login_pass_input)
                    if user:
                        if user.get("disabled") or user.get("status") == "Disabled":
                            st.error("Account Deactivated: This account has been disabled by an administrator. Please contact system support.")
                        else:
                            clear_active_analysis_session()
                            clear_comparison_state()
                            st.session_state.authenticated = True
                            st.session_state.user_data = user
                            st.session_state["username"] = user["username"]
                            st.session_state["last_authenticated_user"] = user["username"]
                            st.session_state["last_activity_ts"] = time.time()
                            st.session_state.theme_preference = user.get("theme_preference", "System")
                            st.session_state.analysis_history = {user["username"]: auth.get_user_history(user["username"])}
                            if user.get("role") == "Administrator":
                                st.session_state.current_page = "🛡️ Admin Dashboard"
                            elif user.get("role") == "Researcher":
                                st.session_state.current_page = "🏠 Researcher Dashboard"
                            else:
                                st.session_state.current_page = "🏠 Student Dashboard"
                            st.success("Signed in successfully!")
                            st.rerun()
                    else:
                        locked, rem_sec, locked_until = auth.check_account_lockout(login_user_input)
                        if locked:
                            mins_rem = max(1, (int(rem_sec) + 59) // 60)
                            st.error(f"🔒 Account Locked: Too many failed sign-in attempts. Please wait {mins_rem} minute(s) (until {locked_until}) before trying again.")
                        else:
                            st.error("Invalid username or password. Please verify your credentials.")
                else:
                    st.warning("Please enter your username and password.")

            # 6. Secondary Demo Action (Explicit Student Demo Access)
            if st.button("🎓 Continue as Demo Student", use_container_width=True, key="btn_demo"):
                demo_username = "Student_Demo"
                demo_role = "Student (Demo Access)"
                auth.signup_user(demo_username, "demo123", demo_role)
                clear_active_analysis_session()
                clear_comparison_state()
                st.session_state.authenticated = True
                st.session_state.user_data = {"username": demo_username, "role": demo_role}
                st.session_state["username"] = demo_username
                st.session_state["last_authenticated_user"] = demo_username
                st.session_state["last_activity_ts"] = time.time()
                st.session_state.theme_preference = auth.get_user_theme(demo_username) or "System"
                st.session_state.analysis_history = {demo_username: auth.get_user_history(demo_username)}
                st.session_state.current_page = "🏠 Student Dashboard"
                st.rerun()

            # 7. Card Footer: Switch to Registration & Admin Portal
            render_html(
                f"<div style='height: 1px; background: {'rgba(214, 225, 234, 0.5)' if eff_auth_theme == 'Light' else 'rgba(56, 189, 248, 0.18)'}; margin: 8px 0 4px 0;'></div>"
            )
            col_sw1, col_sw2 = st.columns([1.1, 1.1])
            with col_sw1:
                if st.button("Create Account →", key="btn_switch_signup", use_container_width=True):
                    st.session_state.auth_mode = "signup"
                    st.rerun()
            with col_sw2:
                if st.button("🛡️ Admin Portal", key="btn_switch_admin_portal", use_container_width=True):
                    st.session_state.auth_mode = "admin_login"
                    st.rerun()

        elif st.session_state.auth_mode == "admin_login":
            render_html([
                '<div style="text-align: center; margin-bottom: 14px;">',
                f'<div style="font-size: 19px; font-weight: 700; color: {"#0E2A47" if eff_auth_theme == "Light" else "#FFFFFF"};">🛡️ Administrator Sign-In</div>',
                f'<div style="font-size: 12px; color: {"#62758A" if eff_auth_theme == "Light" else "#94A3B8"}; margin-top: 3px;">PaperIQ System Management &amp; Auditing Portal</div>',
                '</div>',
            ])
            admin_user_input = st.text_input("Administrator Username", placeholder="Enter administrator username", key="auth_admin_user")
            admin_pass_input = st.text_input("Administrator Password", type="password", placeholder="Enter administrator password", key="auth_admin_pass")

            if st.button("LOG IN AS ADMINISTRATOR", type="primary", use_container_width=True, key="btn_admin_signin"):
                if admin_user_input and admin_pass_input:
                    user = auth.login_user(admin_user_input, admin_pass_input)
                    if user:
                        if user.get("disabled") or user.get("status") == "Disabled":
                            st.error("Account Deactivated: This administrator account has been disabled. Please contact system management.")
                        elif user.get("role") == "Administrator":
                            clear_active_analysis_session()
                            clear_comparison_state()
                            st.session_state.authenticated = True
                            st.session_state.user_data = user
                            st.session_state["username"] = user["username"]
                            st.session_state["last_authenticated_user"] = user["username"]
                            st.session_state["last_activity_ts"] = time.time()
                            st.session_state.theme_preference = user.get("theme_preference", "System")
                            st.session_state.analysis_history = {user["username"]: auth.get_user_history(user["username"])}
                            st.session_state.current_page = "🛡️ Admin Dashboard"
                            st.success("Administrator session verified.")
                            st.rerun()
                        else:
                            st.error("Access Denied: This account does not possess administrator privileges. Please sign in via the Student or Researcher portal.")
                    else:
                        locked, rem_sec, locked_until = auth.check_account_lockout(admin_user_input)
                        if locked:
                            mins_rem = max(1, (int(rem_sec) + 59) // 60)
                            st.error(f"🔒 Account Locked: Too many failed sign-in attempts. Please wait {mins_rem} minute(s) (until {locked_until}) before trying again.")
                        else:
                            st.error("Invalid administrator credentials. Please verify your credentials.")
                else:
                    st.warning("Please enter your administrator username and password.")

            render_html("<div style='height: 1px; background: rgba(214, 225, 234, 0.5); margin: 14px 0 10px 0;'></div>")
            if st.button("← Back to Student / Researcher Sign In", key="btn_back_from_admin", use_container_width=True):
                st.session_state.auth_mode = "login"
                st.rerun()

        else:
            # CREATE ACCOUNT (REGISTRATION) VIEW
            render_html([
                '<div style="text-align: center; margin-bottom: 12px;">',
                f'<div style="font-size: 18px; font-weight: 700; color: {"#0E2A47" if eff_auth_theme == "Light" else "#FFFFFF"};">Create Your Account</div>',
                f'<div style="font-size: 12px; color: {"#62758A" if eff_auth_theme == "Light" else "#94A3B8"}; margin-top: 2px;">Join PaperIQ to analyze and evaluate research papers</div>',
                '</div>',
            ])

            reg_enabled = str(auth.get_platform_setting("registration_enabled", "true") or "true").strip().lower() in ("true", "1", "yes")
            def_role_cfg = str(auth.get_platform_setting("default_role", "Student") or "Student").strip()
            if not reg_enabled:
                st.warning("🚫 New account self-registration is currently disabled by the platform administrator.")

            signup_user_input = st.text_input("Username", placeholder="Choose a username", key="auth_reg_user", disabled=not reg_enabled)

            signup_pass_input = st.text_input(
                "Password",
                type="password",
                placeholder="Create a password (min 4 chars)",
                key="auth_reg_pass",
                disabled=not reg_enabled
            )

            signup_pass_confirm = st.text_input(
                "Confirm Password",
                type="password",
                placeholder="Re-enter password",
                key="auth_reg_confirm",
                disabled=not reg_enabled
            )

            signup_role_input = st.selectbox(
                "Academic Role",
                ["Student", "Researcher"],
                index=1 if def_role_cfg == "Researcher" else 0,
                key="auth_reg_role",
                disabled=not reg_enabled
            )

            if st.button("REGISTER ACCOUNT", type="primary", use_container_width=True, key="btn_register", disabled=not reg_enabled):
                if not signup_user_input or not signup_pass_input:
                    st.warning("Please complete all required fields.")
                elif len(signup_pass_input) < 4:
                    st.warning("Password must be at least 4 characters.")
                elif signup_pass_input != signup_pass_confirm:
                    st.error("Passwords do not match.")
                else:
                    success, msg = auth.signup_user(signup_user_input, signup_pass_input, signup_role_input)
                    if success:
                        st.success("Account created successfully! Please sign in with your credentials.")
                        st.session_state["auth_mode"] = "login"
                        st.rerun()
                    else:
                        st.error(msg)

            render_html("<div style='height: 1px; background: rgba(214, 225, 234, 0.5); margin: 12px 0 8px 0;'></div>")
            col_b1, col_b2 = st.columns([1.1, 1.3])
            with col_b1:
                render_html(
                    f'<div style="font-size: 12.5px; color: {"#62758A" if eff_auth_theme == "Light" else "#94A3B8"}; line-height: 36px; white-space: nowrap;">Already have an account?</div>'
                )
            with col_b2:
                if st.button("← Back to Sign In", key="btn_switch_login", use_container_width=True):
                    st.session_state.auth_mode = "login"
                    st.rerun()

    st.stop()

# ----------------------------------------------------
# Authenticated State & Role-Based Sidebar Navigation
# ----------------------------------------------------
username = st.session_state.user_data.get("username", "Student")
if st.session_state.get("last_authenticated_user") not in [None, username]:
    clear_active_analysis_session()
    clear_comparison_state()
    st.session_state.analysis_history = {}
st.session_state["last_authenticated_user"] = username
st.session_state["username"] = username
raw_role = st.session_state.user_data.get("role", "Student")
role = "Researcher" if str(raw_role).lower() == "teacher" else raw_role
is_admin = (role == "Administrator")
is_researcher = (role == "Researcher")
is_student = (role == "Student" or not (is_admin or is_researcher))
theme = get_theme_tokens(st.session_state.get("theme_preference", "System"))

# Enforce platform session_timeout_minutes
try:
    timeout_mins = int(auth.get_platform_setting("session_timeout_minutes", 60) or 60)
except Exception:
    timeout_mins = 60
now_ts = time.time()
last_act_ts = st.session_state.get("last_activity_ts")
if last_act_ts and timeout_mins > 0 and (now_ts - float(last_act_ts)) > (timeout_mins * 60):
    try:
        auth.record_audit_event(
            username=username,
            action_type="SESSION_TIMEOUT",
            target="session",
            details=f"Session timed out after {timeout_mins} minutes of inactivity",
            status="Info",
        )
    except Exception:
        pass
    saved_theme = st.session_state.get("theme_preference", "System")
    for k in list(st.session_state.keys()):
        del st.session_state[k]
    st.session_state["theme_preference"] = saved_theme
    st.session_state["authenticated"] = False
    st.session_state["user_data"] = {}
    st.session_state["auth_mode"] = "login"
    st.session_state["session_expired_notice"] = True
    st.rerun()
st.session_state["last_activity_ts"] = now_ts

# Enforce mandatory temporary password change (Part 2.2)
if st.session_state.user_data.get("must_change_password"):
    with st.container(border=True):
        render_html([
            f'<div style="margin-bottom: 12px;">',
            f'<h2 style="margin: 0; color: {theme["text_primary"]};">🔑 Password Change Required</h2>',
            f'<p style="color: {theme["text_muted"]}; font-size: 14px; margin-top: 4px;">Your account is using a temporary administrator-reset password. Please set a new password to continue to your workspace.</p>',
            f'</div>',
        ])
        new_pw_in = st.text_input("New Password (minimum 6 characters)", type="password", key="force_pw_new")
        conf_pw_in = st.text_input("Confirm New Password", type="password", key="force_pw_conf")
        c_fp1, c_fp2 = st.columns([2, 1])
        with c_fp1:
            if st.button("Update Password & Continue", type="primary", use_container_width=True, key="btn_force_pw_save"):
                if not new_pw_in or len(new_pw_in) < 6:
                    st.warning("New password must be at least 6 characters long.")
                elif new_pw_in != conf_pw_in:
                    st.error("New password and confirmation do not match.")
                else:
                    ok_pw, msg_pw = auth.change_user_password(username, new_pw_in)
                    if ok_pw:
                        st.session_state.user_data["must_change_password"] = False
                        st.success("✓ Password updated successfully!")
                        st.rerun()
                    else:
                        st.error(msg_pw)
        with c_fp2:
            if st.button("🚪 Sign Out", use_container_width=True, key="btn_force_pw_logout"):
                st.session_state["authenticated"] = False
                st.session_state["user_data"] = {}
                st.rerun()
    st.stop()

# Ensure user theme preference is loaded if user was already authenticated
if username and username != "Student" and st.session_state.get("theme_preference") in [None, ""]:
    st.session_state["theme_preference"] = auth.get_user_theme(username) or "System"

# Keep history synchronized (strictly isolated to active user)
if username not in st.session_state.analysis_history:
    st.session_state.analysis_history[username] = auth.get_user_history(username)
user_history = st.session_state.analysis_history.get(username, [])

with st.sidebar:
    # Sidebar Logo Header
    render_html([
        '<div class="sidebar-logo-card">',
        '<div class="sidebar-logo-title">PAPERIQ</div>',
        '<div class="sidebar-logo-sub">AI Research Intelligence</div>',
        '</div>',
    ])
    
    # User Profile Card with Dynamic Role Icon
    role_icon = "🛡️" if is_admin else ("🔬" if is_researcher else "🎓")
    render_html([
        '<div class="sidebar-user-card">',
        f'<div style="font-size: 11px; color: {theme["text_muted"]}; text-transform: uppercase; font-weight: 600;">Active Account</div>',
        f'<div style="font-size: 16px; font-weight: 700; color: {theme["text_primary"]}; margin-top: 2px;">{safe_text(username)}</div>',
        f'<span class="sidebar-user-role">{role_icon} {safe_text(role)}</span>',
        '</div>',
    ])
    
    current_pg = st.session_state.get("current_page", "🏠 Dashboard")

    if is_admin:
        render_html(f"<div style='font-size: 11px; color: {theme['text_muted']}; text-transform: uppercase; font-weight: 700; margin: 12px 0 6px 4px; letter-spacing: 0.08em;'>Administration</div>")
        admin_nav = [
            ("🛡️ Admin Dashboard", "nav_admin_dash"),
            ("👥 All Users / User Management", "nav_admin_users"),
            ("🕒 Login History & Audit Logs", "nav_admin_login_hist"),
        ]
        for label, key in admin_nav:
            is_active = (current_pg == label or (label == "👥 All Users / User Management" and current_pg in ["👥 All Users", "👥 User Management", "👥 All Users / User Management"]) or (label == "🕒 Login History & Audit Logs" and current_pg in ["🕒 Login History", "🕒 Login & Audit Logs", "🕒 Login History & Audit Logs"]))
            st.button(label, key=key, on_click=navigate_to, args=(label,), use_container_width=True, type="primary" if is_active else "secondary")

        render_html(f"<div style='font-size: 11px; color: {theme['text_muted']}; text-transform: uppercase; font-weight: 700; margin: 14px 0 6px 4px; letter-spacing: 0.08em;'>Research Repository</div>")
        repo_nav = [
            ("📁 Saved Papers", "nav_admin_saved_papers"),
            ("⬆️ Upload History", "nav_admin_upload_hist"),
            ("📤 Upload Research Paper", "nav_admin_upload_paper"),
        ]
        for label, key in repo_nav:
            is_active = (current_pg == label)
            st.button(label, key=key, on_click=navigate_to, args=(label,), use_container_width=True, type="primary" if is_active else "secondary")

        render_html(f"<div style='font-size: 11px; color: {theme['text_muted']}; text-transform: uppercase; font-weight: 700; margin: 14px 0 6px 4px; letter-spacing: 0.08em;'>Reporting</div>")
        rep_nav = [
            ("📈 Analytics & Reports", "nav_admin_analytics"),
        ]
        for label, key in rep_nav:
            is_active = (current_pg == label)
            st.button(label, key=key, on_click=navigate_to, args=(label,), use_container_width=True, type="primary" if is_active else "secondary")

        render_html(f"<div style='font-size: 11px; color: {theme['text_muted']}; text-transform: uppercase; font-weight: 700; margin: 14px 0 6px 4px; letter-spacing: 0.08em;'>System &amp; Governance</div>")
        sys_nav = [
            ("🩺 System Health", "nav_admin_health"),
            ("💾 Storage & Data Management", "nav_admin_storage"),
            ("⚙️ Platform Settings", "nav_admin_settings"),
            ("ℹ️ About & System Info", "nav_admin_about"),
        ]
        for label, key in sys_nav:
            is_active = (
                current_pg == label
                or (label == "⚙️ Platform Settings" and current_pg in ["⚙️ Platform Settings", "⚙ Settings", "⚙️ Settings"])
                or (label == "ℹ️ About & System Info" and current_pg in ["ℹ️ About & System Info", "ℹ About", "ℹ️ About"])
            )
            st.button(label, key=key, on_click=navigate_to, args=(label,), use_container_width=True, type="primary" if is_active else "secondary")

    elif is_researcher:
        import core.registry as page_registry
        resolved_res_pg = page_registry.resolve_researcher_page(current_pg)
        render_html(f"<div style='font-size: 11px; color: {theme['text_muted']}; font-weight: 700; margin: 12px 0 6px 4px; letter-spacing: 0.05em;'>Researcher Navigation</div>")
        for label, key in page_registry.RESEARCHER_NAV_ITEMS:
            is_active = (resolved_res_pg == label)
            st.button(label, key=key, on_click=navigate_to, args=(label,), use_container_width=True, type="primary" if is_active else "secondary")

    else:
        # Student Workspace
        render_html(f"<div style='font-size: 11px; color: {theme['text_muted']}; font-weight: 700; margin: 12px 0 6px 4px; letter-spacing: 0.05em;'>Workspace</div>")
        stu_workspace_nav = [
            ("🏠 Student Dashboard", "nav_dashboard"),
            ("📄 Analyze Paper", "nav_analyze"),
        ]
        for label, key in stu_workspace_nav:
            is_active = (current_pg == label or (label == "🏠 Student Dashboard" and current_pg in ["🏠 Dashboard", "🏠 Student Dashboard"]))
            st.button(label, key=key, on_click=navigate_to, args=(label,), use_container_width=True, type="primary" if is_active else "secondary")
            
        render_html(f"<div style='font-size: 11px; color: {theme['text_muted']}; font-weight: 700; margin: 14px 0 6px 4px; letter-spacing: 0.05em;'>Research Tools</div>")
        stu_results_nav = [
            ("📊 Writing Quality & Feedback", "nav_analysis"),
            ("🔄 Compare Papers", "nav_compare"),
            ("💬 Ask PaperIQ", "nav_ask"),
            ("📚 My History", "nav_history"),
        ]
        for label, key in stu_results_nav:
            is_active = (
                current_pg == label
                or (label == "📊 Writing Quality & Feedback" and current_pg in ["📊 Analysis", "📊 Writing Quality & Feedback"])
                or (label == "🔄 Compare Papers" and current_pg in ["🔄 Compare Papers", "🔄 Compare Manuscripts"])
                or (label == "💬 Ask PaperIQ" and current_pg in ["💬 Ask PaperIQ", "💬 Ask PaperIQ Chat"])
                or (label == "📚 My History" and current_pg in ["📚 History", "📚 My History", "🕘 History"])
            )
            st.button(label, key=key, on_click=navigate_to, args=(label,), use_container_width=True, type="primary" if is_active else "secondary")

    if is_student:
        render_html(f"<div style='font-size: 11px; color: {theme['text_muted']}; font-weight: 700; margin: 14px 0 6px 4px; letter-spacing: 0.05em;'>Account</div>")
        account_nav = [
            ("⚙ Settings", "nav_settings"),
            ("ℹ About", "nav_about"),
        ]
        for label, key in account_nav:
            is_active = (current_pg == label or (label == "⚙ Settings" and current_pg in ["⚙ Settings", "⚙️ Settings"]) or (label == "ℹ About" and current_pg in ["ℹ About", "ℹ️ About"]))
            st.button(label, key=key, on_click=navigate_to, args=(label,), use_container_width=True, type="primary" if is_active else "secondary")
        
    render_html(f"<div style='height: 12px; border-top: 1px solid {theme['border_divider']}; margin: 14px 0 8px 0;'></div>")
    if st.button("🚪 Logout", use_container_width=True, key="btn_logout"):
        saved_theme = st.session_state.get("theme_preference", "System")
        for k in list(st.session_state.keys()):
            del st.session_state[k]
        st.session_state["theme_preference"] = saved_theme
        st.session_state["authenticated"] = False
        st.session_state["user_data"] = {}
        st.session_state["auth_mode"] = "login"
        st.rerun()

# Display maintenance banner if configured in app_settings
try:
    active_maint_banner = str(auth.get_platform_setting("maintenance_banner", "") or "").strip()
    if active_maint_banner:
        render_html([
            f'<div style="background: {theme["warning_bg"]}; border: 1px solid {theme["warning_border"]}; border-left: 4px solid {theme["warning_text"]}; border-radius: 8px; padding: 10px 16px; margin-bottom: 14px; font-size: 13px; color: {theme["text_primary"]};">',
            f'<b style="color: {theme["warning_text"]};">📢 Platform Notice:</b> {safe_text(active_maint_banner)}',
            f'</div>',
        ])
except Exception:
    pass

# ----------------------------------------------------
def require_role(allowed_roles):
    """
    Centralized server-side authorization check.
    Verifies that the currently authenticated user's database role
    matches one of the allowed_roles. If unauthorized, records a security audit event,
    halts execution and safely redirects.
    """
    if not st.session_state.get("authenticated", False):
        st.error("Authentication required. Please sign in.")
        st.session_state.current_page = "🔑 Login"
        st.stop()
        
    u_name = st.session_state.get("username")
    db_role = auth.get_user_role(u_name)
    if not db_role:
        st.error("Your user account could not be verified or is deactivated. Please log in again.")
        st.session_state.authenticated = False
        st.stop()
        
    allowed_standard = [r.lower() for r in allowed_roles]
    if db_role.lower() not in allowed_standard:
        auth.record_audit_event(
            username=u_name,
            action_type="UNAUTHORIZED_ACCESS",
            target=st.session_state.get("current_page", "unknown"),
            details=f"User with database role '{db_role}' attempted to access section requiring {allowed_roles}",
            status="Denied"
        )
        st.error(f"🚫 Access Denied: This section requires {', '.join(allowed_roles)} privileges.")
        target_dash = "🏠 Researcher Dashboard" if db_role.lower() == "researcher" else "🏠 Student Dashboard"
        st.session_state.current_page = target_dash
        st.rerun()

admin_pages = [
    "🛡️ Admin Dashboard", "👥 All Users", "👥 User Management", "👥 All Users / User Management",
    "🕒 Login History", "🕒 Login & Audit Logs", "🕒 Login History & Audit Logs",
    "📁 Saved Papers", "⬆️ Upload History", "📤 Upload Research Paper",
    "📈 Analytics & Reports", "🩺 System Health", "💾 Storage & Data Management",
    "⚙️ Platform Settings", "ℹ️ About & System Info"
]
if st.session_state.current_page in admin_pages:
    require_role(["Administrator", "Admin"])

# Enforce role isolation: Student cannot access Researcher-only routes
import core.registry as _page_reg_guard
if is_student and _page_reg_guard.is_researcher_only_page(st.session_state.current_page):
    require_role(["Researcher"])

# Part 3 Privacy Guard: Administrator cannot access Student/Researcher private analysis pages
if is_admin and st.session_state.current_page not in admin_pages and st.session_state.current_page not in [
    "🏠 Dashboard", "🏠 Student Dashboard", "🏠 Researcher Dashboard",
    "⚙ Settings", "⚙️ Settings", "ℹ About", "ℹ️ About"
]:
    st.session_state.current_page = "🛡️ Admin Dashboard"
    st.rerun()

# ----------------------------------------------------
# EVALUATION & REPORTING WRAPPERS
# ----------------------------------------------------
def evaluate_academic_writing(text, filename="Paper", domain_override=None):
    cleaned = clean_text(text)
    res = analyze_full_document(cleaned)
    if not res:
        res = {
            "scores": {"Language Quality": 78, "Structural Coherence": 76, "Argumentation": 74, "Academic Style": 77, "Readability": 75},
            "overall_score": 76.0, "grade": "B", "full_text": cleaned
        }
    if domain_override:
        res["domain"] = domain_override
    res["filename"] = filename
    return res

def generate_pdf_report_wrapper(eval_res, sum_res=None, original_filename="Paper", title=None):
    if title:
        eval_res["title"] = title
    return generate_pdf_report(eval_res, original_filename)

# ----------------------------------------------------
# MODULAR RESEARCHER WORKSPACE DISPATCH (via core/registry.py)
# ----------------------------------------------------
if is_researcher:
    require_role(["Researcher"])
    try:
        _page_reg_guard.dispatch_researcher_page(
            canonical_page=st.session_state.current_page,
            theme=theme,
            username=username,
            user_data=st.session_state.get("user_data") or {},
            navigate_to=navigate_to,
            load_analysis_into_session=load_analysis_into_session,
            extract_text_from_file=extract_text_from_file,
            clean_text=clean_text,
            analyze_full_document=analyze_full_document,
            get_clean_display_title=get_clean_display_title,
        )
    except Exception as _res_pg_exc:
        auth.log_app_error(_res_pg_exc, page=st.session_state.get("current_page", "Researcher"), username=username, role=role)
        st.error("An unexpected error occurred while rendering this section. The incident has been logged in System Health.")
    st.stop()

# ----------------------------------------------------
# PAGE 1: DASHBOARD (STUDENT & ADMIN)
# ----------------------------------------------------
if st.session_state.current_page in ["🏠 Dashboard", "🏠 Student Dashboard", "🏠 Researcher Dashboard", "🛡️ Admin Dashboard"]:
    if is_admin:
        admin_panel.render_admin_dashboard(theme, navigate_to)
    else:
        workspaces.render_student_dashboard(
            theme, username, user_history, navigate_to,
            get_overall_score, get_grade, get_clean_display_title,
            get_clean_report_filename, load_analysis_into_session
        )

# ----------------------------------------------------
# ADMINISTRATIVE ROUTING
# ----------------------------------------------------
elif st.session_state.current_page in ["👥 All Users", "👥 User Management", "👥 All Users / User Management"]:
    admin_panel.render_all_users_view(theme, username)

elif st.session_state.current_page in ["🕒 Login History", "🕒 Login & Audit Logs", "🕒 Login History & Audit Logs"]:
    admin_panel.render_login_history_view(theme, username)

elif st.session_state.current_page == "📁 Saved Papers":
    admin_panel.render_saved_papers_view(theme, navigate_to, load_analysis_into_session, username)

elif st.session_state.current_page == "⬆️ Upload History":
    admin_panel.render_upload_history_view(theme, username)

elif st.session_state.current_page == "📤 Upload Research Paper":
    admin_panel.render_upload_research_paper_view(
        theme, navigate_to, username, extract_text_from_file,
        evaluate_academic_writing, summarizer.summarize_research_paper,
        generate_pdf_report_wrapper
    )

elif st.session_state.current_page == "📈 Analytics & Reports":
    admin_panel.render_analytics_view(theme, username)

elif st.session_state.current_page == "🩺 System Health":
    admin_panel.render_system_health_view(theme, username)

elif st.session_state.current_page == "💾 Storage & Data Management":
    admin_panel.render_storage_management_view(theme, username)

elif st.session_state.current_page == "⚙️ Platform Settings" or (is_admin and st.session_state.current_page in ["⚙ Settings", "⚙️ Settings"]):
    admin_panel.render_platform_settings_view(theme, username)

elif st.session_state.current_page == "ℹ️ About & System Info" or (is_admin and st.session_state.current_page in ["ℹ About", "ℹ️ About"]):
    admin_panel.render_admin_about_view(theme, username)
elif st.session_state.current_page in ["📄 Analyze Paper", "📄 Analyze Manuscript"]:
    pg_title = "Analyze Manuscript" if is_researcher else "Analyze Your Research Paper"
    pg_desc = "Upload an academic manuscript for in-depth evaluation across language quality, structural coherence, argumentation density, and readability." if is_researcher else "Upload a research paper for academic evaluation across language, structure, argumentation, and readability."
    render_html([
        '<div style="margin-bottom: 20px;">',
        f'<h1 style="margin: 0; font-size: 28px; color: {theme["text_primary"]};">{safe_text(pg_title)}</h1>',
        f'<p style="margin: 4px 0 0 0; color: {theme["text_muted"]}; font-size: 15px;">{safe_text(pg_desc)}</p>',
        '</div>',
    ])
    
    upload_card_title = "📄 Academic Manuscript Upload" if is_researcher else "📄 Research Paper Upload"
    render_html([
        f'<div style="background: {theme["bg_card_inner"]}; border: 1px dashed {theme["accent_primary_border"]}; border-radius: 12px; padding: 18px 22px 14px 22px; margin-bottom: 14px;">',
        '<div style="display: flex; justify-content: space-between; align-items: center; margin-bottom: 6px;">',
        f'<h3 style="margin: 0; font-size: 17px; color: {theme["text_primary"]}; font-weight: 600;">{safe_text(upload_card_title)}</h3>',
        f'<span style="background: {theme["badge_role_bg"]}; border: 1px solid {theme["badge_role_border"]}; border-radius: 6px; padding: 3px 9px; font-size: 11px; color: {theme["badge_role_text"]}; font-weight: 600;">Max 30 MB • PDF, DOCX, TXT</span>',
        '</div>',
        f'<p style="color: {theme["text_muted"]}; font-size: 13px; margin: 0;">Drag and drop your manuscript below or browse files. In-memory processing ensures full privacy and no permanent external storage.</p>',
        '</div>',
    ])
    
    uploaded_file = st.file_uploader(
        "Upload research document",
        type=["pdf", "docx", "txt"],
        accept_multiple_files=False,
        label_visibility="collapsed",
        key=f"file_uploader_widget_{st.session_state.get('uploader_key', 0)}"
    )
    
    if uploaded_file:
        file_bytes_raw = uploaded_file.getvalue()
        file_size_mb = len(file_bytes_raw) / (1024 * 1024)
        ext = uploaded_file.name.rsplit(".", 1)[-1].upper()

        # Invalidate previous summary if new file selected
        if st.session_state.get("summary_doc_name") != uploaded_file.name:
            st.session_state.summary_result = None
            st.session_state.summary_doc_text = ""
            st.session_state.summary_doc_name = uploaded_file.name

        # Extract & cache cleaned text upfront so title and word count are consistent everywhere
        cached_text = st.session_state.get("summary_doc_text", "")
        extraction_error = None
        if not cached_text and hasattr(uploaded_file, "getvalue"):
            try:
                pre_extracted = extract_text_from_file(uploaded_file)
                if pre_extracted and pre_extracted.startswith("__ERROR_"):
                    extraction_error = pre_extracted
                elif pre_extracted:
                    cached_text = clean_text(pre_extracted)
                    st.session_state.summary_doc_text = cached_text
            except Exception as ex:
                extraction_error = f"__ERROR_EXCEPTION__: {ex}"

        preview_title = get_clean_display_title(uploaded_file.name, text=cached_text, file=uploaded_file)
        doc_counts = compute_text_counts(cached_text) if cached_text else {"word_count": 0, "sentence_count": 0}
        doc_word_count = doc_counts["word_count"]

        render_html([
            f'<div class="iq-card" style="margin-top: 14px; margin-bottom: 14px; padding: 18px 22px !important; border-left: 4px solid {theme["success_text"]};">',
            '<div style="display: flex; justify-content: space-between; align-items: center; margin-bottom: 10px;">',
            f'<span style="font-size: 11px; font-weight: 700; color: {theme["success_text"]}; letter-spacing: 0.08em; text-transform: uppercase;">✓ FILE UPLOADED &amp; VALIDATED</span>',
            f'<span style="background: {theme["success_bg"]}; color: {theme["success_text"]}; font-size: 11px; font-weight: 600; padding: 2px 8px; border-radius: 4px; border: 1px solid {theme["success_border"]};">Ready for analysis</span>',
            '</div>',
            '<div style="display: grid; grid-template-columns: repeat(auto-fit, minmax(180px, 1fr)); gap: 14px;">',
            f'<div><span style="color: {theme["text_muted"]}; font-size: 12px;">Document:</span><br><b style="color: {theme["text_primary"]}; font-size: 14px; word-break: break-word;">📄 {safe_text(preview_title)}</b></div>',
            f'<div><span style="color: {theme["text_muted"]}; font-size: 12px;">Original file:</span><br><b style="color: {theme["text_secondary"]}; font-size: 13px; word-break: break-all;">{safe_text(uploaded_file.name)}</b></div>',
            f'<div><span style="color: {theme["text_muted"]}; font-size: 12px;">Format:</span><br><b style="color: {theme["accent_purple"]}; font-size: 14px;">{safe_text(ext)}</b></div>',
            f'<div><span style="color: {theme["text_muted"]}; font-size: 12px;">Size:</span><br><b style="color: {theme["text_primary"]}; font-size: 14px;">{file_size_mb:.2f} MB</b></div>',
            '</div>',
            '</div>',
        ])

        # Phase 4: Warn clearly on very short documents (<300 words) or non-English text
        if 0 < doc_word_count < 300:
            st.warning(f"⚠️ **Short Document Notice ({doc_word_count} words):** Documents with fewer than 300 words provide limited context; heuristic writing scores and section summaries may be less reliable.")
        if cached_text and doc_word_count >= 20:
            en_func_words = {"the", "and", "of", "to", "in", "is", "for", "that", "with", "on", "as", "are", "this", "by", "be", "at", "or", "from", "an", "was", "were"}
            sample_tokens = re.findall(r"\b[a-zA-Z]{2,}\b", cached_text[:4000].lower())
            en_hits = sum(1 for w in sample_tokens if w in en_func_words)
            non_ascii_chars = sum(1 for ch in cached_text[:4000] if ord(ch) > 127)
            if (len(sample_tokens) >= 20 and (en_hits / max(1, len(sample_tokens))) < 0.04) or (non_ascii_chars / max(1, len(cached_text[:4000])) > 0.25):
                st.warning("⚠️ **Language Notice:** This document appears to contain predominantly non-English text. PaperIQ's heuristic scoring, readability formulas, and section detectors are calibrated for English academic prose.")

        tab_summarize, tab_eval = st.tabs(["📝 Summarize Document", "🚀 Academic Quality Evaluation"])

        with tab_summarize:
            render_html([
                '<div style="margin-bottom: 12px; margin-top: 6px;">',
                f'<h3 style="margin: 0; font-size: 19px; color: {theme["text_primary"]}; font-weight: 700;">Executive Document Summarization</h3>',
                f'<p style="margin: 4px 0 0 0; color: {theme["text_muted"]}; font-size: 13.5px;">Generate source-grounded academic summaries of the complete paper across Short, Medium, or Long detail levels.</p>',
                '</div>',
            ])
            
            if doc_word_count > 0:
                doc_page_count = max(1, math.ceil(doc_word_count / 450))
            elif file_size_mb > 0 and not extraction_error:
                doc_page_count = max(1, math.ceil(file_size_mb * 2.5))
            else:
                doc_page_count = 0
            
            if extraction_error == "__ERROR_SCANNED_PDF__":
                status_badge_text = "⚠️ Scanned / Image PDF"
                status_badge_color = theme['error_text']
                status_badge_bg = theme['error_bg']
                status_badge_border = theme['error_border']
            elif extraction_error == "__ERROR_PASSWORD_PROTECTED__":
                status_badge_text = "⚠️ Encrypted Document"
                status_badge_color = theme['error_text']
                status_badge_bg = theme['error_bg']
                status_badge_border = theme['error_border']
            elif extraction_error == "__ERROR_EMPTY_DOCUMENT__" or (file_size_mb == 0 and doc_word_count == 0):
                status_badge_text = "⚠️ Empty Document"
                status_badge_color = theme['error_text']
                status_badge_bg = theme['error_bg']
                status_badge_border = theme['error_border']
            elif doc_word_count > 40:
                status_badge_text = "✓ Text Extracted & Ready"
                status_badge_color = theme['success_text']
                status_badge_bg = theme['success_bg']
                status_badge_border = theme['success_border']
            elif doc_word_count > 0:
                status_badge_text = "⚠️ Insufficient Academic Text"
                status_badge_color = theme['warning_text']
                status_badge_bg = theme['warning_bg']
                status_badge_border = theme['warning_border']
            else:
                status_badge_text = "⚠️ Ready to Extract"
                status_badge_color = theme['warning_text']
                status_badge_bg = theme['warning_bg']
                status_badge_border = theme['warning_border']
            
            render_html([
                f'<div style="background: {theme["bg_card_inner"]}; border: 1px solid {theme["border_color"]}; border-radius: 10px; padding: 14px 18px; margin-bottom: 16px;">',
                '<div style="display: flex; justify-content: space-between; align-items: center; margin-bottom: 8px;">',
                f'<span style="font-size: 11.5px; font-weight: 700; color: {theme["info_text"]}; letter-spacing: 0.05em; text-transform: uppercase;">📋 Document Extraction Profile</span>',
                f'<span style="background: {status_badge_bg}; color: {status_badge_color}; font-size: 11px; font-weight: 600; padding: 2px 8px; border-radius: 4px; border: 1px solid {status_badge_border};">{status_badge_text}</span>',
                '</div>',
                '<div style="display: grid; grid-template-columns: repeat(auto-fit, minmax(160px, 1fr)); gap: 10px; font-size: 13px;">',
                f'<div><span style="color: {theme["text_muted"]};">Document:</span> <b style="color: {theme["text_primary"]};">{safe_text(preview_title)}</b></div>',
                f'<div><span style="color: {theme["text_muted"]};">Format:</span> <b style="color: {theme["accent_purple"]};">{safe_text(ext)} ({file_size_mb:.2f} MB)</b></div>',
                f'<div><span style="color: {theme["text_muted"]};">Extracted Words:</span> <b style="color: {theme["success_text"]};">{f"{doc_word_count:,}" if doc_word_count > 0 else ("0 words (Empty)" if (file_size_mb == 0 or extraction_error == "__ERROR_EMPTY_DOCUMENT__") else "Ready to parse")}</b></div>',
                f'<div><span style="color: {theme["text_muted"]};">Estimated Pages:</span> <b style="color: {theme["text_primary"]};">{f"~{doc_page_count} pages" if doc_page_count > 0 else "0 pages"}</b></div>',
                '</div>',
                '</div>',
            ])
            
            # Step 4: Choose the summary type
            col_mode, col_info = st.columns([1.3, 2.1])
            with col_mode:
                mode_options = ["Short", "Medium", "Long"]
                current_mode_idx = mode_options.index(st.session_state.get("summary_length", "Medium")) if st.session_state.get("summary_length", "Medium") in mode_options else 1
                selected_mode = st.radio(
                    "Choose Summary Detail Level",
                    options=mode_options,
                    index=current_mode_idx,
                    horizontal=True,
                    key="radio_summary_length_choice"
                )
                st.session_state["summary_length"] = selected_mode
            
            with col_info:
                if selected_mode == "Short":
                    st.info("💡 **Short Summary (~100–150 words):** Quick executive overview covering research problem, primary objective, core methodology, key quantitative finding, and main conclusion.")
                elif selected_mode == "Medium":
                    st.info("💡 **Medium Summary (~250–400 words):** Structured thematic synthesis covering background, research objectives, study design/datasets, principal empirical findings, practical implications, and reported limitations.")
                else:
                    st.info("💡 **Long Summary (~600–900 words):** In-depth section-by-section breakdown across all 8 academic pillars: motivation, problem statement, methodology, experimental setup, quantitative results, discussion, validity threats, and future directions.")
            
            btn_sum_click = st.button("📝 Generate Summary", type="primary", use_container_width=True, key="btn_execute_summarization")
            
            if btn_sum_click:
                with st.status("Processing document and generating summary...", expanded=True) as sum_status:
                    try:
                        sum_status.write("✓ Stage 1: Validating document and extracting readable text...")
                        if not st.session_state.get("summary_doc_text"):
                            raw_text = extract_text_from_file(uploaded_file)
                        else:
                            raw_text = st.session_state.summary_doc_text
                            
                        if raw_text == "__ERROR_SCANNED_PDF__":
                            sum_status.update(label="Scanned PDF detected", state="error")
                            st.error("This document appears to be a scanned or image-based PDF with no extractable text layer. Please use an OCR tool (e.g. Adobe Acrobat OCR) to convert the document into searchable text, or upload in DOCX or TXT format.")
                        elif raw_text == "__ERROR_PASSWORD_PROTECTED__":
                            sum_status.update(label="Password-protected document", state="error")
                            st.error("This document is password-protected or encrypted. Please remove password protection before uploading.")
                        elif raw_text == "__ERROR_FILE_TOO_LARGE__":
                            sum_status.update(label="File exceeds size limit", state="error")
                            st.error("The uploaded file exceeds the 30 MB size limit.")
                        elif raw_text == "__ERROR_EMPTY_DOCUMENT__" or not raw_text:
                            sum_status.update(label="Empty document", state="error")
                            st.error("The uploaded document is empty or does not contain readable text.")
                        else:
                            cleaned_doc = clean_text(raw_text)
                            st.session_state.summary_doc_text = cleaned_doc
                            
                            sum_status.write("✓ Stage 2: Segmenting into academic sections & tracking boundaries...")
                            sum_status.write("✓ Stage 3: Scoring sentence salience & empirical evidence...")
                            sum_status.write("✓ Stage 4: Synthesizing cross-section discourse & removing redundancies...")
                            sum_status.write("✓ Stage 5: Generating source-grounded summary...")
                            
                            sum_result = summarizer.summarize_research_paper(
                                text=cleaned_doc,
                                summary_type=selected_mode,
                                filename=uploaded_file.name,
                                doc_title=preview_title
                            )
                            
                            if sum_result.get("success"):
                                st.session_state.summary_result = sum_result
                                st.session_state.summary_doc_name = uploaded_file.name
                                st.session_state.summary_length = selected_mode
                                sum_status.update(label="Summary generated successfully!", state="complete", expanded=False)
                                st.success(f"{selected_mode} summary generated successfully! ({sum_result['word_count']} words)")
                            else:
                                sum_status.update(label="Summarization failed", state="error")
                                st.error(sum_result.get("error", "Summarization failed."))
                    except Exception as sum_err:
                        print(f"[ERROR] Summarization error: {sum_err}")
                        sum_status.update(label="Summarization could not be completed", state="error")
                        st.error("Summarization could not be completed. Please ensure the document contains readable academic text and retry.")

            # Display generated summary if available
            cur_sum = st.session_state.get("summary_result")
            if cur_sum and cur_sum.get("success") and st.session_state.get("summary_doc_name") == uploaded_file.name:
                render_html("<br>")
                with st.container(border=True):
                    sum_title_str = cur_sum.get("document_info", {}).get("title") or preview_title
                    # Summary Header
                    render_html([
                        f'<div style="display: flex; justify-content: space-between; align-items: center; flex-wrap: wrap; gap: 10px; margin-bottom: 12px; border-bottom: 1px solid {theme["border_divider"]}; padding-bottom: 10px;">',
                        f'<div>',
                        f'<span style="font-size: 11px; font-weight: 700; color: {theme["info_text"]}; letter-spacing: 0.08em; text-transform: uppercase;">📄 DOCUMENT SUMMARY</span>',
                        f'<h3 style="margin: 2px 0 0 0; font-size: 20px; color: {theme["text_primary"]};">{safe_text(sum_title_str)}</h3>',
                        f'</div>',
                        f'<div style="display: flex; gap: 6px; flex-wrap: wrap; align-items: center;">',
                        f'<span style="background: {theme["accent_primary_subtle"]}; border: 1px solid {theme["accent_primary_border"]}; border-radius: 6px; padding: 3px 8px; font-size: 11.5px; color: {theme["accent_purple"]}; font-weight: 700;">{safe_text(cur_sum["summary_type"])} Summary</span>',
                        f'<span style="background: {theme["success_bg"]}; border: 1px solid {theme["success_border"]}; border-radius: 6px; padding: 3px 8px; font-size: 11.5px; color: {theme["success_text"]}; font-weight: 600;">{cur_sum["word_count"]} words</span>',
                        f'<span style="background: {theme["info_bg"]}; border: 1px solid {theme["info_border"]}; border-radius: 6px; padding: 3px 8px; font-size: 11.5px; color: {theme["info_text"]}; font-weight: 600;">~{cur_sum["reading_time_min"]} min read</span>',
                        f'<span title="Coverage % measures the proportion of detected core manuscript sections represented in the generated summary attributions." style="cursor: help; background: {theme["warning_bg"]}; border: 1px solid {theme["warning_border"]}; border-radius: 6px; padding: 3px 8px; font-size: 11.5px; color: {theme["warning_text"]}; font-weight: 600;">{cur_sum["engine_info"].get("coverage_percentage", 100)}% Coverage ⓘ</span>',
                        f'</div>',
                        f'</div>',
                    ])
                    
                    if cur_sum.get("warnings"):
                        for warn in cur_sum["warnings"]:
                            st.caption(f"ℹ️ Note: {unescape_text(warn)}")
                    
                    # Formatted Summary Content
                    st.markdown(unescape_text(cur_sum["summary"]))
                    
                    # Grounding & Source Attribution Expander
                    if cur_sum.get("source_attribution"):
                        with st.expander("🔍 Source Evidence Attribution & Grounding", expanded=False):
                            render_html(f"<p style='font-size: 12.5px; color: {theme['text_muted']}; margin-bottom: 8px;'>The statements in this summary were grounded directly in the following sections of the uploaded manuscript:</p>")
                            for attr in cur_sum["source_attribution"]:
                                render_html([
                                    f'<div style="background: {theme["bg_card_inner"]}; border: 1px solid {theme["border_color"]}; border-left: 3px solid {theme["info_text"]}; padding: 6px 12px; margin-bottom: 6px; border-radius: 4px;">',
                                    f'<b style="color: {theme["info_text"]}; font-size: 12px;">[{safe_text(attr.get("pillar", "Key Evidence"))}]</b> &nbsp;',
                                    f'<span style="color: {theme["text_secondary"]}; font-size: 12px;">Section: <i>{safe_text(attr.get("section", "General"))}</i> (Page ~{attr.get("page", 1)})</span><br>',
                                    f'<span style="color: {theme["text_muted"]}; font-size: 11.5px; font-style: italic;">"{safe_text(attr.get("excerpt", ""))}"</span>',
                                    f'</div>',
                                ])

                    render_html(f"<hr style='border: none; border-top: 1px solid {theme['border_divider']}; margin: 16px 0 12px 0;'>")
                    
                    # Action Bar (Download PDF, Download TXT, Copy)
                    c_act1, c_act2, c_act3 = st.columns([1.2, 1.2, 1.2])
                    with c_act1:
                        sum_pdf_bytes = get_cached_summary_pdf(cur_sum, paper_title=sum_title_str, metadata={"filename": uploaded_file.name})
                        sum_pdf_name = pdf_export.get_pdf_filename(sum_title_str, f"Summary_{cur_sum['summary_type']}")
                        st.download_button(
                            "⬇ Download Summary (PDF)",
                            data=sum_pdf_bytes,
                            file_name=sum_pdf_name,
                            mime="application/pdf",
                            use_container_width=True,
                            key="btn_download_summary_pdf"
                        )
                    with c_act2:
                        txt_export_content = summarizer.export_summary_as_txt(cur_sum)
                        txt_export_name = summarizer.get_summary_export_filename(uploaded_file.name, cur_sum["summary_type"])
                        st.download_button(
                            "⬇ Download Summary (.txt)",
                            data=txt_export_content.encode("utf-8"),
                            file_name=txt_export_name,
                            mime="text/plain",
                            use_container_width=True,
                            key="btn_download_summary_txt"
                        )
                    with c_act3:
                        with st.popover("📋 Copy Full Summary Text", use_container_width=True):
                            st.caption("Click the copy icon in the top right corner of the box below:")
                            st.code(unescape_text(cur_sum["summary"]), language=None)

        with tab_eval:
            if st.session_state.get("filename") and uploaded_file.name != st.session_state.get("filename"):
                anl_btn_hint = "🚀 Analyze Manuscript" if is_researcher else "🚀 Analyze Paper"
                doc_term = "manuscript" if is_researcher else "paper"
                st.info(f"ℹ️ New document selected ('{uploaded_file.name}'). Click '{anl_btn_hint}' below to analyze this {doc_term} and replace previous results for '{st.session_state.get('filename')}'.")
            
            has_matching_analysis = has_matching_analysis_for_file(uploaded_file.name)
            
            run_analysis_clicked = False
            if has_matching_analysis:
                succ_msg = "🎉 **Manuscript evaluated successfully!** Scholarly evaluation metrics, structural coherence, and grounded insights are ready." if is_researcher else "🎉 **Paper analyzed successfully!** Scholarly evaluation metrics, structural coherence, and grounded insights are ready."
                st.success(succ_msg)
                c_act1, c_act2, c_act3 = st.columns([2.0, 1.8, 1.6])
                with c_act1:
                    target_results_pg = "📊 Manuscript Evaluation" if is_researcher else ("📊 Writing Quality & Feedback" if is_student else "📊 Analysis")
                    btn_view_label = "📊 View Manuscript Evaluation →" if is_researcher else "📊 View Analysis Results →"
                    btn_view_main = st.button(btn_view_label, type="primary", use_container_width=True, key="btn_view_analysis_results", on_click=navigate_to, args=(target_results_pg,))
                    if btn_view_main:
                        navigate_to(target_results_pg)
                        st.rerun()
                with c_act2:
                    active_res_for_pdf = st.session_state.get("results") or st.session_state.get("active_analysis")
                    saved_pdf = st.session_state.get("pdf_bytes") or (get_cached_analysis_pdf(active_res_for_pdf, uploaded_file.name, st.session_state.get("current_analysis_id")) if active_res_for_pdf else None)
                    saved_clean_fn = st.session_state.get("report_filename") or pdf_export.get_pdf_filename(st.session_state.get("paper_title") or uploaded_file.name, "Analysis")
                    if saved_pdf:
                        st.download_button("⬇ Download Report", data=saved_pdf, file_name=saved_clean_fn, mime="application/pdf", key="btn_direct_down_rep", use_container_width=True)
                    else:
                        st.button("⬇ Download Report", disabled=True, use_container_width=True, key="btn_direct_down_disabled", help="Analysis report not generated yet.")
                with c_act3:
                    re_label = "🔄 Re-analyze Manuscript" if is_researcher else "🔄 Re-analyze Paper"
                    run_analysis_clicked = st.button(re_label, use_container_width=True, key="btn_reanalyze_paper")
            else:
                # Before analysis: Analyze Paper is enabled, View Analysis Results is disabled
                c_btn1, c_btn2 = st.columns([1.6, 1.8])
                with c_btn1:
                    run_label = "🚀 Analyze Manuscript" if is_researcher else "🚀 Analyze Paper"
                    run_analysis_clicked = st.button(run_label, type="primary", use_container_width=True, key="btn_run_analysis")
                with c_btn2:
                    disabled_view_label = "📊 View Manuscript Evaluation →" if is_researcher else "📊 View Analysis Results →"
                    st.button(disabled_view_label, disabled=True, use_container_width=True, key="btn_redirect_top_disabled", help="Analyze this paper first to view results.")
        
            if run_analysis_clicked:
                # Invalidate previous document analysis state to prevent stale data leakage
                st.session_state.pop("results", None)
                st.session_state.pop("active_analysis", None)
                st.session_state.pop("analysis_result", None)
                st.session_state.pop("current_analysis_id", None)
                st.session_state.pop("overall_score", None)
                st.session_state.pop("pdf_bytes", None)
                st.session_state.pop("report_filename", None)
                st.session_state.pop("analysis_just_completed", None)
                st.session_state["chat_history"] = []
            
                status_label = "Evaluating academic manuscript..." if is_researcher else "Analyzing research paper..."
                with st.status(status_label, expanded=True) as status:
                    try:
                        st.write("✓ Document extraction")
                        raw_text = extract_text_from_file(uploaded_file)
                    
                        if raw_text == "__ERROR_SCANNED_PDF__":
                            status.update(label="Scanned PDF detected", state="error")
                            st.error("This document appears to be a scanned or image-based PDF with no extractable text layer. Please use an OCR tool (such as Adobe Acrobat OCR or Tesseract) to convert the document into searchable text, or upload the source manuscript in DOCX or TXT format.")
                        elif raw_text == "__ERROR_PASSWORD_PROTECTED__":
                            status.update(label="Password-protected document", state="error")
                            st.error("This document is password-protected or encrypted. Please remove password protection or encryption before uploading.")
                        elif raw_text == "__ERROR_FILE_TOO_LARGE__":
                            status.update(label="File exceeds size limit", state="error")
                            st.error("The uploaded file exceeds the 30 MB size limit. Please upload a smaller or compressed document.")
                        elif raw_text == "__ERROR_CORRUPTED_PDF__":
                            status.update(label="Corrupted PDF", state="error")
                            st.error("The PDF file appears to be corrupted or could not be read. Please check the file and try again.")
                        elif raw_text == "__ERROR_CORRUPTED_DOCX__":
                            status.update(label="Corrupted Word document", state="error")
                            st.error("The Word document (.docx) could not be parsed. Please ensure it is a valid, uncorrupted DOCX file.")
                        elif raw_text == "__ERROR_EMPTY_DOCUMENT__" or not raw_text:
                            status.update(label="Analysis failed: empty document", state="error")
                            st.error("The uploaded document is empty or does not contain readable text. Please upload a valid research paper.")
                        else:
                            st.write("✓ Text preprocessing")
                            cleaned = clean_text(raw_text)
                        
                            if not cleaned or len(cleaned.strip()) < 50:
                                status.update(label="Analysis failed: insufficient text", state="error")
                                st.error("This document does not contain enough extractable academic text for scholarly analysis (minimum 50 characters required). Please ensure the file contains research content.")
                            else:
                                st.write("✓ Structural analysis")
                                sections_data = extract_sections(cleaned)
                            
                                st.write("✓ Language analysis")
                                st.write("✓ Argumentation analysis")
                                st.write("✓ Readability analysis")
                                results = analyze_full_document(cleaned)
                            
                                if results:
                                    st.write("✓ Research insight generation")
                                    keywords, _ = extract_keywords_and_domain(cleaned)
                                    clf = classify_research_domain(cleaned)
                                    publisher = detect_publisher(cleaned)
                                    citations = analyze_citations(cleaned)
                                    semantic = calculate_semantic_strength(cleaned)
                                    novelty = novelty_score(cleaned)
                                    display_title = get_clean_display_title(uploaded_file.name, text=cleaned, file=uploaded_file)
                                    clean_fn = pdf_export.get_pdf_filename(display_title, "Analysis")
                                    overall_score = get_overall_score(results)
                                    grade = get_grade(overall_score)
                                
                                    results["title"] = display_title
                                    results["display_title"] = display_title
                                    results["overall_score"] = overall_score
                                    results["keywords"] = keywords
                                    results["domains"] = clf["all_candidates"]
                                    results["domain"] = clf["primary"]
                                    results["domain_confidence"] = clf["confidence"]
                                    results["domain_secondary"] = clf["secondary"]
                                    results["publisher"] = publisher
                                    results["full_text"] = cleaned
                                    results["citation_analysis"] = citations
                                    results["semantic_strength"] = semantic
                                    results["novelty_score"] = novelty
                                    results["ai_feedback"] = generate_research_feedback(results)
                                    results["potential_research_areas"] = clf["all_candidates"]
                                    results["future_directions"] = extract_future_directions(cleaned)
                                    results["methodology_and_findings"] = extract_methodology_and_findings(cleaned, results["sections"])
                                    results["structured_gaps"] = extract_structured_research_gaps(cleaned)
                                    results["structured_contributions"] = extract_key_contributions_structured(cleaned)
                                    results["recommended_journal"] = recommend_journal(clf["all_candidates"])
                                    results["issues"] = detect_paper_issues(cleaned, results["sections"], results["scores"], results["stats"], citations)
                                    results["vocab_improvements"] = extract_vocabulary_improvements(cleaned)
                                    results["section_sentiments"] = compute_section_sentiments(results["sections"])
                                
                                    st.write("✓ Report generation")
                                    pdf_bytes = generate_pdf_report(results, uploaded_file.name)
                                
                                    # Serialize safely for persistent viewing
                                    serializable_res = dict(results)
                                    serializable_res.pop("blob", None)
                                    res_json = json.dumps(serializable_res)
                                
                                    doc_hash = ""
                                    try:
                                        if file_bytes_raw:
                                            doc_hash = hashlib.sha256(file_bytes_raw).hexdigest()[:16]
                                        elif cleaned:
                                            doc_hash = hashlib.sha256(cleaned.encode("utf-8", errors="ignore")).hexdigest()[:16]
                                    except Exception:
                                        doc_hash = f"hash-{int(time.time())}"
                                    
                                    try:
                                        anl_uuid = f"ANL-{int(time.time())}-{uuid.uuid4().hex[:6].upper()}"
                                    except Exception:
                                        anl_uuid = f"ANL-{int(time.time())}-REC"
                                
                                    rec_id, ts, anl_id = auth.save_analysis_record(
                                        username=username,
                                        filename=uploaded_file.name,
                                        domain=results["domain"],
                                        score=overall_score,
                                        pdf_bytes=pdf_bytes,
                                        results_json=res_json,
                                        analysis_id=anl_uuid,
                                        display_title=display_title,
                                        report_filename=clean_fn,
                                        document_hash=doc_hash,
                                        grade=grade,
                                        dimension_scores=results["scores"]
                                    )

                                    try:
                                        f_kb = len(file_bytes_raw) / 1024.0 if file_bytes_raw else 0.0
                                        f_type = uploaded_file.name.split(".")[-1].upper() if "." in uploaded_file.name else "PDF"
                                        auth.record_upload(username, uploaded_file.name, f_type, f_kb, status="Analyzed", target="Personal")
                                    except Exception:
                                        pass
                                
                                    # Authoritative session state synchronization
                                    st.session_state["results"] = results
                                    st.session_state["active_analysis"] = results
                                    st.session_state["analysis_result"] = results
                                    st.session_state["analysis_results"] = results
                                    st.session_state["analysis_complete"] = True
                                    st.session_state["filename"] = uploaded_file.name
                                    st.session_state["analyzed_filename"] = uploaded_file.name
                                    st.session_state["paper_title"] = display_title
                                    st.session_state["overall_score"] = overall_score
                                    st.session_state["current_analysis_id"] = anl_id
                                    st.session_state["pdf_bytes"] = pdf_bytes
                                    st.session_state["report_filename"] = clean_fn
                                    st.session_state["chat_history"] = []
                                    st.session_state["analysis_just_completed"] = uploaded_file.name
                                    st.session_state["active_analysis_source"] = "uploaded"
                                    st.session_state.analysis_history[username] = auth.get_user_history(username)
                                
                                    status.update(label="Paper analyzed successfully!", state="complete", expanded=False)
                                    st.rerun()
                                else:
                                    status.update(label="Error analyzing document", state="error")
                                    st.error("Something went wrong while analyzing the document. Please check the file and try again.")
                    except Exception as ex:
                        print(f"[ERROR] Analysis failed: {ex}")
                        status.update(label="Analysis could not be completed", state="error")
                        st.error("Analysis could not be completed. Please retry the analysis. If the problem continues, check the uploaded document format.")

    # Compact Active Analysis Status Card (Non-intrusive)
    elif has_any_valid_analysis():
        res = st.session_state.get("results") or st.session_state.get("active_analysis")
        overall_score = get_overall_score(res)
        composite = overall_score
        composite_score = overall_score
        grade = get_grade(overall_score)
        filename = st.session_state.get("filename", "Paper")
        clean_disp = st.session_state.get("paper_title") or get_clean_display_title(filename, stored_title=res.get("title") if isinstance(res, dict) else None)
        clean_fn = pdf_export.get_pdf_filename(clean_disp, "Analysis")
        
        is_from_hist = (st.session_state.get("active_analysis_source") == "history")
        badge_txt = "📚 Saved History" if is_from_hist else "📄 Active Session"
        badge_bg = theme['badge_role_bg']
        badge_color = theme['accent_purple'] if is_from_hist else theme['success_text']

        render_html("<br>")
        with st.container(border=True):
            col_res1, col_res2, col_res3, col_res4 = st.columns([3.4, 1.8, 1.6, 1.2])
            with col_res1:
                dom_str = res.get('domain', 'General Academic') if isinstance(res, dict) else 'General Academic'
                rd_str = res.get('stats', {}).get('reading_difficulty', 'Standard Academic') if isinstance(res, dict) else 'Standard Academic'
                render_html([
                    f'<div style="font-weight: 700; color: {theme["text_primary"]}; font-size: 15px;">',
                    f'✓ Active Analysis: 📄 {safe_text(clean_disp)} &nbsp;',
                    f'<span style="background: {badge_bg}; color: {badge_color}; font-size: 10.5px; font-weight: 700; padding: 2px 7px; border-radius: 6px; border: 1px solid {theme["border_color"]};">{badge_txt}</span>',
                    f'</div>',
                    f'<div style="color: {theme["text_muted"]}; font-size: 12px; margin-top: 3px;">',
                    f'Domain: <b style="color: {theme["info_text"]};">{safe_text(dom_str)}</b> &nbsp;•&nbsp; ',
                    f'Reading Level: <b style="color: {theme["text_secondary"]};">{safe_text(rd_str)}</b>',
                    f'</div>',
                ])
            with col_res2:
                render_html([
                    f'<div style="font-size: 15px; font-weight: 800; color: {theme["accent_purple"]}; margin-top: 6px;">',
                    f'Score: {overall_score:.1f} / 100 &nbsp;',
                    f'<span style="background: {theme["badge_role_bg"]}; color: {theme["badge_role_text"]}; font-size: 11px; font-weight: 700; padding: 2px 7px; border-radius: 8px; border: 1px solid {theme["badge_role_border"]};">Grade {safe_text(grade)}</span>',
                    f'</div>',
                ])
            with col_res3:
                btn_goto_act = st.button("📊 View Results →", type="primary", use_container_width=True, key="btn_goto_analysis_act", on_click=navigate_to, args=("📊 Analysis",))
                if btn_goto_act:
                    navigate_to("📊 Analysis")
                    st.rerun()
            with col_res4:
                if st.button("🗑️ Clear", use_container_width=True, key="btn_clear_act_on_analyze", help="Clear active analysis from workspace"):
                    clear_active_analysis_session()
                    st.rerun()
            
            with st.expander("📝 Generate / View Executive Summary for this Paper", expanded=False):
                cur_sum = st.session_state.get("summary_result")
                if cur_sum and cur_sum.get("success") and (cur_sum.get("document_info", {}).get("filename") == filename or cur_sum.get("document_info", {}).get("title") == clean_disp):
                    render_html([
                        f'<div style="display: flex; gap: 8px; flex-wrap: wrap; margin-bottom: 10px;">',
                        f'<span style="background: {theme["accent_primary_subtle"]}; border: 1px solid {theme["accent_primary_border"]}; border-radius: 6px; padding: 2px 8px; font-size: 11px; color: {theme["accent_purple"]}; font-weight: 700;">{safe_text(cur_sum["summary_type"])} Summary</span>',
                        f'<span style="background: {theme["success_bg"]}; border: 1px solid {theme["success_border"]}; border-radius: 6px; padding: 2px 8px; font-size: 11px; color: {theme["success_text"]}; font-weight: 600;">{cur_sum["word_count"]} words</span>',
                        f'<span style="background: {theme["info_bg"]}; border: 1px solid {theme["info_border"]}; border-radius: 6px; padding: 2px 8px; font-size: 11px; color: {theme["info_text"]}; font-weight: 600;">~{cur_sum["reading_time_min"]} min read</span>',
                        f'</div>',
                    ])
                    st.markdown(unescape_text(cur_sum["summary"]))
                    c_dl_pdf, c_dl_txt, c_cp = st.columns([1.2, 1.2, 1.2])
                    with c_dl_pdf:
                        act_sum_pdf = get_cached_summary_pdf(cur_sum, paper_title=clean_disp, metadata={"filename": filename})
                        act_sum_pdf_name = pdf_export.get_pdf_filename(clean_disp, f"Summary_{cur_sum['summary_type']}")
                        st.download_button("⬇ Download Summary (PDF)", data=act_sum_pdf, file_name=act_sum_pdf_name, mime="application/pdf", use_container_width=True, key="btn_dl_active_sum_pdf")
                    with c_dl_txt:
                        txt_c = summarizer.export_summary_as_txt(cur_sum)
                        txt_n = summarizer.get_summary_export_filename(filename, cur_sum["summary_type"])
                        st.download_button("⬇ Download Summary (TXT)", data=txt_c.encode("utf-8"), file_name=txt_n, mime="text/plain", use_container_width=True, key="btn_dl_active_sum")
                    with c_cp:
                        with st.popover("📋 Copy Summary Text", use_container_width=True):
                            st.caption("Click the copy icon in the top right corner of the box below:")
                            st.code(unescape_text(cur_sum["summary"]), language=None)
                else:
                    st.info("No executive summary has been generated for this active paper yet. Choose a summary mode below to generate one immediately without re-uploading:")
                    col_m1, col_m2 = st.columns([1.3, 2.1])
                    with col_m1:
                        sel_act_mode = st.radio("Summary Mode", ["Short", "Medium", "Long"], index=1, horizontal=True, key="sel_act_paper_mode")
                    with col_m2:
                        if sel_act_mode == "Short":
                            st.caption("Quick executive overview (~100–150 words).")
                        elif sel_act_mode == "Medium":
                            st.caption("Structured thematic explanation (~250–400 words).")
                        else:
                            st.caption("Comprehensive section-by-section breakdown (~600–900 words).")
                    if st.button("📝 Generate Summary for Active Paper", type="primary", use_container_width=True, key="btn_gen_active_sum"):
                        with st.spinner("Generating summary..."):
                            full_txt = res.get("full_text", "")
                            if full_txt:
                                act_sum = summarizer.summarize_research_paper(full_txt, summary_type=sel_act_mode, filename=filename, doc_title=clean_disp)
                                st.session_state.summary_result = act_sum
                                st.session_state.summary_doc_text = full_txt
                                st.session_state.summary_doc_name = filename
                                st.session_state.summary_length = sel_act_mode
                                st.rerun()

# ----------------------------------------------------
# PAGE 3: DETAILED ANALYSIS
# ----------------------------------------------------
elif st.session_state.current_page in ["📊 Analysis", "📊 Writing Quality & Feedback", "📊 Manuscript Evaluation", "📊 Analysis Results"]:
    anl_heading = "Manuscript Evaluation" if is_researcher else ("Student Writing Quality & Feedback" if is_student else "Academic Analysis & Performance")
    render_html([
        f'<div style="margin-bottom: 20px;">',
        f'<h1 style="margin: 0; font-size: 28px; color: {theme["text_primary"]};">{safe_text(anl_heading)}</h1>',
        f'<p style="margin: 4px 0 0 0; color: {theme["text_muted"]}; font-size: 15px;">Clear, explainable evaluation of textual quality, logical structure, and argumentative density.</p>',
        f'</div>',
    ])
    
    if not has_any_valid_analysis():
        empty_title = "No Manuscript Evaluated Yet" if is_researcher else "No Paper Analyzed Yet"
        empty_text = "Upload and evaluate an academic manuscript in <b>Analyze Manuscript</b> to inspect detailed analytics." if is_researcher else "Upload and analyze a research paper in <b>Analyze Paper</b> to inspect detailed analytics."
        btn_label = "Go to Analyze Manuscript" if is_researcher else "Go to Analyze Paper"
        target_anl_pg = "📄 Analyze Manuscript" if is_researcher else "📄 Analyze Paper"
        render_html([
            f'<div class="iq-card" style="text-align: center; padding: 40px !important;">',
            f'<div style="font-size: 32px; margin-bottom: 8px;">📊</div>',
            f'<h4 style="color: {theme["text_primary"]}; margin: 0;">{safe_text(empty_title)}</h4>',
            f'<p style="color: {theme["text_muted"]}; margin-top: 4px;">{empty_text}</p>',
            f'</div>',
        ])
        btn_goto_anl = st.button(btn_label, type="primary", key="btn_goto_analyze", on_click=navigate_to, args=(target_anl_pg,))
        if btn_goto_anl:
            navigate_to(target_anl_pg)
            st.rerun()
    else:
        res = st.session_state.get("results") or st.session_state.get("active_analysis")
        scores = res.get("scores", {})
        stats = res.get("stats", {})
        filename = st.session_state.get("filename", "Paper")
        clean_display_title = st.session_state.get("paper_title") or get_clean_display_title(filename, text=res.get("full_text"), stored_title=res.get("title"))
        clean_fn = pdf_export.get_pdf_filename(clean_display_title, "Analysis")
        feedback_fn = pdf_export.get_pdf_filename(clean_display_title, "Feedback")
        comp = get_overall_score(res)
        grade = get_grade(comp)
        
        is_from_hist = (st.session_state.get("active_analysis_source") == "history")
        badge_txt = "📚 Archived Historical Report" if is_from_hist else "✨ Current Session Active Analysis"
        badge_bg = theme['badge_role_bg']
        badge_color = theme['accent_purple'] if is_from_hist else theme['success_text']

        # Cached full analysis/feedback PDF bytes
        anl_pdf_bytes = get_cached_analysis_pdf(res, filename, st.session_state.get("current_analysis_id"))

        # Action Toolbar
        c_tb1, c_tb_rep, c_tb_fb, c_tb2, c_tb3 = st.columns([2.6, 1.5, 1.6, 1.6, 1.3])
        with c_tb1:
            st.caption(f"Currently inspecting: **{clean_display_title}** ({badge_txt})")
        with c_tb_rep:
            st.download_button(
                "⬇ Download Report",
                data=anl_pdf_bytes,
                file_name=clean_fn,
                mime="application/pdf",
                use_container_width=True,
                key="btn_anl_page_download_report"
            )
        with c_tb_fb:
            st.download_button(
                "⬇ Download Feedback (PDF)",
                data=anl_pdf_bytes,
                file_name=feedback_fn,
                mime="application/pdf",
                use_container_width=True,
                key="btn_anl_page_download_feedback"
            )
        with c_tb2:
            btn_another_label = "📄 Analyze Another Manuscript" if is_researcher else "📄 Analyze Another Paper"
            target_another_pg = "📄 Analyze Manuscript" if is_researcher else "📄 Analyze Paper"
            if st.button(btn_another_label, key="btn_anl_another", use_container_width=True, on_click=navigate_to, args=(target_another_pg,)):
                navigate_to(target_another_pg)
                st.rerun()
        with c_tb3:
            if st.button("🗑️ Clear Active Analysis", key="btn_anl_clear_active", use_container_width=True, help="Remove this paper from active workspace"):
                clear_active_analysis_session()
                st.rerun()

        # A. PAPER OVERVIEW
        areas = res.get("potential_research_areas") or res.get("domains") or [res.get("domain", "General Academic")]
        area_badges_html = "".join([f"<span class='domain-badge'>{safe_text(a)}</span>" for a in areas[:3]])
        rd_diff_str = safe_text(res.get('stats', {}).get('reading_difficulty', 'Standard Academic'))
        render_html([
            f'<div class="iq-card" style="padding: 18px 22px !important; margin-bottom: 18px !important;">',
            f'<div style="display: flex; justify-content: space-between; align-items: flex-start; flex-wrap: wrap; gap: 10px;">',
            f'<div>',
            f'<div style="font-size: 12px; color: {theme["text_muted"]}; text-transform: uppercase; font-weight: 700; letter-spacing: 0.05em;">',
            f'Analyzed Document &nbsp;•&nbsp; <span style="background: {badge_bg}; color: {badge_color}; font-size: 11px; padding: 2px 7px; border-radius: 6px; border: 1px solid {theme["border_color"]};">{badge_txt}</span>',
            f'</div>',
            f'<div style="font-size: 20px; font-weight: 800; color: {theme["text_primary"]}; margin-top: 4px;">📄 {safe_text(clean_display_title)}</div>',
            f'<div style="margin-top: 8px;">',
            f'<span style="font-size: 12px; color: {theme["text_muted"]}; font-weight: 600;">Research Domain:</span>&nbsp;{area_badges_html}',
            f'<span style="font-size: 12px; color: {theme["text_muted"]}; font-weight: 600; margin-left: 12px;">Reading Level:</span>&nbsp;',
            f'<span style="color: {theme["text_secondary"]}; font-size: 13px; font-weight: 600;">{rd_diff_str}</span>',
            f'</div>',
            f'</div>',
            f'<div style="text-align: right;">',
            f'<div style="font-size: 26px; font-weight: 800; color: {theme["accent_purple"]};">{comp:.1f} <span style="font-size: 14px; color: {theme["text_muted"]};">/ 100</span></div>',
            f'<div style="background: {theme["badge_role_bg"]}; color: {theme["badge_role_text"]}; font-size: 13px; font-weight: 700; padding: 2px 10px; border-radius: 10px; border: 1px solid {theme["badge_role_border"]}; display: inline-block; margin-top: 4px;">Grade {safe_text(grade)}</div>',
            f'</div>',
            f'</div>',
            f'<div style="border-top: 1px solid {theme["border_divider"]}; margin-top: 14px; padding-top: 10px; font-size: 13px; color: {theme["text_muted"]};">',
            f'Overall evaluation based on language quality, structural coherence, argumentation, academic style, and readability.',
            f'</div>',
            f'</div>',
        ])
        
        # Executive Summary Quick View / Generation Expander
        with st.expander("📝 Executive Paper Summary (PaperIQ Summarizer)", expanded=False):
            cur_sum = st.session_state.get("summary_result")
            if cur_sum and cur_sum.get("success") and (cur_sum.get("document_info", {}).get("filename") == filename or cur_sum.get("document_info", {}).get("title") == clean_display_title):
                render_html([
                    f'<div style="display: flex; gap: 8px; flex-wrap: wrap; margin-bottom: 10px;">',
                    f'<span style="background: {theme["accent_primary_subtle"]}; border: 1px solid {theme["accent_primary_border"]}; border-radius: 6px; padding: 2px 8px; font-size: 11px; color: {theme["accent_purple"]}; font-weight: 700;">{safe_text(cur_sum["summary_type"])} Summary</span>',
                    f'<span style="background: {theme["success_bg"]}; border: 1px solid {theme["success_border"]}; border-radius: 6px; padding: 2px 8px; font-size: 11px; color: {theme["success_text"]}; font-weight: 600;">{cur_sum["word_count"]} words</span>',
                    f'<span style="background: {theme["info_bg"]}; border: 1px solid {theme["info_border"]}; border-radius: 6px; padding: 2px 8px; font-size: 11px; color: {theme["info_text"]}; font-weight: 600;">~{cur_sum["reading_time_min"]} min read</span>',
                    f'</div>',
                ])
                st.markdown(unescape_text(cur_sum["summary"]))
                c_dl_pdf, c_dl_txt, c_cp = st.columns([1.2, 1.2, 1.2])
                with c_dl_pdf:
                    anl_sum_pdf = get_cached_summary_pdf(cur_sum, paper_title=clean_display_title, metadata={"filename": filename})
                    anl_sum_pdf_name = pdf_export.get_pdf_filename(clean_display_title, f"Summary_{cur_sum['summary_type']}")
                    st.download_button("⬇ Download Summary (PDF)", data=anl_sum_pdf, file_name=anl_sum_pdf_name, mime="application/pdf", use_container_width=True, key="btn_dl_anl_summary_pdf")
                with c_dl_txt:
                    txt_c = summarizer.export_summary_as_txt(cur_sum)
                    txt_n = summarizer.get_summary_export_filename(filename, cur_sum["summary_type"])
                    st.download_button("⬇ Download Summary (TXT)", data=txt_c.encode("utf-8"), file_name=txt_n, mime="text/plain", use_container_width=True, key="btn_dl_anl_summary")
                with c_cp:
                    with st.popover("📋 Copy Summary Text", use_container_width=True):
                        st.caption("Click the copy icon in the top right corner of the box below:")
                        st.code(unescape_text(cur_sum["summary"]), language=None)
            else:
                st.info("No executive summary has been generated for this paper yet.")
                col_sm1, col_sm2 = st.columns([1.5, 3.0])
                with col_sm1:
                    sel_anl_len = st.selectbox("Summary Mode", ["Short", "Medium", "Long"], index=1, key="sb_anl_sum_mode")
                with col_sm2:
                    st.write("")
                    st.write("")
                    if st.button("📝 Generate Executive Summary Now", type="primary", key="btn_anl_run_quick_sum"):
                        with st.spinner("Generating summary..."):
                            full_txt = res.get("full_text", "")
                            if full_txt:
                                gen_sum = summarizer.summarize_research_paper(full_txt, summary_type=sel_anl_len, filename=filename, doc_title=clean_display_title)
                                st.session_state.summary_result = gen_sum
                                st.session_state.summary_doc_text = full_txt
                                st.session_state.summary_doc_name = filename
                                st.session_state.summary_length = sel_anl_len
                                st.rerun()
        
        # B. QUALITY DIMENSIONS (5 Main Cards Only)
        render_html("<h3 style='font-size: 18px; margin: 0 0 12px 0;'>Quality Dimensions</h3>")
        metrics_keys = ["Language Quality", "Structural Coherence", "Argumentation", "Academic Style", "Readability"]
        dim_explanations = res.get("dimension_explanations") or {}
        card_cols = st.columns(5)
        for c, m in zip(card_cols, metrics_keys):
            val = scores.get(m, scores.get(m.split()[0], 0.0))
            rating, color, interp = interpret_score(val, m)
            metric_driver = safe_text(dim_explanations.get(m) or get_dimension_metric_explanation(m, res))
            with c:
                render_html([
                    f'<div class="score-card">',
                    f'<div class="score-card-title">{safe_text(m)}</div>',
                    f'<div class="score-card-val">{val:.1f} <span style="font-size: 13px; color: #64748B;">/ 100</span></div>',
                    f'<div style="font-size: 13px; font-weight: 700; color: {color}; margin: 2px 0;">{safe_text(rating)}</div>',
                    f'<div class="score-card-interp">"{safe_text(interp)}"</div>',
                    f'<div style="font-size: 11px; color: {theme["text_muted"]}; margin-top: 6px; padding-top: 6px; border-top: 1px dashed {theme["border_divider"]}; line-height: 1.4;">{metric_driver}</div>',
                    f'</div>',
                ])
                
        render_html("<br>")
        
        # C & D. KEY FINDINGS & AREAS REQUIRING ATTENTION
        col_findings, col_attention = st.columns(2)
        
        l_val = scores.get("Language Quality", scores.get("Language", 0.0))
        c_val = scores.get("Structural Coherence", scores.get("Coherence", 0.0))
        a_val = scores.get("Argumentation", scores.get("Reasoning", 0.0))
        s_val = scores.get("Academic Style", scores.get("Sophistication", 0.0))
        r_val = scores.get("Readability", 0.0)
        
        with col_findings:
            with st.container(border=True):
                render_html(f"<h4 style='color: {theme['success_text']}; margin: 0 0 12px 0;'>✓ Key Findings</h4>")
                findings_list = []
                if l_val >= 70:
                    findings_list.append("<b>Strong academic vocabulary:</b> Clear sentence structure with formal scholarly cadence.")
                else:
                    findings_list.append("<b>Consistent prose rhythm:</b> Sentences maintain an understandable scholarly cadence.")
                    
                if c_val >= 65:
                    findings_list.append("<b>Consistent structural organization:</b> Logical transitions between sections and topic sentences.")
                else:
                    findings_list.append("<b>Structured thematic blocks:</b> Paper maintains recognizable topical organization.")
                    
                if a_val >= 65:
                    findings_list.append("<b>Evidence-supported arguments:</b> Causal propositions are linked to empirical observations.")
                if s_val >= 65:
                    findings_list.append("<b>Lexical sophistication:</b> Effective use of domain-specific scholarly nomenclature.")
                if r_val >= 60:
                    findings_list.append("<b>Accessible academic readability:</b> Phrasing is well-calibrated for scholarly peer review.")
                    
                for item in findings_list[:4]:
                    render_html(f"<p style='font-size: 13px; margin: 0 0 8px 0; color: {theme['text_secondary']};'>✓ {item}</p>")
            
        with col_attention:
            with st.container(border=True):
                render_html(f"<h4 style='color: {theme['error_text']}; margin: 0 0 12px 0;'>⚠️ Areas Requiring Attention</h4>")
                issues_found = res.get("issues", [])
                if issues_found:
                    for idx, iss in enumerate(issues_found[:3]):
                        category = safe_text(iss.get("category", "General"))
                        prob_text = safe_text(iss.get("problem", iss.get("issue", "Issue identified")))
                        why_text = safe_text(iss.get("why_it_matters", "Essential for peer-review clarity and scholarly rigor."))
                        act_text = safe_text(iss.get("recommendation", "Review and revise phrasing."))
                        border_div = f"<div style='border-bottom: 1px solid {theme['border_divider']}; margin: 10px 0;'></div>" if idx < len(issues_found[:3]) - 1 else ""
                        render_html([
                            f"<div style='margin-bottom: 4px;'>",
                            f"<div style='font-size: 11px; font-weight: 700; color: {theme['warning_text']}; text-transform: uppercase; letter-spacing: 0.05em;'>{category}</div>",
                            f"<div style='font-size: 13px; font-weight: 700; color: {theme['text_primary']}; margin: 2px 0 4px 0;'>{prob_text}</div>",
                            f"<div style='font-size: 12px; color: {theme['text_muted']}; margin-bottom: 3px;'><b style='color: {theme['text_secondary']};'>Why it matters:</b> {why_text}</div>",
                            f"<div style='font-size: 12px; color: {theme['text_secondary']};'><b style='color: {theme['accent_purple']};'>Action:</b> {act_text}</div>",
                            f"</div>",
                            border_div,
                        ])
                else:
                    render_html(f"<p style='font-size: 13px; color: {theme['success_text']}; margin: 0;'>✓ No critical structural, citation, or readability issues detected in the manuscript.</p>")
            
        render_html("<br>")
        
        # E. VISUAL ANALYSIS (Radar + Score Breakdown)
        render_html("<h3 style='font-size: 18px; margin: 0 0 12px 0;'>Visual Analysis</h3>")
        col_c1, col_c2 = st.columns(2)
        chart_vals = [scores.get(k, scores.get(k.split()[0], 0)) for k in metrics_keys]
        
        with col_c1:
            render_html([
                f'<div style="margin-bottom: 8px;">',
                f'<h4 style="color: {theme["text_primary"]}; margin: 0 0 4px 0;">🎯 Academic Performance Radar</h4>',
                f'<p style="color: {theme["text_muted"]}; font-size: 12px; margin: 0;">Relative score distribution across the 5 evaluation dimensions.</p>',
                f'</div>',
            ])
            ctheme = get_chart_theme(st.session_state.get("theme_preference", "System"))
            fig_radar = go.Figure()
            fig_radar.add_trace(go.Scatterpolar(
                r=chart_vals + [chart_vals[0]],
                theta=metrics_keys + [metrics_keys[0]],
                fill='toself',
                fillcolor=ctheme["radar_fill"],
                line=dict(color=ctheme["radar_line"], width=2.5),
                hoverinfo="theta+r"
            ))
            fig_radar.update_layout(
                polar=dict(
                    bgcolor=ctheme["polar_bg"],
                    radialaxis=dict(visible=True, range=[0, 100], gridcolor=ctheme["grid_color"], tickfont=dict(color=ctheme["tick_color"], size=10)),
                    angularaxis=dict(tickfont=dict(color=ctheme["font_color"], size=11), linecolor=ctheme["line_color"])
                ),
                paper_bgcolor=ctheme["paper_bgcolor"],
                plot_bgcolor=ctheme["plot_bgcolor"],
                font=dict(color=ctheme["font_color"]),
                margin=dict(t=35, b=35, l=45, r=45),
                height=360,
                showlegend=False
            )
            st.plotly_chart(fig_radar, use_container_width=True)
            
        with col_c2:
            render_html([
                f'<div style="margin-bottom: 8px;">',
                f'<h4 style="color: {theme["text_primary"]}; margin: 0 0 4px 0;">📊 Score Breakdown</h4>',
                f'<p style="color: {theme["text_muted"]}; font-size: 12px; margin: 0;">Component dimension scores compared against composite score.</p>',
                f'</div>',
            ])
            bar_labels = ["Language<br>Quality", "Structural<br>Coherence", "Argumentation", "Academic<br>Style", "Readability", "Overall<br>Score"]
            bar_vals = chart_vals + [comp]
            fig_bar = go.Figure(data=[
                go.Bar(
                    x=bar_labels,
                    y=bar_vals,
                    marker=dict(color=["#8B5CF6", "#6366F1", "#3B82F6", "#06B6D4", "#10B981", "#A855F7"]),
                    text=[f"{v:.1f}" for v in bar_vals],
                    textposition="auto",
                    textfont=dict(color="#FFFFFF", size=11)
                )
            ])
            fig_bar.update_layout(
                paper_bgcolor=ctheme["paper_bgcolor"],
                plot_bgcolor=ctheme["plot_bgcolor"],
                font=dict(color=ctheme["font_color"]),
                yaxis=dict(range=[0, 105], gridcolor=ctheme["grid_color"], linecolor=ctheme["line_color"], title="Score (0-100)"),
                xaxis=dict(linecolor=ctheme["line_color"], tickfont=dict(color=ctheme["tick_color"], size=11)),
                margin=dict(t=35, b=35, l=35, r=35),
                height=360
            )
            st.plotly_chart(fig_bar, use_container_width=True)
            
        render_html("<br>")
        
        # F. DOCUMENT STATISTICS (Collapsible / Expandable Section)
        canonical_counts = compute_text_counts(res.get("full_text", "")) if res.get("full_text") else None
        stat_words = canonical_counts["word_count"] if canonical_counts and canonical_counts["word_count"] > 0 else stats.get("word_count", 0)
        stat_sents = canonical_counts["sentence_count"] if canonical_counts and canonical_counts["sentence_count"] > 0 else stats.get("sentence_count", 0)
        with st.expander("📊 Document Statistics", expanded=False):
            dm1, dm2, dm3, dm4, dm5, dm6 = st.columns(6)
            dm1.metric("Total Words", f"{stat_words:,}")
            dm2.metric("Total Sentences", f"{stat_sents:,}")
            dm3.metric("Avg Sentence Len", f"{stats.get('avg_sentence_len', 0):.1f}")
            dm4.metric("Avg Word Len", f"{stats.get('avg_word_len', 0):.1f}")
            dm5.metric("Vocab Diversity", f"{stats.get('vocab_diversity', 0):.2f}")
            dm6.metric("Complex Word Ratio", f"{stats.get('complex_word_ratio', 0):.2f}")
            
        # G. CLEAN SECTION SUMMARIES (Only if Reliable Academic Sections Exist)
        reliable_secs = get_reliable_academic_sections(res.get("sections", {}))
        if reliable_secs:
            render_html("<br>")
            st.subheader("Section Summaries")
            for sec_name, sec_content in reliable_secs.items():
                if sec_name == "References":
                    ref_entries = extract_reference_entries(sec_content)
                    with st.expander(f"📚 References ({len(ref_entries)} entries)", expanded=False):
                        render_html(f"<p style='font-size: 13px; font-weight: 700; color: {theme['text_primary']}; margin-bottom: 6px;'>Extracted Reference Entries ({len(ref_entries)} detected):</p>")
                        for ref_item in ref_entries:
                            st.markdown(f"- {safe_text(ref_item)}")
                else:
                    with st.expander(f"📑 {sec_name}", expanded=True if sec_name in ["Abstract", "Introduction"] else False):
                        summary_text = summarize_text(sec_content, 3)
                        key_points = get_important_sentences(summary_text, 2)
                        render_html(f"<div style='font-size: 13.5px; color: {theme['text_secondary']}; line-height: 1.6;'><b>Executive Summary:</b><br>{safe_text(summary_text)}</div>")
                        if key_points:
                            render_html(f"<div style='margin-top: 8px; font-size: 13px; font-weight: 700; color: {theme['text_primary']};'>Key Observations:</div>")
                            for kp in key_points:
                                st.markdown(f"• {unescape_text(kp)}")

            missing_std_secs = get_missing_standard_sections(reliable_secs)
            if missing_std_secs:
                missing_notes_html = " &nbsp;&bull;&nbsp; ".join(
                    [f"<b>{safe_text(ms)}:</b> <i>Not detected in this manuscript</i>" for ms in missing_std_secs]
                )
                render_html(
                    f"<div style='background: {theme['bg_card_inner']}; border: 1px solid {theme['border_color']}; border-radius: 8px; padding: 10px 14px; margin-top: 8px; font-size: 12.5px; color: {theme['text_muted']};'>"
                    f"ℹ️ <b>Standard Section Status:</b> {missing_notes_html}"
                    f"</div>"
                )

# ----------------------------------------------------
# PAGE 4: INSIGHTS & IMPROVEMENTS
# ----------------------------------------------------
elif st.session_state.current_page in ["💡 Insights", "💡 Research Insights"]:
    ins_heading = "Research Insights" if is_researcher else "Research Insights & Writing Improvements"
    render_html([
        f'<div style="margin-bottom: 20px;">',
        f'<h1 style="margin: 0; font-size: 28px; color: {theme["text_primary"]};">{safe_text(ins_heading)}</h1>',
        f'<p style="margin: 4px 0 0 0; color: {theme["text_muted"]}; font-size: 15px;">Actionable evaluation of research opportunities, contributions, and structural writing refinements.</p>',
        f'</div>',
    ])
    
    if not has_any_valid_analysis():
        ins_empty_title = "No Manuscript Evaluated Yet" if is_researcher else "No Insights Generated Yet"
        ins_empty_text = "Upload and evaluate an academic manuscript in <b>Analyze Manuscript</b> to generate research insights and evidence extraction." if is_researcher else "Analyze a research paper in <b>Analyze Paper</b> to generate scholarly insights and recommendations."
        ins_btn_label = "Go to Analyze Manuscript" if is_researcher else "Go to Analyze Paper"
        ins_btn_target = "📄 Analyze Manuscript" if is_researcher else "📄 Analyze Paper"
        render_html([
            f'<div class="iq-card" style="text-align: center; padding: 40px !important;">',
            f'<div style="font-size: 32px; margin-bottom: 8px;">💡</div>',
            f'<h4 style="color: {theme["text_primary"]}; margin: 0;">{safe_text(ins_empty_title)}</h4>',
            f'<p style="color: {theme["text_muted"]}; margin-top: 4px;">{ins_empty_text}</p>',
            f'</div>',
        ])
        btn_goto_ins = st.button(ins_btn_label, type="primary", key="btn_insights_goto_analyze", on_click=navigate_to, args=(ins_btn_target,))
        if btn_goto_ins:
            navigate_to(ins_btn_target)
            st.rerun()
    else:
        res = st.session_state.get("results") or st.session_state.get("active_analysis")
        full_text = res.get("full_text", "")
        scores = res.get("scores", {})
        stats = res.get("stats", {})
        filename_ins = st.session_state.get("filename", "Paper")
        title_ins = st.session_state.get("paper_title") or get_clean_display_title(filename_ins, text=full_text, stored_title=res.get("title"))
        fb_pdf_bytes_ins = get_cached_analysis_pdf(res, filename_ins, st.session_state.get("current_analysis_id"))
        fb_pdf_fn_ins = pdf_export.get_pdf_filename(title_ins, "Feedback")

        c_ins_hdr1, c_ins_hdr2 = st.columns([4, 1.8])
        with c_ins_hdr1:
            st.caption(f"Grounded insights and writing feedback for: **{title_ins}**")
        with c_ins_hdr2:
            st.download_button(
                "⬇ Download Feedback (PDF)",
                data=fb_pdf_bytes_ins,
                file_name=fb_pdf_fn_ins,
                mime="application/pdf",
                use_container_width=True,
                key="btn_insights_download_feedback_pdf"
            )
        
        # 1. Potential Research Areas & Domain Context
        areas_list = res.get("potential_research_areas") or res.get("domains") or [res.get("domain", "General Academic")]
        areas_badges_str = "".join([f"<span class='domain-badge'>{safe_text(d)}</span>" for d in areas_list])
        render_html([
            f'<div class="journal-hero-card">',
            f'<h4>📚 Potential Research Areas</h4>',
            f'<div style="margin: 8px 0 10px 0;">{areas_badges_str}</div>',
            f'<div style="font-size: 12px; color: {theme["text_muted"]}; border-top: 1px solid {theme["border_divider"]}; padding-top: 8px;">',
            f'🛡️ <b>Disclaimer:</b> Inferred based on document terminology and computational keyword modeling. Not an official journal endorsement or indexing guarantee.',
            f'</div>',
            f'</div>',
        ])
        
        # RESEARCH INSIGHTS (Gaps & Contributions)
        st.subheader("Research Insights")
        col_gaps, col_contrib = st.columns(2)
        
        with col_gaps:
            with st.container(border=True):
                render_html(f"<h4 style='color: {theme['info_text']}; margin: 0 0 12px 0;'>🔍 Research Gaps & Opportunities</h4>")
                structured_gaps = res.get("structured_gaps") or extract_structured_research_gaps(full_text)
                if structured_gaps:
                    for g in structured_gaps:
                        conf_color = theme['success_text'] if g.get("confidence") == "High" else theme['warning_text']
                        conf_bg = theme['success_bg'] if g.get("confidence") == "High" else theme['warning_bg']
                        conf_border = theme['success_border'] if g.get("confidence") == "High" else theme['warning_border']
                        conf_label = safe_text(g.get("confidence", "Moderate"))
                        lim_label = safe_text(g.get("limitation", "Identified Constraint"))
                        ev_text = safe_text(g.get("evidence", ""))
                        opp_text = safe_text(g.get("opportunity", ""))
                        render_html([
                            f"<div style='background: {theme['bg_card_inner']}; border: 1px solid {theme['border_color']}; border-radius: 8px; padding: 12px 14px; margin-bottom: 10px; word-break: break-word;'>",
                            f"<div style='display: flex; justify-content: space-between; align-items: center; margin-bottom: 6px;'>",
                            f"<span style='font-size: 11px; font-weight: 700; color: {theme['info_text']}; text-transform: uppercase;'>{lim_label}</span>",
                            f"<span style='font-size: 10px; font-weight: 700; color: {conf_color}; background: {conf_bg}; padding: 2px 7px; border-radius: 4px; border: 1px solid {conf_border};'>Confidence: {conf_label}</span>",
                            f"</div>",
                            f"<div style='color: {theme['text_secondary']}; font-style: italic; font-size: 12px; margin: 6px 0; line-height: 1.5;'>&ldquo;{ev_text}&rdquo;</div>",
                            f"<div style='color: {theme['text_muted']}; font-size: 12px; margin-top: 6px;'><b style='color: {theme['accent_purple']};'>Research Opportunity:</b> {opp_text}</div>",
                            f"</div>",
                        ])
                else:
                    render_html(f"<div style='color: {theme['text_muted']}; font-size: 13px; font-style: italic; padding: 8px 0;'>A specific research gap could not be established from the available document content. The manuscript does not contain explicit limitation statements or unresolved inquiry markers.</div>")
            
        with col_contrib:
            with st.container(border=True):
                render_html(f"<h4 style='color: {theme['success_text']}; margin: 0 0 12px 0;'>🎯 Key Contributions</h4>")
                structured_contribs = res.get("structured_contributions") or extract_key_contributions_structured(full_text)
                if structured_contribs:
                    for c in structured_contribs:
                        is_author = c.get("type") == "Author-Stated Contribution"
                        badge_color = theme['success_text'] if is_author else theme['info_text']
                        type_label = safe_text(c.get("type", "Key Contribution"))
                        conf_label = safe_text(c.get("confidence", "Moderate"))
                        c_text = safe_text(c.get("text", ""))
                        render_html([
                            f"<div style='background: {theme['bg_card_inner']}; border: 1px solid {theme['border_color']}; border-radius: 8px; padding: 12px 14px; margin-bottom: 10px; word-break: break-word;'>",
                            f"<div style='display: flex; justify-content: space-between; align-items: center; margin-bottom: 6px;'>",
                            f"<span style='font-size: 11px; font-weight: 700; color: {badge_color}; text-transform: uppercase;'>{type_label}</span>",
                            f"<span style='font-size: 10px; font-weight: 700; color: {theme['text_secondary']}; background: {theme.get('input_bg', theme['bg_card_inner'])}; padding: 2px 7px; border-radius: 4px; border: 1px solid {theme['border_color']};'>Confidence: {conf_label}</span>",
                            f"</div>",
                            f"<div style='color: {theme['text_primary']}; font-size: 13px; margin-top: 4px; line-height: 1.5;'>&ldquo;{c_text}&rdquo;</div>",
                            f"</div>",
                        ])
                else:
                    render_html(f"<div style='color: {theme['text_muted']}; font-size: 13px; font-style: italic; padding: 8px 0;'>No explicit author contribution statements could be established from the analyzed manuscript text.</div>")
            
        # METHODOLOGY, DATA & EMPIRICAL FINDINGS
        render_html("<br>")
        st.subheader("Methodology, Data & Empirical Findings")
        meth_findings = res.get("methodology_and_findings") or extract_methodology_and_findings(full_text, res.get("sections", {}))
        c_m1, c_m2, c_m3 = st.columns(3)
        with c_m1:
            with st.container(border=True):
                render_html(f"<h4 style='color: {theme['info_text']}; margin: 0 0 10px 0;'>🔬 Methodology & Study Design</h4>")
                m_txt = safe_text(meth_findings.get('methodology', 'Not explicitly reported in the analyzed text.'))
                render_html(f"<p style='font-size: 13px; color: {theme['text_secondary']}; line-height: 1.6; margin: 0;'>{m_txt}</p>")
        with c_m2:
            with st.container(border=True):
                render_html(f"<h4 style='color: {theme['warning_text']}; margin: 0 0 10px 0;'>📦 Data Sources & Materials</h4>")
                d_txt = safe_text(meth_findings.get('data_and_materials', 'Not explicitly reported in the analyzed text.'))
                render_html(f"<p style='font-size: 13px; color: {theme['text_secondary']}; line-height: 1.6; margin: 0;'>{d_txt}</p>")
        with c_m3:
            with st.container(border=True):
                render_html(f"<h4 style='color: {theme['success_text']}; margin: 0 0 10px 0;'>📊 Empirical Findings & Outcomes</h4>")
                f_txt = safe_text(meth_findings.get('findings', 'Not explicitly reported in the analyzed text.'))
                render_html(f"<p style='font-size: 13px; color: {theme['text_secondary']}; line-height: 1.6; margin: 0;'>{f_txt}</p>")
            
        # FUTURE RESEARCH DIRECTIONS
        future_dirs = res.get("future_directions") or extract_future_directions(full_text)
        has_author_stated = any("not enough evidence" not in fd.lower() for fd in future_dirs)
        
        with st.container(border=True):
            render_html(f"<h4 style='color: {theme['accent_purple']}; margin: 0 0 12px 0;'>🔭 Future Research Directions</h4>")
            if has_author_stated:
                for fd in future_dirs:
                    if "not enough evidence" not in fd.lower():
                        fd_clean = safe_text(fd)
                        render_html([
                            f"<div style='background: {theme['bg_card_inner']}; border: 1px solid {theme['border_color']}; border-radius: 8px; padding: 12px 14px; margin-bottom: 8px; word-break: break-word;'>",
                            f"<span style='font-size: 11px; font-weight: 700; color: {theme['accent_purple']}; text-transform: uppercase;'>Author-Stated Future Work:</span>",
                            f"<div style='color: {theme['text_primary']}; font-size: 13px; margin-top: 4px; line-height: 1.5;'>&ldquo;{fd_clean}&rdquo;</div>",
                            f"</div>",
                        ])
            else:
                render_html([
                    f"<div style='background: {theme['bg_card_inner']}; border: 1px solid {theme['border_color']}; border-radius: 8px; padding: 12px 14px; margin-bottom: 8px;'>",
                    f"<span style='font-size: 11px; font-weight: 700; color: {theme['text_muted']}; text-transform: uppercase;'>Analytical Suggestion (Derived from evaluation, not claimed by authors):</span>",
                    f"<div style='color: {theme['text_secondary']}; font-size: 13px; margin-top: 4px; line-height: 1.5;'>Conduct external validation experiments and benchmark against standardized baselines to establish broader domain generalizability.</div>",
                    f"</div>",
                ])
            
        render_html("<br>")
        
        # WRITING IMPROVEMENTS (Readability, Structural, Argumentation)
        st.subheader("Writing Improvements")
        
        # Readability Improvements
        long_sentences = extract_overly_long_sentences(full_text, max_words=35, limit=2)
        with st.container(border=True):
            render_html(f"<h4 style='color: {theme['text_primary']}; margin: 0 0 12px 0;'>📖 Readability Improvements</h4>")
            if long_sentences:
                for s_text, word_cnt in long_sentences:
                    clean_sample = safe_text(s_text)
                    render_html([
                        f"<div style='background: {theme['bg_card_inner']}; border: 1px solid {theme['border_color']}; border-radius: 8px; padding: 12px 14px; margin-bottom: 10px; word-break: break-word;'>",
                        f"<div style='color: {theme['warning_text']}; font-weight: 700; font-size: 12px; text-transform: uppercase; letter-spacing: 0.05em; margin-bottom: 4px;'>Complex Sentence Identified ({word_cnt} words):</div>",
                        f"<div style='color: {theme['text_secondary']}; font-style: italic; font-size: 13px; margin: 6px 0; line-height: 1.5;'>&ldquo;{clean_sample}&rdquo;</div>",
                        f"<div style='color: {theme['text_muted']}; font-size: 12px; margin-top: 6px;'><b style='color: {theme['accent_purple']};'>Recommendation:</b> Break this compound sentence into two concise statements to improve readability and reviewer comprehension.</div>",
                        f"</div>",
                    ])
            else:
                avg_len = stats.get('avg_sentence_len', 22) if isinstance(stats, dict) else 22
                render_html(f"<p style='color: {theme['success_text']}; font-size: 13px; margin: 0;'>✓ Sentence lengths are well-calibrated (average: {avg_len:.1f} words). Phrasing remains accessible for scholarly evaluation.</p>")
        
        # Structural & Argumentation Improvements
        col_struct, col_arg = st.columns(2)
        with col_struct:
            c_score = scores.get("Structural Coherence", scores.get("Coherence", 0.0))
            with st.container(border=True):
                render_html(f"<h4 style='color: {theme['text_primary']}; margin: 0 0 6px 0;'>🏗️ Structural Improvements</h4>")
                render_html(f"<p style='font-size: 13px; color: {theme['text_secondary']}; margin-bottom: 8px;'>Coherence Index: <b>{c_score:.1f} / 100</b></p>")
                rec_struct = 'Incorporate directional discourse markers (e.g., "Consequently", "In contrast", "Furthermore") between subsection transitions to improve thematic flow.' if c_score < 70 else 'Discourse connectors are consistently deployed across major argument blocks; ensure paragraph openers state clear thematic topic sentences.'
                render_html(f"<p style='font-size: 13px; color: {theme['text_muted']}; margin: 0;'>{safe_text(rec_struct)}</p>")
            
        with col_arg:
            a_score = scores.get("Argumentation", scores.get("Reasoning", 0.0))
            with st.container(border=True):
                render_html(f"<h4 style='color: {theme['text_primary']}; margin: 0 0 6px 0;'>⚖️ Argumentation Improvements</h4>")
                render_html(f"<p style='font-size: 13px; color: {theme['text_secondary']}; margin-bottom: 8px;'>Argumentation Index: <b>{a_score:.1f} / 100</b></p>")
                rec_arg = 'Directly couple causal propositions with explicit empirical evidence indicators (e.g., "demonstrates that", "substantiates", "implies").' if a_score < 70 else 'Causal assertions are well-grounded by supporting evidence and empirical outcomes throughout findings.'
                render_html(f"<p style='font-size: 13px; color: {theme['text_muted']}; margin: 0;'>{safe_text(rec_arg)}</p>")
            
        render_html("<br>")
        
        # RECOMMENDED ACTIONS (Clear, Practical Next Steps)
        st.subheader("Recommended Next Actions")
        actions = []
        if any(iss.get("category") == "Structure" for iss in res.get("issues", [])):
            actions.append("<b>Clarify Methodology:</b> Add or expand a dedicated Methodology section outlining research design, dataset specifications, and evaluation procedure.")
        if any(iss.get("category") == "Citations" for iss in res.get("issues", [])):
            actions.append("<b>Improve Citation Coverage:</b> Anchor empirical findings and related work within authoritative peer-reviewed literature.")
        if long_sentences:
            actions.append("<b>Split Overly Long Sentences:</b> Segment compound sentences exceeding 35 words into concise propositions to improve readability.")
        if c_score < 70:
            actions.append("<b>Improve Structural Transitions:</b> Use directional discourse markers between consecutive sections to guide reader flow.")
        if a_score < 70:
            actions.append("<b>Strengthen Empirical Evidence:</b> Ensure all primary claims link directly to experimental data, benchmark metrics, or baseline comparisons.")
        if not actions:
            actions.append("<b>Camera-Ready Polish:</b> Polish abstract and conclusion sections to highlight principal empirical findings and future research directions.")
            actions.append("<b>Consistency Check:</b> Verify consistent notation and formatting across all equations, figures, and table captions.")
            
        with st.container(border=True):
            items_html = "".join([f"<li style='margin-bottom: 6px;'>{act}</li>" for act in actions[:4]])
            render_html(f"<ol style='color: {theme['text_secondary']}; font-size: 14px; padding-left: 20px; margin: 0; line-height: 1.8;'>{items_html}</ol>")

# ----------------------------------------------------
# PAGE 5: ASK PAPERIQ (DOCUMENT Q&A)
# ----------------------------------------------------
elif st.session_state.current_page == "💬 Ask PaperIQ":
    ask_desc_str = "Ask questions about your research manuscript. PaperIQ answers questions grounded exclusively in the currently analyzed document." if is_researcher else "Ask questions about your research paper. PaperIQ answers questions grounded exclusively in the currently analyzed document."
    render_html([
        f'<div style="margin-bottom: 20px;">',
        f'<h1 style="margin: 0; font-size: 28px; color: {theme["text_primary"]};">Ask PaperIQ</h1>',
        f'<p style="margin: 4px 0 0 0; color: {theme["text_muted"]}; font-size: 15px;">{safe_text(ask_desc_str)}</p>',
        f'</div>',
    ])
    
    if not has_any_valid_analysis():
        ask_empty_title = "Analyze a Manuscript to Start Q&A" if is_researcher else "Analyze a Paper to Start Q&A"
        ask_empty_text = "Upload and evaluate an academic manuscript in <b>Analyze Manuscript</b> first to chat with its contents." if is_researcher else "Upload and analyze a research paper first to chat with its contents."
        ask_btn_label = "Go to Analyze Manuscript" if is_researcher else "Go to Analyze Paper"
        ask_btn_target = "📄 Analyze Manuscript" if is_researcher else "📄 Analyze Paper"
        render_html([
            f'<div class="iq-card" style="text-align: center; padding: 40px !important;">',
            f'<div style="font-size: 32px; margin-bottom: 8px;">💬</div>',
            f'<h4 style="color: {theme["text_primary"]}; margin: 0;">{safe_text(ask_empty_title)}</h4>',
            f'<p style="color: {theme["text_muted"]}; margin-top: 4px;">{ask_empty_text}</p>',
            f'</div>',
        ])
        btn_goto_ask = st.button(ask_btn_label, type="primary", key="btn_ask_goto_analyze", on_click=navigate_to, args=(ask_btn_target,))
        if btn_goto_ask:
            navigate_to(ask_btn_target)
            st.rerun()
    else:
        res = st.session_state.get("results") or st.session_state.get("active_analysis")
        # Auto-clear chat history if switching to a different manuscript
        active_doc_id = str(st.session_state.get("filename") or st.session_state.get("current_analysis_id") or "")
        if st.session_state.get("last_chat_doc_id") != active_doc_id:
            st.session_state["chat_history"] = []
            st.session_state["last_chat_doc_id"] = active_doc_id

        full_text = res.get("full_text", "")
        sections = res.get("sections", {})
        active_doc = st.session_state.get("paper_title") or get_clean_display_title(st.session_state.get("filename", "Current Document"), stored_title=res.get("title"))
        domain_name = res.get("domain", "General Academic")
        overall_score = get_overall_score(res)
        doc_badge_title = "Current Manuscript" if is_researcher else "Current Paper"
        
        col_ctx1, col_ctx_dl, col_ctx2 = st.columns([3.8, 1.4, 1.1])
        with col_ctx1:
            render_html([
                f'<div class="iq-card" style="padding: 14px 18px !important; margin-bottom: 16px !important;">',
                f'<div style="font-size: 11px; text-transform: uppercase; color: {theme["text_muted"]}; font-weight: 700; letter-spacing: 0.05em;">{safe_text(doc_badge_title)}</div>',
                f'<div style="font-size: 17px; font-weight: 700; color: {theme["text_primary"]}; margin-top: 2px;">📄 {safe_text(active_doc)}</div>',
                f'<div style="font-size: 12px; color: {theme["text_muted"]}; margin-top: 4px;">',
                f'Domain: <b style="color: {theme["info_text"]};">{safe_text(domain_name)}</b> &nbsp;•&nbsp; ',
                f'Score: <b style="color: {theme["accent_purple"]};">{overall_score:.1f} / 100</b> &nbsp;•&nbsp; ',
                f'<span style="color: {theme["success_text"]}; font-weight: 600;">● Document loaded</span>',
                f'</div>',
                f'</div>',
            ])
        with col_ctx_dl:
            render_html("<div style='height: 18px;'></div>")
            chat_pdf_bytes = get_cached_chat_pdf(active_doc, st.session_state.get("chat_history", []))
            chat_pdf_fn = pdf_export.get_pdf_filename(active_doc, "Chat")
            st.download_button(
                "⬇ Download Chat (PDF)",
                data=chat_pdf_bytes,
                file_name=chat_pdf_fn,
                mime="application/pdf",
                key="btn_download_chat_pdf",
                use_container_width=True
            )
        with col_ctx2:
            render_html("<div style='height: 18px;'></div>")
            if st.button("🗑️ Clear Chat", key="btn_clear_chat", use_container_width=True):
                st.session_state.chat_history = []
                st.rerun()
        
        # 11 Standardized Suggested Questions
        render_html(f"<p style='color: {theme['text_muted']}; font-size: 13px; margin-bottom: 8px; font-weight: 600;'>Suggested Questions:</p>")
        suggested = [
            "What is the main objective of this paper?",
            "What methodology was used?",
            "What are the key findings?",
            "What are the limitations?",
            "What research gap does this paper address?",
            "What datasets or materials were used?",
            "How is this work evaluated?",
            "What are the core contributions?",
            "What future work is suggested?",
            "What theoretical framework is applied?",
            "Summarize the conclusion."
        ]
        
        sug_cols = st.columns(3)
        for i, q_text in enumerate(suggested):
            target_col = sug_cols[i % 3]
            with target_col:
                if st.button(q_text, key=f"sug_btn_{i}", use_container_width=True):
                    ans_text, sources = grounded_paper_qa(q_text, full_text, sections)
                    st.session_state.chat_history.append({"role": "user", "content": q_text})
                    st.session_state.chat_history.append({"role": "assistant", "answer": ans_text, "evidence": sources})
                    st.rerun()
                    
        render_html(f"<hr style='border: none; border-top: 1px solid {theme['border_divider']}; margin: 16px 0 14px 0;'>")
        
        # Render Chat History
        if not st.session_state.chat_history:
            render_html(f"<div style='color: {theme['text_muted']}; font-size: 13px; font-style: italic; margin-bottom: 12px;'>No questions asked yet. Click a suggested question above or enter a question below.</div>")
        else:
            for msg in st.session_state.chat_history:
                if msg.get("role") == "user":
                    user_c = safe_text(msg.get('content', ''))
                    render_html(f"<div class='chat-msg-user'><b>👤 You:</b><div style='margin-top: 4px; color: {theme['text_primary']};'>{user_c}</div></div>")
                else:
                    raw_ans_str = str(msg.get("answer") or msg.get("content", ""))
                    raw_ans_clean = re.split(r"\n\nSupporting evidence \(from ", raw_ans_str, maxsplit=1)[0].strip()
                    ans_body = safe_text(raw_ans_clean)
                    ev_items = msg.get("evidence", [])
                    
                    evidence_html = ""
                    if ev_items and not raw_ans_clean.startswith(("The uploaded paper does not provide", "I could not find this information")):
                        ev_cards = []
                        for ev in ev_items:
                            sec_lbl = safe_text(ev.get("section", "Document Excerpt"))
                            exc_txt = safe_text(str(ev.get("excerpt", "")).strip())
                            if exc_txt:
                                ev_cards.append(
                                    f"<div style='background: {theme['bg_card_inner']}; border-left: 3px solid {theme['accent_purple']}; border-radius: 6px; padding: 10px 14px; margin-top: 10px;'>"
                                    f"<div style='font-size: 11.5px; font-weight: 700; color: {theme['accent_purple']}; letter-spacing: 0.04em;'>📌 Supporting evidence (from {sec_lbl})</div>"
                                    f"<div style='font-style: italic; font-size: 13px; color: {theme['text_secondary']}; margin-top: 5px; line-height: 1.5;'>\"{exc_txt}\"</div>"
                                    f"</div>"
                                )
                        evidence_html = "".join(ev_cards)
                        
                    render_html([
                        f"<div class='chat-msg-bot'>",
                        f"<div style='font-weight: 700; color: {theme['accent_purple']}; font-size: 14px; margin-bottom: 6px;'>🤖 PaperIQ</div>",
                        f"<div style='color: {theme['text_primary']}; font-size: 14px; line-height: 1.6; margin-bottom: 6px; word-break: break-word;'>{ans_body}</div>",
                        evidence_html,
                        f"</div>",
                    ])
                    
        # Compact Inline Chat Input (Eliminates huge blank bottom gap!)
        with st.form(key="ask_form", clear_on_submit=True):
            col_inp, col_snd = st.columns([5, 1])
            with col_inp:
                user_query = st.text_input("Ask about this paper...", placeholder="Ask a question about the analyzed paper...", label_visibility="collapsed")
            with col_snd:
                send_pressed = st.form_submit_button("Send", type="primary", use_container_width=True)
                
            if send_pressed and user_query and user_query.strip():
                ans_text, sources = grounded_paper_qa(user_query.strip(), full_text, sections)
                st.session_state.chat_history.append({"role": "user", "content": user_query.strip()})
                st.session_state.chat_history.append({"role": "assistant", "answer": ans_text, "evidence": sources})
                st.rerun()

# ----------------------------------------------------
# PAGE 6: PAPER COMPARISON
# ----------------------------------------------------
elif st.session_state.current_page in ["🔄 Compare Papers", "🔄 Compare Manuscripts"]:
    cmp_heading = "Compare Manuscripts" if is_researcher else "Compare Research Papers"
    cmp_desc = "Compare research manuscripts across writing quality, structure, argumentation, and readability." if is_researcher else "Compare research papers across writing quality, structure, argumentation, and readability."
    render_html([
        f'<div style="margin-bottom: 20px;">',
        f'<h1 style="margin: 0; font-size: 28px; color: {theme["text_primary"]};">{safe_text(cmp_heading)}</h1>',
        f'<p style="margin: 4px 0 0 0; color: {theme["text_muted"]}; font-size: 15px;">{safe_text(cmp_desc)}</p>',
        f'</div>',
    ])
    
    comp_card_title = "Upload Manuscripts for Comparison" if is_researcher else "Upload Papers for Comparison"
    comp_card_desc = "Upload 2 or 3 research manuscripts (PDF, DOCX, TXT, max 30 MB each) to evaluate writing clarity and structural consistency." if is_researcher else "Upload 2 or 3 research papers (PDF, DOCX, TXT, max 30 MB each) to evaluate writing clarity and structural consistency."
    render_html([
        f'<div class="iq-card">',
        f'<h4 style="margin: 0 0 8px 0; color: {theme["text_primary"]};">{safe_text(comp_card_title)}</h4>',
        f'<p style="color: {theme["text_muted"]}; font-size: 13px; margin: 0 0 14px 0;">{safe_text(comp_card_desc)}</p>',
        f'</div>',
    ])
    
    c_up1, c_up2, c_up3 = st.columns(3)
    with c_up1:
        f1_label = "Manuscript 1" if is_researcher else "Paper 1"
        f1 = st.file_uploader(f1_label, type=["pdf", "docx", "txt"], key="comp_file_1")
    with c_up2:
        f2_label = "Manuscript 2" if is_researcher else "Paper 2"
        f2 = st.file_uploader(f2_label, type=["pdf", "docx", "txt"], key="comp_file_2")
    with c_up3:
        f3_label = "Manuscript 3 (Optional)" if is_researcher else "Paper 3 (Optional)"
        f3 = st.file_uploader(f3_label, type=["pdf", "docx", "txt"], key="comp_file_3")
        
    run_cmp_label = "🚀 Compare Manuscripts" if is_researcher else "🚀 Compare Papers"
    if st.button(run_cmp_label, type="primary", use_container_width=True, key="btn_run_compare"):
        uploaded_comp = [f for f in [f1, f2, f3] if f is not None]
        if len(uploaded_comp) < 2:
            st.warning("Please upload at least 2 manuscripts to perform comparison." if is_researcher else "Please upload at least 2 papers to perform comparison.")
        else:
            with st.spinner("Analyzing and comparing papers..."):
                compared_list = []
                failed_files = []
                for f in uploaded_comp:
                    text_c = extract_text_from_file(f)
                    if isinstance(text_c, str) and text_c.startswith("__ERROR_"):
                        failed_files.append((f.name, text_c))
                        continue
                    cleaned_c = clean_text(text_c)
                    if not cleaned_c or len(cleaned_c.strip()) < 50:
                        failed_files.append((f.name, "__ERROR_EMPTY_DOCUMENT__"))
                        continue
                    res_c = analyze_full_document(cleaned_c)
                    if res_c:
                        clf_c = classify_research_domain(cleaned_c)
                        cits = analyze_citations(cleaned_c)
                        res_c["filename"] = f.name
                        res_c["clean_title"] = get_clean_display_title(f.name, text=cleaned_c, file=f)
                        res_c["domain"] = clf_c["primary"]
                        res_c["domain_confidence"] = clf_c["confidence"]
                        res_c["citation_analysis"] = cits
                        res_c["full_text"] = cleaned_c
                        w_cnt = compute_text_counts(cleaned_c)["word_count"]
                        res_c["word_count"] = w_cnt
                        res_c["doc_page_count"] = max(1, math.ceil(w_cnt / 450))
                        compared_list.append(res_c)
                    else:
                        failed_files.append((f.name, "Analysis failed"))
                
                if failed_files:
                    for fname, reason in failed_files:
                        err_msg = "insufficient extractable text"
                        if "SCANNED" in reason:
                            err_msg = "scanned/image-only PDF (OCR required)"
                        elif "PASSWORD" in reason:
                            err_msg = "password-protected document"
                        elif "LARGE" in reason:
                            err_msg = "file too large (>30 MB)"
                        elif "CORRUPTED" in reason:
                            err_msg = "corrupted or unreadable file"
                        st.warning(f"⚠️ Could not process '{fname}': {err_msg}")
                        
                if len(compared_list) < 2:
                    st.error("Comparison requires at least 2 successfully extracted research papers. Please provide valid documents.")
                else:
                    st.session_state["comparison_data"] = compared_list
                    succ_comp_msg = f"Successfully compared {len(compared_list)} manuscripts!" if is_researcher else f"Successfully compared {len(compared_list)} papers!"
                    st.success(succ_comp_msg)

    comp_data = st.session_state.get("comparison_data", [])
    if comp_data and len(comp_data) >= 2:
        # Check for identical uploaded documents
        texts = [p.get("full_text", "").strip() for p in comp_data]
        is_duplicate = (len(texts) != len(set(texts)))
        if is_duplicate:
            st.info("ℹ️ **Identical Document Detection**: Two or more uploaded papers contain identical text. As expected under consistent deterministic scoring, identical documents receive identical scores across all dimensions.")

        render_html("<br>")
        col_hdr1, col_hdr2 = st.columns([5, 1.2])
        with col_hdr1:
            render_html([
                f'<div style="background: {theme["accent_primary_subtle"]}; border: 1px solid {theme["accent_primary_border"]}; border-radius: 8px; padding: 10px 14px; margin-bottom: 16px;">',
                f'<span style="color: {theme["text_primary"]}; font-size: 13px; font-weight: 600;">Comparative Summary</span>',
                f'<div style="color: {theme["text_muted"]}; font-size: 12px; margin-top: 2px;">Side-by-side comparison across key academic dimensions. Evaluates writing clarity and structure without claiming scientific superiority.</div>',
                f'</div>',
            ])
        with col_hdr2:
            render_html("<div style='height: 4px;'></div>")
            if st.button("🗑️ Reset", key="btn_reset_comp_view", use_container_width=True, help="Clear comparison results"):
                clear_comparison_state()
                st.rerun()
        
        # 1. Paper Overview Metadata Cards
        meta_cols = st.columns(len(comp_data))
        palette_colors = ["#8B5CF6", "#06B6D4", "#F59E0B"]
        for idx_m, p in enumerate(comp_data):
            with meta_cols[idx_m]:
                p_label = f"Manuscript {idx_m + 1}" if is_researcher else f"Paper {idx_m + 1}"
                p_title = unescape_text(p.get("clean_title", p["filename"]))
                p_short = truncate_title_word_boundary(p_title, 44)
                p_words = p.get("word_count") or compute_text_counts(p.get("full_text", ""))["word_count"]
                p_pages = p.get("doc_page_count", max(1, math.ceil(p_words / 450)))
                p_dom = p.get("domain", "General Academic")
                p_ovr = get_overall_score(p)
                p_grd = get_grade(p_ovr)
                accent_c = palette_colors[idx_m % len(palette_colors)]
                
                render_html([
                    f'<div class="iq-card" style="border-top: 3px solid {accent_c} !important; margin-bottom: 14px; padding: 14px 16px !important; height: 100%;">',
                    f'<div style="display: flex; align-items: center; justify-content: space-between; margin-bottom: 6px;">',
                    f'<span style="font-size: 11px; font-weight: 800; color: {accent_c}; text-transform: uppercase; letter-spacing: 0.05em;">{safe_text(p_label)}</span>',
                    f'<span class="domain-badge" style="margin: 0; font-size: 10.5px;">{safe_text(p_dom)}</span>',
                    f'</div>',
                    f'<div style="font-weight: 700; font-size: 13.5px; color: {theme["text_primary"]}; margin-bottom: 6px; line-height: 1.4; overflow-wrap: break-word;" title="{safe_text(p_title)}">{safe_text(p_short)}</div>',
                    f'<div style="display: flex; justify-content: space-between; align-items: center; margin-bottom: 6px;">',
                    f'<span style="font-size: 13px; font-weight: 800; color: {accent_c};">Score: {p_ovr:.1f} / 100</span>',
                    f'<span style="background: {theme["badge_role_bg"]}; color: {theme["badge_role_text"]}; font-size: 10.5px; font-weight: 700; padding: 2px 7px; border-radius: 6px; border: 1px solid {theme["badge_role_border"]};">Grade {safe_text(p_grd)}</span>',
                    f'</div>',
                    f'<div style="font-size: 11.5px; color: {theme["text_muted"]}; display: flex; gap: 12px;">',
                    f'<span>📄 ~{p_pages} pages</span>',
                    f'<span>📝 {p_words:,} words</span>',
                    f'</div>',
                    f'</div>',
                ])

        # 2. Theme-Aware Responsive Comparison Table
        dimensions_list = [
            ("Language Quality", "Language Quality"),
            ("Structural Coherence", "Structural Coherence"),
            ("Argumentation", "Argumentation"),
            ("Academic Style", "Academic Style"),
            ("Readability", "Readability"),
            ("Overall (Composite)", "Composite"),
        ]
        
        # Build semantic HTML table
        th_headers = ["<th style='width: 28%; text-align: left;'>Academic Dimension</th>"]
        for idx_h, p in enumerate(comp_data):
            p_label = f"Paper {idx_h + 1}"
            p_title = unescape_text(p.get("clean_title", p["filename"]))
            p_short = truncate_title_word_boundary(p_title, 32)
            accent_c = palette_colors[idx_h % len(palette_colors)]
            th_headers.append(
                f"<th style='text-align: center;' title='{safe_text(p_title)}'>"
                f"<span style='color: {accent_c}; font-weight: 800;'>{safe_text(p_label)}</span> — {safe_text(p_short)}"
                f"</th>"
            )
            
        tr_rows = []
        for label, key in dimensions_list:
            vals = []
            for p in comp_data:
                s = p["scores"]
                if key == "Composite":
                    val = get_overall_score(p)
                else:
                    val = s.get(key, s.get(key.split()[0], 0.0))
                vals.append(val)
            max_v = max(vals) if vals else 0.0
            
            td_cells = [f"<td class='iq-comp-dim-name'><b>{safe_text(label)}</b></td>"]
            for v in vals:
                is_max = (v == max_v and max_v > 0)
                badge_style = f"background: {theme['success_bg']}; color: {theme['success_text']}; border: 1px solid {theme['success_border']};" if is_max else f"background: {theme['bg_card_inner']}; color: {theme['text_primary']}; border: 1px solid {theme['border_subtle']};"
                star = " ★" if is_max else ""
                td_cells.append(
                    f"<td class='iq-comp-val-cell'>"
                    f"<span class='iq-comp-val-pill' style='{badge_style}'>{v:.1f}{star}</span>"
                    f"</td>"
                )
            tr_rows.append(f"<tr>{''.join(td_cells)}</tr>")
            
        render_html([
            '<div class="iq-comparison-table-wrapper">',
            '<table class="iq-comp-table">',
            f'<thead><tr>{"".join(th_headers)}</tr></thead>',
            f'<tbody>{"".join(tr_rows)}</tbody>',
            '</table>',
            '</div>',
        ])
        
        # 3. Winner / Strongest Dimension & Largest Score Disparity Summary Cards
        best_p = max(comp_data, key=lambda p: get_overall_score(p))
        best_title = unescape_text(best_p.get("clean_title", best_p["filename"]))
        best_score = get_overall_score(best_p)
        best_p_scores = best_p.get("scores", {})
        best_p_dim_map = {
            "Language Quality": float(best_p_scores.get("Language Quality", best_p_scores.get("Language", 0.0))),
            "Structural Coherence": float(best_p_scores.get("Structural Coherence", best_p_scores.get("Coherence", 0.0))),
            "Argumentation": float(best_p_scores.get("Argumentation", best_p_scores.get("Reasoning", 0.0))),
            "Academic Style": float(best_p_scores.get("Academic Style", best_p_scores.get("Sophistication", 0.0))),
            "Readability": float(best_p_scores.get("Readability", 0.0)),
        }
        strongest_dim_name, strongest_dim_val = max(best_p_dim_map.items(), key=lambda x: x[1])

        # Calculate dimension score differences / gaps across papers
        dim_gaps = []
        for dim_name in ["Language Quality", "Structural Coherence", "Argumentation", "Academic Style", "Readability"]:
            vals = []
            for p in comp_data:
                s = p.get("scores", {})
                if dim_name == "Language Quality":
                    v = s.get("Language Quality", s.get("Language", 0.0))
                elif dim_name == "Structural Coherence":
                    v = s.get("Structural Coherence", s.get("Coherence", 0.0))
                elif dim_name == "Argumentation":
                    v = s.get("Argumentation", s.get("Reasoning", 0.0))
                elif dim_name == "Academic Style":
                    v = s.get("Academic Style", s.get("Sophistication", 0.0))
                elif dim_name == "Readability":
                    v = s.get("Readability", 0.0)
                else:
                    v = s.get(dim_name, 0.0)
                vals.append(float(v))
            gap = max(vals) - min(vals) if vals else 0.0
            dim_gaps.append((dim_name, gap, max(vals) if vals else 0.0, min(vals) if vals else 0.0))
        
        dim_gaps.sort(key=lambda x: x[1], reverse=True)
        largest_gap_dim, max_gap, top_val, low_val = dim_gaps[0]

        col_comp_c1, col_comp_c2 = st.columns(2)
        with col_comp_c1:
            render_html([
                f'<div style="background: {theme["success_bg"]}; border: 1px solid {theme["success_border"]}; border-radius: 8px; padding: 12px 18px; margin: 14px 0; height: 100%;" title="{safe_text(best_title)}">',
                f'<div style="color: {theme["success_text"]}; font-size: 11px; font-weight: 700; text-transform: uppercase; letter-spacing: 0.05em;">🏆 Overall Winner &amp; Strongest Dimension</div>',
                f'<div style="color: {theme["text_primary"]}; font-size: 15px; font-weight: 700; margin-top: 2px;">{safe_text(best_title)}</div>',
                f'<div style="color: {theme["text_secondary"]}; font-size: 13px; margin-top: 2px;">Overall Score: <b>{best_score:.1f} / 100</b> (Grade {safe_text(get_grade(best_score))}) &nbsp;&bull;&nbsp; Strongest Dimension: <b>{safe_text(strongest_dim_name)} ({strongest_dim_val:.1f})</b></div>',
                f'</div>',
            ])
        with col_comp_c2:
            render_html([
                f'<div style="background: {theme["info_bg"]}; border: 1px solid {theme["info_border"]}; border-radius: 8px; padding: 12px 18px; margin: 14px 0; height: 100%;">',
                f'<div style="color: {theme["info_text"]}; font-size: 11px; font-weight: 700; text-transform: uppercase; letter-spacing: 0.05em;">⚡ Largest Writing Quality Gap</div>',
                f'<div style="color: {theme["text_primary"]}; font-size: 15px; font-weight: 700; margin-top: 2px;">{safe_text(largest_gap_dim)} (Δ {max_gap:.1f} pts)</div>',
                f'<div style="color: {theme["text_secondary"]}; font-size: 13px; margin-top: 2px;">Spread from <b>{low_val:.1f}</b> to <b>{top_val:.1f}</b> across manuscripts</div>',
                f'</div>',
            ])
        
        # 4. Detailed Strengths AND Weaknesses per Paper
        render_html("<h3 style='font-size: 18px; margin: 18px 0 12px 0;'>Strengths &amp; Areas for Improvement</h3>")
        
        comp_cols = st.columns(len(comp_data))
        for idx_p, p in enumerate(comp_data):
            with comp_cols[idx_p]:
                p_title = unescape_text(p.get("clean_title", p["filename"]))
                p_display_title = truncate_title_word_boundary(p_title, 72)
                s = p["scores"]
                p_ovr = get_overall_score(p)
                l_v = s.get("Language Quality", s.get("Language", 0.0))
                c_v = s.get("Structural Coherence", s.get("Coherence", 0.0))
                a_v = s.get("Argumentation", s.get("Reasoning", 0.0))
                st_v = s.get("Academic Style", s.get("Sophistication", 0.0))
                r_v = s.get("Readability", 0.0)
                
                # Identify strengths (top 2) and areas for improvement (lowest 2)
                dim_dict = {
                    "Language Quality": l_v, "Structural Coherence": c_v,
                    "Argumentation": a_v, "Academic Style": st_v, "Readability": r_v
                }
                sorted_by_score = sorted(dim_dict.items(), key=lambda x: x[1], reverse=True)
                strengths = sorted_by_score[:2]
                weaknesses = sorted_by_score[-2:]
                accent_c = palette_colors[idx_p % len(palette_colors)]
                
                render_html([
                    f'<div class="iq-card" style="border-top: 3px solid {accent_c} !important; height: 100%; min-height: 310px; display: flex; flex-direction: column; justify-content: space-between; overflow-wrap: break-word;">',
                    f'<div>',
                    f'<div style="font-weight: 700; color: {theme["text_primary"]}; font-size: 14px; margin-bottom: 4px; line-height: 1.4; min-height: 42px;" title="{safe_text(p_title)}">📄 Paper {idx_p + 1}: {safe_text(p_display_title)}</div>',
                    f'<div style="font-size: 12px; color: {theme["accent_purple"]}; font-weight: 700; margin-bottom: 10px;">Overall: {p_ovr:.1f} / 100 (Grade {safe_text(get_grade(p_ovr))})</div>',
                    f'<div style="font-size: 12px; color: {theme["success_text"]}; font-weight: 700; text-transform: uppercase;">✔ Strengths</div>',
                    f'<ul style="font-size: 12px; color: {theme["text_secondary"]}; padding-left: 18px; margin: 4px 0 10px 0; line-height: 1.5;">',
                    f'<li><b>{safe_text(strengths[0][0])}</b> ({strengths[0][1]:.1f}/100) — Strongest dimension.</li>',
                    f'<li><b>{safe_text(strengths[1][0])}</b> ({strengths[1][1]:.1f}/100) — Strong academic consistency.</li>',
                    f'</ul>',
                    f'<div style="font-size: 12px; color: {theme["error_text"]}; font-weight: 700; text-transform: uppercase;">▲ Areas for Improvement</div>',
                    f'<ul style="font-size: 12px; color: {theme["text_secondary"]}; padding-left: 18px; margin: 4px 0 4px 0; line-height: 1.5;">',
                    f'<li><b>{safe_text(weaknesses[0][0])}</b> ({weaknesses[0][1]:.1f}/100) — Needs editorial attention.</li>',
                    f'<li><b>{safe_text(weaknesses[1][0])}</b> ({weaknesses[1][1]:.1f}/100) — Primary improvement focus.</li>',
                    f'</ul>',
                    f'</div>',
                    f'</div>',
                ])
                
        # 5. Cross-Manuscript Similarities and Differences Synthesis
        render_html("<h3 style='font-size: 18px; margin: 20px 0 12px 0;'>Comparative Synthesis: Similarities &amp; Differences</h3>")
        col_syn1, col_syn2 = st.columns(2)
        with col_syn1:
            render_html([
                f'<div class="iq-card" style="height: 100%;">',
                f'<div style="font-size: 13.5px; font-weight: 700; color: {theme["info_text"]}; margin-bottom: 6px;">🤝 Structural &amp; Stylistic Similarities</div>',
                f'<ul style="font-size: 12.5px; color: {theme["text_secondary"]}; padding-left: 16px; margin: 0; line-height: 1.6;">',
                f'<li><b>Academic Register</b>: All compared manuscripts maintain formal scholarly tone without colloquial phrasing.</li>',
                f'<li><b>Disciplinary Grounding</b>: The manuscripts demonstrate familiarity with conventions in {safe_text(comp_data[0].get("domain", "academic"))} research.</li>',
                f'<li><b>Readability Alignment</b>: Reading ease metrics fall within standard peer-reviewed scholarly tolerances across the documents.</li>',
                f'</ul>',
                f'</div>',
            ])
        with col_syn2:
            best_title_syn = truncate_title_word_boundary(best_title, 44)
            render_html([
                f'<div class="iq-card" style="height: 100%;">',
                f'<div style="font-size: 13.5px; font-weight: 700; color: {theme["accent_purple"]}; margin-bottom: 6px;">⚡ Key Divergences &amp; Differences</div>',
                f'<ul style="font-size: 12.5px; color: {theme["text_secondary"]}; padding-left: 16px; margin: 0; line-height: 1.6;">',
                f'<li><b>Greatest Disparity</b>: The largest divergence appears in <b>{safe_text(largest_gap_dim)}</b> with a delta of <b>{max_gap:.1f} points</b>.</li>',
                f'<li><b>Composite Ranking</b>: <b title="{safe_text(best_title)}">{safe_text(best_title_syn)}</b> achieves the leading writing score ({best_score:.1f}/100).</li>',
                f'<li><b>Rhetorical Complexity</b>: Structural coherence and argumentation depth show distinct variations in paragraph transitions and supporting evidence.</li>',
                f'</ul>',
                f'</div>',
            ])

        # 6. Actionable Writing Recommendations Specific to Each Paper
        render_html("<h3 style='font-size: 18px; margin: 20px 0 12px 0;'>Actionable Recommendations per Manuscript</h3>")
        rec_cols = st.columns(len(comp_data))
        for idx_r, p in enumerate(comp_data):
            with rec_cols[idx_r]:
                p_title = unescape_text(p.get("clean_title", p["filename"]))
                s_map = p.get("scores", {})
                dim_items = sorted([
                    ("Language Quality", s_map.get("Language Quality", s_map.get("Language", 0.0))),
                    ("Structural Coherence", s_map.get("Structural Coherence", s_map.get("Coherence", 0.0))),
                    ("Argumentation", s_map.get("Argumentation", s_map.get("Reasoning", 0.0))),
                    ("Academic Style", s_map.get("Academic Style", s_map.get("Sophistication", 0.0))),
                    ("Readability", s_map.get("Readability", 0.0)),
                ], key=lambda x: x[1])
                lowest_dim, lowest_val = dim_items[0]
                
                tips = {
                    "Language Quality": "Enhance lexical variation and reduce repetitive sentence openings.",
                    "Structural Coherence": "Add explicit transition signposts between major subsection paragraphs.",
                    "Argumentation": "Strengthen empirical claims with direct literature cross-references.",
                    "Academic Style": "Replace passive nominalizations with concise, active scholarly verbs.",
                    "Readability": "Decompose multi-clause sentences (>35 words) into digestible units."
                }
                tip_text = tips.get(lowest_dim, "Refine sectional clarity and evidence grounding.")
                
                render_html([
                    f'<div class="iq-card" style="border-top: 3px solid {palette_colors[idx_r % len(palette_colors)]} !important; height: 100%;" title="{safe_text(p_title)}">',
                    f'<div style="font-weight: 700; color: {theme["text_primary"]}; font-size: 13.5px; margin-bottom: 4px;">Paper {idx_r + 1} Guidance</div>',
                    f'<div style="font-size: 12px; color: {theme["info_text"]}; font-weight: 600; margin-bottom: 6px;">Primary Focus: {safe_text(lowest_dim)} ({lowest_val:.1f}/100)</div>',
                    f'<p style="font-size: 12px; color: {theme["text_secondary"]}; margin: 0; line-height: 1.5;">{safe_text(tip_text)}</p>',
                    f'</div>',
                ])
                
        # 7. Visual Comparison: Radar Overlay + Grouped Bar Chart
        render_html("<br>")
        render_html("<h3 style='font-size: 18px; margin: 10px 0 12px 0;'>Visual Comparison: Radar Overlay &amp; Score Breakdown</h3>")
        col_viz_radar, col_viz_bar = st.columns(2)
        ctheme_comp = get_chart_theme(st.session_state.get("theme_preference", "System"))
        radar_dims = ["Language Quality", "Structural Coherence", "Argumentation", "Academic Style", "Readability"]
        dims = radar_dims + ["Overall"]
        fill_rgba_palette = [
            "rgba(139, 92, 246, 0.20)",
            "rgba(6, 182, 212, 0.20)",
            "rgba(245, 158, 11, 0.20)",
        ]

        with col_viz_radar:
            fig_comp_radar = go.Figure()
            for idx_p, p in enumerate(comp_data):
                s = p["scores"]
                r_vals = [
                    float(s.get("Language Quality", s.get("Language", 0))),
                    float(s.get("Structural Coherence", s.get("Coherence", 0))),
                    float(s.get("Argumentation", s.get("Reasoning", 0))),
                    float(s.get("Academic Style", s.get("Sophistication", 0))),
                    float(s.get("Readability", 0)),
                ]
                p_label = f"Paper {idx_p + 1}"
                p_title = unescape_text(p.get("clean_title", p["filename"]))
                p_short = truncate_title_word_boundary(p_title, 26)
                trace_name = f"{p_label} ({p_short})"
                fig_comp_radar.add_trace(go.Scatterpolar(
                    name=trace_name,
                    r=r_vals + [r_vals[0]],
                    theta=radar_dims + [radar_dims[0]],
                    fill='toself',
                    fillcolor=fill_rgba_palette[idx_p % len(fill_rgba_palette)],
                    line=dict(color=palette_colors[idx_p % len(palette_colors)], width=2.5),
                    hoverinfo="name+theta+r"
                ))
            fig_comp_radar.update_layout(
                polar=dict(
                    bgcolor=ctheme_comp["polar_bg"],
                    radialaxis=dict(visible=True, range=[0, 100], gridcolor=ctheme_comp["grid_color"], tickfont=dict(color=ctheme_comp["tick_color"], size=10)),
                    angularaxis=dict(tickfont=dict(color=ctheme_comp["font_color"], size=11), linecolor=ctheme_comp["line_color"])
                ),
                paper_bgcolor=ctheme_comp["paper_bgcolor"],
                plot_bgcolor=ctheme_comp["plot_bgcolor"],
                font=dict(color=ctheme_comp["font_color"]),
                height=400,
                margin=dict(t=45, b=40, l=45, r=45),
                legend=dict(
                    orientation="h",
                    y=1.06,
                    yanchor="bottom",
                    x=0.5,
                    xanchor="center",
                    font=dict(color=ctheme_comp["font_color"], size=11.5),
                    bgcolor="rgba(0,0,0,0)"
                )
            )
            st.plotly_chart(fig_comp_radar, use_container_width=True)

        with col_viz_bar:
            fig_comp = go.Figure()
            for idx_p, p in enumerate(comp_data):
                s = p["scores"]
                y_vals = [
                    s.get("Language Quality", s.get("Language", 0)),
                    s.get("Structural Coherence", s.get("Coherence", 0)),
                    s.get("Argumentation", s.get("Reasoning", 0)),
                    s.get("Academic Style", s.get("Sophistication", 0)),
                    s.get("Readability", 0),
                    get_overall_score(p)
                ]
                p_label = f"Paper {idx_p + 1}"
                p_title = unescape_text(p.get("clean_title", p["filename"]))
                p_short = truncate_title_word_boundary(p_title, 26)
                trace_name = f"{p_label} ({p_short})"
                
                fig_comp.add_trace(go.Bar(
                    name=trace_name,
                    x=dims,
                    y=y_vals,
                    marker_color=palette_colors[idx_p % len(palette_colors)]
                ))
            fig_comp.update_layout(
                barmode='group',
                bargap=0.22,
                bargroupgap=0.08,
                height=400,
                margin=dict(t=45, b=40, l=45, r=20),
                paper_bgcolor=ctheme_comp["paper_bgcolor"],
                plot_bgcolor=ctheme_comp["plot_bgcolor"],
                font=dict(color=ctheme_comp["font_color"]),
                yaxis=dict(range=[0, 105], gridcolor=ctheme_comp["grid_color"], linecolor=ctheme_comp["line_color"], title="Score (0-100)"),
                xaxis=dict(linecolor=ctheme_comp["line_color"], tickfont=dict(color=ctheme_comp["tick_color"])),
                legend=dict(
                    orientation="h",
                    y=1.06,
                    yanchor="bottom",
                    x=0.5,
                    xanchor="center",
                    font=dict(color=ctheme_comp["font_color"], size=11.5),
                    bgcolor="rgba(0,0,0,0)",
                    bordercolor="rgba(0,0,0,0)"
                )
            )
            st.plotly_chart(fig_comp, use_container_width=True)
        
        # 8. Downloadable Comparison Report Export (PDF & TXT)
        comp_report_lines = [
            "=" * 70,
            "PAPERIQ RESEARCH INTELLIGENCE — COMPARATIVE EVALUATION REPORT",
            "=" * 70,
            f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
            f"Number of Compared Manuscripts: {len(comp_data)}",
            f"Leading Manuscript: {best_title} ({best_score:.1f}/100 - Grade {get_grade(best_score)})",
            f"Largest Score Disparity: {largest_gap_dim} (Delta of {max_gap:.1f} pts)",
            "-" * 70,
            "DETAILED SCORES PER MANUSCRIPT:",
        ]
        for idx_p, p in enumerate(comp_data, 1):
            p_title = unescape_text(p.get("clean_title", p["filename"]))
            ovr = get_overall_score(p)
            comp_report_lines.append(f"\n[{idx_p}] {p_title}")
            comp_report_lines.append(f"    Overall Score: {ovr:.1f} / 100 (Grade {get_grade(ovr)})")
            comp_report_lines.append(f"    Domain: {p.get('domain', 'General Academic')}")
            for d_name in ["Language Quality", "Structural Coherence", "Argumentation", "Academic Style", "Readability"]:
                s_map = p.get("scores", {})
                if d_name == "Language Quality":
                    sc = s_map.get("Language Quality", s_map.get("Language", 0.0))
                elif d_name == "Structural Coherence":
                    sc = s_map.get("Structural Coherence", s_map.get("Coherence", 0.0))
                elif d_name == "Argumentation":
                    sc = s_map.get("Argumentation", s_map.get("Reasoning", 0.0))
                elif d_name == "Academic Style":
                    sc = s_map.get("Academic Style", s_map.get("Sophistication", 0.0))
                elif d_name == "Readability":
                    sc = s_map.get("Readability", 0.0)
                else:
                    sc = s_map.get(d_name, 0.0)
                comp_report_lines.append(f"    - {d_name}: {float(sc):.1f} / 100")
        comp_report_lines.append("\n" + "=" * 70)
        comp_report_lines.append("Evaluated deterministically by PaperIQ AI Research Intelligence Platform.")
        comp_report_lines.append("Disclaimer: Evaluates writing clarity and structure; does not claim scientific superiority.")
        comp_report_lines.append("=" * 70)
        comp_report_txt = "\n".join(comp_report_lines)

        comp_pdf_bytes = get_cached_comparison_pdf(comp_data)
        comp_pdf_fn = pdf_export.get_pdf_filename(best_title, "Comparison")

        col_dl_c1, col_dl_c2, _ = st.columns([2, 2, 1])
        with col_dl_c1:
            st.download_button(
                label="⬇ Download Comparison Report (PDF)",
                data=comp_pdf_bytes,
                file_name=comp_pdf_fn,
                mime="application/pdf",
                key="btn_download_comparison_pdf",
                use_container_width=True
            )
        with col_dl_c2:
            st.download_button(
                label="⬇ Download Comparison Report (.txt)",
                data=comp_report_txt,
                file_name=f"PaperIQ_Comparison_{len(comp_data)}_Manuscripts.txt",
                mime="text/plain",
                key="btn_download_comparison_txt",
                use_container_width=True
            )
        
        # 9. Recommendation & Limitation Notice
        render_html([
            f'<div class="iq-card" style="margin-top: 14px; border-left: 4px solid {theme["info_text"]} !important;">',
            f'<div style="font-size: 13px; font-weight: 700; color: {theme["text_primary"]}; margin-bottom: 4px;">ℹ️ Academic Evaluation Disclaimer</div>',
            f'<p style="font-size: 12px; color: {theme["text_muted"]}; margin: 0; line-height: 1.5;">',
            f'PaperIQ evaluates writing clarity, structural coherence, argumentation flow, academic style, and readability. A higher writing-quality score indicates clearer academic presentation and does not imply greater scientific validity, empirical novelty, or methodological superiority.',
            f'</p>',
            f'</div>',
        ])

# ----------------------------------------------------
# PAGE 7: ANALYSIS HISTORY
# ----------------------------------------------------
elif st.session_state.current_page in ["📚 History", "📚 My History", "📚 Research History", "🕘 History"]:
    hist_heading = "Research History" if is_researcher else ("My History" if is_student else "Analysis History")
    render_html([
        f'<div style="margin-bottom: 20px;">',
        f'<h1 style="margin: 0; font-size: 28px; color: {theme["text_primary"]};">{safe_text(hist_heading)}</h1>',
        f'<p style="margin: 4px 0 0 0; color: {theme["text_muted"]}; font-size: 15px;">Persistently saved research evaluations and downloadable analysis reports from your local workspace.</p>',
        f'</div>',
    ])
    
    if not user_history:
        hist_empty_title = "No Research History Yet" if is_researcher else "No Analysis History Yet"
        hist_empty_desc = "Evaluate your first manuscript in <b>Analyze Manuscript</b> to archive evaluations and reports here." if is_researcher else "Analyze your first paper in <b>Analyze Paper</b> to archive evaluations and reports here."
        render_html([
            f'<div class="iq-card" style="text-align: center; padding: 40px !important;">',
            f'<div style="font-size: 32px; margin-bottom: 8px;">📚</div>',
            f'<h4 style="color: {theme["text_primary"]}; margin: 0;">{safe_text(hist_empty_title)}</h4>',
            f'<p style="color: {theme["text_muted"]}; margin-top: 6px; margin-bottom: 20px;">{hist_empty_desc}</p>',
            f'</div>',
        ])
        c_emp1, c_emp2, c_emp3 = st.columns([1, 1.2, 1])
        with c_emp2:
            target_analyze = "📄 Analyze Manuscript" if is_researcher else "📄 Analyze Paper"
            btn_first_label = "📄 Evaluate Your First Manuscript" if is_researcher else "📄 Analyze Your First Paper"
            st.button(
                btn_first_label,
                type="primary",
                key="btn_hist_empty_goto_analyze",
                on_click=navigate_to,
                args=(target_analyze,),
                use_container_width=True
            )
    else:
        # Search and Domain Filtering Controls
        f_col1, f_col2, f_col3 = st.columns([3, 1.8, 1.2])
        with f_col1:
            search_query = st.text_input("🔍 Search papers by title", "", key="hist_search_box")
        with f_col2:
            all_domains = sorted(list(set([item.get("domain", "General Academic") for item in user_history])))
            domain_filter = st.selectbox("Filter by Domain", ["All Domains"] + all_domains, key="hist_domain_filter")
        with f_col3:
            st.write("")
            st.write("")
            if st.button("🗑️ Clear All", key="hist_btn_clear_all_top", use_container_width=True):
                st.session_state["confirm_clear_all_hist_page"] = True
                st.rerun()

        if st.session_state.get("confirm_clear_all_hist_page", False):
            render_html([
                f'<div style="background: {theme["error_bg"]}; border: 1px solid {theme["error_border"]}; border-radius: 8px; padding: 14px 18px; margin: 12px 0;">',
                f'<div style="color: {theme["error_text"]}; font-weight: 700; font-size: 14px; margin-bottom: 6px;">⚠ Clear All Analysis History</div>',
                f'<div style="color: {theme["text_secondary"]}; font-size: 13px;">',
                f'Permanently delete all <b>{len(user_history)}</b> saved analysis records for account <b>{safe_text(username)}</b>?',
                f'</div>',
                f'</div>',
            ])
            c_cl1, c_cl2, _ = st.columns([1, 1.6, 4])
            with c_cl1:
                if st.button("Cancel", key="btn_cancel_clear_all_hist_page"):
                    st.session_state["confirm_clear_all_hist_page"] = False
                    st.rerun()
            with c_cl2:
                if st.button("Confirm Delete All", type="primary", key="btn_confirm_clear_all_hist_page"):
                    success, msg = auth.clear_user_history(username)
                    if success:
                        st.session_state.analysis_history[username] = []
                        clear_active_analysis_session()
                        st.session_state["confirm_clear_all_hist_page"] = False
                        st.session_state["hist_cleared_success"] = True
                        st.rerun()
                    else:
                        st.error(f"Failed to clear history: {msg}")

        if st.session_state.get("hist_cleared_success", False):
            st.success("✓ All analysis history and active workspace analysis cleared successfully.")
            st.session_state["hist_cleared_success"] = False

        filtered_history = [
            item for item in user_history
            if (not search_query or search_query.lower() in f"{item.get('display_title', '')} {item.get('original_filename', '')} {item.get('filename', '')} {item.get('domain', '')}".lower())
            and (domain_filter == "All Domains" or item.get("domain") == domain_filter)
        ]
        
        render_html(f"<p style='color: {theme['text_muted']}; font-size: 13px; margin: 6px 0 12px 0;'>Showing <b>{len(filtered_history)}</b> of <b>{len(user_history)}</b> analyses</p>")
        
        if not filtered_history:
            st.info("No analyses match the current search or filter criteria.")
        else:
            for idx, item in enumerate(filtered_history):
                item_score = get_overall_score(item)
                grade = get_grade(item_score)
                res_meta = None
                if item.get("results_json"):
                    try:
                        res_meta = json.loads(item["results_json"])
                    except Exception:
                        pass
                stored_title = item.get("display_title") or (res_meta.get("title") if res_meta else None)
                clean_title = get_clean_display_title(item.get("filename", "Paper"), stored_title=stored_title)
                truncated_title = truncate_title_word_boundary(clean_title, 48)
                clean_fn = pdf_export.get_pdf_filename(clean_title, "Analysis")
                
                with st.container(border=True):
                    col_info, col_score, col_v, col_d, col_del = st.columns([3.8, 1.8, 1.6, 1.6, 1.2])
                    with col_info:
                        render_html([
                            f'<div style="font-size: 15px; font-weight: 700; color: {theme["text_primary"]};" title="{safe_text(clean_title)}">📄 {safe_text(truncated_title)}</div>',
                            f'<div style="font-size: 12px; color: {theme["text_muted"]}; margin-top: 3px;">',
                            f'Domain: <b style="color: {theme["info_text"]};">{safe_text(item.get("domain", "General Academic"))}</b> &nbsp;•&nbsp; ',
                            f'Date: {safe_text(item.get("created_at", ""))}',
                            f'</div>',
                        ])
                    with col_score:
                        render_html([
                            f'<div style="font-size: 15px; font-weight: 800; color: {theme["accent_purple"]}; margin-top: 6px;">',
                            f'Score: {item_score:.1f} / 100 &nbsp;',
                            f'<span style="background: {theme["badge_role_bg"]}; color: {theme["badge_role_text"]}; font-size: 11px; font-weight: 700; padding: 2px 7px; border-radius: 8px; border: 1px solid {theme["badge_role_border"]};">Grade {safe_text(grade)}</span>',
                            f'</div>',
                        ])
                    with col_v:
                        view_hist_btn_label = "View Evaluation" if is_researcher else "View Analysis"
                        target_v_pg = "📊 Manuscript Evaluation" if is_researcher else ("📊 Writing Quality & Feedback" if is_student else "📊 Analysis")
                        if st.button(view_hist_btn_label, key=f"hist_v_{item.get('id', idx)}", use_container_width=True):
                            load_analysis_into_session(item, source="history")
                            st.session_state.current_page = target_v_pg
                            st.rerun()
                    with col_d:
                        rec_owner_ok = auth.verify_analysis_ownership(item.get("id") or item.get("analysis_id"), username)
                        if rec_owner_ok:
                            hist_pdf_bytes = item.get("pdf")
                            if res_meta and isinstance(res_meta, dict):
                                hist_pdf_bytes = get_cached_analysis_pdf(res_meta, item.get("filename", "Paper"), item.get("analysis_id") or item.get("id"))
                            if hist_pdf_bytes:
                                st.download_button(
                                    "Download Report",
                                    data=hist_pdf_bytes,
                                    file_name=clean_fn,
                                    mime="application/pdf",
                                    key=f"hist_d_{item.get('id', idx)}",
                                    use_container_width=True
                                )
                    with col_del:
                        if st.button("Delete", key=f"hist_del_{item.get('id', idx)}", use_container_width=True):
                            st.session_state[f"confirm_del_{item.get('id', idx)}"] = True
                            st.rerun()
                            
                    # Confirmation Step to Prevent Accidental Deletion
                    if st.session_state.get(f"confirm_del_{item.get('id', idx)}", False):
                        render_html([
                            f"<div style='background: {theme['error_bg']}; border: 1px solid {theme['error_border']}; border-radius: 8px; padding: 10px 14px; margin-top: 8px;'>",
                            f"<span style='color: {theme['error_text']}; font-weight: 700; font-size: 13px;'>Delete Analysis?</span>",
                            f"<div style='color: {theme['text_secondary']}; font-size: 12px; margin-top: 2px;'>This will remove this analysis record and its generated report snapshot.</div>",
                            f"</div>",
                        ])
                        cd_c1, cd_c2, _ = st.columns([1, 1.2, 4])
                        with cd_c1:
                            if st.button("Cancel", key=f"cancel_del_{item.get('id', idx)}"):
                                st.session_state[f"confirm_del_{item.get('id', idx)}"] = False
                                st.rerun()
                        with cd_c2:
                            if st.button("Delete", type="primary", key=f"confirm_del_btn_{item.get('id', idx)}"):
                                ok, msg = auth.delete_analysis_record(item.get("id"), username)
                                if ok:
                                    st.session_state.analysis_history[username] = [
                                        h for h in st.session_state.analysis_history[username] if h.get("id") != item.get("id")
                                    ]
                                    cur_id = str(st.session_state.get("current_analysis_id", ""))
                                    del_id = str(item.get("id", ""))
                                    del_anl_id = str(item.get("analysis_id", ""))
                                    del_fn = str(item.get("filename", ""))
                                    cur_fn = str(st.session_state.get("filename", ""))
                                    if (cur_id and (cur_id == del_id or cur_id == del_anl_id)) or (del_fn and cur_fn and del_fn == cur_fn):
                                        clear_active_analysis_session()
                                    st.session_state[f"confirm_del_{item.get('id', idx)}"] = False
                                    st.success("✓ Analysis deleted successfully.")
                                    st.rerun()
                                else:
                                    st.error(f"Failed to delete analysis: {msg}")

# ----------------------------------------------------
# PAGE 8: SETTINGS
# ----------------------------------------------------
elif st.session_state.current_page in ["⚙ Settings", "⚙️ Settings"]:
    render_html([
        f'<div style="margin-bottom: 20px;">',
        f'<h1 style="margin: 0; font-size: 28px; color: {theme["text_primary"]};">Settings &amp; Workspace</h1>',
        f'<p style="margin: 4px 0 0 0; color: {theme["text_muted"]}; font-size: 15px;">Manage your account details, workspace preferences, and application data.</p>',
        f'</div>',
    ])
    
    st.subheader("Account Profile")
    render_html([
        '<div class="iq-card">',
        f'<p><b>Username:</b> {safe_text(username)}</p>',
        f'<p><b>Academic Role:</b> {safe_text(role)}</p>',
        f'<p><b>Total Analyses in Database:</b> {len(user_history)}</p>',
        '</div>',
    ])
    
    st.subheader("Preferences")
    
    current_pref = st.session_state.get("theme_preference", "System")
    if current_pref not in ["Light", "Dark", "System"]:
        current_pref = "System"
    eff_theme = get_effective_theme(current_pref)
    
    render_html([
        '<div class="iq-card" style="margin-bottom: 14px;">',
        '<div style="display: flex; justify-content: space-between; align-items: center; margin-bottom: 6px;">',
        '<div>',
        '<div style="font-weight: 700; font-size: 15px;">Interface Theme</div>',
        '<div style="font-size: 13px; opacity: 0.8; margin-top: 2px;">Customize application workspace styling across all views</div>',
        '</div>',
        f'<span style="background: {theme["badge_role_bg"]}; color: {theme["accent_purple"]}; font-size: 12px; font-weight: 700; padding: 4px 12px; border-radius: 8px; border: 1px solid {theme["badge_role_border"]};">',
        f'Active: {safe_text(current_pref)} {f"({safe_text(eff_theme)} Mode)" if current_pref == "System" else "Mode"}',
        '</span>',
        '</div>',
        '</div>',
    ])
    
    theme_options = ["Light", "Dark", "System"]
    theme_idx = theme_options.index(current_pref)
    
    col_t1, col_t2 = st.columns([3, 2])
    with col_t1:
        selected_theme = st.radio(
            "Appearance Mode",
            options=theme_options,
            index=theme_idx,
            horizontal=True,
            key="theme_mode_radio",
            help="Select Light for daytime reading, Dark for nighttime focus, or System to follow your OS setting automatically."
        )
    with col_t2:
        desc_map = {
            "Light": "Clean, crisp light interface with dark typography.",
            "Dark": "Signature dark navy & purple analytical workspace.",
            "System": f"Automatically tracks your operating system appearance (currently {eff_theme})."
        }
        render_html([
            f'<div style="padding: 10px 14px; background: {theme["bg_card_inner"]}; border-radius: 8px; border: 1px solid {theme["border_color"]}; margin-top: 6px;">',
            f'<div style="font-size: 11px; font-weight: 700; color: {theme["text_muted"]}; text-transform: uppercase;">Selected Mode</div>',
            f'<div style="font-size: 14px; font-weight: 700; color: {theme["accent_purple"]}; margin-top: 2px;">{safe_text(selected_theme)}</div>',
            f'<div style="font-size: 11px; color: {theme["text_muted"]}; margin-top: 2px;">{safe_text(desc_map.get(selected_theme, ""))}</div>',
            f'</div>',
        ])
        
    if selected_theme != current_pref:
        st.session_state["theme_preference"] = selected_theme
        if username:
            auth.update_user_theme(username, selected_theme)
        st.rerun()
        
    render_html("<br>")
    st.subheader("Data Management")
    if st.session_state.get("history_cleared_msg"):
        st.success(st.session_state["history_cleared_msg"])
        st.session_state["history_cleared_msg"] = False
        
    has_active = has_any_valid_analysis()
    col_dm1, col_dm2 = st.columns(2)
    with col_dm1:
        render_html([
            '<div class="iq-card" style="margin-bottom: 10px;">',
            f'<div style="font-weight: 700; color: {theme["text_primary"]}; font-size: 14px; margin-bottom: 6px;">📄 Active Workspace Analysis</div>',
            f'<p style="color: {theme["text_muted"]}; font-size: 13px; line-height: 1.5; margin-bottom: 4px;">',
            f'{"Active manuscript loaded in workspace." if has_active else "No active manuscript currently loaded in memory."} Resetting clears the current paper, scores, executive summary, and chat without deleting saved database history.',
            '</p>',
            '</div>',
        ])
        if st.button("🗑️ Clear Active Analysis", key="btn_clear_active_only", disabled=not has_active, use_container_width=True):
            clear_active_analysis_session()
            st.session_state["history_cleared_msg"] = "✓ Active workspace analysis has been cleared."
            st.rerun()

    with col_dm2:
        render_html([
            '<div class="iq-card" style="margin-bottom: 10px;">',
            f'<div style="font-weight: 700; color: {theme["text_primary"]}; font-size: 14px; margin-bottom: 6px;">📚 Saved History Storage</div>',
            f'<p style="color: {theme["text_muted"]}; font-size: 13px; line-height: 1.5; margin-bottom: 4px;">',
            f'<b>{len(user_history)}</b> archived paper evaluations in database. Clearing history permanently removes all archived score records and PDF snapshots for account <b>{safe_text(username)}</b>.',
            '</p>',
            '</div>',
        ])
        if st.button("🗑️ Clear All Analysis History", key="btn_clear_all_hist", disabled=len(user_history) == 0, use_container_width=True):
            st.session_state["confirm_clear_all"] = True
            st.rerun()
        
    if st.session_state.get("confirm_clear_all", False):
        render_html([
            f'<div style="background: {theme["error_bg"]}; border: 1px solid {theme["error_border"]}; border-radius: 8px; padding: 14px 18px; margin: 12px 0;">',
            f'<div style="color: {theme["error_text"]}; font-weight: 700; font-size: 14px; margin-bottom: 6px;">⚠ Clear Analysis History</div>',
            f'<div style="color: {theme["text_secondary"]}; font-size: 13px; line-height: 1.6;">',
            'This will permanently delete:',
            '<ul style="margin: 6px 0 0 18px; padding: 0;">',
            f'<li>Archived analysis records for account <b>{safe_text(username)}</b></li>',
            '<li>Stored quality scores &amp; metrics</li>',
            '<li>Generated PDF report snapshots</li>',
            '<li>Historical evaluations and audit metadata</li>',
            '</ul>',
            '</div>',
            '</div>',
        ])
        clear_active_too = st.checkbox(
            "Also reset active analysis and uploaded document from workspace",
            value=True,
            key="chk_clear_active_settings"
        )
        c_cancel, c_confirm, _ = st.columns([1, 1.8, 4])
        with c_cancel:
            if st.button("Cancel", key="btn_cancel_clear_all"):
                st.session_state["confirm_clear_all"] = False
                st.rerun()
        with c_confirm:
            if st.button("Delete Everything", type="primary", key="btn_confirm_clear_all_exec"):
                success, msg = auth.clear_user_history(username)
                if success:
                    st.session_state.analysis_history[username] = []
                    if clear_active_too:
                        clear_active_analysis_session()
                    st.session_state["confirm_clear_all"] = False
                    st.session_state["history_cleared_msg"] = "✓ Analysis history and active workspace analysis cleared successfully." if clear_active_too else "✓ Analysis history cleared. Active manuscript retained in current session."
                    st.rerun()
                else:
                    st.error(f"Failed to clear history: {msg}")
                
    render_html("<br>")
    st.subheader("Privacy Information & Data Processing Architecture")
    render_html([
        '<div class="iq-card">',
        '<p><b>What Data is Stored:</b> PaperIQ stores only the following information scoped to your authenticated account in the local SQLite database (<code>users.db</code>) and active session state:</p>',
        '<ul style="margin: 6px 0 12px 18px; line-height: 1.65;">',
        '<li><b>Account Identity:</b> <code>username</code>, hashed credentials, and assigned academic <code>role</code> (Student, Researcher, or Administrator).</li>',
        '<li><b>Uploaded Paper Metadata:</b> original filename, detected display title, inferred research domain, and SHA-256 document hash.</li>',
        '<li><b>Extracted Text &amp; Evaluation Artifacts:</b> cleaned extracted text (where retained in saved analysis JSON snapshots or active session memory), composite and 5-dimension <code>scores</code>, structured writing <code>feedback</code>, and section summaries.</li>',
        '<li><b>Q&amp;A Session State:</b> Ask PaperIQ <code>chat history</code> (retained transiently in active session memory only and automatically cleared when switching papers or logging out).</li>',
        '<li><b>Audit &amp; Report Records:</b> analysis <code>timestamps</code> and downloadable PDF <code>generated reports</code>.</li>',
        '</ul>',
        '<p><b>How to Delete Your Data:</b></p>',
        '<ul style="margin: 6px 0 12px 18px; line-height: 1.65;">',
        '<li><b>Clear Active Session Memory:</b> Click <i>🗑️ Clear Active Analysis</i> above (or on the Analyze / Evaluation pages) to immediately purge active extracted text, scores, summaries, and chat history from memory.</li>',
        '<li><b>Delete Individual Saved Records:</b> Navigate to <i>📚 My History / Research History</i> and click <i>Delete</i> on any individual record (with confirmation) to remove its database row, extracted text snapshot, scores, and generated PDF report.</li>',
        '<li><b>Purge All Saved History:</b> Click <i>🗑️ Clear All Analysis History</i> above (or <i>Clear All</i> on the History page) and confirm to permanently delete all stored analyses, scores, feedback, timestamps, and generated PDF reports for your account.</li>',
        '</ul>',
        '<p><b>In-Memory Linguistic Processing &amp; Zero External AI Transmission:</b> All document text extraction (PDF, DOCX, TXT), statistical computational linguistics, scoring algorithms, and question-answering routines execute locally using deterministic NLP models (TextBlob, TF-IDF vectorization, cosine similarity). No manuscript text or user queries are transmitted to external generative AI APIs.</p>',
        '</div>',
    ])

# ----------------------------------------------------
# PAGE 9: ABOUT & METHODOLOGY
# ----------------------------------------------------
elif st.session_state.current_page in ["ℹ About", "ℹ️ About"]:
    render_html([
        '<div style="margin-bottom: 20px;">',
        f'<h1 style="margin: 0; font-size: 28px; color: {theme["text_primary"]};">About PaperIQ</h1>',
        f'<p style="margin: 4px 0 0 0; color: {theme["text_muted"]}; font-size: 15px;">AI-Powered Research Paper Analysis &amp; Evaluation System.</p>',
        '</div>',
    ])
    
    render_html([
        '<div class="iq-card">',
        '<h4>What is PaperIQ?</h4>',
        '<p>PaperIQ is an intelligent academic paper evaluation platform that analyzes research manuscripts for language quality, structural coherence, argumentation density, academic style, and readability. It empowers students, researchers, faculty, and peer evaluators to identify structural gaps, understand manuscript writing quality, and generate comprehensive review reports.</p>',
        '</div>',
    ])
    
    render_html([
        '<div class="iq-card">',
        '<h4>Role-Tailored Research Experiences</h4>',
        '<p>PaperIQ provides three purpose-built workspaces configured to distinct academic workflows:</p>',
        '<div style="display: grid; grid-template-columns: repeat(auto-fit, minmax(260px, 1fr)); gap: 14px; margin-top: 14px;">',
        f'<div style="background: {theme["bg_card_inner"]}; border: 1px solid {theme["border_color"]}; border-radius: 8px; padding: 14px 16px;">',
        f'<div style="font-weight: 700; color: {theme["text_primary"]}; margin-bottom: 6px;">🎓 Student Workspace</div>',
        f'<p style="font-size: 13px; color: {theme["text_muted"]}; margin: 0; line-height: 1.5;">Focused on understanding academic papers, learning scholarly terminology, improving thesis and paper drafting, receiving actionable sentence-level writing tips, and tracking writing quality across five core dimensions.</p>',
        '</div>',
        f'<div style="background: {theme["bg_card_inner"]}; border: 1px solid {theme["border_color"]}; border-radius: 8px; padding: 14px 16px;">',
        f'<div style="font-weight: 700; color: {theme["text_primary"]}; margin-bottom: 6px;">🔬 Researcher Workspace</div>',
        f'<p style="font-size: 13px; color: {theme["text_muted"]}; margin: 0; line-height: 1.5;">Engineered for deep manuscript intelligence, comparative multi-paper evaluation, objective-methodology-finding synthesis, evidence-grounded research Q&amp;A, and pre-submission peer-review preparation.</p>',
        '</div>',
        f'<div style="background: {theme["bg_card_inner"]}; border: 1px solid {theme["border_color"]}; border-radius: 8px; padding: 14px 16px;">',
        f'<div style="font-weight: 700; color: {theme["text_primary"]}; margin-bottom: 6px;">🛡️ Administrator Workspace</div>',
        f'<p style="font-size: 13px; color: {theme["text_muted"]}; margin: 0; line-height: 1.5;">Institutional administration portal featuring live database metrics, role governance and user management, complete login audit logs, an institutional saved papers repository, and document upload history.</p>',
        '</div>',
        '</div>',
        '</div>',
    ])
    
    render_html([
        '<div class="iq-card">',
        '<h4>Analytical Scoring Framework</h4>',
        '<p>PaperIQ evaluates papers across five weighted quality dimensions to calculate a transparent composite score (0 – 100):</p>',
        f'<div style="background: {theme["bg_card_inner"]}; border: 1px solid {theme["border_color"]}; border-radius: 8px; padding: 14px 18px; margin: 12px 0 16px 0; font-size: 13px; color: {theme["accent_purple"]}; line-height: 1.8;">',
        '<b>Overall Writing Score =</b><br>',
        '&nbsp;&nbsp;Language Quality × 25%<br>',
        '+ Structural Coherence × 25%<br>',
        '+ Argumentation × 20%<br>',
        '+ Academic Style × 15%<br>',
        '+ Readability × 15%',
        '</div>',
        '<ul>',
        '<li><b>Language Quality (25%):</b> Evaluates syntactic cadence, sentence length calibration, and formal academic phrasing.</li>',
        '<li><b>Structural Coherence (25%):</b> Measures the density and placement of transitional discourse markers between paragraphs and sections.</li>',
        '<li><b>Argumentation (20%):</b> Evaluates causal links and evidentiary propositions supporting analytical assertions.</li>',
        '<li><b>Academic Style (15%):</b> Measures lexical richness, technical terminology, and scholarly vocabulary diversity.</li>',
        '<li><b>Readability (15%):</b> Evaluates reading ease calibrated for academic literature and peer-reviewed manuscripts.</li>',
        '</ul>',
        '</div>',
    ])
    
    render_html([
        '<div class="iq-card">',
        '<h4>Limitations</h4>',
        '<ul style="margin: 6px 0 0 18px; line-height: 1.7;">',
        '<li><b>Heuristic Evaluation Only (Not Peer Review):</b> All composite and dimension scores are deterministic computational heuristics based on surface lexical, structural, and syntactic signals. They are <b>not a substitute for expert human peer review</b> and cannot verify mathematical proofs, empirical truth, experimental reproducibility, or scientific novelty.</li>',
        '<li><b>Domain &amp; Language Calibration Limits:</b> Scoring rubrics, discourse transition lexicons, and readability indices are calibrated for standard English-language academic prose. Non-English manuscripts, formula-heavy theoretical notes, or unconventional humanities formats may yield less reliable scores.</li>',
        '<li><b>Document Extraction Constraints:</b> Text extraction relies on embedded digital text streams. Extraction can fail on scanned or image-only PDFs without an OCR text layer, heavily corrupted files, or encrypted/password-protected documents.</li>',
        '</ul>',
        '</div>',
    ])

    render_html([
        '<div class="iq-card">',
        '<h4>Responsible Use</h4>',
        '<p>PaperIQ is designed as an educational and pre-submission writing-assistance tool for students and researchers. Users should treat all scores, extracted summaries, and revision recommendations as supportive diagnostic feedback rather than definitive editorial or grading judgments. Always verify extracted claims and citations directly against the primary manuscript text.</p>',
        '</div>',
    ])
    
    render_html([
        '<div class="iq-card">',
        '<h4>Technology Stack Architecture</h4>',
        '<p><b>Application Core:</b> Python, Streamlit</p>',
        '<p><b>NLP &amp; Computational Linguistics:</b> TextBlob, Scikit-learn (TF-IDF &amp; Cosine Similarity), NLTK</p>',
        '<p><b>Document Processing:</b> pdfplumber, python-docx, ReportLab, Kaleido</p>',
        '<p><b>Data Visualization:</b> Plotly Graph Objects (Radar &amp; Comparative Charts)</p>',
        '<p><b>Storage &amp; Security:</b> Local SQLite database, SHA-256 password hashing</p>',
        '</div>',
    ])
